"""Run folder: log file, Excel report, error screenshots."""
from __future__ import annotations

import logging
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Outcome, TicketResult

FILLS = {
    Outcome.DONE: "C6EFCE", Outcome.REJECTED: "FFEB9C", Outcome.SKIPPED: "D9D9D9",
    Outcome.PLANNED: "DDEBF7", Outcome.ERROR: "FFC7CE",
}
COLUMNS = [
    ("#", 5), ("Time", 9), ("Result", 10), ("Ticket type", 34), ("Type", 11), ("Request date", 18),
    ("User id", 10), ("User name", 26), ("Branch (portal)", 22), ("Branch id", 9), ("Job code", 11),
    ("Category", 9), ("User in CEM", 11), ("Action", 22), ("Decision", 14), ("Reason", 14),
    ("Details", 70), ("Steps done", 60), ("Seconds", 8), ("Screenshot", 30),
]


def setup_logging(run_dir: Path, verbose: bool = False) -> logging.Logger:
    log = logging.getLogger("cem_bot")
    log.setLevel(logging.DEBUG)
    log.handlers.clear()
    fmt = logging.Formatter("%(asctime)s %(levelname)-7s %(message)s", "%H:%M:%S")
    fh = logging.FileHandler(run_dir / "run.log", encoding="utf-8")
    fh.setLevel(logging.DEBUG)
    fh.setFormatter(fmt)
    log.addHandler(fh)
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    ch = logging.StreamHandler(sys.stdout)
    ch.setLevel(logging.DEBUG if verbose else logging.INFO)
    ch.setFormatter(fmt)
    log.addHandler(ch)
    log.propagate = False
    return log


def new_run_dir(base: Path, dry_run: bool) -> Path:
    stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    d = base / (stamp + ("_dry-run" if dry_run else ""))
    (d / "screenshots").mkdir(parents=True, exist_ok=True)
    return d


class Report:
    def __init__(self, run_dir: Path, meta: dict):
        self.run_dir = run_dir
        self.meta = meta
        self.results: list[TicketResult] = []
        self.path = run_dir / "report.xlsx"

    def add(self, r: TicketResult):
        self.results.append(r)
        self.save()

    def counts(self) -> dict[str, int]:
        out = {o.value: 0 for o in Outcome}
        for r in self.results:
            out[r.outcome.value] += 1
        return out

    def save(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Tickets"
        ws.sheet_view.rightToLeft = False
        head_fill = PatternFill("solid", fgColor="1F4E78")
        for i, (name, width) in enumerate(COLUMNS, start=1):
            c = ws.cell(row=1, column=i, value=name)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = width
        for n, r in enumerate(self.results, start=1):
            t, p = r.ticket, r.plan
            row = [n, r.started, r.outcome.value, t.ticket_type_raw, t.ticket_type.value if t.ticket_type else "",
                   t.request_date, t.user_id, t.user_name, t.branch_raw, p.branch_id if p else "",
                   p.job_code if p else "", p.category if p else "",
                   "" if r.user_existed is None else ("yes" if r.user_existed else "no"),
                   p.action.value if p else "", p.decision if p and p.decision else "",
                   p.reason if p and p.reason else "", r.message, " > ".join(r.steps), round(r.seconds, 1),
                   r.screenshot]
            for i, v in enumerate(row, start=1):
                c = ws.cell(row=n + 1, column=i, value=v)
                c.alignment = Alignment(vertical="top", wrap_text=i in (17, 18))
            ws.cell(row=n + 1, column=3).fill = PatternFill("solid", fgColor=FILLS[r.outcome])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(1, len(self.results) + 1)}"

        s = wb.create_sheet("Summary")
        s.column_dimensions["A"].width = 24
        s.column_dimensions["B"].width = 60
        rows = list(self.meta.items()) + [("", "")] + [(k, v) for k, v in self.counts().items()]
        for i, (k, v) in enumerate(rows, start=1):
            s.cell(row=i, column=1, value=k).font = Font(bold=True)
            s.cell(row=i, column=2, value=str(v))
        wb.save(self.path)
