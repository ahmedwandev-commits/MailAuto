"""Generate the NBE Operations Hub user guides (EN + AR) and the employee quick card as .docx, into docs/out/.

    python docs/build/make_guides.py            (needs: pip install python-docx)

Pure python-docx (no node, no LibreOffice, no COM) so nothing real on the machine is touched - this is the generator
used on a Windows PC that does not have the Node/LibreOffice toolchain. The guides describe the Hub pages (Home, Email,
SLA, ITSM, My work, System) and the collector setup. Fake/example values only. The ZIP builder picks up docs/out/*.docx
as guides/. (The older Node generator - guide_en.js / guide_ar.js / card.js - built the Call Log guides.)"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches

ROOT = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "out"
OUT.mkdir(parents=True, exist_ok=True)

GREEN = RGBColor(0x0B, 0x6E, 0x4F)
GOLD = RGBColor(0x9A, 0x7A, 0x2E)
GREY = RGBColor(0x55, 0x55, 0x60)
NOTE_FILL = "EAF3EE"
WARN_FILL = "FBECEC"
TIP_FILL = "FEF6E0"
CODE_FILL = "F0F0F3"
HEAD_FILL = "0B6E4F"


def _set(el, tag, **attrs):
    child = OxmlElement(tag)
    for k, v in attrs.items():
        child.set(qn(k), v)
    el.append(child)
    return child


def _shade(paragraph, fill):
    _set(paragraph.paragraph_format.element.get_or_add_pPr(), "w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill})


def _border(paragraph, side, color, size="18"):
    pPr = paragraph.paragraph_format.element.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        pPr.append(pbdr)
    _set(pbdr, f"w:{side}", **{"w:val": "single", "w:sz": size, "w:space": "6", "w:color": color})


def _rtl_para(paragraph):
    _set(paragraph.paragraph_format.element.get_or_add_pPr(), "w:bidi")


def _rtl_run(run):
    _set(run._element.get_or_add_rPr(), "w:rtl")


class Guide:
    def __init__(self, rtl=False, cs_font="Arial"):
        self.doc = Document()
        self.rtl = rtl
        self.cs_font = cs_font
        self._base_styles()

    def _base_styles(self):
        d = self.doc
        for section in d.sections:
            section.page_width, section.page_height = Inches(8.5), Inches(11)
            section.left_margin = section.right_margin = Inches(0.85)
            section.top_margin = section.bottom_margin = Inches(0.8)
            if self.rtl:
                _set(section._sectPr, "w:bidi")
        normal = d.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        rpr = normal.element.get_or_add_rPr()
        _set(rpr, "w:rFonts", **{"w:ascii": "Calibri", "w:hAnsi": "Calibri", "w:cs": self.cs_font})
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.12
        for lvl, sz, col in (("Heading 1", 18, GREEN), ("Heading 2", 14, GREEN), ("Heading 3", 12, GOLD)):
            st = d.styles[lvl]
            st.font.size, st.font.color.rgb, st.font.bold = Pt(sz), col, True
            st.font.name = "Calibri"
            _set(st.element.get_or_add_rPr(), "w:rFonts", **{"w:ascii": "Calibri", "w:hAnsi": "Calibri", "w:cs": self.cs_font})

    def _p(self, paragraph):
        if self.rtl:
            _rtl_para(paragraph)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for r in paragraph.runs:
                _rtl_run(r)
        return paragraph

    def heading(self, text, level=1, page_break=False):
        if page_break:
            self.doc.add_page_break()
        p = self.doc.add_heading(level=level)
        run = p.add_run(text)
        if self.rtl:
            _rtl_run(run)
            _rtl_para(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        return p

    def para(self, text, bold=False, italic=False, color=None, size=None):
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        run.bold, run.italic = bold, italic
        if color:
            run.font.color.rgb = color
        if size:
            run.font.size = Pt(size)
        return self._p(p)

    def bullets(self, items, numbered=False):
        style = "List Number" if numbered else "List Bullet"
        for it in items:
            p = self.doc.add_paragraph(style=style)
            run = p.add_run(it)
            self._p(p)
            if self.rtl:
                _rtl_run(run)

    def box(self, label, text, fill):
        p = self.doc.add_paragraph()
        lead = p.add_run(f"{label}  ")
        lead.bold = True
        body = p.add_run(text)
        _shade(p, fill)
        _border(p, "left", HEAD_FILL if fill == NOTE_FILL else ("C0392B" if fill == WARN_FILL else "B8860B"))
        p.paragraph_format.space_before = Pt(4)
        if self.rtl:
            _rtl_para(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _rtl_run(lead)
            _rtl_run(body)
        return p

    def note(self, text):
        self.box("Note" if not self.rtl else "ملاحظة", text, NOTE_FILL)

    def warn(self, text):
        self.box("Important" if not self.rtl else "مهم", text, WARN_FILL)

    def tip(self, text):
        self.box("Tip" if not self.rtl else "نصيحة", text, TIP_FILL)

    def code(self, text):
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = "Consolas"
        run.font.size = Pt(10)
        _set(run._element.get_or_add_rPr(), "w:rFonts", **{"w:ascii": "Consolas", "w:hAnsi": "Consolas"})
        _shade(p, CODE_FILL)
        p.paragraph_format.left_indent = Inches(0.2)
        return p

    def table(self, headers, rows, widths=None):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        if self.rtl:
            _set(t._tbl.tblPr, "w:bidiVisual")
        hdr = t.rows[0].cells
        for i, htext in enumerate(headers):
            cell = hdr[i]
            p = cell.paragraphs[0]
            run = p.add_run(htext)
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _shade(p, HEAD_FILL)
            _set(cell._tc.get_or_add_tcPr(), "w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": HEAD_FILL})
            if self.rtl:
                _rtl_para(p); _rtl_run(run); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                p = cells[i].paragraphs[0]
                run = p.add_run(val)
                if self.rtl:
                    _rtl_para(p); _rtl_run(run); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Inches(w)
        self.doc.add_paragraph()
        return t

    def cover(self, title, subtitle, contents, version):
        for _ in range(2):
            self.doc.add_paragraph()
        t = self.doc.add_paragraph()
        r = t.add_run(title); r.bold = True; r.font.size = Pt(30); r.font.color.rgb = GREEN
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if self.rtl:
            _rtl_para(t); _rtl_run(r)
        s = self.doc.add_paragraph()
        rs = s.add_run(subtitle); rs.font.size = Pt(16); rs.font.color.rgb = GREY
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if self.rtl:
            _rtl_para(s); _rtl_run(rs)
        self.doc.add_paragraph()
        c = self.doc.add_paragraph()
        rc = c.add_run("  ·  ".join(contents)); rc.font.size = Pt(11); rc.italic = True
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if self.rtl:
            _rtl_para(c); _rtl_run(rc)
        for _ in range(2):
            self.doc.add_paragraph()
        v = self.doc.add_paragraph()
        rv = v.add_run(version); rv.font.size = Pt(10); rv.font.color.rgb = GREY
        v.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if self.rtl:
            _rtl_para(v); _rtl_run(rv)
        self.doc.add_page_break()

    def save(self, name):
        path = OUT / name
        self.doc.save(str(path))
        return path


SERVER = r"\\10.104.46.34\calllog"


# ===================================================================================== ENGLISH
def build_en():
    g = Guide(rtl=False)
    g.cover("NBE Operations Hub", "User Guide",
            ["Part 1 - Setup", "Part 2 - Employee", "Part 3 - Admin & Manager", "Part 4 - Troubleshooting"],
            "Version 1.0  ·  October 2026")

    g.heading("About the Operations Hub", 1)
    g.para("The NBE Operations Hub follows the whole sector from one place. It brings four things together in one "
           "application that runs from the shared folder, next to the Call Log you already use:")
    g.table(["Module", "What it shows"], [
        ["Calls", "Incoming calls each employee logs (the existing Call Log)."],
        ["Email", "Which business e-mails got a team reply, the SLA, teams and reply quality."],
        ["SLA", "Security requests waiting at the \"security user\" for more than 1 hour (reminder) / 2 hours (breach)."],
        ["ITSM", "Implementation and Rollback request counts per role (from CA SDM)."],
    ], widths=[1.4, 5.3])
    g.heading("Who sees what", 2)
    g.table(["Role", "Home page", "Can do"], [
        ["Administrator", "/hub (Admin dashboard)", "Everything: all modules, all teams, settings, collectors, users, audit, \"view as employee\"."],
        ["Manager", "/hub", "The same dashboards for their own team(s); assign e-mails, mark no-action, change team, review reply quality. No settings or users."],
        ["Employee", "/my (My work)", "Only what is theirs: their calls, e-mails assigned to/answered by them, their team(s)' e-mails and their security requests. Logs calls on /dashboard."],
    ], widths=[1.3, 1.8, 3.6])
    g.para("An employee can belong to more than one team and sees the union. An employee never sees an \"Unclassified\" "
           "e-mail. Admin and employee pages share one \"Team status\" strip with the same counts.")
    g.heading("How the parts fit together", 2)
    g.para("The collector runs on a Windows PC that stays logged in (Outlook / Chrome). It writes plain JSON files into "
           "the shared folder. The web app reads those files and lets people act - it never talks to Outlook or the portals.")
    g.code("collector.py  (Outlook / SSO portal / CA SDM)  --writes-->  " + SERVER + r"\data\ops\*.json  <--reads--  run.py (web app)")
    g.note("No database and no files that end in .bat / .cmd / .ps1 / .vbs / .exe. Everything starts with "
           "\"python run.py\" or \"python collector.py\".")

    # --------------------------------------------------------------- PART 1
    g.heading("Part 1 - Setup", 1, page_break=True)
    g.para("Two machines are involved: the web-server PC (anyone opens the site in a browser) and the collector PC "
           "(the one that has Outlook and the browser for the portals). They can be the same PC.")
    g.heading("A. The web server", 2)
    g.bullets([
        "Copy the contents of the \"program\" folder (from NBE-CallLog.zip) into the shared folder " + SERVER + ". The file run.py must sit directly inside it.",
        "Install Python 3.9 or newer on the PC from python.org (tick \"Add python.exe to PATH\").",
        "First start: open Command Prompt and run the command below. The first start downloads Flask into a \"libs\" folder (internet needed once).",
    ], numbered=True)
    g.code("python " + SERVER + r"\run.py --share")
    g.bullets([
        "Write down the administrator password it prints once (User ID: admin). You choose a new one at first sign-in.",
        "Allow python.exe through the Windows firewall, or ask IT to open TCP port 8080.",
        "To run it automatically after a reboot, use Task Scheduler with: python run.py --service  (it keeps the server up and restarts it if it stops or when the files are updated).",
    ])
    g.tip("A quick health check without signing in: open  http://<server>:8080/healthz  - it returns the collectors' online state as plain JSON. A monitoring tool can poll it.")
    g.heading("B. The collector (Outlook + portals)", 2)
    g.para("The collector PC must stay logged in (a locked screen is fine) because Outlook and the browser need an "
           "interactive Windows session. Install the collector packages on first run - they download automatically.")
    g.bullets([
        "Copy collector_config.example.json to collector_config.json (same folder as collector.py) and open it in Notepad.",
        "Fill in the user name and password of the SSO portal (sla) and CA SDM (itsm), and the Outlook archive name for the e-mail scan:",
    ], numbered=True)
    g.code('{\n  "sla":   {"username": "1001", "password": "..."},\n  "itsm":  {"username": "1001", "password": "..."},\n  "email": {"target_archive_name": "2026_Part3"}\n}')
    g.bullets([
        "Instead of writing a password in the file you can run:  python collector.py --set-password sla   (it is kept in Windows Credential Manager).",
        "Check everything is in place:  python collector.py --check",
        "Start it in the background (no window), restarted automatically:  python collector.py --start   - or keep it in a window with:  python collector.py --service",
    ])
    g.warn("collector_config.json holds passwords. Never e-mail it, never put it in the ZIP or in git. Only the "
           "collector reads it - the web server never does. After moving a real password into it, change that password once.")
    g.note("A wrong or expired password pauses that job's scheduled runs (so the account is not locked out). Fix the "
           "password, or press \"Run now\" on the System page, to resume.")

    # --------------------------------------------------------------- PART 2
    g.heading("Part 2 - Employee guide (My work)", 1, page_break=True)
    g.para("After you sign in, you land on My work (/my). It shows only your own work. Call logging stays on the "
           "\"Log calls\" page (/dashboard) - the link is in the top bar.")
    g.heading("Team status strip", 2)
    g.para("At the top is the shared \"Team status\" strip: calls today, your team's e-mails pending, security requests "
           "open, ITSM totals and how fresh the data is. These are counts only - the same numbers your administrator sees.")
    g.heading("My calls", 2)
    g.para("Today / yesterday / last 7 days totals, split into Resolved, Transferred and Ticket inquiry, a small chart "
           "of your calls per hour today, and your share of the team's calls.")
    g.heading("My e-mails", 2)
    g.bullets([
        "Assigned to me: e-mails a manager assigned to you, pending first, with the waiting time and the SLA colour.",
        "Replied today: the e-mails you answered today, the average first-reply time and how many met / breached the SLA.",
        "My team's e-mails: the open e-mails of your team(s) - never the Unclassified ones.",
    ])
    g.heading("My quality reviews", 2)
    g.para("Ratings a manager gave your replies (Excellent / Good / Needs improvement), with the reasons. If you "
           "disagree with a rating you can press Appeal and explain why; the decision shows here once a manager decides.")
    g.heading("My security requests", 2)
    g.para("Security requests open under your name (matched by the \"security name\" on your account), with the live age "
           "and the state. If nothing shows and you handle security requests, ask your administrator to set your security name.")
    g.note("If a section says your e-mail address or security name is not set, ask your administrator to add it to your account.")

    # --------------------------------------------------------------- PART 3
    g.heading("Part 3 - Administrator & manager guide", 1, page_break=True)
    g.heading("Home (Admin dashboard)", 2)
    g.para("/hub shows the whole sector at a glance: a \"needs attention\" list (newest and most serious first), the "
           "four module cards with headline numbers, collector health, today's timeline of what changed, and the "
           "e-mails-by-team table. Managers see the same, scoped to their team. \"View as employee\" opens any "
           "employee's My work page read-only.")
    g.heading("Email", 2)
    g.para("/email has six tabs:")
    g.table(["Tab", "What it is for"], [
        ["Overview", "KPI tiles, a status donut, the per-team table, arrivals per hour, top applications, slowest replies and the escalated list."],
        ["Pending", "The working list. Filter by team, application, carry-over, escalated or assignment; search; click a row for details. Actions: Assign to a person, Mark no action, Change team."],
        ["Replied", "The answered e-mails with the first-reply time and SLA. Managers can Review the reply quality (rating + reasons + note)."],
        ["Quality", "Per-person scorecard, team quality and the review counts."],
        ["Trends", "Replied / pending / past-SLA per day over the last days."],
        ["Unclassified", "E-mails no team could be assigned to. Assign a team to one mail, or Create a rule (keyword to team) that applies from the next scan. The current rules are listed with Remove."],
    ], widths=[1.3, 5.4])
    g.note("Unclassified e-mails are tuned in System > Settings > Classification: keyword-to-team rules, mailbox-to-team "
           "rules, and where mails with no detected team go. An employee never sees an Unclassified e-mail.")
    g.heading("SLA (security requests)", 2)
    g.para("/sla shows the open security requests: totals, per-security-user table, the request list with the live age "
           "and status (within / reminder / breached), and today's timeline. \"Run scan now\" asks the collector for a "
           "fresh scan. The collector sends the reminder / breach e-mails, as the old script did.")
    g.heading("ITSM", 2)
    g.para("/itsm shows the Implementation and Rollback counts per role, what is new or closed since the last scan, and "
           "any role the collector could not read.")
    g.heading("System", 2)
    g.table(["Tab", "What it is for"], [
        ["Health", "Data folder (free disk, writable), snapshots, heartbeats, users by role and today's audit."],
        ["Collectors", "Each job's state (online / offline / running / login failed - paused), the accounts and pinned config keys (no passwords), and \"Run now\" per job."],
        ["Settings", "All module settings, validated. The Classification editor is here. Keys set in collector_config.json show as read-only (pinned)."],
        ["Logs", "The tail of each collector job log."],
        ["Reports", "The Excel / HTML reports the collector produced, to download."],
    ], widths=[1.3, 5.4])
    g.heading("Users and teams", 2)
    g.para("On the Employees tab you add or edit accounts. Pick the role (employee / manager / administrator), the "
           "e-mail address (links replies and assignments), the security name (links SLA requests) and tick the "
           "team(s). An account can hold several teams; a manager may follow \"All teams\". You can import users from a "
           "CSV/Excel file with a \"teams\" column (names separated by ;).")

    # --------------------------------------------------------------- PART 4
    g.heading("Part 4 - Troubleshooting", 1, page_break=True)
    g.table(["Symptom", "What to do"], [
        ["A collector shows Offline", "Check the collector PC is on and logged in and collector.py is running (python collector.py --check). The heartbeat must be fresh."],
        ["\"Login failed - retries paused\"", "The SSO or ITSM password is wrong or expired. Fix it in collector_config.json (or --set-password), then press Run now."],
        ["No e-mail data", "The e-mail scan needs Outlook on the collector PC (pywin32). Run python collector.py --check; the archive name must match an Outlook folder."],
        ["The page says data is old", "The collector is behind or outside working hours. Outside the active window the last scan is shown on purpose."],
        ["Forgot the admin password", "On the server PC: python manage.py reset-password admin  (prints a new temporary password)."],
        ["Other PCs cannot open the site", "Start with --share and open TCP port 8080 in the firewall. Use the address the server window prints."],
    ], widths=[2.3, 4.4])
    g.note("Everything is stored as files under " + SERVER + r"\data. To back up, copy that folder. Reports, history and "
           "logs are pruned automatically (housekeeping runs once a day outside business hours).")
    g.save("NBE-OpsHub-User-Guide-EN.docx")


# ===================================================================================== ARABIC
def build_ar():
    g = Guide(rtl=True)
    g.cover("مركز عمليات البنك الأهلي", "دليل المستخدم",
            ["الجزء 1 - التركيب", "الجزء 2 - الموظف", "الجزء 3 - الأدمن والمدير", "الجزء 4 - حل المشاكل"],
            "الإصدار 1.0  ·  أكتوبر 2026")

    g.heading("عن مركز العمليات", 1)
    g.para("مركز عمليات البنك الأهلي بيتابع القطاع كله من مكان واحد. بيجمع أربع حاجات في برنامج واحد شغّال من الشير فولدر، "
           "جنب الـ Call Log اللي بتستخدمه دلوقتي:")
    g.table(["الموديول", "بيعرض إيه"], [
        ["المكالمات", "المكالمات الواردة اللي كل موظف بيسجلها (الـ Call Log الحالي)."],
        ["الإيميل", "أنهي إيميلات العمل اتردّ عليها من التيم، الـ SLA، التيمات، وجودة الرد."],
        ["SLA", "طلبات السرية اللي مستنية عند مستخدم السرية أكتر من ساعة (تذكير) / ساعتين (تجاوز)."],
        ["ITSM", "أعداد طلبات التنفيذ والـ Rollback لكل دور (من CA SDM)."],
    ], widths=[1.4, 5.3])
    g.heading("مين بيشوف إيه", 2)
    g.table(["الدور", "الصفحة الرئيسية", "بيعمل إيه"], [
        ["الأدمن", "‏/hub (لوحة الأدمن)", "كل حاجة: كل الموديولات، كل التيمات، الإعدادات، الـ collectors، اليوزرز، الـ audit، و\"اعرض كموظف\"."],
        ["المدير", "‏/hub", "نفس اللوحات لتيمه هو بس؛ يوزّع إيميلات، no-action، يغيّر التيم، يراجع جودة الرد. من غير إعدادات ولا يوزرز."],
        ["الموظف", "‏/my (شغلي)", "حاجته بس: مكالماته، الإيميلات المسندة له أو اللي ردّ عليها، إيميلات تيمه، وطلبات السرية بتاعته. بيسجّل المكالمات في ‏/dashboard."],
    ], widths=[1.1, 1.7, 3.9])
    g.para("الموظف ممكن يبقى في أكتر من تيم وبيشوف الاتنين. الموظف عمره ما بيشوف إيميل \"Unclassified\". صفحة الأدمن "
           "وصفحة الموظف بيشتركوا في شريط \"حالة التيم\" بنفس الأرقام.")
    g.heading("إزاي الأجزاء بتشتغل مع بعض", 2)
    g.para("الـ collector بيشتغل على جهاز ويندوز فاضل مفتوح (فيه Outlook و Chrome). بيكتب ملفات JSON بسيطة في الشير فولدر. "
           "الـ web app بيقرا الملفات دي ويخلّي الناس تتصرّف - وعمره ما بيكلّم Outlook ولا البوابات.")
    g.code("collector.py  (Outlook / SSO / CA SDM)  --writes-->  " + SERVER + r"\data\ops\*.json  <--reads--  run.py")
    g.note("مفيش داتابيز ومفيش أي ملف .bat / .cmd / .ps1 / .vbs / .exe. كل حاجة بتبدأ بـ \"python run.py\" أو \"python collector.py\".")

    g.heading("الجزء 1 - التركيب", 1, page_break=True)
    g.para("فيه جهازين: جهاز الـ web server (أي حد بيفتح الموقع من المتصفح) وجهاز الـ collector (اللي عليه Outlook والمتصفح "
           "للبوابات). ممكن يكونوا نفس الجهاز.")
    g.heading("أ. الـ web server", 2)
    g.bullets([
        "انسخ محتويات فولدر \"program\" (من NBE-CallLog.zip) جوّه الشير فولدر " + SERVER + ". لازم run.py يبقى جوّه على طول.",
        "نصّب Python 3.9 أو أحدث من python.org (علّم على \"Add python.exe to PATH\").",
        "أول تشغيل: افتح Command Prompt وشغّل الأمر اللي تحت. أول مرة بينزّل Flask (محتاج إنترنت مرة واحدة).",
    ], numbered=True)
    g.code("python " + SERVER + r"\run.py --share")
    g.bullets([
        "اكتب باسورد الأدمن اللي بيظهر مرة واحدة (اليوزر: admin). هتختار واحد جديد أول ما تسجّل دخول.",
        "اسمح لـ python.exe في الـ firewall، أو اطلب من الـ IT يفتح بورت TCP 8080.",
        "عشان يشتغل لوحده بعد الريستارت، استخدم Task Scheduler مع:  python run.py --service  (بيفضّل شغّال ويعيد نفسه لو وقف أو لو الملفات اتحدّثت).",
    ])
    g.tip("فحص سريع من غير تسجيل دخول: افتح  http://<server>:8080/healthz  - بيرجّع حالة الـ collectors كـ JSON. أي أداة مراقبة تقدر تسأله.")
    g.heading("ب. الـ collector (Outlook والبوابات)", 2)
    g.para("جهاز الـ collector لازم يفضل مسجّل دخول (الشاشة مقفولة عادي) عشان Outlook والمتصفح محتاجين جلسة ويندوز مفتوحة.")
    g.bullets([
        "انسخ collector_config.example.json لـ collector_config.json (نفس فولدر collector.py) وافتحه بالـ Notepad.",
        "املا اسم اليوزر والباسورد بتوع بوابة SSO (sla) و CA SDM (itsm)، واسم أركيف Outlook لفحص الإيميل:",
    ], numbered=True)
    g.code('{\n  "sla":   {"username": "1001", "password": "..."},\n  "itsm":  {"username": "1001", "password": "..."},\n  "email": {"target_archive_name": "2026_Part3"}\n}')
    g.bullets([
        "بدل ما تكتب الباسورد في الملف تقدر تشغّل:  python collector.py --set-password sla   (بيتحفظ في Windows Credential Manager).",
        "اتأكد إن كل حاجة تمام:  python collector.py --check",
        "شغّله في الخلفية من غير نافذة ويعيد نفسه تلقائيًا:  python collector.py --start   - أو خلّيه في نافذة بـ:  python collector.py --service",
    ])
    g.warn("ملف collector_config.json فيه باسوردات. متبعتهوش بإيميل، متحطّهوش في الـ ZIP ولا في git. الـ collector بس اللي "
           "بيقراه. بعد ما تنقل باسورد حقيقي فيه، غيّر الباسورد ده مرة.")
    g.note("الباسورد الغلط أو المنتهي بيوقف تشغيل الـ job المجدول (عشان الحساب ما يتقفلش). صلّح الباسورد أو اضغط \"Run now\" من صفحة System عشان يرجع.")

    g.heading("الجزء 2 - دليل الموظف (شغلي)", 1, page_break=True)
    g.para("بعد ما تسجّل دخول هتلاقي نفسك في \"شغلي\" (/my). بتعرض شغلك إنت بس. تسجيل المكالمات في صفحة \"Log calls\" (/dashboard) - اللينك في الشريط فوق.")
    g.heading("شريط حالة التيم", 2)
    g.para("فوق فيه شريط \"حالة التيم\" المشترك: مكالمات النهارده، إيميلات تيمك المعلّقة، طلبات السرية المفتوحة، إجماليات ITSM، "
           "وحداثة البيانات. دي أعداد بس - نفس الأرقام اللي الأدمن بيشوفها.")
    g.heading("مكالماتي", 2)
    g.para("إجماليات النهارده / إمبارح / آخر 7 أيام، مقسومة لـ محلولة / محوّلة / استعلام تذكرة، ورسم صغير لمكالماتك في الساعة النهارده، ونصيبك من مكالمات التيم.")
    g.heading("إيميلاتي", 2)
    g.bullets([
        "المسندة ليّا: الإيميلات اللي المدير سندها لك، المعلّق الأول، بزمن الانتظار ولون الـ SLA.",
        "ردّيت عليها النهارده: الإيميلات اللي رديت عليها النهارده، متوسط أول رد، وكام التزم / تجاوز الـ SLA.",
        "إيميلات تيمي: الإيميلات المفتوحة بتاعة تيمك - من غير الـ Unclassified أبدًا.",
    ])
    g.heading("مراجعات الجودة بتاعتي", 2)
    g.para("التقييمات اللي المدير حطّها على ردودك (ممتاز / جيد / يحتاج تحسين) بالأسباب. لو مش موافق على تقييم تقدر تضغط "
           "\"Appeal\" وتوضّح ليه؛ والقرار بيظهر هنا أول ما المدير يقرر.")
    g.heading("طلبات السرية بتاعتي", 2)
    g.para("طلبات السرية المفتوحة باسمك (بتتطابق بـ \"الاسم في السرية\" على حسابك)، بالعمر الحي والحالة. لو مفيش حاجة "
           "وإنت بتتعامل مع طلبات السرية، اطلب من الأدمن يحط الاسم ده على حسابك.")
    g.note("لو قسم قال إن إيميلك أو اسمك في السرية مش متظبط، اطلب من الأدمن يضيفه لحسابك.")

    g.heading("الجزء 3 - دليل الأدمن والمدير", 1, page_break=True)
    g.heading("الرئيسية (لوحة الأدمن)", 2)
    g.para("‏/hub بتعرض القطاع كله بنظرة: قائمة \"محتاج انتباه\" (الأحدث والأخطر الأول)، كروت الموديولات الأربعة بالأرقام، "
           "حالة الـ collectors، تايم‌لاين اللي حصل النهارده، وجدول الإيميلات بالتيم. المدير بيشوف نفس الحاجة بس على تيمه. "
           "\"اعرض كموظف\" بتفتح صفحة \"شغلي\" لأي موظف للقراءة بس.")
    g.heading("الإيميل", 2)
    g.para("‏/email فيها ست تابات:")
    g.table(["التاب", "بتستخدم في إيه"], [
        ["Overview", "كروت KPI، دونَت الحالة، جدول التيمات، الوصول بالساعة، أكتر التطبيقات، أبطأ الردود، وقائمة المصعّد."],
        ["Pending", "قائمة الشغل. فلتر بالتيم/التطبيق/المرحّل/المصعّد/الإسناد؛ بحث؛ اضغط صف للتفاصيل. أكشنز: Assign لشخص، No action، تغيير التيم."],
        ["Replied", "الإيميلات اللي اتردّ عليها بوقت أول رد والـ SLA. المدير يقدر يراجع جودة الرد (تقييم + أسباب + ملاحظة)."],
        ["Quality", "scorecard لكل شخص، جودة التيم، وأعداد المراجعات."],
        ["Trends", "ردّيت/معلّق/تجاوز SLA لكل يوم على مدار الأيام الأخيرة."],
        ["Unclassified", "إيميلات ماقدرش تتحدّد لتيم. اسند تيم لإيميل، أو اعمل قاعدة (كلمة → تيم) بتطبّق من الفحص الجاي. القواعد الحالية بزرار حذف."],
    ], widths=[1.3, 5.4])
    g.note("الـ Unclassified بتتظبط في System > Settings > Classification: قواعد كلمة→تيم، ميل‌بوكس→تيم، وفين تروح الإيميلات "
           "اللي مالهاش تيم. الموظف عمره ما بيشوف إيميل Unclassified.")
    g.heading("SLA (طلبات السرية)", 2)
    g.para("‏/sla بتعرض الطلبات المفتوحة: الإجماليات، جدول لكل مستخدم سرية، قائمة الطلبات بالعمر الحي والحالة "
           "(within / reminder / breached)، وتايم‌لاين النهارده. \"Run scan now\" بتطلب فحص جديد من الـ collector. "
           "الـ collector بيبعت إيميلات التذكير/التجاوز زي السكربت القديم.")
    g.heading("ITSM", 2)
    g.para("‏/itsm بتعرض أعداد التنفيذ والـ Rollback لكل دور، اللي جديد أو اتقفل من آخر فحص، وأي دور ماقدرش الـ collector يقراه.")
    g.heading("System", 2)
    g.table(["التاب", "بتستخدم في إيه"], [
        ["Health", "فولدر الداتا (المساحة، قابل للكتابة)، الـ snapshots، الـ heartbeats، اليوزرز بالدور، والـ audit النهارده."],
        ["Collectors", "حالة كل job (online/offline/running/login failed - paused)، الحسابات ومفاتيح الكونفج المثبّتة (من غير باسوردات)، و\"Run now\" لكل job."],
        ["Settings", "كل الإعدادات متحقّق منها. محرّر الـ Classification هنا. المفاتيح اللي في collector_config.json بتظهر للقراءة بس (مثبّتة)."],
        ["Logs", "آخر سطور لوج كل job."],
        ["Reports", "تقارير الإكسل/HTML اللي الـ collector عملها، للتنزيل."],
    ], widths=[1.3, 5.4])
    g.heading("اليوزرز والتيمات", 2)
    g.para("في تاب Employees بتضيف أو تعدّل الحسابات. اختار الدور (موظف/مدير/أدمن)، الإيميل (بيربط الردود والإسناد)، "
           "الاسم في السرية (بيربط طلبات SLA)، وعلّم التيم/التيمات. الحساب يقدر يبقى في أكتر من تيم؛ المدير ممكن يتابع \"All teams\". "
           "تقدر تستورد يوزرز من ملف CSV/Excel فيه عمود \"teams\" (الأسماء مفصولة بـ ;).")

    g.heading("الجزء 4 - حل المشاكل", 1, page_break=True)
    g.table(["المشكلة", "تعمل إيه"], [
        ["collector ظاهر Offline", "اتأكد إن جهاز الـ collector شغّال ومسجّل دخول و collector.py شغّال (python collector.py --check). الـ heartbeat لازم يكون حديث."],
        ["\"Login failed - retries paused\"", "باسورد SSO أو ITSM غلط أو منتهي. صلّحه في collector_config.json (أو --set-password)، وبعدين اضغط Run now."],
        ["مفيش بيانات إيميل", "فحص الإيميل محتاج Outlook على جهاز الـ collector (pywin32). شغّل python collector.py --check؛ اسم الأركيف لازم يطابق فولدر في Outlook."],
        ["الصفحة بتقول البيانات قديمة", "الـ collector متأخّر أو بره ساعات العمل. بره النافذة النشطة بيتعرض آخر فحص بالقصد."],
        ["نسيت باسورد الأدمن", "على جهاز السيرفر: python manage.py reset-password admin  (بيطبع باسورد مؤقت جديد)."],
        ["أجهزة تانية مش بتفتح الموقع", "شغّل بـ --share وافتح بورت TCP 8080 في الـ firewall. استخدم العنوان اللي نافذة السيرفر بتطبعه."],
    ], widths=[2.3, 4.4])
    g.note("كل حاجة متخزّنة كملفات تحت " + SERVER + r"\data. للنسخ الاحتياطي، انسخ الفولدر ده. التقارير والتاريخ واللوجات "
           "بتتنضّف تلقائيًا (housekeeping مرة في اليوم بره ساعات العمل).")
    g.save("NBE-OpsHub-User-Guide-AR.docx")


# ===================================================================================== QUICK CARD
def build_card():
    g = Guide(rtl=False)
    g.para("NBE Operations Hub - Employee Quick Card", bold=True, color=GREEN, size=16)
    g.para("Everything that is yours is on one page: My work (/my).", italic=True, color=GREY)
    g.doc.add_paragraph()
    g.heading("Every day", 2)
    g.bullets([
        "Sign in, then log each incoming call on \"Log calls\" (/dashboard): the IP phone, the branch (if a branch called) and the reason.",
        "Open \"My work\" (/my) to see your calls today, your e-mails and your security requests.",
    ])
    g.heading("On My work you see", 2)
    g.bullets([
        "My calls: today / yesterday / last 7 days, split by result.",
        "My e-mails: assigned to you, replied today, and your team's open e-mails.",
        "My quality reviews: your ratings - press Appeal if you disagree.",
        "My security requests: what is open under your name, with the live age.",
    ])
    g.note("You see only your own work and your team's e-mails - never another person's rows and never an Unclassified e-mail.")
    g.doc.add_paragraph()
    g.para("بطاقة سريعة للموظف", bold=True, color=GREEN, size=16)
    pcard = g.doc.paragraphs[-1]
    _rtl_para(pcard); pcard.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in pcard.runs:
        _rtl_run(run)

    def ar_line(text, bold=False):
        p = g.doc.add_paragraph()
        r = p.add_run(text); r.bold = bold
        _rtl_para(p); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; _rtl_run(r)
        return p
    ar_line("كل يوم:", bold=True)
    for t in ["سجّل دخول، وسجّل كل مكالمة واردة في \"Log calls\" (/dashboard): رقم الـ IP، الفرع (لو فرع اتصل)، والسبب.",
              "افتح \"شغلي\" (/my) تشوف مكالماتك النهارده، إيميلاتك، وطلبات السرية بتاعتك."]:
        p = g.doc.add_paragraph(style="List Bullet"); r = p.add_run(t); _rtl_para(p); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; _rtl_run(r)
    ar_line("في \"شغلي\" بتشوف: مكالماتي، إيميلاتي (المسندة ليّا، ردّيت عليها، وإيميلات تيمي)، مراجعات الجودة (اضغط Appeal لو مش موافق)، وطلبات السرية بتاعتي.")
    ar_line("بتشوف شغلك إنت بس وإيميلات تيمك - مش صفوف حد تاني ولا إيميل Unclassified أبدًا.")
    g.save("NBE-OpsHub-Employee-Quick-Card.docx")


if __name__ == "__main__":
    build_en()
    build_ar()
    build_card()
    for f in sorted(OUT.glob("*.docx")):
        print("wrote", f.name, f"({f.stat().st_size:,} bytes)")
