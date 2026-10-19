import io
import multiprocessing
import threading

import openpyxl
import pytest

from calllog.storage import Store
from conftest import Browser


# ------------------------------------------------------------------- login
def test_first_run_creates_admin_who_must_change_password(app):
    pw = app.config["INITIAL_ADMIN_PASSWORD"]
    assert pw and len(pw) >= 10
    b = Browser(app)
    resp = b.login("admin", pw)
    assert resp.status_code == 302
    assert b.get("/api/me").status_code == 403           # blocked until the password changes
    assert b.client.get("/admin").headers["Location"].endswith("/account/password")
    resp = b.client.post("/account/password", data={
        "current_password": pw, "new_password": "short", "confirm_password": "short",
        "csrf_token": b.csrf})
    assert resp.status_code == 400
    resp = b.client.post("/account/password", data={
        "current_password": pw, "new_password": "BrandNewPass1", "confirm_password": "BrandNewPass1",
        "csrf_token": b.csrf})
    assert resp.status_code == 302
    assert b.get("/api/me").get_json()["user"]["role"] == "admin"
    assert app.test_client().post("/login", data={"user_id": "admin", "password": pw}).status_code == 400


def test_wrong_password_is_rejected_and_throttled(app):
    b = Browser(app)
    for _ in range(5):
        assert b.login("admin", "nope").status_code == 401
    assert b.login("admin", app.config["INITIAL_ADMIN_PASSWORD"]).status_code == 429


def test_login_needs_csrf_token(app):
    resp = app.test_client().post("/login", data={"user_id": "admin", "password": "x"})
    assert resp.status_code == 400


def test_api_post_without_csrf_header_is_rejected(admin):
    resp = admin.client.post("/api/admin/users", json={"user_id": "1", "name": "x"})
    assert resp.status_code == 400


def test_anonymous_gets_401_and_redirect(app):
    c = app.test_client()
    assert c.get("/api/me").status_code == 401
    assert c.get("/admin").headers["Location"].endswith("/login")


# ---------------------------------------------------------------- records
def log_call(b, ip="75942", reason="solved", note="", branch=None):
    body = {"ip_phone": ip, "reason": reason, "note": note}
    if branch is not None:
        body["branch"] = branch
    return b.post("/api/records", body)


def test_employee_logs_edits_and_deletes_calls(make_employee, clock):
    emp = make_employee("230966", "Khaled")
    assert log_call(emp, "75942", "solved").status_code == 201
    clock.advance(minutes=30)
    assert log_call(emp, "٩٥٠٧٧", "answered").status_code == 201      # Arabic-Indic digits
    clock.advance(hours=1)
    rid = log_call(emp, "75071", "followup", "call back tomorrow").get_json()["record"]["id"]

    day = emp.get("/api/records").get_json()
    assert [r["ip_phone"] for r in day["records"]] == ["75071", "95077", "75942"]   # newest first
    s = day["stats"]
    assert (s["total"], s["resolved"], s["transferred"], s["ticket_inquiry"]) == (3, 2, 0, 1)
    assert s["first_call"] == "2026-10-11 09:00:00" and s["last_call"] == "2026-10-11 10:30:00"
    assert {h["key"]: h["total"] for h in s["hourly"] if h["total"]} == {9: 2, 10: 1}

    r = emp.put(f"/api/records/{rid}", {"reason": "transferred", "note": "sent to network team"})
    assert r.status_code == 200 and r.get_json()["record"]["status"] == "transferred"
    assert r.get_json()["record"]["edited"] is True
    assert emp.get("/api/records").get_json()["stats"]["transferred"] == 1

    assert emp.delete(f"/api/records/{rid}").status_code == 200
    assert emp.get("/api/records").get_json()["stats"]["total"] == 2
    assert emp.delete(f"/api/records/{rid}").status_code == 404


@pytest.mark.parametrize("payload", [
    {"ip_phone": "", "reason": "solved"},
    {"ip_phone": "abc", "reason": "solved"},
    {"ip_phone": "12", "reason": "solved"},
    {"ip_phone": "75942", "reason": "nonsense"},
    {"ip_phone": "75942", "reason": "solved", "note": "x" * 301},
])
def test_invalid_calls_are_rejected(make_employee, payload):
    assert make_employee("1001").post("/api/records", payload).status_code == 400


def test_employee_cannot_touch_other_employees_records(make_employee):
    a, b = make_employee("1001"), make_employee("1002")
    rid = log_call(a).get_json()["record"]["id"]
    assert b.put(f"/api/records/{rid}", {"note": "hacked"}).status_code == 404
    assert b.delete(f"/api/records/{rid}").status_code == 404
    assert a.get("/api/records").get_json()["records"][0]["note"] == ""
    assert b.get("/api/records").get_json()["stats"]["total"] == 0


@pytest.mark.parametrize("rid", ["../../users", "2026-10-11_../../x_abc", "garbage", "2026-13-40_1001_abc"])
def test_malformed_record_ids_are_404(make_employee, rid):
    emp = make_employee("1001")
    assert emp.delete(f"/api/records/{rid}").status_code == 404
    assert emp.put(f"/api/records/{rid}", {"note": "x"}).status_code == 404


def test_edit_window_locks_old_records_for_employees_not_admins(make_app, clock):
    app = make_app(user_edit_days=1)
    adm = Browser(app)
    adm.login("admin", app.config["INITIAL_ADMIN_PASSWORD"])
    adm.client.post("/account/password", data={"current_password": app.config["INITIAL_ADMIN_PASSWORD"],
                    "new_password": "AdminPass123", "confirm_password": "AdminPass123", "csrf_token": adm.csrf})
    adm.post("/api/admin/users", {"user_id": "1001", "name": "E", "password": "TempPass123"})
    emp = Browser(app)
    emp.login("1001", "TempPass123")
    emp.client.post("/account/password", data={"current_password": "TempPass123", "new_password": "MyOwnPass123",
                    "confirm_password": "MyOwnPass123", "csrf_token": emp.csrf})
    rid = log_call(emp).get_json()["record"]["id"]
    clock.advance(days=2)
    assert emp.put(f"/api/records/{rid}", {"note": "late"}).status_code == 403
    assert emp.delete(f"/api/records/{rid}").status_code == 403
    assert adm.put(f"/api/records/{rid}", {"note": "fixed by admin"}).status_code == 200
    day = emp.get("/api/records?date=2026-10-11").get_json()
    assert day["records"][0]["editable"] is False and day["records"][0]["note"] == "fixed by admin"


def test_cannot_view_future_dates(make_employee):
    assert make_employee("1001").get("/api/records?date=2026-10-12").status_code == 400


def test_admin_account_cannot_log_calls(admin):
    assert log_call(admin).status_code == 403


# -------------------------------------------------------------- admin views
def test_admin_overview_records_and_filters(admin, make_employee, clock):
    a, b = make_employee("1001", "Aya", "Identity"), make_employee("1002", "Omar", "Billing")
    make_employee("1003", "Idle", "Billing")
    log_call(a, "11111", "solved")
    clock.advance(minutes=1)
    log_call(a, "22222", "transferred")
    clock.advance(hours=2)
    log_call(b, "33333", "followup", "ticket 55")

    ov = admin.get("/api/admin/overview").get_json()
    assert ov["totals"] == {"total": 3, "resolved": 1, "transferred": 1, "ticket_inquiry": 1,
                            "active_employees": 2, "employees": 3}
    assert ov["chart"]["kind"] == "hour" and len(ov["chart"]["buckets"]) == 24
    assert [e["user_id"] for e in ov["employees"]][:2] == ["1001", "1002"]
    assert {e["user_id"]: e["total"] for e in ov["employees"]}["1003"] == 0   # idle staff are listed
    assert ov["recent"][0]["ip_phone"] == "33333" and ov["recent"][0]["name"] == "Omar"
    assert ov["departments"] == ["Billing", "Identity"]

    q = lambda s: admin.get("/api/admin/records?" + s).get_json()
    assert q("")["total"] == 3
    assert [r["ip_phone"] for r in q("user_id=1001")["items"]] == ["22222", "11111"]
    assert q("reason=followup")["total"] == 1
    assert q("q=ticket")["items"][0]["user_id"] == "1002"
    assert q("department=Billing")["total"] == 1
    assert q("page_size=2")["pages"] == 2

    week = admin.get("/api/admin/overview?from=2026-10-05&to=2026-10-11").get_json()
    assert week["chart"]["kind"] == "day" and len(week["chart"]["buckets"]) == 7


def test_admin_can_edit_and_delete_any_record(admin, make_employee):
    emp = make_employee("1001")
    rid = log_call(emp).get_json()["record"]["id"]
    assert admin.put(f"/api/records/{rid}", {"reason": "answered"}).status_code == 200
    assert admin.delete(f"/api/records/{rid}").status_code == 200


def test_employee_is_forbidden_from_admin_endpoints(make_employee):
    emp = make_employee("1001")
    for url in ("/api/admin/overview", "/api/admin/records", "/api/admin/users",
                "/api/admin/audit", "/export/records.csv", "/export/report.xlsx"):
        assert emp.get(url).status_code == 403, url
    assert emp.client.get("/admin").headers["Location"].endswith("/")


def test_bad_date_ranges(admin):
    assert admin.get("/api/admin/overview?from=2026-10-12&to=2026-10-01").status_code == 400
    assert admin.get("/api/admin/overview?from=nope").status_code == 400
    assert admin.get("/api/admin/overview?from=2020-01-01&to=2026-01-01").status_code == 400


# ---------------------------------------------------------- user management
def test_user_management_flow(admin, app):
    r = admin.post("/api/admin/users", {"user_id": "230966", "name": "خالد محمد", "department": "إدارة اعمال الهوية"})
    assert r.status_code == 201
    temp = r.get_json()["temp_password"]
    assert len(temp) >= 10
    assert admin.post("/api/admin/users", {"user_id": "230966", "name": "dup"}).status_code == 409
    assert admin.post("/api/admin/users", {"user_id": "bad id!", "name": "x"}).status_code == 400
    assert admin.post("/api/admin/users", {"user_id": "5", "name": ""}).status_code == 400

    emp = Browser(app)
    assert emp.login("230966", temp).status_code == 302
    assert emp.get("/api/me").status_code == 403       # must change the temporary password

    assert admin.put("/api/admin/users/230966", {"active": False}).status_code == 200
    assert Browser(app).login("230966", temp).status_code == 403
    assert emp.get("/api/records").status_code == 401   # an open session dies when disabled
    admin.put("/api/admin/users/230966", {"active": True, "department": "Billing"})

    new_pw = admin.post("/api/admin/users/230966/reset-password", {}).get_json()["temp_password"]
    assert Browser(app).login("230966", temp).status_code == 401
    assert Browser(app).login("230966", new_pw).status_code == 302
    users = {u["user_id"]: u for u in admin.get("/api/admin/users").get_json()["users"]}
    assert users["230966"]["department"] == "Billing" and "password_hash" not in users["230966"]


def test_admin_cannot_lock_themselves_out(admin):
    assert admin.put("/api/admin/users/admin", {"active": False}).status_code == 409
    assert admin.put("/api/admin/users/admin", {"role": "user"}).status_code == 409
    assert admin.put("/api/admin/users/ghost", {"name": "x"}).status_code == 404


def test_import_users_from_csv(admin, app):
    data = "user_id,name,department,password\n230966,خالد,Identity,\n230967,Sara,Billing,Passw0rdX\n,NoId,x,\n230966,Dup,x,\n"
    r = admin.post("/api/admin/users/import", data={"file": (io.BytesIO(data.encode("utf-8")), "u.csv")},
                   content_type="multipart/form-data")
    body = r.get_json()
    assert r.status_code == 200 and body["created"] == 2 and len(body["skipped"]) == 2
    creds = {c["user_id"]: c["password"] for c in body["credentials"]}
    assert creds["230967"] == "Passw0rdX" and len(creds["230966"]) >= 10
    assert Browser(app).login("230966", creds["230966"]).status_code == 302
    assert admin.post("/api/admin/users/import", data={"file": (io.BytesIO(b"x"), "u.txt")},
                      content_type="multipart/form-data").status_code == 400


# ------------------------------------------------------------------- audit
def test_audit_log_records_who_did_what(admin, make_employee):
    emp = make_employee("1001", "Aya")
    rid = log_call(emp, "75942", "solved").get_json()["record"]["id"]
    emp.put(f"/api/records/{rid}", {"ip_phone": "75943", "note": "typo"})
    emp.delete(f"/api/records/{rid}")
    Browser(admin.client.application).login("1001", "wrong-password")

    entries = admin.get("/api/admin/audit").get_json()
    actions = [e["action"] for e in entries["items"]]
    for expected in ("login", "record_created", "record_updated", "record_deleted",
                     "login_failed", "user_created", "password_changed"):
        assert expected in actions, expected
    upd = next(e for e in entries["items"] if e["action"] == "record_updated")
    assert upd["actor"] == "1001" and upd["details"]["before"]["ip_phone"] == "75942"
    assert upd["details"]["after"] == {"ip_phone": "75943", "note": "typo"}
    deleted = next(e for e in entries["items"] if e["action"] == "record_deleted")
    assert deleted["details"]["ip_phone"] == "75943"                  # the deleted data is kept
    only = admin.get("/api/admin/audit?action=record_created&actor=aya").get_json()
    assert only["total"] == 1


# ----------------------------------------------------------------- exports
def test_exports(admin, make_employee):
    emp = make_employee("1001", "Aya", "Identity")
    log_call(emp, "75942", "solved", "=HYPERLINK(\"http://evil\")")
    log_call(emp, "75943", "followup")

    r = admin.get("/export/report.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.mimetype
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Summary", "By branch", "Calls"]
    summary = list(wb["Summary"].iter_rows(values_only=True))
    assert summary[0][:4] == ("User ID", "Employee", "Department", "Total calls")
    assert summary[1][:7] == ("1001", "Aya", "Identity", 2, 1, 0, 1)
    rows = list(wb["Calls"].iter_rows(values_only=True))
    assert rows[0][5:12] == ("IP Phone", "Branch No.", "Branch", "Region", "Action Taken", "Action (English)", "Status")
    by_ip = {r[5]: r for r in rows[1:]}
    assert by_ip["75942"][9] == "تم حل المشكلة" and by_ip["75942"][11] == "Resolved"
    assert by_ip["75943"][11] == "Ticket Inquiry"
    formula_cells = [c for row in wb["Calls"].iter_rows() for c in row if c.data_type == "f"]
    assert by_ip["75942"][12].startswith("=HYPERLINK") and not formula_cells    # stored as text, never a live formula

    csv_resp = admin.get("/export/records.csv")
    text = csv_resp.data.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("Date,Time,User ID") and "'=HYPERLINK" in text
    assert "تم حل المشكلة" in text
    assert admin.get("/export/audit.xlsx").status_code == 200
    assert admin.get("/export/audit.csv").status_code == 200
    assert admin.get("/export/records.pdf").status_code == 404

    mine = emp.get("/export/my.xlsx")
    assert mine.status_code == 200
    assert len(list(openpyxl.load_workbook(io.BytesIO(mine.data)).active.iter_rows())) == 3
    actions = [e["action"] for e in admin.get("/api/admin/audit").get_json()["items"]]
    assert actions.count("export") >= 4


# ----------------------------------------------------------------- storage
def test_concurrent_writes_do_not_lose_records(tmp_path):
    store = Store(tmp_path / "d")
    store.init()

    def work(n):
        for i in range(25):
            store.add_record({"id": f"2026-10-11_u{n}_{i:06x}", "user_id": f"u{n}", "date": "2026-10-11",
                              "time": "09:00:00", "ip_phone": "12345", "reason": "solved", "note": ""})
            # also hammer one shared file
            store.add_record({"id": f"2026-10-11_shared_{n}{i:05x}", "user_id": "shared",
                              "date": "2026-10-11", "time": "09:00:01", "ip_phone": "1", "reason": "solved", "note": ""})

    threads = [threading.Thread(target=work, args=(n,)) for n in range(8)]
    [t.start() for t in threads]
    [t.join() for t in threads]
    assert len(store.records_between("2026-10-11", "2026-10-11")) == 8 * 25 * 2
    assert len(store.day_records("2026-10-11", "shared")) == 200


def test_unreachable_data_folder_gives_a_clear_error(tmp_path):
    from calllog.storage import StorageError
    blocker = tmp_path / "file"
    blocker.write_text("x")
    with pytest.raises(StorageError, match="Cannot use the data folder"):
        Store(blocker / "sub").init()


def test_security_headers(app):
    r = app.test_client().get("/login")
    assert "script-src 'self'" in r.headers["Content-Security-Policy"]
    assert r.headers["X-Frame-Options"] == "DENY"


# ---------------------------------------------------------------- branches
def test_branch_list_is_seeded_and_cached(make_employee):
    emp = make_employee("1001")
    r = emp.get("/api/branches")
    items = r.get_json()["branches"]
    assert len(items) == 597 and items[0] == {"code": "100", "name": "مصر الجديدة", "region": "شمال القاهرة الكبرى",
                                                "type": "كبير", "id": "100"}
    assert len({b["id"] for b in items}) == 597                 # ids are unique even where numbers repeat
    assert {"183/شمال القاهرة الكبرى", "183/غرب القاهرة الكبرى"} <= {b["id"] for b in items}
    again = emp.get("/api/branches", headers={"If-None-Match": r.headers["ETag"]})
    assert again.status_code == 304                              # unchanged list costs almost nothing


def test_branch_is_optional_and_stored_as_a_snapshot(make_employee):
    emp = make_employee("1001")
    rec = log_call(emp, "75942", "solved", branch="100").get_json()["record"]
    assert (rec["branch"], rec["branch_name"], rec["region"]) == ("100", "مصر الجديدة", "شمال القاهرة الكبرى")
    plain = log_call(emp, "75943", "answered").get_json()["record"]          # a sector / other caller
    assert (plain["branch"], plain["branch_name"], plain["region"]) == ("", "", "")
    assert log_call(emp, "75944", "answered", branch="").status_code == 201   # empty = no branch
    assert log_call(emp, "75945", "answered", branch="999999").status_code == 400
    day = emp.get("/api/records").get_json()["records"]
    assert sorted(r["branch"] for r in day) == ["", "", "100"]

    r = emp.put(f"/api/records/{plain['id']}", {"branch": "183/غرب القاهرة الكبرى"})
    assert r.get_json()["record"]["region"] == "غرب القاهرة الكبرى"
    r = emp.put(f"/api/records/{plain['id']}", {"branch": ""})                # clear it again
    assert r.get_json()["record"]["branch"] == ""
    assert emp.put(f"/api/records/{plain['id']}", {"branch": "nope"}).status_code == 400


def test_branch_import_replaces_the_list_but_keeps_old_calls(admin, make_employee):
    emp = make_employee("1001")
    rid = log_call(emp, "75942", "solved", branch="100").get_json()["record"]["id"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "الفروع"
    ws.append(["اسم المنطقة التابع لها", "رقم الفرع", "اسم الفرع", "نوع الفرع"])
    ws.append(["الإسكندرية", 5, "سموحة", "كبير"])                       # a number typed in Excel loses its zeros
    ws.append(["الإسكندرية", "007", "البنيــــان", "صغير"])               # tatweel is stripped
    ws.append(["الإسكندرية", "007", "البنيــــان", "متوسط"])              # same branch listed twice -> merged
    ws.append(["", "", "", ""])
    other = wb.create_sheet("ملخص")
    other.append(["x"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    r = admin.post("/api/admin/branches/import", data={"file": (buf, "branches.xlsx")}, content_type="multipart/form-data")
    assert r.status_code == 200 and r.get_json()["count"] == 2 and r.get_json()["duplicates_merged"] == 1
    items = admin.get("/api/branches").get_json()["branches"]
    assert [(b["id"], b["name"]) for b in items] == [("005", "سموحة"), ("007", "البنيان")]
    assert log_call(emp, "75943", "solved", branch="100").status_code == 400     # old branch no longer offered
    assert log_call(emp, "75944", "solved", branch="005").status_code == 201
    old = next(r for r in emp.get("/api/records").get_json()["records"] if r["id"] == rid)
    assert old["branch_name"] == "مصر الجديدة"                                    # history is unchanged
    assert admin.get("/api/admin/audit?action=branches_imported").get_json()["total"] == 1
    bad = io.BytesIO(b"a,b\n1,2\n")
    assert admin.post("/api/admin/branches/import", data={"file": (bad, "x.csv")}, content_type="multipart/form-data").status_code == 400
    assert make_employee("1002").post("/api/admin/branches/import", data={"file": (io.BytesIO(b"x"), "x.csv")},
                                      content_type="multipart/form-data").status_code == 403


def test_branch_analytics_filters_and_exports(admin, make_employee, clock):
    a, b = make_employee("1001", "Aya"), make_employee("1002", "Omar")
    log_call(a, "11111", "solved", branch="100")
    clock.advance(minutes=1)
    log_call(a, "22222", "transferred", branch="100")
    clock.advance(minutes=1)
    log_call(b, "33333", "followup", branch="183/غرب القاهرة الكبرى")
    clock.advance(minutes=1)
    log_call(b, "44444", "answered")                                            # no branch

    ov = admin.get("/api/admin/overview").get_json()["branches"]
    assert ov["with_branch"] == 3 and ov["without_branch"] == 1 and ov["distinct"] == 2
    assert ov["top"][0]["code"] == "100" and ov["top"][0]["total"] == 2 and ov["top"][0]["resolved"] == 1
    assert [r["region"] for r in ov["by_region"]] == ["شمال القاهرة الكبرى", "غرب القاهرة الكبرى"]

    q = lambda s: admin.get("/api/admin/records?" + s).get_json()
    assert q("region=غرب القاهرة الكبرى")["total"] == 1
    assert q("q=مصر")["total"] == 2                                               # matches the branch name
    assert q("q=183")["items"][0]["ip_phone"] == "33333"

    wb = openpyxl.load_workbook(io.BytesIO(admin.get("/export/report.xlsx").data))
    rows = list(wb["By branch"].iter_rows(values_only=True))
    assert rows[0][:4] == ("Branch No.", "Branch", "Region", "Total calls")
    assert rows[1][:7] == ("100", "مصر الجديدة", "شمال القاهرة الكبرى", 2, 1, 1, 0)
    calls = list(wb["Calls"].iter_rows(values_only=True))
    assert {(r[5], r[6] or "", r[8] or "") for r in calls[1:]} == {("11111", "100", "شمال القاهرة الكبرى"), ("22222", "100", "شمال القاهرة الكبرى"),
                                                       ("33333", "183", "غرب القاهرة الكبرى"), ("44444", "", "")}
    created = [e for e in admin.get("/api/admin/audit?action=record_created").get_json()["items"]]
    assert any(e["details"]["branch"] == "100 · مصر الجديدة" for e in created)


def test_admin_overview_sees_a_new_call_immediately(admin, make_employee):
    emp = make_employee("1001")
    assert admin.get("/api/admin/overview").get_json()["totals"]["total"] == 0     # result now cached ~2 s ...
    log_call(emp)
    assert admin.get("/api/admin/overview").get_json()["totals"]["total"] == 1     # ... but a write clears it
    assert admin.get("/api/admin/records").get_json()["total"] == 1


def test_parse_branch_rows():
    from calllog import branches
    entries, stats = branches.parse_rows([["Region", "Branch No", "Branch Name", "Type"], ["R1", "95.0", " Giza  Branch ", "x"],
                                          ["R1", "095", "Giza Branch", "y"], ["R2", "095", "Giza Branch", "x"], ["R2", "", "no code", ""]])
    assert stats == {"rows": 4, "skipped": 1, "duplicates_merged": 1}
    assert [e["id"] for e in entries] == ["095/R1", "095/R2"]                       # same number in two regions
    with pytest.raises(ValueError):
        branches.parse_rows([["a", "b"], ["1", "2"]])


# ----------------------------------------------------- several PCs, one folder
def test_separate_processes_sharing_one_folder_do_not_lose_data(tmp_path):
    store = Store(tmp_path / "d")
    store.init()
    ctx = multiprocessing.get_context("spawn")           # fresh interpreters, like separate PCs
    procs = [ctx.Process(target=__import__("mp_worker").hammer, args=(str(tmp_path / "d"), n, 20)) for n in range(4)]
    [p.start() for p in procs]
    [p.join(60) for p in procs]
    assert [p.exitcode for p in procs] == [0, 0, 0, 0]
    assert len(Store(tmp_path / "d").day_records("2026-10-11", "shared")) == 80
    assert len(Store(tmp_path / "d").read_audit("2026-10-11", "2026-10-11")) == 80
    lines = (tmp_path / "d" / "logs" / "audit-2026-10-11.jsonl").read_text(encoding="utf-8").splitlines()
    assert len(lines) == 80 and all(line.startswith("{") and line.endswith("}") for line in lines)   # no torn lines


def test_data_folder_defaults_to_next_to_the_program(monkeypatch, tmp_path):
    from calllog import config
    monkeypatch.delenv("CALLLOG_DATA_DIR", raising=False)
    monkeypatch.setenv("CALLLOG_CONFIG", str(tmp_path / "none.json"))
    assert config.load_config()["data_dir"] == str(config.APP_DIR / "data")
    cfg_file = tmp_path / "c.json"
    cfg_file.write_text('{"data_dir": "' + str(tmp_path / "elsewhere").replace("\\", "\\\\") + '"}')
    monkeypatch.setenv("CALLLOG_CONFIG", str(cfg_file))
    assert config.load_config()["data_dir"] == str(tmp_path / "elsewhere")


# --------------------------------------------- per-day summaries vs. plain counting
def test_overview_from_daily_summaries_matches_brute_force(admin, make_employee, app, clock):
    import random
    from collections import Counter, defaultdict
    from calllog.storage import new_record_id
    store = app.extensions["store"]
    uids = ["1001", "1002", "1003"]
    for u in uids:
        make_employee(u, f"Emp {u}", "Billing" if u == "1002" else "Cards")
    rnd = random.Random(5)
    branches = store.branches()
    keys = ["answered", "solved", "transferred", "followup", "retired-reason"]       # one reason the list does not know
    all_records = []
    for back in range(6):                                                           # six days ending at the fake "today"
        day = (clock().date() - __import__("datetime").timedelta(days=back)).isoformat()
        for u in uids[: 3 - (back % 2)]:
            recs = []
            for _ in range(rnd.randint(0, 25)):
                b = rnd.choice(branches) if rnd.random() < 0.7 else None
                recs.append({"id": new_record_id(day, u), "user_id": u, "date": day, "ip_phone": "12345",
                             "time": f"{rnd.randint(7, 15):02d}:{rnd.randint(0, 59):02d}:{rnd.randint(0, 59):02d}",
                             "reason": rnd.choice(keys), "note": "", "branch": b["code"] if b else "",
                             "branch_name": b["name"] if b else "", "region": b["region"] if b else ""})
            for r in recs:
                store.add_record(r)
            all_records += recs
    status = {r["key"]: r["status"] for r in store.reasons()}

    def brute(d_from, d_to):
        recs = [r for r in all_records if d_from <= r["date"] <= d_to]
        out = {"total": len(recs), "by_reason": dict(Counter(r["reason"] for r in recs))}
        users = defaultdict(lambda: {"total": 0, "resolved": 0, "transferred": 0, "ticket_inquiry": 0, "first": None, "last": None})
        branch = defaultdict(lambda: {"total": 0, "resolved": 0, "transferred": 0, "ticket_inquiry": 0})
        hours = defaultdict(int)
        for r in recs:
            st, stamp = status.get(r["reason"]), f"{r['date']} {r['time']}"
            u = users[r["user_id"]]
            u["total"] += 1
            u["first"] = min(u["first"] or stamp, stamp)
            u["last"] = max(u["last"] or stamp, stamp)
            if st:
                u[st] += 1
                hours[int(r["time"][:2])] += 1
            if r["branch"]:
                b = branch[(r["branch"], r["branch_name"], r["region"])]
                b["total"] += 1
                if st:
                    b[st] += 1
        out.update(users=users, branch=branch, hours=hours, status=Counter(status.get(r["reason"]) for r in recs))
        return out

    today = clock().date().isoformat()
    for d_from, d_to in [(today, today), ("2026-10-08", "2026-10-09"), ("2026-10-06", today)]:
        ov = admin.get(f"/api/admin/overview?from={d_from}&to={d_to}").get_json()
        want = brute(d_from, d_to)
        t = ov["totals"]
        assert t["total"] == want["total"]
        assert (t["resolved"], t["transferred"], t["ticket_inquiry"]) == tuple(want["status"].get(k, 0) for k in ("resolved", "transferred", "ticket_inquiry"))
        assert ov["by_reason"] == want["by_reason"]
        for row in ov["employees"]:
            u = want["users"].get(row["user_id"])
            if u is None:
                assert row["total"] == 0
                continue
            assert (row["total"], row["resolved"], row["transferred"], row["ticket_inquiry"]) == (u["total"], u["resolved"], u["transferred"], u["ticket_inquiry"])
            assert (row["first_call"], row["last_call"]) == (u["first"], u["last"])
        expected_top = sorted(want["branch"].items(), key=lambda kv: (-kv[1]["total"], kv[0][0], kv[0][2], kv[0][1]))[:10]
        assert [((b["code"], b["name"], b["region"]), {k: b[k] for k in ("total", "resolved", "transferred", "ticket_inquiry")})
                for b in ov["branches"]["top"]] == [(k, dict(v)) for k, v in expected_top]
        assert ov["branches"]["with_branch"] == sum(b["total"] for b in want["branch"].values())
        if d_from == d_to:
            assert {h["key"]: h["total"] for h in ov["chart"]["buckets"] if h["total"]} == dict(want["hours"])
        else:
            assert sum(b["total"] for b in ov["chart"]["buckets"]) == sum(want["status"].get(k, 0) for k in ("resolved", "transferred", "ticket_inquiry"))

    # editing a call on a PAST day must show up at once (its stored summary is dropped)
    old = next(r for r in all_records if r["date"] != today)
    before = admin.get(f"/api/admin/overview?from={old['date']}&to={old['date']}").get_json()["totals"]
    other = "solved" if status.get(old["reason"]) != "resolved" else "followup"
    assert admin.put(f"/api/records/{old['id']}", {"reason": other}).status_code == 200
    after = admin.get(f"/api/admin/overview?from={old['date']}&to={old['date']}").get_json()["totals"]
    assert after["total"] == before["total"] and after != before


def test_calls_saved_in_the_same_second_are_listed_newest_first(admin, make_employee):
    emp = make_employee("1001")
    for ip in ("11111", "22222", "33333"):                  # the fake clock does not move: identical timestamps
        log_call(emp, ip)
    assert [r["ip_phone"] for r in emp.get("/api/records").get_json()["records"]] == ["33333", "22222", "11111"]
    assert [r["ip_phone"] for r in admin.get("/api/admin/records").get_json()["items"]][:3] == ["33333", "22222", "11111"]
    assert [r["ip_phone"] for r in admin.get("/api/admin/overview").get_json()["recent"]][:3] == ["33333", "22222", "11111"]
