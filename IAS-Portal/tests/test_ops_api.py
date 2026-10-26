"""Operations Hub server side: roles, /api/ops endpoints, overrides overlay, exports, user profile fields."""
import datetime as dt
import io
import json

import openpyxl
import pytest

from calllog.ops import overlay, schemas
from conftest import Browser

SNAP_AT = "2026-10-11T09:00:00"          # same moment as the `clock` fixture


# ---------------------------------------------------------------------------- builders
def email_row(mid, **kw):
    row = {
        "subject": f"Subject {mid}", "entry_id": f"E-{mid}", "message_id": mid, "detected_systems": "Flexcube",
        "team": "Flex Cube", "all_teams": ["Flex Cube"], "teams_pending": ["Flex Cube"], "teams_replied": [],
        "customer_name": f"Customer {mid}", "customer_email": f"cust-{mid}@example.test",
        "received_time": "2026-10-11T08:00:00", "mailbox": "", "folder": "", "match_method": "No Match",
        "first_reply_by": "", "first_reply_email": "", "first_reply_time": None, "first_response_time": "",
        "response_minutes": None, "duration_minutes": 60, "reply_count": 0, "current_waiting_time": "1 H",
        "status": "Pending Reply", "sla_minutes": 60, "sla_status": "Within SLA", "escalated": False,
        "reminder_count": 0, "is_carryover": False, "age_days": 0, "no_action_reason": "", "reply_by_manager": False,
        "assigned_to_email": "", "assigned_to_name": "", "quality_rating": "", "quality_note": "",
        "quality_reasons": "", "quality_appeal": "", "appeal_decision": "", "rating_withdrawn": False,
        "body_preview": "", "reply_body_preview": "",
    }
    row.update(kw)
    return row


def replied_row(mid, by_email, minutes=20, **kw):
    received = dt.datetime(2026, 10, 11, 8, 0, 0)
    fields = dict(status="Replied", first_reply_by=f"Replier {by_email.split('@')[0]}", first_reply_email=by_email,
                  first_reply_time=(received + dt.timedelta(minutes=minutes)).isoformat(), response_minutes=minutes,
                  first_response_time=f"{minutes} Min", duration_minutes=minutes, reply_count=1,
                  sla_status="Met" if minutes <= 60 else "Breached", teams_pending=[], teams_replied=["Flex Cube"],
                  current_waiting_time="")
    fields.update(kw)
    return email_row(mid, **fields)


def email_data(rows, **stats_extra):
    stats = overlay.recount(rows, {"sla_target_minutes": 60, "team_stats": {
        t: {"team_emails_sent": 0} for t in ("Flex Cube", "Card Management", "Unclassified")}, **stats_extra})
    return {"date": "2026-10-11", "sla_target_minutes": 60, "stats": stats, "rows": rows,
            "teams": ["MFA Team", "Flex Cube", "Card Management", "Third Party", "Open System", "Unclassified"],
            "quality": {"scorecard": [], "team_index": {}}}


def sla_request(key, user, status="within", age=10, **fields):
    """An open security request ``age`` minutes old at the test clock (09:00), as the scan writes it."""
    sent = (dt.datetime(2026, 10, 11, 9, 0) - dt.timedelta(minutes=age)).isoformat(timespec="seconds") if age is not None else None
    return {"key": key, "security_user": user, "sent_at": sent, "age_minutes": age, "status": status,
            "fields": {"النظام": "Flexcube", **fields}}


def sla_data(requests):
    totals = {"open": len(requests), "oldest_minutes": max((r["age_minutes"] or 0 for r in requests), default=0)}
    for s in schemas.SLA_STATUSES:
        totals[s] = sum(1 for r in requests if r["status"] == s)
    users = sorted({r["security_user"] for r in requests})
    return {"thresholds": {"warning_minutes": 60, "breach_minutes": 120}, "totals": totals,
            "by_user": [{"user": u, "total": sum(1 for r in requests if r["security_user"] == u),
                         "within": 0, "reminder": 0, "breached": 0, "no_date": 0, "oldest_minutes": 0} for u in users],
            "requests": requests}


def itsm_data(imp=3, rollback=1, errors=()):
    roles = [{"role": "IMP User provisioning", "error": None, "categories": {
        "Implementation Requests": {"count": imp, "rows": [{"number": "100001", "summary": "Add user"}],
                                    "warning": None, "note": None},
        "Rollback Requests": {"count": rollback, "rows": [], "warning": None, "note": None}}}]
    roles += [{"role": r, "error": "Set Role button not found", "categories": {}} for r in errors]
    return {"categories": ["Implementation Requests", "Rollback Requests"], "roles": roles,
            "totals": {"Implementation Requests": imp, "Rollback Requests": rollback},
            "changes": {"new": [{"role": "IMP User provisioning", "category": "Rollback Requests",
                                 "number": "100002", "summary": "Rollback"}], "closed": []}}


def put_snapshot(ops, module, data, at=SNAP_AT):
    return ops.write_snapshot(module, data, generated_at=at, host="PC-TEST")


def beat(ops, job, at="2026-10-11T08:58:00", **fields):
    ops.write_json(ops.heartbeat_dir / f"{job}.json", {
        "job": job, "host": "PC-TEST", "running": False, "enabled": True, "last_status": "ok", "last_error": None,
        "last_ok_at": at, "updated_at": at, **fields})


def audit_actions(app, clock):
    day = clock().date().isoformat()
    return [(e["action"], e.get("target"), e.get("actor")) for e in app.extensions["store"].read_audit(day, day)]


# ---------------------------------------------------------------------------- fixtures
@pytest.fixture
def ops(app):
    return app.extensions["ops"]


@pytest.fixture
def make_account(app, admin):
    """Create an account through the API and return a signed-in Browser for it."""
    def factory(uid, role="user", name=None, **profile):
        body = {"user_id": uid, "name": name or uid.title(), "department": "IT", "role": role,
                "password": "TempPass123", **profile}
        r = admin.post("/api/admin/users", body)
        assert r.status_code == 201, r.get_json()
        b = Browser(app)
        b.login(uid, "TempPass123")
        resp = b.client.post("/account/password", data={
            "current_password": "TempPass123", "new_password": "MyOwnPass123",
            "confirm_password": "MyOwnPass123", "csrf_token": b.csrf})
        assert resp.status_code == 302
        return b
    return factory


@pytest.fixture
def mgr(make_account):
    """Manager of every team."""
    return make_account("mona", "manager", "Mona Manager", email="mona@example.test", team="__all__")


@pytest.fixture
def flex_mgr(make_account):
    return make_account("fadi", "manager", "Fadi Flex", email="fadi@example.test", team="Flex Cube")


@pytest.fixture
def emp(make_account):
    return make_account("1001", "user", "Alice Employee", email="alice@example.test",
                        security_name="احمد سالم")


@pytest.fixture
def assignees(admin):
    """Accounts of Admin › Users a mail can be assigned to: Bob and Dina of Flex Cube, Carol of Card Management."""
    for uid, name, email, team in (("2101", "Bob B", "bob@example.test", "Flex Cube"),
                                   ("2102", "Carol C", "carol@example.test", "Card Management"),
                                   ("2103", "Dina Dev", "dina.dev@example.test", "Flex Cube")):
        r = admin.post("/api/admin/users", {"user_id": uid, "name": name, "department": "IT", "role": "user",
                                            "password": "TempPass123", "email": email, "team": team})
        assert r.status_code == 201, r.get_json()


@pytest.fixture
def hub(ops):
    """Email snapshot with one row per situation, plus matching collector heartbeats."""
    rows = [
        email_row("<p1@x>", subject="Reset Flexcube password"),
        email_row("<p2@x>", subject="Card limit change", team="Card Management", all_teams=["Card Management"],
                  teams_pending=["Card Management"], received_time="2026-10-11T07:00:00", current_waiting_time="2 H",
                  sla_status="Breached", escalated=True, duration_minutes=120),
        email_row("<u1@x>", subject="Strange request", team="Unclassified", all_teams=["Unclassified"],
                  teams_pending=["Unclassified"], detected_systems=""),
        replied_row("<r1@x>", "bob@example.test", 20, subject="Replied mail one"),
        replied_row("<r2@x>", "carol@example.test", 90, subject="Replied mail two", team="Card Management",
                    all_teams=["Card Management"], teams_replied=["Card Management"]),
        email_row("<old@x>", subject="Old pending", is_carryover=True, age_days=2,
                  received_time="2026-10-09T08:00:00"),
    ]
    put_snapshot(ops, "email", email_data(rows))
    put_snapshot(ops, "sla", sla_data([sla_request("k1", "احمد سالم", "breached", 150),
                                       sla_request("k2", "منى علي", "reminder", 70)]))
    put_snapshot(ops, "itsm", itsm_data(errors=["IMP UAT I-Flex SMS"]))
    for job in schemas.JOBS:
        beat(ops, job)
    return rows


# ---------------------------------------------------------------------------- access matrix
GET_ENDPOINTS = {
    # path: (anonymous, user, manager, admin)
    "/api/ops/shared": (401, 200, 200, 200),
    "/api/ops/my": (401, 200, 200, 200),
    "/api/ops/overview": (401, 403, 200, 200),
    "/api/ops/email": (401, 403, 200, 200),
    "/api/ops/sla": (401, 403, 200, 200),
    "/api/ops/itsm": (401, 403, 200, 200),
    "/api/ops/email/history": (401, 403, 200, 200),
    "/api/ops/events": (401, 403, 200, 200),
    "/api/ops/reports": (401, 403, 200, 200),
    "/api/ops/jobs": (401, 403, 403, 200),
    "/api/ops/jobs/sla/log": (401, 403, 403, 200),
    "/api/ops/settings": (401, 403, 403, 200),
    "/api/ops/health": (401, 403, 403, 200),
    "/export/ops/email.xlsx": (401, 403, 200, 200),
    "/export/ops/sla.csv": (401, 403, 200, 200),
}


@pytest.mark.parametrize("path", sorted(GET_ENDPOINTS))
def test_get_endpoint_access_by_role(app, admin, mgr, emp, hub, path):
    anon, user, manager, adm = GET_ENDPOINTS[path]
    assert app.test_client().get(path).status_code == anon
    assert emp.get(path).status_code == user
    assert mgr.get(path).status_code == manager
    assert admin.get(path).status_code == adm


WRITE_ENDPOINTS = [
    ("POST", "/api/ops/email/assign", {"message_id": "<p1@x>", "assigned_to_email": "bob@example.test"}),
    ("POST", "/api/ops/email/no-action", {"message_id": "<p1@x>"}),
    ("DELETE", "/api/ops/email/no-action", {"message_id": "<p1@x>"}),
    # /api/ops/email/team is NOT here: it is no longer ops-only - an employee may re-file a mail assigned to them
    # (see test_done_flow.py::test_assignee_can_change_their_mails_team); only anonymous is denied, checked below.
    ("POST", "/api/ops/email/quality", {"message_id": "<r1@x>", "rating": "Good"}),
    ("POST", "/api/ops/email/quality/decision", {"message_id": "<r1@x>", "decision": "Accepted"}),
    ("POST", "/api/ops/email/classification/rule", {"keyword": "swift", "team": "Flex Cube"}),
    ("DELETE", "/api/ops/email/classification/rule", {"keyword": "swift"}),
    ("POST", "/api/ops/email/classification/suggestion", {"message_id": "<p1@x>", "keyword": "swift", "team": "Flex Cube"}),
    ("DELETE", "/api/ops/email/classification/suggestion", {"message_id": "<p1@x>"}),
    ("POST", "/api/ops/jobs/sla/run", {}),
    ("PUT", "/api/ops/settings", {"general": {}}),
]


@pytest.mark.parametrize("method,path,body", WRITE_ENDPOINTS)
def test_write_endpoints_deny_employees_and_anonymous(app, emp, hub, method, path, body):
    assert emp.api(method, path, json=body).status_code == 403
    anon = app.test_client()
    with anon.session_transaction() as s:
        s["csrf"] = "tok"
    r = anon.open(path, method=method, json=body, headers={"X-CSRF-Token": "tok"})
    assert r.status_code == 401


@pytest.mark.parametrize("method,path,body", WRITE_ENDPOINTS)
def test_write_endpoints_need_the_csrf_header(admin, hub, method, path, body):
    r = admin.client.open(path, method=method, json=body)
    assert r.status_code == 400
    assert "session expired" in r.get_json()["error"].lower()


# ---------------------------------------------------------------------------- module answers
def test_absent_snapshot_is_200_with_meta(admin):
    r = admin.get("/api/ops/sla")
    assert r.status_code == 200
    body = r.get_json()
    assert body["snapshot"] is None and body["meta"]["online"] is False and body["meta"]["generated_at"] is None
    email = admin.get("/api/ops/email").get_json()
    assert email["snapshot"] is None and email["options"]["ratings"] == ["Excellent", "Very Good", "Good", "Needs Improvement"]


def test_module_answer_carries_envelope_meta_and_options(admin, hub):
    body = admin.get("/api/ops/email").get_json()
    assert body["module"] == "email" and body["generated_at"] == SNAP_AT and body["host"] == "PC-TEST"
    assert len(body["data"]["rows"]) == 6
    assert body["meta"]["online"] is True and body["meta"]["stale"] is False and body["meta"]["interval_minutes"] == 10
    opts = body["options"]
    assert "Flex Cube" in opts["teams"] and "Incomplete answer" in opts["quality_reasons"]
    assert opts["unclassified_team"] == "" and opts["system_team_map"] == {}
    assert "options" not in admin.get("/api/ops/sla").get_json()


def test_unchanged_answer_when_since_matches(admin, hub):
    first = admin.get("/api/ops/sla").get_json()
    again = admin.get(f"/api/ops/sla?since={first['generated_at']}").get_json()
    assert again["unchanged"] is True and "data" not in again and again["meta"]["online"] is True
    other = admin.get("/api/ops/sla?since=2026-10-11T08:00:00").get_json()
    assert "unchanged" not in other and other["data"]["totals"]["open"] == 2


def test_unchanged_is_not_returned_after_someone_acted(admin, hub, ops):
    since = admin.get("/api/ops/email").get_json()
    assert admin.get(f"/api/ops/email?since={SNAP_AT}").get_json()["unchanged"] is True
    assert admin.post("/api/ops/email/no-action", {"message_id": "<p1@x>"}).status_code == 200
    # an override is newer than the scan: the page must refetch
    assert "unchanged" not in admin.get(f"/api/ops/email?since={SNAP_AT}").get_json()
    # a page that sends the revision it already shows gets "unchanged" only while that revision is current
    rev = admin.get("/api/ops/email").get_json()["meta"]["rev"]
    assert admin.get(f"/api/ops/email?since={SNAP_AT}&rev={rev}").get_json()["unchanged"] is True
    assert "unchanged" not in admin.get(f"/api/ops/email?since={SNAP_AT}&rev={since['meta']['rev']}").get_json()


def test_meta_offline_and_stale(app, admin, hub, ops, clock):
    clock.value = clock.value.replace(hour=11, minute=0)          # snapshot and heartbeats are now 2 h old
    meta = admin.get("/api/ops/sla").get_json()["meta"]
    assert meta["online"] is False and meta["stale"] is True and meta["age_minutes"] == 120
    beat(ops, "sla", at="2026-10-11T10:58:00")
    meta = admin.get("/api/ops/sla").get_json()["meta"]
    assert meta["online"] is True and meta["stale"] is True            # alive but its data is old
    clock.value = clock.value.replace(hour=20)                          # outside the 08:00-17:00 window: not a fault
    assert admin.get("/api/ops/sla").get_json()["meta"]["stale"] is False


def test_meta_reports_last_error_from_heartbeat(admin, hub, ops):
    beat(ops, "sla", last_status="error", last_error="Login failed")
    meta = admin.get("/api/ops/sla").get_json()["meta"]
    assert meta["last_status"] == "error" and meta["last_error"] == "Login failed"


def test_damaged_snapshot_is_reported_not_fatal(admin, ops):
    (ops.snapshots_dir / "sla.json").write_text("{not json", encoding="utf-8")
    body = admin.get("/api/ops/sla").get_json()
    assert body["snapshot"] is None and "damaged" in body["meta"]["problem"]


def test_history_endpoint(admin, ops):
    for i in range(5):
        ops.add_history_point("sla", "2026-10-11", {"ts": f"2026-10-11T09:0{i}:00", "open": i})
    ops.add_history_point("sla", "2026-10-10", {"ts": "2026-10-10T09:00:00", "open": 9})
    body = admin.get("/api/ops/sla/history?days=14").get_json()
    assert [d["day"] for d in body["days"]] == ["2026-10-10", "2026-10-11"]
    assert len(body["days"][1]["points"]) == 5
    assert [d["day"] for d in admin.get("/api/ops/sla/history?days=1").get_json()["days"]] == ["2026-10-11"]
    assert admin.get("/api/ops/sla/history?days=x").status_code == 400
    assert admin.get("/api/ops/nope/history").status_code == 404


def test_events_endpoint_and_manager_scope(admin, flex_mgr, ops):
    ops.add_events([
        {"ts": "2026-10-11T09:00:00", "module": "sla", "level": "danger", "kind": "sla_breached", "text": "1 breached"},
        {"ts": "2026-10-11T09:05:00", "module": "email", "level": "danger", "kind": "email_escalated", "text": "2 e-mails",
         "items": [{"subject": "Flex thing", "team": "Flex Cube"}, {"subject": "Card thing", "team": "Card Management"}]},
        {"ts": "2026-10-11T09:06:00", "module": "email", "level": "danger", "kind": "email_escalated", "text": "1 e-mail",
         "items": [{"subject": "Only card", "team": "Card Management"}]}])
    body = admin.get("/api/ops/events?day=2026-10-11").get_json()
    assert [e["ts"][11:16] for e in body["events"]] == ["09:06", "09:05", "09:00"]        # newest first
    mine = flex_mgr.get("/api/ops/events?day=2026-10-11").get_json()["events"]
    flat = json.dumps(mine)
    assert "Flex thing" in flat and "Card thing" not in flat and "Only card" not in flat
    assert [(e["text"], e["count"]) for e in mine] == [("1 e-mail", 1)]      # the count is theirs too, not the sector's
    assert admin.get("/api/ops/events?day=bad").status_code == 400
    assert admin.get("/api/ops/events").get_json()["events"] == body["events"]            # default day = today


# ---------------------------------------------------------------------------- overlay (pure)
def test_overlay_no_action_team_assignment_quality():
    rows = [email_row("<a>", escalated=True, sla_status="Breached"), replied_row("<b>", "bob@example.test"),
            email_row("<c>", team="Unclassified", all_teams=["Unclassified"])]
    ovr = {"no_action": {"<a>": {"marked_by": "Mona", "marked_at": "2026-10-11T09:01:00"}},
           "team": {"<c>": {"team": "Card Management"}},
           "assignments": {"<c>": {"assigned_to_email": "Bob@Example.test", "assigned_to_name": "Bob", "note": "urgent",
                                   "marked_by": "Mona", "marked_at": "2026-10-11T09:02:00"}},
           "quality": {"<b>": {"rating": "Good", "note": "ok", "reasons": "", "appeal": "hmm"}}}
    out = overlay.apply_overrides(rows, ovr)
    a, b, c = out
    assert a["status"] == "No Action Needed" and a["sla_status"] == "N/A" and a["escalated"] is False
    assert a["no_action_reason"] == 'Manually marked "No Action Needed" by Mona on 2026-10-11T09:01:00'
    assert a["teams_pending"] == []
    assert c["team"] == "Card Management" and c["all_teams"] == ["Card Management"] and c["teams_pending"] == ["Card Management"]
    assert c["assigned_to_email"] == "bob@example.test" and c["assigned_to_name"] == "Bob" and c["assigned_note"] == "urgent"
    assert b["quality_rating"] == "Good" and b["quality_appeal"] == "hmm"
    assert rows[0]["status"] == "Pending Reply" and rows[2]["team"] == "Unclassified"      # input untouched


def test_overlay_no_action_does_not_touch_replied_mails_and_drops_old_ones():
    rows = [replied_row("<b>", "bob@example.test"),
            email_row("<old>", is_carryover=True, age_days=3, received_time="2026-10-08T08:00:00")]
    out = overlay.apply_overrides(rows, {"no_action": {"<b>": {"marked_by": "x"}, "<old>": {"marked_by": "x"}}})
    assert [r["message_id"] for r in out] == ["<b>"] and out[0]["status"] == "Replied"


def test_overlay_removed_no_action_mark_shows_the_mail_pending_again():
    marked = overlay.apply_overrides([email_row("<a>")], {"no_action": {"<a>": {"marked_by": "Mona", "marked_at": "t"}}})
    import datetime as dt
    again = overlay.apply_overrides(marked, {}, now=dt.datetime(2026, 10, 11, 9, 30))
    assert again[0]["status"] == "Pending Reply" and again[0]["sla_status"] == "Breached"        # waiting 90 min > 60
    assert again[0]["escalated"] is True and again[0]["current_waiting_time"] == "1 H 30 Min"
    # an automatic no-action (no manual text) stays as it is
    auto = email_row("<z>", status="No Action Needed", no_action_reason="keyword thanks", sla_status="N/A")
    assert overlay.apply_overrides([auto], {})[0]["status"] == "No Action Needed"


def test_overlay_assignment_and_quality_follow_the_override_files():
    stale = [email_row("<a>", assigned_to_email="old@example.test"), replied_row("<b>", "bob@example.test", quality_rating="Good")]
    out = overlay.apply_overrides(stale, {})
    assert out[0]["assigned_to_email"] == "" and out[1]["quality_rating"] == ""


def test_overlay_quality_only_on_replied_rows():
    out = overlay.apply_overrides([email_row("<a>")], {"quality": {"<a>": {"rating": "Good"}}})
    assert out[0]["quality_rating"] == ""


def test_overlay_team_override_collapses_split_children():
    shared = {"message_id": "<s>", "subject": "Two teams", "is_split_child": True,
              "split_sibling_teams": ["Flex Cube", "Card Management"]}
    flex = replied_row("<s>", "bob@example.test", 15, **shared)
    card = email_row("<s>", team="Card Management", all_teams=["Card Management"], **shared)
    out = overlay.apply_overrides([flex, card], {"team": {"<s>": {"team": "Third Party"}}})
    assert len(out) == 1 and out[0]["team"] == "Third Party" and out[0]["status"] == "Replied"
    assert "is_split_child" not in out[0] and out[0]["teams_replied"] == ["Third Party"]
    # no override: both children stay
    assert len(overlay.apply_overrides([flex, card], {})) == 2


def test_overlay_no_action_collapses_pending_split_children():
    kw = {"message_id": "<s>", "is_split_child": True, "split_sibling_teams": ["Flex Cube", "Card Management"]}
    rows = [email_row("<s>", **kw), email_row("<s>", team="Card Management", all_teams=["Card Management"], **kw)]
    out = overlay.apply_overrides(rows, {"no_action": {"<s>": {"marked_by": "x", "marked_at": "t"}}})
    assert len(out) == 1 and out[0]["status"] == "No Action Needed" and out[0]["all_teams"] == ["Flex Cube", "Card Management"]


def test_recount_is_idempotent_without_overrides_and_math_after_overrides():
    rows = [email_row("<p1>"), email_row("<p2>", escalated=True, sla_status="Breached", received_time="2026-10-11T06:00:00"),
            replied_row("<r1>", "bob@example.test", 20), replied_row("<r2>", "bob@example.test", 90),
            email_row("<n1>", status="No Action Needed", sla_status="N/A", teams_pending=[]),
            email_row("<c1>", is_carryover=True, age_days=2)]
    base = overlay.recount(rows, {"team_stats": {"Flex Cube": {"team_emails_sent": 7}}})
    assert overlay.recount(rows, base) == base
    assert (base["business_emails"], base["replied"], base["pending"], base["no_action_needed"]) == (5, 2, 3, 1)
    assert base["needed_reply"] == 4 and base["reply_rate"] == 50.0
    assert base["sla_breached"] == 2 and base["escalated_count"] == 1 and base["carryover_count"] == 1
    assert base["carryover_oldest_days"] == 2 and base["average_response"] == 55 and base["fastest_response"] == 20
    assert base["slowest_response"] == 90 and base["team_stats"]["Flex Cube"]["team_emails_sent"] == 7
    assert base["top_slowest_replies"][0]["response_minutes"] == 90

    new = overlay.apply_overrides(rows, {
        "no_action": {"<p1>": {"marked_by": "x", "marked_at": "t"}},
        "assignments": {"<p2>": {"assigned_to_email": "bob@example.test"}},
        "team": {"<r2>": {"team": "Card Management"}},
        "quality": {"<r1>": {"rating": "Excellent"}, "<r2>": {"rating": "Needs Improvement", "appeal": "unfair"}}})
    stats = overlay.recount(new, base)
    assert (stats["pending"], stats["no_action_needed"], stats["needed_reply"]) == (2, 2, 3)
    assert stats["reply_rate"] == round(2 / 3 * 100, 2)
    assert stats["assigned_count"] == 1 and stats["escalated_count"] == 1
    flex, card = stats["team_stats"]["Flex Cube"], stats["team_stats"]["Card Management"]
    assert (flex["business_emails"], flex["replied"], flex["no_action_needed"]) == (5, 1, 2)       # like the collector, a team's count includes its carry-over rows
    assert (card["business_emails"], card["replied"]) == (1, 1) and card["fastest_response_minutes"] == 90
    assert stats["quality"]["reviewed"] == 2 and stats["quality"]["excellent"] == 1
    assert stats["quality"]["needs_improvement"] == 1 and stats["quality"]["appeals"] == 1
    assert stats["quality"]["coverage_rate"] == 100.0


def test_scope_rows_and_rows_for_email():
    rows = [email_row("<a>"), email_row("<b>", team="Card Management", all_teams=["Card Management"]),
            email_row("<c>", team="Open System", all_teams=["Open System", "Flex Cube"]),
            email_row("<d>", assigned_to_email="x@example.test"),
            replied_row("<e>", "x@example.test")]
    flex = {"role": "manager", "team": "Flex Cube"}
    assert [r["message_id"] for r in overlay.scope_rows(rows, flex)] == ["<a>", "<c>", "<d>", "<e>"]
    assert len(overlay.scope_rows(rows, {"role": "manager", "team": "__all__"})) == 5
    assert len(overlay.scope_rows(rows, {"role": "manager", "team": ""})) == 5
    assert len(overlay.scope_rows(rows, {"role": "admin", "team": "Flex Cube"})) == 5
    assert overlay.scope_rows(rows, {"role": "user", "team": ""}) == []
    assert [r["message_id"] for r in overlay.rows_for_email(rows, "X@Example.test")] == ["<d>", "<e>"]
    assert overlay.rows_for_email(rows, "") == []


def test_display_name_from_email():
    assert overlay.display_name_from_email("flexcube.Alice.Smith2@nbe.test") == "Alice Smith2"
    assert overlay.display_name_from_email("alice_smith@example.test") == "Alice Smith"


# ---------------------------------------------------------------------------- e-mail answers with overrides
def test_email_answer_applies_overrides_and_recounts(admin, hub, assignees):
    before = admin.get("/api/ops/email").get_json()["data"]["stats"]
    assert admin.post("/api/ops/email/no-action", {"message_id": "<p1@x>"}).status_code == 200
    assert admin.post("/api/ops/email/assign", {"message_id": "<p2@x>", "assigned_to_email": "Bob@Example.test",
                                                "note": "please"}).status_code == 200
    after = admin.get("/api/ops/email").get_json()["data"]
    rows = {r["message_id"]: r for r in after["rows"]}
    assert rows["<p1@x>"]["status"] == "No Action Needed" and rows["<p2@x>"]["assigned_to_email"] == "bob@example.test"
    assert after["stats"]["pending"] == before["pending"] - 1 and after["stats"]["no_action_needed"] == 1
    assert after["stats"]["assigned_count"] == 1
    assert after["stats"]["needed_reply"] == before["needed_reply"] - 1


def test_manager_sees_only_their_team_with_recounted_stats(admin, flex_mgr, mgr, hub, assignees):
    full = admin.get("/api/ops/email").get_json()
    mine = flex_mgr.get("/api/ops/email").get_json()
    subjects = {r["subject"] for r in mine["data"]["rows"]}
    assert "Card limit change" not in subjects and "Reset Flexcube password" in subjects
    assert all(r["team"] == "Flex Cube" for r in mine["data"]["rows"])
    stats = mine["data"]["stats"]
    assert stats["business_emails"] == sum(1 for r in mine["data"]["rows"] if not r["is_carryover"])
    assert list(stats["team_stats"]) == ["Flex Cube"] and stats["escalated_count"] == 0
    assert stats["business_address_volume"] == {} and stats["trend"] == {}
    assert "Card limit change" not in json.dumps(mine) and "Card limit change" in json.dumps(full)
    assert len(mgr.get("/api/ops/email").get_json()["data"]["rows"]) == 6          # team "__all__" = everything
    # a manager assigns a mail only to the people of their own team; the administrator to anybody
    assert {p["email"] for p in mine["options"]["people"]} == {"fadi@example.test", "bob@example.test", "dina.dev@example.test"}
    assert "carol@example.test" in {p["email"] for p in full["options"]["people"]} and mine["options"]["scope"] == "Flex Cube"


def test_manager_quality_sections_carry_nothing_about_other_teams(admin, flex_mgr, hub, ops):
    data = email_data([email_row("<q1>"), replied_row("<q2>", "bob@example.test")])
    data["quality"] = {
        "scorecard": [{"name": "Bob F", "team": "Flex Cube", "score": 90}, {"name": "Cora C", "team": "Card Management", "score": 40}],
        "team_index": {"Flex Cube": 90, "Card Management": 40},
        "leaderboard_top": [{"name": "Bob F", "team": "Flex Cube"}, {"name": "Cora C", "team": "Card Management"}],
        "person_trends": [{"name": "No Team Given"}], "repeat_issue_alerts": [{"issue": "x", "team": "Card Management"}]}
    put_snapshot(ops, "email", data)
    mine = flex_mgr.get("/api/ops/email").get_json()["data"]
    assert [p["name"] for p in mine["quality"]["scorecard"]] == ["Bob F"]
    assert mine["quality"]["team_index"] == {"Flex Cube": 90}
    assert [p["name"] for p in mine["quality"]["leaderboard_top"]] == ["Bob F"]
    assert mine["quality"]["person_trends"] == [] and mine["quality"]["repeat_issue_alerts"] == []
    assert len(admin.get("/api/ops/email").get_json()["data"]["quality"]["scorecard"]) == 2


def test_options_people_are_the_accounts_of_admin_users_and_a_manager_gets_only_their_team(admin, flex_mgr, hub, assignees):
    """The people a mail can be assigned to are the active accounts of Admin › Users with an e-mail address (named and
    teamed from the account) - never the Settings team-member addresses without an account. A manager gets only the
    people of their own team."""
    cfg = admin.get("/api/ops/settings").get_json()["settings"]
    cfg["email"]["team_members_by_team"] = {"Flex Cube": ["flexcube.Alice.Smith2@example.test"],
                                           "Card Management": ["carla.cardholder@example.test"]}
    cfg["email"]["team_members"] = ["flexcube.Alice.Smith2@example.test", "solo.person@example.test"]
    assert admin.put("/api/ops/settings", cfg).status_code == 200
    people = admin.get("/api/ops/email").get_json()["options"]["people"]
    assert {p["email"]: (p["name"], p["team"]) for p in people} == {
        "fadi@example.test": ("Fadi Flex", "Flex Cube"), "bob@example.test": ("Bob B", "Flex Cube"),
        "carol@example.test": ("Carol C", "Card Management"), "dina.dev@example.test": ("Dina Dev", "Flex Cube")}
    mine = flex_mgr.get("/api/ops/email").get_json()["options"]["people"]
    assert [p["email"] for p in mine] == ["bob@example.test", "dina.dev@example.test", "fadi@example.test"]


def test_a_manager_assigns_only_to_an_account_of_their_own_team(flex_mgr, hub, ops, assignees):
    ok = flex_mgr.post("/api/ops/email/assign", {"message_id": "<p1@x>", "assigned_to_email": "Bob@Example.test",
                                                 "assigned_to_name": "Somebody Else"})
    assert ok.status_code == 200
    assert ops.overrides("assignments")["<p1@x>"]["assigned_to_name"] == "Bob B"          # the account's own name
    for address in ("carol@example.test", "ghost@example.test"):                           # another team / no account
        r = flex_mgr.post("/api/ops/email/assign", {"message_id": "<p1@x>", "assigned_to_email": address})
        assert r.status_code == 403 and "your own team" in r.get_json()["error"], address
    assert ops.overrides("assignments")["<p1@x>"]["assigned_to_email"] == "bob@example.test"


# ---------------------------------------------------------------------------- e-mail actions
def test_assign_validates_and_audits(app, admin, hub, ops, clock, assignees):
    ok = admin.post("/api/ops/email/assign", {"message_id": "<p1@x>", "assigned_to_email": "Bob@Example.test",
                                              "assigned_to_name": "Bob B", "note": "today please"})
    assert ok.status_code == 200
    stored = ops.overrides("assignments")["<p1@x>"]
    assert stored["assigned_to_email"] == "bob@example.test" and stored["assigned_to_name"] == "Bob B"
    assert stored["note"] == "today please" and stored["notified"] is False and stored["marked_by"] == "Administrator"
    assert stored["subject"] == "Reset Flexcube password" and stored["marked_at"] == "2026-10-11T09:00:00"
    assert ("ops_assign", "<p1@x>", "admin") in audit_actions(app, clock)
    # the name is the account's
    admin.post("/api/ops/email/assign", {"message_id": "<p2@x>", "assigned_to_email": "dina.dev@example.test"})
    assert ops.overrides("assignments")["<p2@x>"]["assigned_to_name"] == "Dina Dev"
    bad = [({"message_id": "<p1@x>", "assigned_to_email": "nope"}, 400),
           ({"message_id": "<p1@x>", "assigned_to_email": ""}, 400),
           ({"message_id": "<missing@x>", "assigned_to_email": "bob@example.test"}, 404),
           ({"message_id": "", "assigned_to_email": "bob@example.test"}, 400),
           ({"message_id": "<p1@x>", "assigned_to_email": "bob@example.test", "note": "x" * 301}, 400),
           ({"message_id": "<p1@x>", "assigned_to_email": "no.account@example.test"}, 403)]
    for body, status in bad:
        r = admin.post("/api/ops/email/assign", body)
        assert r.status_code == status, (body, r.get_json())
        assert "error" in r.get_json()
    assert admin.post("/api/ops/email/assign", [1, 2]).status_code == 400


def test_assign_keeps_notified_flag_only_for_the_same_assignee(admin, hub, ops, assignees):
    admin.post("/api/ops/email/assign", {"message_id": "<p1@x>", "assigned_to_email": "bob@example.test", "note": "a"})
    entry = ops.overrides("assignments")["<p1@x>"]
    ops.set_override("assignments", "<p1@x>", {**entry, "notified": True})
    admin.post("/api/ops/email/assign", {"message_id": "<p1@x>", "assigned_to_email": "bob@example.test", "note": "a"})
    assert ops.overrides("assignments")["<p1@x>"]["notified"] is True
    admin.post("/api/ops/email/assign", {"message_id": "<p1@x>", "assigned_to_email": "carol@example.test", "note": "a"})
    assert ops.overrides("assignments")["<p1@x>"]["notified"] is False


def test_unassign(admin, hub, ops, assignees):
    admin.post("/api/ops/email/assign", {"message_id": "<p1@x>", "assigned_to_email": "bob@example.test"})
    assert admin.delete("/api/ops/email/assign", json={"message_id": "<p1@x>"}).status_code == 200
    assert ops.overrides("assignments") == {}
    assert admin.delete("/api/ops/email/assign", json={"message_id": "<p1@x>"}).status_code == 404
    rows = {r["message_id"]: r for r in admin.get("/api/ops/email").get_json()["data"]["rows"]}
    assert rows["<p1@x>"]["assigned_to_email"] == ""


def test_no_action_mark_and_undo(app, admin, hub, ops, clock):
    assert admin.post("/api/ops/email/no-action", {"message_id": "<p1@x>"}).status_code == 200
    mark = ops.overrides("no_action")["<p1@x>"]
    assert mark["marked_by"] == "Administrator" and mark["subject"] == "Reset Flexcube password"
    assert admin.post("/api/ops/email/no-action", {"message_id": "<r1@x>"}).status_code == 409      # already answered
    assert admin.post("/api/ops/email/no-action", {"message_id": "<zzz>"}).status_code == 404
    assert admin.delete("/api/ops/email/no-action", json={"message_id": "<p1@x>"}).status_code == 200
    assert ops.overrides("no_action") == {}
    assert admin.delete("/api/ops/email/no-action", json={"message_id": "<p1@x>"}).status_code == 404
    acts = [a[0] for a in audit_actions(app, clock)]
    assert "ops_no_action" in acts and "ops_no_action_undo" in acts


def test_no_action_on_old_mail_hides_it_and_undo_still_works(admin, hub, ops):
    assert admin.post("/api/ops/email/no-action", {"message_id": "<old@x>"}).status_code == 200
    ids = {r["message_id"] for r in admin.get("/api/ops/email").get_json()["data"]["rows"]}
    assert "<old@x>" not in ids
    assert admin.delete("/api/ops/email/no-action", json={"message_id": "<old@x>"}).status_code == 200
    ids = {r["message_id"] for r in admin.get("/api/ops/email").get_json()["data"]["rows"]}
    assert "<old@x>" in ids


def test_team_override(app, admin, hub, ops, clock):
    assert admin.post("/api/ops/email/team", {"message_id": "<u1@x>", "team": "Open System"}).status_code == 200
    assert ops.overrides("team")["<u1@x>"]["team"] == "Open System"
    rows = {r["message_id"]: r for r in admin.get("/api/ops/email").get_json()["data"]["rows"]}
    assert rows["<u1@x>"]["team"] == "Open System"
    for body in ({"message_id": "<u1@x>", "team": "Nonexistent Team"}, {"message_id": "<u1@x>", "team": ""},
                 {"message_id": "<u1@x>"}):
        assert admin.post("/api/ops/email/team", body).status_code == 400
    assert admin.post("/api/ops/email/team", {"message_id": "<nope>", "team": "Open System"}).status_code == 404
    assert ("ops_team", "<u1@x>", "admin") in audit_actions(app, clock)


def test_quality_review_validation_and_storage(app, admin, mgr, hub, ops, clock):
    body = {"message_id": "<r2@x>", "rating": "Needs Improvement", "note": "Missed a step",
            "reasons": ["Too slow", "Incomplete answer"]}
    assert mgr.post("/api/ops/email/quality", body).status_code == 200
    review = ops.overrides("quality")["<r2@x>"]
    assert review["rating"] == "Needs Improvement" and review["reasons"] == "Too slow, Incomplete answer"
    assert review["replied_by_email"] == "carol@example.test" and review["reviewed_by"] == "Mona Manager"
    assert review["team"] == "Card Management" and review["rating_withdrawn"] is False
    rows = {r["message_id"]: r for r in admin.get("/api/ops/email").get_json()["data"]["rows"]}
    assert rows["<r2@x>"]["quality_rating"] == "Needs Improvement"
    assert admin.get("/api/ops/email").get_json()["data"]["stats"]["quality"]["needs_improvement"] == 1
    bad = [{"message_id": "<r2@x>", "rating": "Superb"}, {"message_id": "<r2@x>", "rating": ""},
           {"message_id": "<r2@x>", "rating": "Good", "reasons": ["Made up reason"]},
           {"message_id": "<r2@x>", "rating": "Good", "note": "n" * 501}]
    for b in bad:
        assert admin.post("/api/ops/email/quality", b).status_code == 400, b
    assert admin.post("/api/ops/email/quality", {"message_id": "<p1@x>", "rating": "Good"}).status_code == 409     # no reply
    assert admin.post("/api/ops/email/quality", {"message_id": "<nope>", "rating": "Good"}).status_code == 404
    assert ("ops_quality", "<r2@x>", "mona") in audit_actions(app, clock)


def test_a_manager_can_rate_a_reply_very_good(admin, mgr, hub, ops):
    """Excellent, Very Good, Good, Needs Improvement: the page offers the four, the API takes "Very Good", the counts
    and the quality statistics carry it on its own."""
    options = admin.get("/api/ops/email").get_json()["options"]
    assert options["ratings"] == ["Excellent", "Very Good", "Good", "Needs Improvement"]
    assert mgr.post("/api/ops/email/quality", {"message_id": "<r2@x>", "rating": "Very Good"}).status_code == 200
    assert ops.overrides("quality")["<r2@x>"]["rating"] == "Very Good"
    quality = admin.get("/api/ops/email").get_json()["data"]["stats"]["quality"]
    assert quality["very_good"] == 1 and quality["reviewed"] == 1
    stats = admin.get("/api/ops/analytics?from=2026-10-01&to=2026-10-11").get_json()["quality"]
    assert stats["very_good"] == 1 and stats["score"] == 90.0
    assert admin.post("/api/ops/email/quality", {"message_id": "<r2@x>", "rating": "Very good!"}).status_code == 400


def test_quality_rules_manager_replies_and_own_reply(admin, mgr, hub, ops):
    put_snapshot(ops, "email", email_data([
        replied_row("<m1>", "boss@example.test", reply_by_manager=True),
        replied_row("<own>", "mona@example.test"), replied_row("<ok>", "bob@example.test")]))
    assert admin.post("/api/ops/email/quality", {"message_id": "<m1>", "rating": "Good"}).status_code == 409
    assert mgr.post("/api/ops/email/quality", {"message_id": "<own>", "rating": "Good"}).status_code == 403
    assert mgr.post("/api/ops/email/quality", {"message_id": "<ok>", "rating": "Good"}).status_code == 200


def test_appeal_and_decision_flow(app, admin, mgr, flex_mgr, make_account, hub, ops, clock):
    bob = make_account("1002", "user", "Bob Employee", email="bob@example.test")
    carol = make_account("1003", "user", "Carol Employee", email="carol@example.test")
    assert mgr.post("/api/ops/email/quality", {"message_id": "<r1@x>", "rating": "Needs Improvement",
                                               "reasons": "Too slow"}).status_code == 200
    # only the replier may appeal
    assert carol.post("/api/ops/email/quality/appeal", {"message_id": "<r1@x>", "appeal": "not mine"}).status_code == 404
    assert admin.post("/api/ops/email/quality/appeal", {"message_id": "<r1@x>", "appeal": "admin cannot"}).status_code == 404
    assert bob.post("/api/ops/email/quality/appeal", {"message_id": "<r1@x>", "appeal": "ok"}).status_code == 400   # too short
    assert bob.post("/api/ops/email/quality/appeal", {"message_id": "<nope>", "appeal": "unfair rating"}).status_code == 404
    assert bob.post("/api/ops/email/quality/appeal", {"message_id": "<r1@x>", "appeal": "The customer asked late"}).status_code == 200
    assert ops.overrides("quality")["<r1@x>"]["appeal"] == "The customer asked late"
    assert ops.overrides("quality")["<r1@x>"]["appeal_by"] == "Bob Employee"
    # decisions: managers/admin only; a Flex Cube manager may decide (the mail is Flex Cube), others not
    assert bob.post("/api/ops/email/quality/decision", {"message_id": "<r1@x>", "decision": "Accepted"}).status_code == 403
    assert admin.post("/api/ops/email/quality/decision", {"message_id": "<r1@x>", "decision": "Maybe"}).status_code == 400
    assert flex_mgr.post("/api/ops/email/quality/decision", {"message_id": "<r1@x>", "decision": "Accepted"}).status_code == 200
    review = ops.overrides("quality")["<r1@x>"]
    assert review["appeal_decision"] == "Accepted" and review["rating_withdrawn"] is True
    assert review["decided_by"] == "Fadi Flex"
    assert bob.post("/api/ops/email/quality/appeal", {"message_id": "<r1@x>", "appeal": "again please"}).status_code == 409
    rows = {r["message_id"]: r for r in admin.get("/api/ops/email").get_json()["data"]["rows"]}
    assert rows["<r1@x>"]["rating_withdrawn"] is True and rows["<r1@x>"]["appeal_decision"] == "Accepted"
    acts = [a[0] for a in audit_actions(app, clock)]
    assert {"ops_quality", "ops_quality_appeal", "ops_quality_decision"} <= set(acts)


def test_appeal_needs_a_signed_in_user_and_csrf(app, emp, hub):
    body = {"message_id": "<r1@x>", "appeal": "please look again"}
    anon = app.test_client()
    with anon.session_transaction() as s:
        s["csrf"] = "tok"
    assert anon.post("/api/ops/email/quality/appeal", json=body, headers={"X-CSRF-Token": "tok"}).status_code == 401
    assert emp.client.post("/api/ops/email/quality/appeal", json=body).status_code == 400
    assert emp.post("/api/ops/email/quality/appeal", body).status_code == 404         # alice has no rating on that mail


def test_decision_needs_an_appeal_and_team_access(admin, mgr, make_account, hub, ops):
    card_mgr = make_account("cora", "manager", "Cora Card", team="Card Management")
    mgr.post("/api/ops/email/quality", {"message_id": "<r1@x>", "rating": "Good"})
    assert admin.post("/api/ops/email/quality/decision", {"message_id": "<r1@x>", "decision": "Accepted"}).status_code == 409
    assert admin.post("/api/ops/email/quality/decision", {"message_id": "<zzz>", "decision": "Accepted"}).status_code == 404
    ops.update_override("quality", "<r1@x>", lambda old: {**old, "appeal": "please check"})
    assert card_mgr.post("/api/ops/email/quality/decision", {"message_id": "<r1@x>", "decision": "Rejected"}).status_code == 403
    assert admin.post("/api/ops/email/quality/decision", {"message_id": "<r1@x>", "decision": "rejected"}).status_code == 200
    assert ops.overrides("quality")["<r1@x>"]["rating_withdrawn"] is False


def test_a_second_team_manager_keeps_handling_a_shared_mail(admin, make_account, hub, ops):
    """A replied mail named for two teams is one row (team = the first). The second team's manager files their
    assignment and their rating under their OWN team, so they can still handle both once the mail leaves the scan.
    A reply is rated only by the managers of the team of the person who wrote it - here Card Management's."""
    cara = make_account("cara", "manager", "Cara Card", email="cara@example.test", team="Card Management")
    bob = make_account("1002", "user", "Bob Employee", email="bob@example.test", team="Card Management")
    shared = replied_row("<s1@x>", "bob@example.test", all_teams=["Flex Cube", "Card Management"])
    put_snapshot(ops, "email", email_data([shared]))
    r = cara.post("/api/ops/email/assign", {"message_id": "<s1@x>", "assigned_to_email": "bob@example.test"})
    assert r.status_code == 200 and r.get_json()["assignment"]["team"] == "Card Management"
    assert cara.post("/api/ops/email/quality", {"message_id": "<s1@x>", "rating": "Needs Improvement"}).status_code == 200
    assert ops.overrides("quality")["<s1@x>"]["team"] == "Card Management"
    assert bob.post("/api/ops/email/quality/appeal", {"message_id": "<s1@x>", "appeal": "The customer asked late"}).status_code == 200
    put_snapshot(ops, "email", email_data([]), at="2026-10-11T09:30:00")           # the mail left the scan
    assert cara.post("/api/ops/email/quality/decision", {"message_id": "<s1@x>", "decision": "Rejected"}).status_code == 200
    assert cara.delete("/api/ops/email/assign", json={"message_id": "<s1@x>"}).status_code == 200


def test_view_as_shows_only_the_ratings_of_the_viewers_teams(admin, flex_mgr, make_account, hub, ops):
    cara = make_account("cara", "manager", "Cara Card", email="cara@example.test", team="Card Management")
    make_account("1002", "user", "Bob Employee", email="bob@example.test", team="Flex Cube")
    put_snapshot(ops, "email", email_data([replied_row("<fx@x>", "bob@example.test", subject="FLEX ONLY SUBJECT")]))
    assert flex_mgr.post("/api/ops/email/quality", {"message_id": "<fx@x>", "rating": "Needs Improvement",
                                                    "note": "private coaching note"}).status_code == 200
    seen = lambda who: [(q["subject"], q["note"]) for q in who.get("/api/ops/my?as=1002").get_json()["quality"]["reviews"]]  # noqa: E731
    assert seen(flex_mgr) == [("FLEX ONLY SUBJECT", "private coaching note")]
    assert seen(admin) == [("FLEX ONLY SUBJECT", "private coaching note")]
    assert cara.get("/api/ops/my?as=1002").status_code == 404   # another team's manager cannot even open the page
    # an employee of both teams: the Card manager opens the page but reads none of the Flex Cube ratings
    make_account("1003", "user", "Dina Both", email="dina@example.test", teams=["Flex Cube", "Card Management"])
    put_snapshot(ops, "email", email_data([replied_row("<fx2@x>", "dina@example.test", subject="FLEX TWO")]))
    assert flex_mgr.post("/api/ops/email/quality", {"message_id": "<fx2@x>", "rating": "Good", "note": "kept private"}).status_code == 200
    both = cara.get("/api/ops/my?as=1003").get_json()["quality"]
    assert both["reviews"] == [] and both["counts"]["reviewed"] == 0


def test_rerating_keeps_the_appeal_text_but_resets_the_decision(admin, mgr, hub, ops):
    mgr.post("/api/ops/email/quality", {"message_id": "<r1@x>", "rating": "Good"})
    ops.update_override("quality", "<r1@x>", lambda old: {**old, "appeal": "text", "appeal_decision": "Accepted",
                                                         "rating_withdrawn": True})
    mgr.post("/api/ops/email/quality", {"message_id": "<r1@x>", "rating": "Excellent"})
    review = ops.overrides("quality")["<r1@x>"]
    assert review["appeal"] == "text" and review["appeal_decision"] == "" and review["rating_withdrawn"] is False


def test_manager_cannot_act_on_another_teams_mail(app, flex_mgr, hub, ops, clock):
    card_mail = "<p2@x>"
    assert flex_mgr.post("/api/ops/email/assign", {"message_id": card_mail, "assigned_to_email": "x@example.test"}).status_code == 403
    assert flex_mgr.post("/api/ops/email/no-action", {"message_id": card_mail}).status_code == 403
    assert flex_mgr.post("/api/ops/email/team", {"message_id": card_mail, "team": "Flex Cube"}).status_code == 403
    assert flex_mgr.post("/api/ops/email/quality", {"message_id": "<r2@x>", "rating": "Good"}).status_code == 403
    assert ops.overrides("assignments") == {} and ops.overrides("no_action") == {} and ops.overrides("team") == {}
    # ... but may work on their own and hand a mail over to another team
    assert flex_mgr.post("/api/ops/email/team", {"message_id": "<p1@x>", "team": "Card Management"}).status_code == 200
    assert flex_mgr.post("/api/ops/email/assign", {"message_id": "<p1@x>", "assigned_to_email": "x@example.test"}).status_code == 403
    assert ("ops_team", "<p1@x>", "fadi") in audit_actions(app, clock)


def test_manager_can_undo_only_own_team_no_action(admin, flex_mgr, hub, ops):
    admin.post("/api/ops/email/no-action", {"message_id": "<p2@x>"})
    assert flex_mgr.delete("/api/ops/email/no-action", json={"message_id": "<p2@x>"}).status_code == 403
    assert "<p2@x>" in ops.overrides("no_action")


# ---------------------------------------------------------------------------- classification rules
def test_classification_rule_roundtrip(app, admin, hub, ops, clock):
    r = admin.post("/api/ops/email/classification/rule", {"keyword": "  SWIFT   gateway ", "team": "Flex Cube"})
    assert r.status_code == 200 and r.get_json()["rule"] == {"keyword": "SWIFT gateway", "team": "Flex Cube"}
    assert ops.read_settings_raw()["email"]["classification"]["system_team_map"] == {"SWIFT gateway": "Flex Cube"}
    opts = admin.get("/api/ops/email").get_json()["options"]
    assert opts["system_team_map"] == {"SWIFT gateway": "Flex Cube"}
    # same keyword (any case) replaces the rule
    admin.post("/api/ops/email/classification/rule", {"keyword": "swift GATEWAY", "team": "Card Management"})
    assert admin.get("/api/ops/settings").get_json()["settings"]["email"]["classification"]["system_team_map"] == {
        "swift GATEWAY": "Card Management"}
    d = admin.delete("/api/ops/email/classification/rule", json={"keyword": "SWIFT gateway"})
    assert d.status_code == 200 and d.get_json()["system_team_map"] == {}
    assert admin.delete("/api/ops/email/classification/rule", json={"keyword": "SWIFT gateway"}).status_code == 404
    assert ("ops_classification_rule", "swift GATEWAY", "admin") in audit_actions(app, clock)


def test_classification_rule_validation(admin, hub):
    for body in ({"keyword": "x", "team": "Flex Cube"}, {"keyword": "y" * 81, "team": "Flex Cube"},
                 {"keyword": "swift", "team": "Nope"}, {"keyword": "swift", "team": "Unclassified"},
                 {"keyword": "swift"}, {"team": "Flex Cube"}):
        assert admin.post("/api/ops/email/classification/rule", body).status_code == 400, body
    assert admin.delete("/api/ops/email/classification/rule", json={}).status_code == 400
    assert admin.delete("/api/ops/email/classification/rule", json=[1, 2]).status_code == 400          # not an object
    assert admin.delete("/api/ops/email/assign", json=[1]).status_code == 400


def test_manager_rules_are_limited_to_their_own_team(admin, flex_mgr, mgr, hub, ops):
    assert flex_mgr.post("/api/ops/email/classification/rule", {"keyword": "swift", "team": "Flex Cube"}).status_code == 200
    assert flex_mgr.post("/api/ops/email/classification/rule", {"keyword": "visa", "team": "Card Management"}).status_code == 403
    admin.post("/api/ops/email/classification/rule", {"keyword": "visa", "team": "Card Management"})
    assert flex_mgr.delete("/api/ops/email/classification/rule", json={"keyword": "visa"}).status_code == 403
    assert flex_mgr.post("/api/ops/email/classification/rule", {"keyword": "visa", "team": "Flex Cube"}).status_code == 403
    assert flex_mgr.delete("/api/ops/email/classification/rule", json={"keyword": "swift"}).status_code == 200
    # the manager's options only list the rules of their team
    assert flex_mgr.get("/api/ops/email").get_json()["options"]["system_team_map"] == {}
    assert mgr.post("/api/ops/email/classification/rule", {"keyword": "mastercard", "team": "Card Management"}).status_code == 200
    assert set(ops.read_settings_raw()["email"]["classification"]["system_team_map"]) == {"visa", "mastercard"}


# ------------------------------------------------------------- learned classifications (pending admin approval)
def uncls_replied_row(mid, suggested="Flex Cube", by="bob@example.test", hint="password reset; vpn access", **kw):
    """An Unclassified mail a team member answered: the scan proposes filing it under the replier's team."""
    return replied_row(mid, by, 20, team="Unclassified", all_teams=["Unclassified"], teams_pending=[],
                       teams_replied=[], classified_by="none", classification_hint=hint, suggested_team=suggested, **kw)


SUG = "/api/ops/email/classification/suggestion"
SUGS = SUG + "s"


def test_classification_suggestion_is_listed_scoped_and_learned_on_approval(app, admin, flex_mgr, make_account, hub, ops, clock):
    cara = make_account("cara", "manager", "Cara Card", email="cara@example.test", team="Card Management")
    rows = [uncls_replied_row("<u1@x>"), uncls_replied_row("<u2@x>", hint="statement request"),
            email_row("<c1@x>", classified_by="system")]
    put_snapshot(ops, "email", email_data(rows))
    beat(ops, "email")
    # admin and the Flex Cube manager see both suggestions; the Card manager sees none (not their team)
    s = admin.get(SUGS).get_json()["suggestions"]
    assert {x["message_id"] for x in s} == {"<u1@x>", "<u2@x>"}
    one = next(x for x in s if x["message_id"] == "<u1@x>")
    assert one["suggested_team"] == "Flex Cube" and one["keyword"] == "password reset" and one["reply_by"] == "Replier bob"
    assert {x["message_id"] for x in flex_mgr.get(SUGS).get_json()["suggestions"]} == {"<u1@x>", "<u2@x>"}
    assert cara.get(SUGS).get_json()["suggestions"] == []
    # the Card manager cannot approve a Flex Cube suggestion; a real team is required
    assert cara.post(SUG, {"message_id": "<u1@x>", "keyword": "password reset", "team": "Flex Cube"}).status_code == 403
    assert admin.post(SUG, {"message_id": "<u1@x>", "keyword": "x", "team": "Flex Cube"}).status_code == 400
    assert admin.post(SUG, {"message_id": "<u1@x>", "keyword": "password reset", "team": "Unclassified"}).status_code == 400
    # approve: the keyword rule is learned AND this mail is filed under Flex Cube now
    r = admin.post(SUG, {"message_id": "<u1@x>", "keyword": "password reset", "team": "Flex Cube"})
    assert r.status_code == 200 and r.get_json()["rule"] == {"keyword": "password reset", "team": "Flex Cube"}
    assert ops.read_settings_raw()["email"]["classification"]["system_team_map"] == {"password reset": "Flex Cube"}
    assert ops.overrides("team")["<u1@x>"]["team"] == "Flex Cube"
    assert ("ops_classification_learn", "<u1@x>", "admin") in audit_actions(app, clock)
    # the approved one is no longer proposed; the other still is
    assert {x["message_id"] for x in admin.get(SUGS).get_json()["suggestions"]} == {"<u2@x>"}


def test_classification_suggestion_approval_cannot_move_a_mail_that_is_not_a_pending_suggestion(
        admin, flex_mgr, make_account, hub, ops):
    cara = make_account("cara", "manager", "Cara Card", email="cara@example.test", team="Card Management")
    rows = [email_row("<card1@x>", team="Card Management", all_teams=["Card Management"], classified_by="system"),
            uncls_replied_row("<u1@x>")]
    put_snapshot(ops, "email", email_data(rows))
    beat(ops, "email")
    # another team's classified mail is not a suggestion: it cannot be pulled into the manager's own team
    assert flex_mgr.post(SUG, {"message_id": "<card1@x>", "keyword": "card stuff", "team": "Flex Cube"}).status_code == 404
    # a suggestion for another team cannot be filed under the manager's own team either
    assert cara.post(SUG, {"message_id": "<u1@x>", "keyword": "password reset", "team": "Card Management"}).status_code == 403
    assert ops.overrides("team") == {}
    assert ops.read_settings_raw().get("email", {}).get("classification", {}).get("system_team_map", {}) == {}
    # approving twice: the second time it is no longer pending
    assert flex_mgr.post(SUG, {"message_id": "<u1@x>", "keyword": "password reset", "team": "Flex Cube"}).status_code == 200
    assert flex_mgr.post(SUG, {"message_id": "<u1@x>", "keyword": "password reset", "team": "Flex Cube"}).status_code == 404


def test_classification_suggestion_can_be_dismissed(app, admin, hub, ops, clock):
    put_snapshot(ops, "email", email_data([uncls_replied_row("<u9@x>")]))
    beat(ops, "email")
    assert admin.delete(SUG, json={"message_id": "<u9@x>"}).status_code == 200
    assert admin.get(SUGS).get_json()["suggestions"] == []           # remembered as rejected, not proposed again
    assert ("ops_classification_learn", "<u9@x>", "admin") in audit_actions(app, clock)


# ---------------------------------------------------------------------------- jobs
def test_run_now_creates_a_visible_request_and_audits(app, admin, hub, ops, clock):
    r = admin.post("/api/ops/jobs/sla/run", {})
    assert r.status_code == 200 and r.get_json()["queued"] is True
    assert len(ops.pending_run_requests("sla")) == 1
    listing = admin.get("/api/ops/jobs").get_json()
    assert [p["job"] for p in listing["pending"]] == ["sla"] and listing["pending"][0]["requested_by"] == "admin"
    assert set(listing["jobs"]) == set(schemas.JOBS) and listing["jobs"]["sla"]["online"] is True
    again = admin.post("/api/ops/jobs/sla/run", {}).get_json()
    assert again["queued"] is True and again["already_pending"] is True
    assert len(ops.pending_run_requests("sla")) == 1
    assert ("ops_run_job", "sla", "admin") in audit_actions(app, clock)


def test_run_now_rules(admin, mgr, emp, hub, ops):
    assert admin.post("/api/ops/jobs/nope/run", {}).status_code == 404
    assert admin.post("/api/ops/jobs/../run", {}).status_code == 404
    assert mgr.post("/api/ops/jobs/itsm/run", {}).status_code == 403          # managers no longer run scans / "Run now"
    assert mgr.post("/api/ops/jobs/email_report/run", {}).status_code == 403
    assert emp.post("/api/ops/jobs/sla/run", {}).status_code == 403
    assert admin.post("/api/ops/jobs/email_report/run", {}).status_code == 200
    assert admin.post("/api/ops/jobs/itsm/run", {}).status_code == 200
    assert {r["job"] for r in ops.pending_run_requests()} == {"itsm", "email_report"}


# ---------------------------------------------------------------------------- auto-assign
def test_auto_assign_scope_and_access(admin, flex_mgr, emp, hub):
    assert emp.get("/api/ops/auto-assign").status_code == 403            # employees have no auto-assign
    body = flex_mgr.get("/api/ops/auto-assign").get_json()
    assert set(body["teams"]) == {"Flex Cube"} and body["is_admin"] is False
    assert admin.get("/api/ops/auto-assign").get_json()["is_admin"] is True
    r = flex_mgr.put("/api/ops/auto-assign/team", {"team": "Card Management", "kind": "email", "enabled": True, "rules": []})
    assert r.status_code == 403                                          # a manager cannot configure another team


def test_auto_assign_admin_manages_applications(admin, flex_mgr):
    assert admin.post("/api/ops/auto-assign/apps", {"name": "Flexcube", "team": "Flex Cube"}).status_code == 200
    r = admin.post("/api/ops/auto-assign/apps", {"name": "Prepaid", "team": "Card Management"}).get_json()
    assert {a["name"] for a in r["applications"]} == {"Flexcube", "Prepaid"}
    assert flex_mgr.get("/api/ops/auto-assign").get_json()["teams"]["Flex Cube"]["applications"] == ["Flexcube"]
    assert flex_mgr.post("/api/ops/auto-assign/apps", {"name": "X", "team": "Flex Cube"}).status_code == 403    # admin only
    left = admin.delete("/api/ops/auto-assign/apps?name=Flexcube").get_json()
    assert {a["name"] for a in left["applications"]} == {"Prepaid"}


def test_auto_assign_distributes_pending_mail_balanced(admin, flex_mgr, make_account, ops):
    make_account("3001", "user", "Emp One", email="one@example.test", team="Flex Cube")
    make_account("3002", "user", "Emp Two", email="two@example.test", team="Flex Cube")
    put_snapshot(ops, "email", email_data([email_row("<m1@x>"), email_row("<m2@x>"), email_row("<m3@x>")]))
    r = flex_mgr.put("/api/ops/auto-assign/team", {"team": "Flex Cube", "kind": "email", "enabled": True,
        "rules": [{"employees": ["one@example.test", "two@example.test"], "applications": [], "from": "", "to": ""}]}).get_json()
    assert r["assigned"] == 3
    who = {m: ops.overrides("assignments")[m]["assigned_to_email"] for m in ("<m1@x>", "<m2@x>", "<m3@x>")}
    assert set(who.values()) == {"one@example.test", "two@example.test"}
    vals = list(who.values())
    assert sorted([vals.count("one@example.test"), vals.count("two@example.test")]) == [1, 2]        # balanced
    assert ops.overrides("assignments")["<m1@x>"]["marked_by"] == "Auto-assign"
    assert flex_mgr.post("/api/ops/auto-assign/apply", {}).get_json()["assigned"] == 0                # idempotent


def test_auto_assign_application_filter(admin, flex_mgr, make_account, ops):
    make_account("4001", "user", "Only Flex", email="flex@example.test", team="Flex Cube")
    admin.post("/api/ops/auto-assign/apps", {"name": "Flexcube", "team": "Flex Cube"})
    put_snapshot(ops, "email", email_data([email_row("<f1@x>"), email_row("<f2@x>")]))     # both detect "Flexcube"
    flex_mgr.put("/api/ops/auto-assign/team", {"team": "Flex Cube", "kind": "email", "enabled": True,
        "rules": [{"employees": ["flex@example.test"], "applications": ["Flexcube"], "from": "", "to": ""}]})
    assigns = ops.overrides("assignments")
    assert assigns["<f1@x>"]["assigned_to_email"] == "flex@example.test"
    assert assigns["<f2@x>"]["assigned_to_email"] == "flex@example.test"


def test_auto_assign_max_per_run_is_saved_and_limits_each_run(admin, flex_mgr, make_account, ops):
    make_account("3101", "user", "Emp Max", email="max@example.test", team="Flex Cube")
    put_snapshot(ops, "email", email_data([email_row(f"<n{i}@x>") for i in range(5)]))
    r = flex_mgr.put("/api/ops/auto-assign/team", {"team": "Flex Cube", "kind": "email", "enabled": True, "max_per_run": 2,
        "rules": [{"employees": ["max@example.test"], "applications": [], "from": "", "to": ""}]}).get_json()
    assert r["block"]["max_per_run"] == 2 and r["assigned"] == 2                  # this run: 2, the rest wait
    assert flex_mgr.post("/api/ops/auto-assign/apply", {}).get_json()["assigned"] == 2      # the next run: 2 more
    got = flex_mgr.get("/api/ops/auto-assign").get_json()["teams"]["Flex Cube"]["email"]
    assert got["max_per_run"] == 2
    bad = flex_mgr.put("/api/ops/auto-assign/team", {"team": "Flex Cube", "kind": "ticketing", "enabled": True,
                                                      "max_per_run": "lots", "rules": []}).get_json()
    assert bad["block"]["max_per_run"] == 0                                         # not a number = no limit


def test_auto_assign_disabled_does_nothing(admin, flex_mgr, make_account, ops):
    make_account("5001", "user", "Somebody", email="some@example.test", team="Flex Cube")
    put_snapshot(ops, "email", email_data([email_row("<z1@x>")]))
    r = flex_mgr.put("/api/ops/auto-assign/team", {"team": "Flex Cube", "kind": "email", "enabled": False,
        "rules": [{"employees": ["some@example.test"], "applications": [], "from": "", "to": ""}]}).get_json()
    assert r["assigned"] == 0 and not ops.overrides("assignments").get("<z1@x>")


# ---------------------------------------------------------------------------- SLA analysis (live recording)
def test_itsm_analysis_records_execution_time_and_conclusion(admin, emp, ops):
    assert emp.get("/api/ops/itsm-analysis").status_code == 403           # the ITSM analysis is admin/manager only
    assert admin.get("/api/ops/itsm-analysis").get_json()["available"] is False      # no scan recorded yet
    data = {"categories": ["Implementation Requests", "Implementation Completed", "Rollback Requests"],
            "roles": [{"role": "IMP User provisioning", "error": None, "categories": {
                "Implementation Completed": {"count": 1, "rows": [{"number": "100001", "summary": "Add user"}], "warning": None, "note": None},
                "Implementation Requests": {"count": 1, "rows": [{"number": "100002", "summary": "Reset password"}], "warning": None, "note": None},
                "Rollback Requests": {"count": 1, "rows": [{"number": "100003", "summary": "Revert change"}], "warning": None, "note": None}}}],
            "totals": {}, "changes": {"new": [], "closed": []}}
    put_snapshot(ops, "itsm", data)
    body = admin.get("/api/ops/itsm-analysis").get_json()
    assert body["available"] is True and body["buckets"] == {"pending": 1, "completed": 1, "rollback": 1}
    done = next(r for r in body["records"] if r["number"] == "100001")
    assert done["conclusion"] == "completed" and done["completed_at"][:10] == SNAP_AT[:10]    # the execution time
    assert body["options"]["roles"] == ["IMP User provisioning"]
    assert {r["number"] for r in admin.get("/api/ops/itsm-analysis?conclusion=completed").get_json()["records"]} == {"100001"}
    assert admin.get("/api/ops/itsm-analysis?from=2020-01-01&to=2020-01-02").get_json()["meta"]["records"] == 0


def test_itsm_recording_stamps_execution_time_across_two_scans(admin, ops):
    """The live point: a CR seen pending in one scan and completed in a LATER scan gets its execution time stamped
    from the scan that first saw it completed (the scoreboard has no explicit completion timestamp)."""
    def snap(completed_rows, generated):
        data = {"categories": ["Implementation Requests", "Implementation Completed", "Rollback Requests"],
                "roles": [{"role": "IMP User provisioning", "error": None, "categories": {
                    "Implementation Requests": {"count": 1, "rows": [{"number": "700001", "summary": "Add user"}], "warning": None, "note": None},
                    "Implementation Completed": {"count": len(completed_rows), "rows": completed_rows, "warning": None, "note": None}}}],
                "totals": {}, "changes": {"new": [], "closed": []}}
        put_snapshot(ops, "itsm", data, at=generated)
    snap([], "2026-10-18T09:00:00")                                       # scan 1: 700001 is only pending
    b1 = admin.get("/api/ops/itsm-analysis").get_json()
    assert b1["buckets"]["pending"] == 1 and b1["buckets"]["completed"] == 0
    snap([{"number": "700001", "summary": "Add user"}], "2026-10-18T10:30:00")   # scan 2: now completed
    b2 = admin.get("/api/ops/itsm-analysis").get_json()
    assert b2["buckets"] == {"pending": 0, "completed": 1, "rollback": 0}
    done = next(r for r in b2["records"] if r["number"] == "700001")
    assert done["conclusion"] == "completed" and done["completed_at"] == "2026-10-18T10:30:00"


def test_sla_config_endpoints(admin, flex_mgr, emp):
    assert emp.get("/api/ops/sla-config").status_code == 403              # employees have no Ticketing config
    body = flex_mgr.get("/api/ops/sla-config").get_json()                 # a manager may view it
    assert body["config"]["sla_minutes"] == 180 and body["is_admin"] is False
    assert flex_mgr.put("/api/ops/sla-config", {"sla_minutes": 120}).status_code == 403    # only an admin edits
    r = admin.put("/api/ops/sla-config", {"sla_minutes": 240, "statuses": {"executed": ["تم التنفيذ"]}}).get_json()
    assert r["ok"] is True and r["config"]["sla_minutes"] == 240
    assert admin.get("/api/ops/sla-config").get_json()["config"]["sla_minutes"] == 240


def _sla_scan(ops, at, reqs):
    """A live SLA snapshot (the Re-Assign scrape) holding ``reqs`` = [(key, security user, sent at)]."""
    put_snapshot(ops, "sla", {"thresholds": {"warning_minutes": 60, "breach_minutes": 120}, "totals": {}, "by_user": [],
                              "requests": [{"key": k, "security_user": u, "sent_at": t, "age_minutes": 0,
                                            "status": "within", "fields": {"النظام": "Flexcube", "اسم الموظف": "x"}}
                                           for k, u, t in reqs]}, at=at)


def test_job_log_tail(admin, ops):
    for i in range(10):
        ops.append_log("sla", f"line {i}")
    lines = admin.get("/api/ops/jobs/sla/log?lines=3").get_json()["lines"]
    assert len(lines) == 3 and lines[-1].endswith("line 9")
    assert admin.get("/api/ops/jobs/nope/log").status_code == 404
    assert admin.get("/api/ops/jobs/sla/log?lines=abc").status_code == 400


def test_collector_info_is_passed_through_without_secrets(admin, ops):
    ops.write_state("collector_info", {"host": "PC1", "pinned_keys": ["sla.username"], "accounts": {
        "sla": {"username": "user-a", "has_password": True, "password": "secret-pw"}}, "token": "t"})
    body = admin.get("/api/ops/jobs").get_json()["collector_info"]
    assert body["pinned_keys"] == ["sla.username"] and body["accounts"]["sla"] == {"username": "user-a", "has_password": True}
    assert "secret-pw" not in json.dumps(admin.get("/api/ops/jobs").get_json())
    settings = admin.get("/api/ops/settings").get_json()
    assert settings["pinned_keys"] == ["sla.username"] and "secret-pw" not in json.dumps(settings)


# ---------------------------------------------------------------------------- settings
def test_settings_roundtrip_audit_and_errors(app, admin, mgr, ops, clock):
    got = admin.get("/api/ops/settings").get_json()
    assert got["settings"]["email"]["interval_minutes"] == 10 and "Mon" in got["weekdays"]
    doc = got["settings"]
    doc["email"]["interval_minutes"] = 5
    doc["email"]["business_emails"] = ["Team.Box@Example.test", "team.box@example.test"]
    r = admin.put("/api/ops/settings", doc)
    assert r.status_code == 200 and r.get_json()["settings"]["email"]["business_emails"] == ["team.box@example.test"]
    assert "email.interval_minutes" in r.get_json()["changed"]
    assert admin.get("/api/ops/settings").get_json()["settings"]["email"]["interval_minutes"] == 5
    assert ops.read_settings_raw()["email"]["interval_minutes"] == 5
    assert admin.put("/api/ops/settings", {"settings": doc}).status_code == 200            # wrapped form
    entry = next(e for e in app.extensions["store"].read_audit("2026-10-11", "2026-10-11") if e["action"] == "ops_settings_saved")
    assert "email.interval_minutes" in entry["details"]["changed"]

    doc["email"]["interval_minutes"] = 0
    doc["sla"]["breach_minutes"] = 10
    doc["email"]["business_emails"] = ["not-an-address"]
    bad = admin.put("/api/ops/settings", doc)
    assert bad.status_code == 400
    paths = {e["path"] for e in bad.get_json()["errors"]}
    assert {"email.interval_minutes", "sla.breach_minutes", "email.business_emails"} <= paths
    assert admin.get("/api/ops/settings").get_json()["settings"]["email"]["interval_minutes"] == 5      # nothing was saved
    assert mgr.put("/api/ops/settings", doc).status_code == 403
    assert mgr.get("/api/ops/settings").status_code == 403
    assert admin.put("/api/ops/settings", [1]).status_code == 400


def test_settings_never_accept_a_password(admin, ops):
    doc = admin.get("/api/ops/settings").get_json()["settings"]
    doc["sla"]["password"] = "secret-pw"
    r = admin.put("/api/ops/settings", doc)
    assert r.status_code == 400 and "collector_config.json" in r.get_json()["error"]
    assert "secret-pw" not in json.dumps(ops.read_settings_raw())
    # the CSS selector of the password box is not a password
    doc["sla"].pop("password")
    doc["sla"]["selectors"] = {"password": "#pwd"}
    assert admin.put("/api/ops/settings", doc).status_code == 200


# ---------------------------------------------------------------------------- reports
def test_reports_list_and_download(app, admin, mgr, emp, ops, clock):
    (ops.report_dir("email") / "daily.xlsx").write_bytes(b"PK-fake-xlsx")
    (ops.report_dir("itsm") / "status.html").write_text("<p>hi</p>", encoding="utf-8")
    (ops.reports_dir / "secret.txt").write_text("outside", encoding="utf-8")
    listing = admin.get("/api/ops/reports").get_json()["reports"]
    assert {(r["module"], r["name"]) for r in listing} == {("email", "daily.xlsx"), ("itsm", "status.html")}
    assert admin.get("/api/ops/reports?module=email").get_json()["reports"][0]["url"] == "/export/ops/report/email/daily.xlsx"
    assert admin.get("/api/ops/reports?module=bad").status_code == 404
    r = mgr.get("/export/ops/report/email/daily.xlsx")
    assert r.status_code == 200 and r.data == b"PK-fake-xlsx" and "attachment" in r.headers["Content-Disposition"]
    assert emp.get("/export/ops/report/email/daily.xlsx").status_code == 403
    assert app.test_client().get("/export/ops/report/email/daily.xlsx").status_code == 401
    assert ("ops_report_download", "email/daily.xlsx", "admin") not in audit_actions(app, clock)
    assert ("ops_report_download", "email/daily.xlsx", "mona") in audit_actions(app, clock)


@pytest.mark.parametrize("path", [
    "/export/ops/report/email/../secret.txt", "/export/ops/report/email/..%2Fsecret.txt",
    "/export/ops/report/email/%2e%2e%2fsecret.txt", "/export/ops/report/email/..%5Csecret.txt",
    "/export/ops/report/email/%2Fetc%2Fpasswd", "/export/ops/report/email//etc/passwd",
    "/export/ops/report/..%2Fsnapshots/email.json", "/export/ops/report/email/.hidden", "/export/ops/report/email/nope.xlsx",
    "/export/ops/report/email/a%00.xlsx", "/export/ops/report/bad/daily.xlsx"])
def test_report_download_rejects_traversal(admin, ops, path):
    (ops.report_dir("email") / "daily.xlsx").write_bytes(b"x")
    (ops.report_dir("email") / ".hidden").write_bytes(b"x")
    (ops.reports_dir / "secret.txt").write_text("outside", encoding="utf-8")
    r = admin.get(path)
    assert r.status_code in (404, 400)
    assert b"outside" not in r.data and b"email.json" not in r.data


# ---------------------------------------------------------------------------- exports
def _book(resp):
    assert resp.status_code == 200, resp.get_json()
    return openpyxl.load_workbook(io.BytesIO(resp.data))


def test_email_export_has_overrides_green_header_and_safe_text(admin, flex_mgr, hub, ops):
    rows = [email_row("<f1>", subject="=HYPERLINK(\"http://evil\")"), email_row("<f2>", subject="Normal")]
    put_snapshot(ops, "email", email_data(rows))
    admin.post("/api/ops/email/no-action", {"message_id": "<f2>"})
    wb = _book(admin.get("/export/ops/email.xlsx"))
    assert wb.sheetnames == ["Summary", "E-mails", "Teams"]
    ws = wb["E-mails"]
    assert ws["A1"].fill.fgColor.rgb.endswith("00573F") and ws["A1"].font.bold
    data = {r[1]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert data["Normal"][5] == "No Action Needed"
    assert ws["B2"].data_type == "s" and ws["B2"].value.startswith("=")
    csv_resp = admin.get("/export/ops/email.csv")
    assert csv_resp.status_code == 200 and csv_resp.data.startswith(b"\xef\xbb\xbfReceived")
    assert b"'=HYPERLINK" in csv_resp.data


def test_exports_drop_control_characters_and_still_block_formulas(admin, hub, ops):
    rows = [email_row("<c1>", subject="Esc\x1bape"), email_row("<c2>", subject="\x01=HYPERLINK(\"http://evil\")")]
    put_snapshot(ops, "email", email_data(rows))
    ws = _book(admin.get("/export/ops/email.xlsx"))["E-mails"]
    subjects = {row[1].value: row[1] for row in ws.iter_rows(min_row=2)}
    assert set(subjects) == {"Escape", "=HYPERLINK(\"http://evil\")"}
    assert subjects["=HYPERLINK(\"http://evil\")"].data_type == "s"
    put_snapshot(ops, "sla", sla_data([sla_request("k1", "Sam\x02", **{"Remark\x03": "a\x0cb"})]))
    head, row = list(_book(admin.get("/export/ops/sla.xlsx"))["Requests"].iter_rows(values_only=True))
    assert row[0] == "Sam" and "Remark" in head and row[head.index("Remark")] == "ab"


def test_email_export_is_team_scoped_for_managers(flex_mgr, hub):
    wb = _book(flex_mgr.get("/export/ops/email.xlsx"))
    subjects = {r[1] for r in wb["E-mails"].iter_rows(min_row=2, values_only=True)}
    assert "Card limit change" not in subjects and "Reset Flexcube password" in subjects
    assert flex_mgr.get("/export/ops/email.csv").data.count(b"Card limit change") == 0


def test_sla_and_itsm_exports(app, admin, hub, clock):
    wb = _book(admin.get("/export/ops/sla.xlsx"))
    assert wb.sheetnames == ["Summary", "Requests", "By user"]
    rows = list(wb["Requests"].iter_rows(values_only=True))
    assert rows[0][:4] == ("Security user", "Sent at", "Age (min)", "Status") and rows[0][4] == "النظام"
    assert {r[3] for r in rows[1:]} == {"SLA breached", "Reminder"}
    wb = _book(admin.get("/export/ops/itsm.xlsx"))
    assert wb.sheetnames == ["Summary", "Requests", "By role", "Changes"]
    assert list(wb["Requests"].iter_rows(min_row=2, values_only=True))[0][2] == "100001"
    assert any("Set Role button not found" in (r[1] or "") for r in wb["By role"].iter_rows(values_only=True))
    assert admin.get("/export/ops/itsm.csv").data.startswith(b"\xef\xbb\xbfRole,Category,Number,Summary")
    assert ("ops_export", "itsm.xlsx", "admin") in audit_actions(app, clock)


def test_export_without_data_is_404(admin):
    assert admin.get("/export/ops/sla.xlsx").status_code == 404
    assert admin.get("/export/ops/nope.xlsx").status_code == 404
    assert admin.get("/export/ops/sla.pdf").status_code == 404


# ---------------------------------------------------------------------------- overview and shared
def test_overview_numbers_and_attention(admin, hub):
    body = admin.get("/api/ops/overview").get_json()
    assert body["email"]["available"] and body["email"]["business_emails"] == 5
    assert body["sla"]["breached"] == 1 and body["sla"]["reminder"] == 1 and body["itsm"]["totals"]["Rollback Requests"] == 1
    assert body["itsm"]["role_errors"][0]["role"] == "IMP UAT I-Flex SMS"
    texts = [(a["level"], a["module"], a["text"]) for a in body["attention"]]
    assert ("danger", "sla", "1 ticket open past the 2 h SLA") in texts
    assert ("warn", "sla", "1 ticket open past the reminder time (1 h)") in texts
    assert ("danger", "email", "1 e-mail waiting past the SLA with no team reply") in texts
    assert ("warn", "email", "1 e-mail carried over from earlier days (oldest 2 days)") in texts
    assert ("info", "email", "1 e-mail not classified to a team") in texts
    assert ("warn", "itsm", "1 rollback request waiting") in texts
    assert any(m == "itsm" and lvl == "danger" and "IMP UAT I-Flex SMS" in t for lvl, m, t in texts)
    assert [a["level"] for a in body["attention"]] == sorted((a["level"] for a in body["attention"]),
                                                             key=["danger", "warn", "info"].index)
    assert all(set(a) == {"level", "module", "text", "link"} for a in body["attention"])
    assert set(body["collectors"]) == set(schemas.JOBS) and body["shared"]["email"]["business_emails"] == 5


def test_overview_attention_for_collector_problems(admin, hub, ops, clock):
    clock.value = clock.value.replace(hour=11)
    beat(ops, "sla", at="2026-10-11T10:59:00", last_status="error", last_error="Login failed")
    beat(ops, "email_report", at="2026-10-11T10:59:00", last_status="error", last_error="Outlook busy")
    body = admin.get("/api/ops/overview").get_json()
    texts = [a["text"] for a in body["attention"] if a["module"] in ("sla", "email", "itsm") and a["link"] == "/system"]
    assert any("itsm" not in t and "E-mail collector is offline" in t for t in texts)
    assert "The last Ticketing scan failed: Login failed" in texts
    assert "The last E-mail daily report failed: Outlook busy" in texts
    assert any("data is 120 minutes old" in t for t in texts)


def test_overview_with_no_data_says_so(admin):
    body = admin.get("/api/ops/overview").get_json()
    assert body["email"]["available"] is False and body["sla"]["available"] is False
    assert sum(1 for a in body["attention"] if "no data yet" in a["text"]) == 3


def test_overview_is_team_scoped_for_managers(flex_mgr, hub):
    body = flex_mgr.get("/api/ops/overview").get_json()
    assert body["scope"] == "Flex Cube" and body["scope_teams"] == ["Flex Cube"] and body["email"]["escalated_count"] == 0
    assert all(t["team"] == "Flex Cube" for t in body["email"]["teams"])
    assert not any("past the SLA" in a["text"] for a in body["attention"])
    # the strip is the same component for everybody, but its e-mail numbers are the viewer's own teams'
    strip = body["shared"]["email"]
    assert strip["business_emails"] == 2 and strip["scope"] == "teams" and strip["teams"] == ["Flex Cube"]
    assert strip["sla_breached"] == 0


def test_the_oldest_request_is_aged_by_the_clock_everywhere_like_the_live_tab(admin, emp, ops, clock):
    """Home, the Team-status strip and My work give the oldest request's age NOW (from its sent date), as the Live
    tickets tab does - not the age frozen at the last scan; with no dated request it is "—" (None), not 0."""
    put_snapshot(ops, "sla", sla_data([sla_request("k1", "احمد سالم", "breached", 150), sla_request("k2", "Bob", "within", 20)]))
    clock.advance(minutes=40)                                                    # 40 minutes after the scan
    assert admin.get("/api/ops/overview").get_json()["sla"]["oldest_minutes"] == 190
    assert admin.get("/api/ops/shared").get_json()["sla"]["oldest_minutes"] == 190
    assert emp.get("/api/ops/my").get_json()["sla"]["totals"]["oldest_minutes"] == 190
    put_snapshot(ops, "sla", sla_data([sla_request("k3", "Bob", "no_date", None)]))
    assert admin.get("/api/ops/overview").get_json()["sla"]["oldest_minutes"] is None


def test_the_oldest_request_and_my_tickets_count_from_the_sent_date_and_show_the_first_assignment(admin, emp, ops, clock):
    """Like the owner's sla_monitor.py, the live SLA counts from «تاريخ الارسال للسرية»: a request sent 150 minutes
    before the scan is 190 minutes old 40 minutes later on Home and My work, even though the scan noted it was first
    assigned 30 minutes before it; My tickets still shows that assignment and since when its holder has it."""
    req = sla_request("k1", "احمد سالم", "breached", 150)
    req.update(sent_at="2026-10-11T06:30:00", assigned_at="2026-10-11T08:30:00", holder_since="2026-10-11T08:50:00")
    put_snapshot(ops, "sla", sla_data([req, sla_request("k2", "Bob", "within", 20)]))
    clock.advance(minutes=40)
    assert admin.get("/api/ops/overview").get_json()["sla"]["oldest_minutes"] == 190
    mine = emp.get("/api/ops/my").get_json()["sla"]
    assert mine["totals"]["oldest_minutes"] == 190 and mine["totals"]["breached"] == 1
    one = mine["requests"][0]
    assert one["sent_at"] == "2026-10-11T06:30:00" and one["age_minutes"] == 150 and one["status"] == "breached"
    assert one["assigned_at"] == "2026-10-11T08:30:00" and one["holder_since"] == "2026-10-11T08:50:00"
    assert "age_from" not in one


def test_a_manager_is_told_when_a_role_of_their_team_could_not_be_read(admin, flex_mgr, ops):
    """Their totals would look complete otherwise. The manager gets the role's name, not the collector's error text."""
    data = itsm_data(errors=["IMP UAT I-Flex SMS"])                        # a Flex Cube role whose read failed
    put_snapshot(ops, "itsm", data)
    body = flex_mgr.get("/api/ops/overview").get_json()
    assert body["itsm"]["role_errors"] == [{"role": "IMP UAT I-Flex SMS", "error": ""}]
    notes = [a for a in body["attention"] if a["module"] == "itsm" and "could not be read" in a["text"]]
    assert len(notes) == 1 and "Set Role button" not in notes[0]["text"] and "missing from the totals" in notes[0]["text"]
    whole = admin.get("/api/ops/overview").get_json()                       # the administrator still gets the reason
    assert whole["itsm"]["role_errors"][0]["error"] == "Set Role button not found"


def test_the_change_requests_tab_says_how_many_scoreboard_requests_it_could_not_list(admin, ops):
    """The scoreboard counts 3 implementation requests but the portal listed one row: the tab says 2 are not in it
    (the Scoreboard's 3 and the tab's 1 are both right - one is a count, the other the rows that could be read)."""
    put_snapshot(ops, "itsm", itsm_data(imp=3, rollback=0))
    body = admin.get("/api/ops/itsm-analysis").get_json()
    assert body["scoreboard"] == {"counted": 3, "listed": 1}


def _ticketing_ready(flex_mgr, make_account):
    make_account("t1", "user", "Tarek Ticket", email="tarek@example.test", team="Flex Cube", portal_id="700201")
    r = flex_mgr.put("/api/ops/auto-assign/team", {"team": "Flex Cube", "kind": "ticketing", "enabled": True,
                                                   "interval_minutes": 1,
                                                   "rules": [{"employees": ["tarek@example.test"], "applications": [],
                                                              "from": "", "to": ""}]})
    assert r.status_code == 200, r.get_json()


def _wait_for(check, seconds=10):
    import time
    end = time.monotonic() + seconds
    while time.monotonic() < end:
        out = check()
        if out:
            return out
        time.sleep(0.05)
    raise AssertionError("timed out")


def test_ticketing_auto_assign_keeps_running_until_stopped_and_never_keeps_the_password(app, flex_mgr, make_account,
                                                                                       monkeypatch):
    """Run now with "keep running": like the bank's bot, a run every N minutes until the manager stops it. The password
    lives only in the running loop's memory: never in the status, the log or any file."""
    from calllog.ops import api as ops_api
    from calllog.ops.engines import assign_engine, common
    _ticketing_ready(flex_mgr, make_account)
    calls = []

    def fake_run(ctx):
        calls.append(ctx.options["password"] == "Portal-Secret-9")
        ctx.log(f"run {len(calls)}")
        return common.JobResult(ok=True, summary="assigned 1 of 1 new request(s)",
                                snapshot_data={"assigned": 1, "found": 1, "by_employee": {"Tarek Ticket": 1}})
    monkeypatch.setattr(assign_engine, "run", fake_run)
    monkeypatch.setattr(ops_api, "TICKETING_LOOP_UNIT_SECONDS", 0.05)      # "1 minute" = 50 ms in the test
    r = flex_mgr.post("/api/ops/auto-assign/run-ticketing", {"team": "Flex Cube", "username": "mgr1",
                                                              "password": "Portal-Secret-9", "repeat": True})
    assert r.status_code == 200, r.get_json()
    st = _wait_for(lambda: (lambda d: d if d["runs"] >= 3 else None)(flex_mgr.get("/api/ops/auto-assign/run-ticketing").get_json()))
    assert st["running"] and st["repeat"] and st["assigned"] >= 3
    assert flex_mgr.post("/api/ops/auto-assign/run-ticketing/stop", {}).status_code == 200
    done = _wait_for(lambda: (lambda d: d if not d["running"] else None)(flex_mgr.get("/api/ops/auto-assign/run-ticketing").get_json()))
    assert done["done"] and "stopped" in done["summary"].lower() and all(calls)
    everything = json.dumps(done) + json.dumps(app.extensions.get("ops_ticketing_run"), default=str)
    assert "Portal-Secret-9" not in everything
    assert not any("Portal-Secret-9" in p.read_text(errors="ignore") for p in app.extensions["ops"].root.rglob("*") if p.is_file())


def test_ticketing_run_gets_the_saved_max_per_run(app, flex_mgr, make_account, monkeypatch):
    from calllog.ops.engines import assign_engine, common
    _ticketing_ready(flex_mgr, make_account)
    block = flex_mgr.get("/api/ops/auto-assign").get_json()["teams"]["Flex Cube"]["ticketing"]
    assert flex_mgr.put("/api/ops/auto-assign/team", {"team": "Flex Cube", "kind": "ticketing", **block,
                                                       "max_per_run": 15}).status_code == 200
    seen = []
    monkeypatch.setattr(assign_engine, "run", lambda ctx: seen.append(ctx.options["teams"]["Flex Cube"]["block"]["max_per_run"])
                        or common.JobResult(ok=True, summary="ok", snapshot_data={"assigned": 0, "found": 0}))
    flex_mgr.post("/api/ops/auto-assign/run-ticketing", {"team": "Flex Cube", "username": "mgr1", "password": "pw"})
    done = _wait_for(lambda: (lambda d: d if not d["running"] else None)(flex_mgr.get("/api/ops/auto-assign/run-ticketing").get_json()))
    assert seen == [15] and done["max_per_run"] == 15


def test_a_wrong_password_stops_the_loop_at_once(app, flex_mgr, make_account, monkeypatch):
    from calllog.ops import api as ops_api
    from calllog.ops.engines import assign_engine, common
    _ticketing_ready(flex_mgr, make_account)
    calls = []
    monkeypatch.setattr(assign_engine, "run", lambda ctx: calls.append(1) or common.JobResult(
        ok=False, error="Login failed - still on the login page.", error_kind="auth"))
    monkeypatch.setattr(ops_api, "TICKETING_LOOP_UNIT_SECONDS", 0.05)
    flex_mgr.post("/api/ops/auto-assign/run-ticketing", {"team": "Flex Cube", "username": "mgr1", "password": "bad",
                                                          "repeat": True})
    done = _wait_for(lambda: (lambda d: d if not d["running"] else None)(flex_mgr.get("/api/ops/auto-assign/run-ticketing").get_json()))
    assert calls == [1] and "Login failed" in done["error"]


def test_shared_is_counts_only(make_account, hub):
    emp = make_account("1005", "user", "Eli Employee", email="eli@example.test", teams=["Flex Cube", "Card Management"])
    body = emp.get("/api/ops/shared").get_json()
    text = json.dumps(body, ensure_ascii=False)
    for secret in ("Subject", "Reset Flexcube", "Card limit", "Customer", "cust-", "example.test", "احمد", "منى",
                   "Replier", "Strange request", "100001", "Add user"):
        assert secret not in text, secret
    # an employee's e-mail numbers are those of their own teams, never the "Unclassified" bucket
    assert body["email"]["business_emails"] == 4 and body["email"]["replied"] == 2 and body["email"]["scope"] == "teams"
    assert body["email"]["teams"] == ["Flex Cube", "Card Management"]
    assert body["sla"]["breached"] == 1 and body["itsm"]["totals"]["Rollback Requests"] == 1
    assert body["collector"]["online"] is True and body["collector"]["last_update"] == SNAP_AT
    assert body["calls"]["total"] == 0 and "employees" in body["calls"]


def test_shared_with_nothing_collected_yet(emp):
    body = emp.get("/api/ops/shared").get_json()
    assert body["email"]["available"] is False and body["collector"]["online"] is False
    assert body["collector"]["last_update"] is None


# ---------------------------------------------------------------------------- health
def test_health_report(app, admin, hub, ops, tmp_path, clock):
    app.extensions["store"].audit({"ts": "2026-10-11 08:00:00", "actor": "x", "action": "login_failed", "details": {}})
    (ops.report_dir("email") / "a.xlsx").write_bytes(b"x" * 100)
    ops.add_run_request("sla", "admin")
    body = admin.get("/api/ops/health").get_json()
    assert body["data_folder"]["path"] == str(tmp_path / "data") and body["data_folder"]["writable"] is True
    assert body["data_folder"]["disk"]["total_bytes"] > 0 and body["data_folder"]["disk"]["free_bytes"] > 0
    assert body["ops_folder"]["files"] >= 5 and body["ops_folder"]["bytes"] > 100
    assert body["snapshots"]["sla"]["present"] and body["snapshots"]["sla"]["age_minutes"] == 0
    assert body["heartbeats"]["sla"]["online"] is True and body["heartbeats"]["sla"]["age_minutes"] == 2
    assert body["pending_requests"]["count"] == 1 and body["pending_requests"]["items"][0]["job"] == "sla"
    assert body["users"]["by_role"]["admin"] == {"active": 1, "disabled": 0} and body["users"]["total"] == 1
    assert body["audit_today"]["failed_sign_ins"] == 1 and body["audit_today"]["entries"] >= 1
    server = body["server"]
    assert server["python"] and server["version"] and server["started_at"] and server["uptime_seconds"] >= 0
    assert body["status"] in ("ok", "warn") and isinstance(body["problems"], list)
    assert not list((tmp_path / "data").glob(".health-probe-*"))              # the probe file is removed again


def test_health_counts_users_by_role_and_flags_problems(admin, make_account, hub, ops, clock):
    make_account("1001", "user")
    make_account("mona", "manager")
    admin.put("/api/admin/users/1001", {"active": False})
    clock.value = clock.value.replace(hour=11, minute=30)
    body = admin.get("/api/ops/health").get_json()
    assert body["users"]["by_role"]["user"] == {"active": 0, "disabled": 1}
    assert body["users"]["by_role"]["manager"]["active"] == 1 and body["users"]["total"] == 3
    assert body["status"] == "error"
    assert any("offline" in p["text"] for p in body["problems"])


# ---------------------------------------------------------------------------- users and roles
def test_create_manager_with_profile_fields(app, admin, clock):
    r = admin.post("/api/admin/users", {"user_id": "mona", "name": "Mona", "role": "manager", "department": "IT",
                                        "email": " Mona@Example.TEST ", "team": "Flex Cube",
                                        "security_name": "  منى   علي ", "password": "TempPass123"})
    assert r.status_code == 201, r.get_json()
    user = r.get_json()["user"]
    assert user["role"] == "manager" and user["email"] == "mona@example.test" and user["team"] == "Flex Cube"
    assert user["security_name"] == "منى علي"
    listing = admin.get("/api/admin/users").get_json()["users"]
    assert [u["user_id"] for u in listing] == ["admin", "mona"]
    created = next(e for e in app.extensions["store"].read_audit("2026-10-11", "2026-10-11") if e["action"] == "user_created")
    assert created["details"]["role"] == "manager" and created["details"]["email"] == "mona@example.test"


def test_profile_field_validation(admin, make_account):
    def create(uid, **kw):
        return admin.post("/api/admin/users", {"user_id": uid, "name": "N", "password": "TempPass123", **kw})
    assert create("a1", email="bad-address").status_code == 400
    assert create("a2", email="x" * 130 + "@example.test").status_code == 400
    assert create("a3", team="No Such Team").status_code == 400
    assert create("a4", security_name="s" * 121).status_code == 400
    assert create("a5", role="superuser").status_code == 400
    assert create("a6", role="manager", team="__all__").status_code == 201
    assert create("a7", team="").status_code == 201
    assert create("a8", email="").status_code == 201                           # empty e-mails never collide
    assert create("a9", email="").status_code == 201
    assert create("b1", email="one@example.test").status_code == 201
    dup = create("b2", email="ONE@example.test")
    assert dup.status_code == 409 and "e-mail" in dup.get_json()["error"]
    assert admin.put("/api/admin/users/b1", {"email": "one@example.test", "name": "Renamed"}).status_code == 200
    assert create("b3", email="three@example.test").status_code == 201
    assert admin.put("/api/admin/users/b3", {"email": "one@example.test"}).status_code == 409
    assert admin.put("/api/admin/users/b3", {"email": ""}).status_code == 200
    assert admin.put("/api/admin/users/b3", {"email": "nope"}).status_code == 400


def test_updating_profile_and_role(app, admin, clock):
    admin.post("/api/admin/users", {"user_id": "u1", "name": "U One", "password": "TempPass123"})
    r = admin.put("/api/admin/users/u1", {"role": "manager", "team": "Card Management", "email": "u1@example.test",
                                          "security_name": "U One"})
    assert r.status_code == 200 and r.get_json()["user"]["role"] == "manager"
    assert admin.put("/api/admin/users/u1", {"role": "user"}).get_json()["user"]["team"] == "Card Management"
    entry = [e for e in app.extensions["store"].read_audit("2026-10-11", "2026-10-11") if e["action"] == "user_updated"][0]
    assert entry["details"]["role"] == {"from": "user", "to": "manager"}
    assert admin.put("/api/admin/users/admin", {"role": "manager"}).status_code == 409           # cannot demote yourself


def test_a_team_removed_from_settings_does_not_block_editing_the_account(admin, ops):
    admin.post("/api/admin/users", {"user_id": "u1", "name": "U", "password": "TempPass123", "team": "Third Party"})
    cfg = admin.get("/api/ops/settings").get_json()["settings"]
    cfg["email"]["teams"] = [t for t in cfg["email"]["teams"] if t != "Third Party"]
    assert admin.put("/api/ops/settings", cfg).status_code == 200
    assert admin.put("/api/admin/users/u1", {"name": "U Renamed", "team": "Third Party"}).status_code == 200
    assert admin.put("/api/admin/users/u1", {"team": "Open System"}).status_code == 200


def test_manager_cannot_manage_users_or_the_audit_log(mgr):
    for method, path in (("GET", "/api/admin/users"), ("POST", "/api/admin/users"), ("PUT", "/api/admin/users/mona"),
                         ("GET", "/api/admin/audit"), ("GET", "/api/admin/overview"), ("GET", "/api/admin/records"),
                         ("POST", "/api/admin/users/import"), ("POST", "/api/admin/branches/import")):
        assert mgr.api(method, path, json={}).status_code == 403, path
    assert mgr.get("/export/audit.xlsx").status_code == 403 and mgr.get("/export/records.csv").status_code == 403


def test_manager_cannot_log_calls_and_is_not_an_employee(admin, mgr, emp):
    r = mgr.post("/api/records", {"ip_phone": "75942", "reason": "solved"})
    assert r.status_code == 403 and "do not log calls" in r.get_json()["error"]
    assert emp.post("/api/records", {"ip_phone": "75942", "reason": "solved"}).status_code == 201
    assert mgr.get("/api/records").get_json()["records"] == []
    ov = admin.get("/api/admin/overview").get_json()
    assert [e["user_id"] for e in ov["employees"]] == ["1001"] and ov["totals"]["employees"] == 1
    assert "mona" not in json.dumps(ov["employees"])
    assert [r["user_id"] for r in admin.get("/api/admin/records").get_json()["items"]] == ["1001"]


def test_me_and_public_user_carry_the_profile(mgr, emp):
    user = mgr.get("/api/me").get_json()["user"]
    assert user["role"] == "manager" and user["team"] == "__all__" and user["email"] == "mona@example.test"
    me = emp.get("/api/me").get_json()["user"]
    assert me["security_name"] == "احمد سالم" and me["email"] == "alice@example.test" and me["team"] == ""


def test_import_users_with_profile_columns(admin):
    csv_text = ("user_id,name,department,role,email,team,security_name\n"
                "5001,Eve,IT,manager,eve@example.test,Flex Cube,Eve Sec\n"
                "5002,Fay,IT,,fay@example.test,,\n"
                "5003,Gus,IT,admin,gus@example.test,,\n"
                "5004,Hal,IT,user,eve@example.test,,\n"
                "5005,Ida,IT,user,not-an-email,,\n"
                "5006,Jon,IT,user,,No Such Team,\n")
    r = admin.post("/api/admin/users/import", data={"file": (io.BytesIO(csv_text.encode()), "users.csv")},
                   content_type="multipart/form-data")
    assert r.status_code == 200, r.get_json()
    body = r.get_json()
    assert body["created"] == 2 and {s["user_id"] for s in body["skipped"]} == {"5003", "5004", "5005", "5006"}
    users = {u["user_id"]: u for u in admin.get("/api/admin/users").get_json()["users"]}
    assert users["5001"]["role"] == "manager" and users["5001"]["team"] == "Flex Cube" and users["5001"]["security_name"] == "Eve Sec"
    assert users["5002"]["role"] == "user" and users["5002"]["email"] == "fay@example.test"
    plain = admin.post("/api/admin/users/import", data={"file": (io.BytesIO(b"user_id,name\n5100,Old Style\n"), "u.csv")},
                       content_type="multipart/form-data")
    assert plain.get_json()["created"] == 1


# ---------------------------------------------------------------------------- landing page per role
def test_index_redirects_by_role(app, admin, mgr, emp):
    assert admin.client.get("/").headers["Location"].endswith("/hub")
    assert mgr.client.get("/").headers["Location"].endswith("/hub")
    assert emp.client.get("/").headers["Location"].endswith("/my")
    assert app.test_client().get("/").headers["Location"].endswith("/login")


def test_hub_blueprints_are_registered(app):
    rules = {r.rule for r in app.url_map.iter_rules()}
    assert "/api/ops/shared" in rules and "/export/ops/<any(email,sla,itsm):module>.<any(xlsx,csv):fmt>" in rules
    assert app.extensions["ops"].root.name == "ops" and app.extensions["ops"].root.is_dir()


# ---------------------------------------------------------------------- a manager sees only their own team, everywhere
FLEX_SEC, CARD_SEC = "فريد فلكسوك", "كارلا كاردليك"
FLEX_ROLE, CARD_ROLE = "IMP UAT I-Flex SMS", "IMP Card Management"


@pytest.fixture
def two_teams(ops, make_account):
    """Data of TWO teams in every module (calls, SLA, ITSM, timeline, history, a report file) and an employee of each."""
    flex = make_account("f1", "user", "Farid Flex", email="farid@example.test", team="Flex Cube", security_name=FLEX_SEC)
    card = make_account("c1", "user", "Carla Card", email="carla@example.test", team="Card Management",
                        security_name=CARD_SEC)
    for _ in range(2):
        flex.post("/api/records", {"ip_phone": "70001", "reason": "solved", "note": "flex call"})
    for _ in range(3):
        card.post("/api/records", {"ip_phone": "70002", "reason": "solved", "note": "CARDLEAK call"})
    reqs = [sla_request("f-1", FLEX_SEC, "breached", 200), sla_request("f-2", FLEX_SEC),
            sla_request("c-1", CARD_SEC, "breached", 300, **{"النظام": "DWP", "اسم الموظف": "CARDLEAK one"}),
            sla_request("c-2", CARD_SEC, "reminder", 70, **{"النظام": "DWP"}),
            sla_request("c-3", CARD_SEC, **{"النظام": "DWP"}),
            sla_request("x-1", "(Unassigned)", **{"النظام": "T24"})]            # no rule gives this one a team
    put_snapshot(ops, "sla", sla_data(reqs))
    cat = lambda n, rows=(): {"count": n, "rows": list(rows), "warning": None, "note": None}            # noqa: E731
    roles = [{"role": FLEX_ROLE, "error": None, "categories": {"Implementation Requests": cat(1, [{"number": "900001", "summary": "flex CR"}]),
                                                                "Rollback Requests": cat(0)}},
             {"role": CARD_ROLE, "error": None, "categories": {"Implementation Requests": cat(2, [{"number": "900101", "summary": "CARDLEAK CR"}]),
                                                                "Rollback Requests": cat(1)}}]
    put_snapshot(ops, "itsm", {"categories": ["Implementation Requests", "Rollback Requests"], "roles": roles,
                               "totals": {"Implementation Requests": 3, "Rollback Requests": 1},
                               "changes": {"new": [{"role": CARD_ROLE, "category": "Rollback Requests", "number": "900102",
                                                    "summary": "CARDLEAK rollback"}], "closed": []}})
    ops.add_events([
        {"ts": "2026-10-11T08:40:00", "module": "sla", "level": "danger", "kind": "sla_breached", "text": "3 tickets passed the 2 h SLA",
         "count": 3, "items": [{"user": FLEX_SEC, "key": "f-1", "system": "Flexcube"}, {"user": CARD_SEC, "key": "c-1", "system": "DWP"}],
         "groups": [{"system": "Flexcube", "user": FLEX_SEC, "n": 1}, {"system": "DWP", "user": CARD_SEC, "n": 2}]},
        {"ts": "2026-10-11T08:45:00", "module": "itsm", "level": "warn", "kind": "itsm_new", "text": "1 new Rollback Requests",
         "count": 1, "items": [{"number": "900102", "summary": "CARDLEAK rollback", "role": CARD_ROLE}],
         "groups": [{"role": CARD_ROLE, "n": 1}]}])
    for module, point in (("sla", {"ts": "2026-10-11T08:30:00", "open": 6, "within": 3, "reminder": 1, "breached": 2,
                                   "groups": [{"system": "Flexcube", "user": FLEX_SEC, "open": 2, "within": 1, "reminder": 0, "breached": 1},
                                              {"system": "DWP", "user": CARD_SEC, "open": 3, "within": 1, "reminder": 1, "breached": 1},
                                              {"system": "T24", "user": "(Unassigned)", "open": 1, "within": 1, "reminder": 0, "breached": 0}]}),
                          ("itsm", {"ts": "2026-10-11T08:30:00", "totals": {"Implementation Requests": 3, "Rollback Requests": 1},
                                    "per_role": {FLEX_ROLE: {"Implementation Requests": 1, "Rollback Requests": 0},
                                                 CARD_ROLE: {"Implementation Requests": 2, "Rollback Requests": 1}}})):
        ops.add_history_point(module, "2026-10-11", point)
    (ops.report_dir("sla") / "SLA_latest.xlsx").write_bytes(b"CARDLEAK report of every team")
    beat(ops, "sla")
    beat(ops, "itsm")
    return flex, card


def test_a_team_manager_sees_only_their_team_in_every_module(app, admin, flex_mgr, two_teams, ops):
    def leaks(resp):
        if resp.mimetype.endswith("sheet"):                       # a workbook: every cell of every sheet
            book = openpyxl.load_workbook(io.BytesIO(resp.data))
            body = " ".join(str(v) for ws in book.worksheets for row in ws.iter_rows(values_only=True) for v in row)
        else:
            body = resp.get_data(as_text=True)
        return [m for m in ("CARDLEAK", CARD_SEC, CARD_ROLE, "T24") if m in body]

    sla = flex_mgr.get("/api/ops/sla").get_json()["data"]
    assert [r["key"] for r in sla["requests"]] == ["f-1", "f-2"] and sla["totals"]["open"] == 2
    assert [u["user"] for u in sla["by_user"]] == [FLEX_SEC]
    itsm = flex_mgr.get("/api/ops/itsm").get_json()["data"]
    assert [r["role"] for r in itsm["roles"]] == [FLEX_ROLE]
    assert itsm["totals"] == {"Implementation Requests": 1, "Rollback Requests": 0} and itsm["changes"]["new"] == []
    overview = flex_mgr.get("/api/ops/overview").get_json()
    assert overview["sla"]["open"] == 2 and [u["user"] for u in overview["sla"]["top_users"]] == [FLEX_SEC]
    assert overview["itsm"]["totals"]["Rollback Requests"] == 0
    assert overview["calls"]["today"]["total"] == 2                                  # Farid's calls, not Carla's
    shared = flex_mgr.get("/api/ops/shared").get_json()
    assert (shared["sla"]["open"], shared["calls"]["total"], shared["itsm"]["totals"]["Rollback Requests"]) == (2, 2, 0)
    events = flex_mgr.get("/api/ops/events?day=2026-10-11").get_json()["events"]
    assert [(e["text"], e.get("count")) for e in events] == [("1 ticket passed the 2 h SLA", 1)]
    sla_day = flex_mgr.get("/api/ops/sla/history?days=1").get_json()["days"][-1]["points"][-1]
    assert (sla_day["open"], sla_day["breached"]) == (2, 1)
    itsm_day = flex_mgr.get("/api/ops/itsm/history?days=1").get_json()["days"][-1]["points"][-1]
    assert itsm_day["totals"] == {"Implementation Requests": 1, "Rollback Requests": 0} and list(itsm_day["per_role"]) == [FLEX_ROLE]
    for rid, system, who in (("h-flex", "Flexcube", FLEX_SEC), ("h-card", "T24", CARD_SEC)):    # the status history
        rec = {"id": rid, "source": "history", "stage": "closed", "first_assigned_at": "2026-10-11T07:00:00",
               "assigned_at": "2026-10-11T07:00:00", "executed_at": "2026-10-11T07:30:00",
               "confirmed_at": "2026-10-11T08:00:00", "system": system, "executor": who, "holder": who,
               "final_outcome": "executed", "exec_outcome": "executed", "outcome": "completed"}
        ops.update_sla_closed_day("2026-10-11", lambda d, rec=rec: d.update({rec["id"]: rec}))
    ops.update_sla_history_state(lambda d: d.update(last_history="2026-10-11T08:55:00"))
    analysis = flex_mgr.get("/api/ops/sla-analysis?from=2026-10-11&to=2026-10-11").get_json()
    assert analysis["options"]["teams"] == ["Flex Cube"] and analysis["options"]["executors"] == [FLEX_SEC]
    assert analysis["meta"]["requests"] == 1
    assert flex_mgr.get("/api/ops/reports").get_json()["reports"] == []
    assert flex_mgr.get("/export/ops/report/sla/SLA_latest.xlsx").status_code == 404
    app.extensions["ops_ticketing_run"] = {"running": False, "done": True, "team": "Card Management",
                                           "log": ["CARDLEAK run"], "summary": "assigned 2 (CARDLEAK)"}
    for path in ("/api/ops/sla", "/api/ops/itsm", "/api/ops/overview", "/api/ops/shared", "/api/ops/events?day=2026-10-11",
                 "/api/ops/sla/history", "/api/ops/itsm/history", "/api/ops/sla-analysis", "/api/ops/itsm-analysis",
                 "/api/ops/auto-assign/run-ticketing", "/export/ops/sla.csv", "/export/ops/sla.xlsx", "/export/ops/itsm.csv",
                 "/export/ops/itsm.xlsx", "/hub"):
        resp = flex_mgr.get(path)
        assert resp.status_code == 200 and leaks(resp) == [], (path, leaks(resp))

    # the administrator still sees everything, including the request no rule gives a team
    everything = admin.get("/api/ops/sla").get_json()["data"]
    assert everything["totals"]["open"] == 6 and admin.get("/api/ops/overview").get_json()["calls"]["today"]["total"] == 5
    assert len(admin.get("/api/ops/reports").get_json()["reports"]) == 1


def test_a_managers_numbers_agree_on_every_dashboard(flex_mgr, two_teams):
    sla = flex_mgr.get("/api/ops/sla").get_json()["data"]["totals"]
    overview = flex_mgr.get("/api/ops/overview").get_json()
    shared = flex_mgr.get("/api/ops/shared").get_json()
    calls = flex_mgr.get("/api/ops/team-calls").get_json()["totals"]
    for key in ("open", "breached", "reminder", "within"):
        assert sla[key] == overview["sla"][key] == shared["sla"][key] == overview["shared"]["sla"][key], key
    assert calls["total"] == overview["calls"]["today"]["total"] == shared["calls"]["total"] == 2
    itsm = flex_mgr.get("/api/ops/itsm").get_json()["data"]["totals"]
    assert itsm == overview["itsm"]["totals"] == shared["itsm"]["totals"]


def test_a_managers_email_answer_counts_only_their_team(flex_mgr, admin, hub, ops):
    data = admin.get("/api/ops/email").get_json()["data"]
    data["stats"]["team_replies"] = 9                                    # what the collector counted in Outlook, every team
    data["stats"]["reply_breakdown"] = [{"name": "Replier bob", "team": "Flex Cube", "count": 2},
                                        {"name": "Replier carol", "team": "Card Management", "count": 7}]
    data["classification"] = {"counts": {"system": 5, "address": 0, "default_team": 0, "override": 0, "none": 1},
                              "unclassified": 1, "unclassified_team": ""}             # as a real scan writes it
    put_snapshot(ops, "email", data)
    mine = flex_mgr.get("/api/ops/email").get_json()["data"]
    assert mine["stats"]["team_replies"] == 2                            # Flex Cube's replies only
    assert mine["classification"]["unclassified"] == 0 and mine["classification"]["counts"]["none"] == 0
    assert admin.get("/api/ops/email").get_json()["data"]["classification"]["unclassified"] == 1
    assert flex_mgr.get("/api/ops/sla-config").get_json().get("recorded", 0) <= admin.get("/api/ops/sla-config").get_json().get("recorded", 0)


def test_best_performers_of_the_day_are_scoped_to_the_managers_team(admin, flex_mgr, ops, assignees):
    """GET /api/ops/performance: each employee's calls answered, e-mails replied and tickets handled, best first; the
    administrator sees every team, a Flex Cube manager only Flex Cube's people."""
    put_snapshot(ops, "email", email_data([replied_row("<p1@x>", "bob@example.test"), replied_row("<p2@x>", "bob@example.test"),
                                           replied_row("<p3@x>", "carol@example.test", team="Card Management",
                                                       all_teams=["Card Management"])]))
    every = admin.get("/api/ops/performance").get_json()
    assert every["day"] == "2026-10-11"
    assert [(r["name"], r["emails"], r["rank"]) for r in every["rows"]] == [("Bob B", 2, 1), ("Carol C", 1, 2)]
    mine = flex_mgr.get("/api/ops/performance?day=2026-10-11").get_json()
    assert [r["name"] for r in mine["rows"]] == ["Bob B"]
    assert admin.get("/api/ops/performance?day=2026-10-12").status_code == 400          # not come yet


def test_the_day_by_day_numbers_are_scoped_to_the_managers_team(admin, flex_mgr, ops):
    """GET /api/ops/sla-days: the kept days - every team's totals for the administrator, a manager's own team(s)."""
    ops.update_history("sla", "2026-10-10", lambda d: d.update(day_log={
        "at": "2026-10-10T17:00:00", "totals": {"arrived": 9, "executed": 7, "not_executed": 1, "rejected": 1, "closed": 8},
        "times": {"execution": {"n": 7, "average_minutes": 30.0, "within_pct": 70.0}},
        "by_team": [{"team": "Flex Cube", "arrived": 5, "executed": 4, "not_executed": 1, "rejected": 0, "closed": 5,
                     "exec_n": 4, "exec_average_minutes": 20.0, "exec_within": 3},
                    {"team": "Card Management", "arrived": 4, "executed": 3, "not_executed": 0, "rejected": 1,
                     "closed": 3, "exec_n": 3, "exec_average_minutes": 43.3, "exec_within": 1}],
        "by_person": [], "people": [{"name": "Carol C", "team": "Card Management", "total": 12},
                                    {"name": "Bob B", "team": "Flex Cube", "total": 10}]}))
    every = admin.get("/api/ops/sla-days?from=2026-10-01&to=2026-10-11").get_json()["days"]
    assert [(d["day"], d["arrived"], d["executed"], d["best"]["name"]) for d in every] == [("2026-10-10", 9, 7, "Carol C")]
    mine = flex_mgr.get("/api/ops/sla-days?from=2026-10-01&to=2026-10-11").get_json()["days"]
    assert [(d["arrived"], d["executed"], d["exec_average_minutes"], d["exec_within_pct"], d["best"]["name"])
            for d in mine] == [(5, 4, 20.0, 75.0, "Bob B")]
    assert flex_mgr.get("/api/ops/sla-days?team=Card Management").get_json()["days"][0]["arrived"] == 5   # not theirs
