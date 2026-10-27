"""SLA engine: parsing, analysis, alert decision, e-mail, Excel, job flow (fake scraper) and an end-to-end run through
real Chromium against tests/fake_portal.py.

Fake data only (example.test addresses, user 1001, password "secret-pw", made-up Arabic names), a frozen clock
(Monday 2026-10-12 10:00) and a throw-away OpsStore. The end-to-end tests skip themselves when Chromium cannot start.
"""
import json
import os
import subprocess
import sys
import tempfile
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
import pytest

from calllog.ops import schemas, settings
from calllog.ops.engines import cdp
from calllog.ops.engines import sla_engine as sla
from calllog.ops.engines.common import CredentialError, FakeMailer, JobContext, norm_arabic
from calllog.ops.store import OpsStore
from tests.fake_portal import HEADERS, sample_requests, table_cells

NOW = datetime(2026, 10, 12, 10, 0, 0)
PASSWORD = "secret-pw"
BOSS = "boss@example.test"
OWNER_A, OWNER_B, OWNER_C = "مستخدم اول", "مستخدم ثان", "مستخدم ثالث"
DETAIL = ["اسم الموظف", "رقم المستخدم", "الفرع / الادارة", "النظام", "نوع الطلب", "موظف الاسناد"]


# --------------------------------------------------------------------------------------------------- builders
def sla_settings(**over):
    raw = settings.deep_merge({"sla": {"username": "1001", "mail": {"to": [BOSS]}}}, {"sla": over})
    clean, errors = settings.validate(raw)
    assert errors == [], errors
    return clean["sla"]


def make_ctx(store, cfg=None, now=NOW, mailer=None, password=PASSWORD, **kw):
    logs = []
    ctx = JobContext(job="sla", settings=cfg or sla_settings(), general={}, store=store, now=lambda: now,
                     log=logs.append, mailer=mailer if mailer is not None else FakeMailer(),
                     data_dir=str(store.root.parent), credentials=lambda kind, user: password, **kw)
    ctx.logs = logs
    return ctx


def sample_table(now=NOW):
    return HEADERS, [table_cells(r) for r in sample_requests(now)]


def scraper_of(headers, rows):
    return lambda ctx: (headers, rows)


def row(name, minutes_ago, owner, now=NOW, user_no="1000"):
    sent = "" if minutes_ago is None else f"{now - timedelta(minutes=minutes_ago):%d/%m/%Y %H:%M}"
    return ["", name, user_no, "فرع", "نظام", "نوع", "مسند", sent, owner]


@pytest.fixture
def store(tmp_path):
    return OpsStore(tmp_path / "data").init()


def run(store, table=None, cfg=None, now=NOW, **kw):
    ctx = make_ctx(store, cfg, now, **kw)
    headers, rows = table or sample_table()
    return sla.run(ctx, scraper=scraper_of(headers, rows)), ctx


def steady():
    """A queue that stays in the same state for hours: two requests that are already breached."""
    return HEADERS, [row("a", 130, OWNER_A), row("b", 200, OWNER_B)]


def by_name(result):
    return {r["fields"]["اسم الموظف"]: r for r in result.snapshot_data["requests"]}


# --------------------------------------------------------------------------------------------------- text and dates
class TestParsing:
    def test_normalisation_matches_the_shared_helper(self):
        assert sla.norm("مُحَمَّد  إأآ ة ى ـ") == norm_arabic("مُحَمَّد  إأآ ة ى ـ") == "محمد ااا ه ي"

    @pytest.mark.parametrize("text, expected", [
        ("11/10/2026 08:12", datetime(2026, 10, 11, 8, 12)),
        ("11/10/2026", datetime(2026, 10, 11)),
        ("1-2-2026 9:05", datetime(2026, 2, 1, 9, 5)),
        ("2026-10-11 08:12:30", datetime(2026, 10, 11, 8, 12, 30)),
        ("١١/١٠/٢٠٢٦ ٠٨:١٢", datetime(2026, 10, 11, 8, 12)),
        ("11/10/2026 01:30 PM", datetime(2026, 10, 11, 13, 30)),
        ("11/10/2026 12:15 AM", datetime(2026, 10, 11, 0, 15)),
        ("11/10/2026 12:15 PM", datetime(2026, 10, 11, 12, 15)),
        ("11/10/2026 08:12 ص", datetime(2026, 10, 11, 8, 12)),
        ("11/10/2026 08:12 م", datetime(2026, 10, 11, 20, 12)),
        ("١١/١٠/٢٠٢٦ ١٢:٠٥ ص", datetime(2026, 10, 11, 0, 5)),
        ("08:12 م 11/10/2026", datetime(2026, 10, 11, 20, 12)),
        ("11/10/2026 08:12:09 pm", datetime(2026, 10, 11, 20, 12, 9)),
    ])
    def test_parse_datetime(self, text, expected):
        assert sla.parse_datetime(text) == expected

    def test_month_first_order_is_configurable(self):
        assert sla.parse_datetime("03/04/2026 10:00", "dmy") == datetime(2026, 4, 3, 10)
        assert sla.parse_datetime("03/04/2026 10:00", "mdy") == datetime(2026, 3, 4, 10)

    @pytest.mark.parametrize("text", ["", None, "soon", "31/02/2026 10:00", "11/13/2026", "--"])
    def test_unparsable_dates_are_none(self, text):
        assert sla.parse_datetime(text) is None

    def test_find_col_prefers_exact_then_contains(self):
        cols = ["تاريخ الارسال", "تاريخ الارسال للسرية", "مستخدم السرية (الحالي)"]
        assert sla.find_col(cols, "تاريخ الإرسال للسرية") == "تاريخ الارسال للسرية"     # hamza ignored
        assert sla.find_col(cols, "مستخدم السرية") == "مستخدم السرية (الحالي)"           # contains
        assert sla.find_col(cols, "غير موجود") is None

    def test_formatting(self):
        assert sla.fmt_age(65) == "1h 05m" and sla.fmt_age(5) == "0h 05m" and sla.fmt_age(None) == "-"
        assert sla.fmt_hours(60) == "1h" and sla.fmt_hours(120) == "2h" and sla.fmt_hours(90) == "1.5h"

    def test_scrub_hides_a_password_in_every_form(self):
        text = "fill('p@ss w0rd') failed at /login?pw=p%40ss%20w0rd and again p@ss w0rd"
        out = sla.scrub(text, ["p@ss w0rd"])
        assert "p@ss" not in out and "p%40ss" not in out and out.count("***") == 3
        assert sla.scrub("nothing here", [None, ""]) == "nothing here"


# --------------------------------------------------------------------------------------------------- analysis
class TestAnalysis:
    def test_build_table_names_pads_and_drops_helper_columns(self):
        cols, rows = sla.build_table(["", "A", "A", " ", "B"], [["x", "1", "2", "", "3"], ["", "4"]])
        assert cols == ["col_1", "A", "A_", "B"]          # col_4 (nameless, empty everywhere) is dropped, col_1 has data
        assert rows == [["x", "1", "2", "3"], ["", "4", "", ""]]
        cols, rows = sla.build_table(["", "A"], [["", "1"]])
        assert cols == ["A"] and rows == [["1"]]

    def test_age_status_boundaries_and_order(self):
        cfg = sla_settings()
        table = [row("n59", 59, OWNER_A), row("n60", 60, OWNER_A), row("n119", 119, OWNER_A), row("n120", 120, OWNER_A),
                 row("none", None, OWNER_A), row("future", -5, OWNER_A), row("n300", 300, OWNER_A)]
        cols, rows = sla.build_table(HEADERS, table)
        reqs = sla.analyze(cols, rows, cfg, NOW)
        got = {r["fields"]["اسم الموظف"]: (r["age_minutes"], r["status"]) for r in reqs}
        assert got == {"n59": (59, "within"), "n60": (60, "reminder"), "n119": (119, "reminder"),
                       "n120": (120, "breached"), "none": (None, "no_date"), "future": (0, "within"),
                       "n300": (300, "breached")}
        assert [r["fields"]["اسم الموظف"] for r in reqs] == ["n300", "n120", "n119", "n60", "n59", "future", "none"]
        assert next(r for r in reqs if r["status"] == "no_date")["sent_at"] is None
        assert reqs[1]["sent_at"] == "2026-10-12T08:00:00"

    def test_thresholds_come_from_the_settings(self):
        cols, rows = sla.build_table(HEADERS, [row("a", 45, OWNER_A), row("b", 90, OWNER_A)])
        reqs = sla.analyze(cols, rows, sla_settings(warning_minutes=30, breach_minutes=60), NOW)
        assert [r["status"] for r in reqs] == ["breached", "reminder"]

    def test_equal_ages_keep_the_page_order(self):
        cols, rows = sla.build_table(HEADERS, [row("first", 90, OWNER_A), row("second", 90, OWNER_B), row("third", 90, OWNER_C)])
        names = [r["fields"]["اسم الموظف"] for r in sla.analyze(cols, rows, sla_settings(), NOW)]
        assert names == ["first", "second", "third"]

    def test_unassigned_and_fields(self):
        cols, rows = sla.build_table(HEADERS, [row("a", 10, "   ")])
        req = sla.analyze(cols, rows, sla_settings(), NOW)[0]
        assert req["security_user"] == "(Unassigned)"
        assert list(req["fields"]) == DETAIL and req["fields"]["رقم المستخدم"] == "1000"

    def test_missing_detail_headers_are_left_out(self):
        headers = ["اسم الموظف", "تاريخ الارسال للسرية", "مستخدم السرية"]
        cols, rows = sla.build_table(headers, [["x", "12/10/2026 09:00", OWNER_A]])
        req = sla.analyze(cols, rows, sla_settings(), NOW)[0]
        assert req["fields"] == {"اسم الموظف": "x"}

    def test_two_detail_names_resolving_to_one_column_appear_once(self):
        headers = ["اسم الموظف", "تاريخ الارسال للسرية", "مستخدم السرية"]
        cols, rows = sla.build_table(headers, [["x", "12/10/2026 09:00", OWNER_A]])
        req = sla.analyze(cols, rows, sla_settings(detail_headers=["اسم الموظف", "الموظف"]), NOW)[0]
        assert list(req["fields"]) == ["اسم الموظف"]

    def test_missing_columns_raise_with_the_headers_seen(self):
        with pytest.raises(sla.PortalError, match="Headers seen on the page: رقم, الحالة"):
            sla.analyze(["رقم", "الحالة"], [["1", "x"]], sla_settings(), NOW)

    def test_key_is_stable_and_ignores_owner_assignee_and_notes(self):
        headers = HEADERS + ["ملاحظات"]
        base = row("a", 70, OWNER_A) + ["note one"]
        moved = row("a", 70, OWNER_B) + ["note two"]
        moved[6] = "مسند اخر"                                    # موظف الاسناد changed too
        other = row("b", 70, OWNER_A) + ["note one"]
        cols, rows = sla.build_table(headers, [base, moved, other])
        keys = [r["key"] for r in sla.analyze(cols, rows, sla_settings(), NOW)]
        assert keys[0] == keys[1] != keys[2]
        again = [r["key"] for r in sla.analyze(cols, rows, sla_settings(), NOW + timedelta(hours=3))]
        assert again == keys                                     # time does not change the key
        assert len(keys[0]) == 16 and int(keys[0], 16) >= 0

    def test_key_changes_when_a_request_is_sent_again(self):
        cols, rows = sla.build_table(HEADERS, [row("a", 70, OWNER_A), row("a", 10, OWNER_A)])
        k1, k2 = [r["key"] for r in sla.analyze(cols, rows, sla_settings(), NOW)]
        assert k1 != k2

    def test_summary_sorting_and_counts(self):
        table = [row("a1", 10, "A"), row("a2", 200, "A"),                    # A: 1 within, 1 breached
                 row("b1", 70, "B"), row("b2", 80, "B"), row("b3", 75, "B"),  # B: 3 reminders
                 row("c1", 130, "C"), row("c2", 20, "C"), row("c3", 30, "C"),  # C: 1 breached, 2 within (total 3)
                 row("d1", None, "D")]                                        # D: no date only
        cols, rows = sla.build_table(HEADERS, table)
        summary = sla.summarize(sla.analyze(cols, rows, sla_settings(), NOW))
        assert [s["user"] for s in summary] == ["C", "A", "B", "D"]       # breached desc, reminder desc, total desc
        c = summary[0]
        assert c == {"user": "C", "total": 3, "within": 2, "reminder": 0, "breached": 1, "no_date": 0, "oldest_minutes": 130}
        assert summary[3] == {"user": "D", "total": 1, "within": 0, "reminder": 0, "breached": 0, "no_date": 1,
                              "oldest_minutes": 0}

    def test_sample_portal_data(self):
        headers, rows = sample_table()
        cols, table = sla.build_table(headers, rows)
        assert cols[0] == "اسم الموظف"                                   # the checkbox column was dropped
        reqs = sla.analyze(cols, table, sla_settings(), NOW)
        assert sla.totals_of(reqs) == {"open": 12, "within": 3, "reminder": 4, "breached": 4, "no_date": 1,
                                       "oldest_minutes": 400}

    def test_snapshot_data_matches_the_contract(self):
        headers, rows = sample_table()
        cols, table = sla.build_table(headers, rows)
        cfg = sla_settings()
        reqs = sla.analyze(cols, table, cfg, NOW)
        data = sla.snapshot_data(reqs, sla.summarize(reqs), cfg)
        env = schemas.make_envelope("sla", data, "2026-10-12T10:00:00")
        assert schemas.validate_snapshot(env) == []
        assert data["thresholds"] == {"warning_minutes": 60, "breach_minutes": 120}
        assert set(data["requests"][0]) == {"key", "security_user", "sent_at", "age_minutes", "status", "fields"}
        assert set(data["by_user"][0]) == {"user", "total", "within", "reminder", "breached", "no_date", "oldest_minutes"}
        json.dumps(data)                                                  # JSON-safe: no datetime objects leaked


# --------------------------------------------------------------------------------------------------- alert decision
class TestDecision:
    STATE = {"alerted": {}, "last_sent": None}

    def test_first_overdue_request_escalates(self):
        d = sla.decide_alert(self.STATE, {"k1": 1}, NOW, 60)
        assert d == {"send": True, "escalated": ["k1"], "repeat": True}

    def test_nothing_overdue_nothing_sent(self):
        assert sla.decide_alert(self.STATE, {}, NOW, 60)["send"] is False

    def test_already_alerted_waits_for_the_repeat_interval(self):
        state = {"alerted": {"k": 1}, "last_sent": NOW - timedelta(minutes=30)}
        assert sla.decide_alert(state, {"k": 1}, NOW, 60)["send"] is False
        state["last_sent"] = NOW - timedelta(minutes=60)
        assert sla.decide_alert(state, {"k": 1}, NOW, 60)["send"] is True

    def test_breach_after_reminder_is_an_escalation_even_inside_the_interval(self):
        state = {"alerted": {"k": 1}, "last_sent": NOW - timedelta(minutes=5)}
        d = sla.decide_alert(state, {"k": 2}, NOW, 60)
        assert d["send"] is True and d["escalated"] == ["k"] and d["repeat"] is False

    def test_repeat_zero_means_only_on_escalation(self):
        state = {"alerted": {"k": 2}, "last_sent": NOW - timedelta(days=3)}
        assert sla.decide_alert(state, {"k": 2}, NOW, 0)["send"] is False
        assert sla.decide_alert(state, {"k": 2, "new": 1}, NOW, 0)["send"] is True

    def test_force_always_sends(self):
        assert sla.decide_alert(self.STATE, {}, NOW, 60, force=True)["send"] is True

    def test_clock_moved_back_does_not_block_the_repeat(self):
        state = {"alerted": {"k": 1}, "last_sent": NOW + timedelta(minutes=50)}
        assert sla.decide_alert(state, {"k": 1}, NOW, 60)["send"] is True

    def test_next_state_after_send_raises_levels_and_forgets_closed(self):
        state = {"alerted": {"gone": 2, "keep": 2, "up": 1}, "last_sent": None}
        out = sla.next_state(state, {"keep": 1, "up": 2, "new": 1}, True, NOW)
        assert out == {"alerted": {"keep": 2, "up": 2, "new": 1}, "last_sent": "2026-10-12T10:00:00"}

    def test_next_state_without_send_only_forgets_closed(self):
        last = datetime(2026, 10, 12, 9, 0)
        state = {"alerted": {"gone": 2, "keep": 2}, "last_sent": last}
        out = sla.next_state(state, {"keep": 1, "new": 1}, False, NOW)
        assert out == {"alerted": {"keep": 2}, "last_sent": "2026-10-12T09:00:00"}


# --------------------------------------------------------------------------------------------------- e-mail html
def analysed(table=None, cfg=None, now=NOW):
    cfg = cfg or sla_settings()
    headers, rows = table or sample_table(now)
    cols, body = sla.build_table(headers, rows)
    reqs = sla.analyze(cols, body, cfg, now)
    return cols, reqs, sla.summarize(reqs), cfg


class TestHtml:
    def test_kpis_users_and_overdue_groups(self):
        _, reqs, summary, cfg = analysed()
        body = sla.html_body(reqs, summary, cfg, NOW, has_logo=True)
        assert 'src="cid:nbe_logo"' in body
        for text in ("Security Requests SLA Reminder", "Open Requests", "Within 1h", "Over 2h (Breached)",
                     "Requests per Security User", "Requests Over 1h (8)", OWNER_A, "(Unassigned)", "SLA BREACHED", "REMINDER",
                     "12 Oct 2026, 10:00"):
            assert text in body, text
        assert ">12<" in body and ">3<" in body and ">4<" in body                # KPI numbers: 12 open, 3 within, 4 / 4
        assert body.count("</tr>") > 12
        assert "#00573F" in body and "#C8A45D" in body                             # NBE colours
        for name in DETAIL:
            assert f">{name}</th>" in body                                         # detail headers are columns
        assert "موظف اختبار 11" in body and "موظف اختبار 12" not in body            # no-date request is not "overdue"

    def test_without_logo_it_falls_back_to_text(self):
        _, reqs, summary, cfg = analysed()
        body = sla.html_body(reqs, summary, cfg, NOW, has_logo=False)
        assert "cid:" not in body and ">NBE</div>" in body

    def test_hostile_text_is_escaped_everywhere(self):
        hostile = '<script>alert(1)</script>'
        img = '"><img src=x onerror=alert(2)>'
        table = [row(hostile, 150, img), row("B&B <b>x</b>", 80, img)]
        table[0][3] = "<iframe src=//evil>"
        _, reqs, summary, cfg = analysed((HEADERS, table))
        body = sla.html_body(reqs, summary, cfg, NOW)
        for raw in (hostile, "<img src=x", "<iframe", "<b>x</b>", "onerror=alert(2)>"):
            assert raw not in body, raw
        for safe in ("&lt;script&gt;alert(1)&lt;/script&gt;", "&quot;&gt;&lt;img src=x onerror=alert(2)&gt;",
                     "B&amp;B &lt;b&gt;x&lt;/b&gt;"):
            assert safe in body, safe

    def test_all_within_says_so(self):
        _, reqs, summary, cfg = analysed((HEADERS, [row("a", 5, OWNER_A)]))
        body = sla.html_body(reqs, summary, cfg, NOW)
        assert "No request is older than 1h" in body and "Requests Over 1h (0)" in body

    def test_long_lists_are_cut_and_say_so(self):
        table = [row(f"n{i}", 130 + i, OWNER_A) for i in range(10)]
        _, reqs, summary, cfg = analysed((HEADERS, table))
        body = sla.html_body(reqs, summary, cfg, NOW, max_rows=4)
        assert "Showing 4 of 10 requests" in body and ">n9<" in body and ">n6<" in body and ">n5<" not in body    # oldest first

    def test_subject_has_no_line_breaks(self):
        counts = {"breached": 4, "reminder": 4}
        subject = sla.mail_subject({"subject": "Hello\r\nBcc: x@example.test"}, sla_settings(), counts, NOW)
        assert "\n" not in subject and "\r" not in subject
        assert subject.endswith("- 4 breached, 4 over 1h - 12-Oct 10:00")
        assert sla.mail_subject({}, sla_settings(), counts, NOW).startswith("Security Requests SLA - 4 breached")

    def test_cc_security_users_matches_names_loosely_and_only_overdue(self):
        table = [row("a", 130, "أحمد سمير"), row("b", 10, OWNER_B), row("c", 90, OWNER_C), row("d", 95, OWNER_C)]
        _, reqs, _, _ = analysed((HEADERS, table))
        owners = {"احمد سمير": "ahmed@example.test", OWNER_B: "bee@example.test", OWNER_C: "cee@example.test"}
        mail = {"to": [BOSS], "cc": ["x@example.test", "cee@example.test"], "cc_security_users": True}
        assert sla.mail_cc(mail, owners, reqs) == ["x@example.test", "cee@example.test", "ahmed@example.test"]
        assert sla.mail_cc({**mail, "cc_security_users": False}, owners, reqs) == ["x@example.test", "cee@example.test"]
        assert sla.mail_cc({"to": ["ahmed@example.test"], "cc_security_users": True}, owners, reqs) == ["cee@example.test"]


# --------------------------------------------------------------------------------------------------- Excel
class TestExcel:
    def build(self, table=None):
        cols, reqs, summary, _ = analysed(table)
        return openpyxl.load_workbook(_save(sla.build_workbook(cols, reqs, summary))), reqs

    def test_three_sheets_with_legacy_styling(self):
        wb, reqs = self.build()
        assert wb.sheetnames == ["By Security User", "Over 1 Hour", "All Requests"]
        for ws in wb.worksheets:
            assert ws.freeze_panes == "A2" and ws.auto_filter.ref == ws.dimensions
            head = ws["A1"]
            assert head.fill.fgColor.rgb.endswith("00573F") and head.font.bold and head.font.color.rgb.endswith("FFFFFF")
            assert ws.sheet_view.rightToLeft is False
            assert ws["A2"].border.left.style == "thin"
        users, over, everything = wb.worksheets
        assert [c.value for c in users[1]] == ["SECURITY_USER", "TOTAL", "WITHIN_SLA", "REMINDER", "BREACHED", "NO_DATE", "OLDEST"]
        assert users.max_row == 5 and everything.max_row == 13 and over.max_row == 9
        heads = [c.value for c in everything[1]]
        assert heads[:6] == ["SECURITY_USER", "SENT_AT", "ASSIGNED_AT", "AGE_MIN", "AGE", "SLA_STATUS"]
        assert "تاريخ الارسال للسرية" in heads

    def test_values_and_status_colours(self):
        wb, _ = self.build()
        _, over, everything = wb.worksheets
        first = [c.value for c in everything[2]]
        assert first[3:6] == [400, "6h 40m", "SLA BREACHED"] and first[1] == "12/10/2026 03:20"
        assert first[2] in (None, "")                        # no SLA history: no assignment to note
        status = everything["F2"]
        assert status.fill.fgColor.rgb.endswith("FDECEA") and status.font.color.rgb.endswith("B71C1C") and status.font.bold
        assert {over.cell(r, 6).value for r in range(2, over.max_row + 1)} == {"SLA BREACHED", "REMINDER"}
        last = [c.value for c in everything[everything.max_row]]
        assert last[1] in (None, "") and last[3] is None and last[4] == "-" and last[5] == "NO SENT DATE"

    def test_count_cells_are_coloured_when_positive(self):
        wb, _ = self.build()
        users = wb["By Security User"]
        names = [c.value for c in users[1]]
        breached = users.cell(2, names.index("BREACHED") + 1)
        assert breached.value == 1 and breached.fill.fgColor.rgb.endswith("FDECEA")
        no_date = users.cell(2, names.index("NO_DATE") + 1)
        assert no_date.fill.fill_type is None

    def test_widths_are_bounded(self):
        wb, _ = self.build()
        ws = wb["All Requests"]
        widths = [d.width for d in ws.column_dimensions.values()]
        assert min(widths) >= 12 and max(widths) <= 45

    def test_scraped_text_never_becomes_a_formula(self, tmp_path):
        table = [row('=HYPERLINK("http://evil.example","x")', 150, "+cmd|' /C calc'!A0"), row("@SUM(1)", 90, "-1")]
        cols, reqs, summary, _ = analysed((HEADERS, table))
        path = _save(sla.build_workbook(cols, reqs, summary))
        wb = openpyxl.load_workbook(path)                         # formulas would load with data_type "f"
        seen = []
        for ws in wb.worksheets:
            for r in ws.iter_rows():
                for c in r:
                    assert c.data_type != "f", (ws.title, c.coordinate, c.value)
                    seen.append(c.value)
        assert '=HYPERLINK("http://evil.example","x")' in seen and "+cmd|' /C calc'!A0" in seen

    def test_control_characters_do_not_break_the_file(self):
        table = [row("bad\x00name\x07", 150, "user\x0b")]
        cols, reqs, summary, _ = analysed((HEADERS, table))
        wb = openpyxl.load_workbook(_save(sla.build_workbook(cols, reqs, summary)))
        assert wb["All Requests"].max_row == 2

    def test_empty_queue_still_gives_a_valid_workbook(self):
        cols, reqs, summary, _ = analysed((HEADERS, []))
        wb = openpyxl.load_workbook(_save(sla.build_workbook(cols, reqs, summary)))
        assert [ws.max_row for ws in wb.worksheets] == [1, 1, 1]


def _save(wb):
    path = Path(tempfile.mkdtemp()) / "x.xlsx"
    sla.save_workbook(wb, path)
    return path


# --------------------------------------------------------------------------------------------------- the job (fake scraper)
class TestRun:
    def test_first_scan_returns_snapshot_and_sends_one_mail(self, store):
        result, ctx = run(store)
        assert result.ok and result.error is None and result.emails_sent == 1
        assert result.summary == "12 open, 4 reminder, 4 breached"
        assert schemas.validate_snapshot(schemas.make_envelope("sla", result.snapshot_data, "2026-10-12T10:00:00")) == []
        assert result.snapshot_data["totals"]["open"] == 12
        point = dict(result.history_point)
        groups = point.pop("groups")
        assert point == {"ts": "2026-10-12T10:00:00", "open": 12, "within": 3, "reminder": 4, "breached": 4}
        # per application + security user, adding up to the totals: what a manager's own history is rebuilt from
        assert {k: sum(g[k] for g in groups) for k in ("open", "within", "reminder", "breached")} == {
            "open": 12, "within": 3, "reminder": 4, "breached": 4}
        assert all(set(g) == {"system", "user", "open", "within", "reminder", "breached"} for g in groups)
        breached_keys = {r["key"] for r in result.snapshot_data["requests"] if r["status"] == "breached"}
        assert set(result.history_extra["first_breach"]) == breached_keys and len(breached_keys) == 4
        assert set(result.history_extra["first_breach"].values()) == {"2026-10-12T10:00:00"}
        assert store.read_snapshot("sla") is None                  # the runner writes snapshots, never the engine

        mail = ctx.mailer.sent[0]
        assert mail["to"] == [BOSS] and mail["cc"] == []
        assert mail["subject"] == "Security Requests SLA - 4 breached, 4 over 1h - 12-Oct 10:00"
        assert mail["inline_images"] == {"nbe_logo": str(sla.LOGO_FILE)} and sla.LOGO_FILE.is_file()
        assert 'cid:nbe_logo' in mail["html"]
        (attachment,) = mail["attachments"]
        assert Path(attachment).name == "SLA_Report_20261012_1000.xlsx" and Path(attachment).is_file()
        assert (store.report_dir("sla") / "SLA_latest.xlsx").is_file()

        state = store.read_state("sla")
        assert state["last_sent"] == "2026-10-12T10:00:00"
        assert sorted(state["alerted"].values()) == [1, 1, 1, 1, 2, 2, 2, 2]
        assert not any(PASSWORD in line for line in ctx.logs)

    def test_second_scan_inside_the_interval_is_quiet_and_keeps_the_state(self, store):
        run(store, steady())
        before = store.read_state("sla")
        result, ctx = run(store, steady(), now=NOW + timedelta(minutes=10))
        assert result.ok and result.emails_sent == 0 and ctx.mailer.sent == []
        assert "No new request crossed a threshold - no e-mail" in ctx.logs
        assert store.read_state("sla") == before
        assert sorted(p.name for p in store.report_dir("sla").iterdir()) == [          # no new copy without a mail
            "SLA_Report_20261012_1000.xlsx", "SLA_latest.xlsx"]

    def test_repeat_after_the_interval(self, store):
        run(store, steady())
        early, _ = run(store, steady(), now=NOW + timedelta(minutes=59))
        result, ctx = run(store, steady(), now=NOW + timedelta(minutes=60))
        assert early.emails_sent == 0 and result.emails_sent == 1
        assert store.read_state("sla")["last_sent"] == "2026-10-12T11:00:00"

    def test_repeat_zero_never_repeats(self, store):
        cfg = sla_settings(repeat_minutes=0)
        run(store, steady(), cfg=cfg)
        result, _ = run(store, steady(), cfg=cfg, now=NOW + timedelta(hours=5))
        assert result.emails_sent == 0

    def test_a_request_crossing_the_breach_time_is_mailed_at_once(self, store):
        reminder_only = [row("r", 100, OWNER_A)]
        run(store, (HEADERS, reminder_only))                                   # reminder at 10:00
        result, ctx = run(store, (HEADERS, reminder_only), now=NOW + timedelta(minutes=25))   # now 125 minutes old
        assert result.emails_sent == 1 and result.snapshot_data["totals"]["breached"] == 1
        assert list(store.read_state("sla")["alerted"].values()) == [2]

    def test_closed_requests_are_forgotten(self, store):
        run(store)
        assert len(store.read_state("sla")["alerted"]) == 8
        run(store, (HEADERS, [row("only", 130, OWNER_A)]), now=NOW + timedelta(minutes=5))
        assert len(store.read_state("sla")["alerted"]) == 1

    def test_run_now_can_force_a_mail(self, store):
        run(store)
        result, ctx = run(store, now=NOW + timedelta(minutes=5), options={"send_mail": True})
        assert result.emails_sent == 1 and len(ctx.mailer.sent) == 1

    def test_cc_security_users_and_display_only(self, store):
        cfg = sla_settings(mail={"to": [BOSS], "cc": ["x@example.test"], "cc_security_users": True, "display_only": True},
                           owner_emails={OWNER_A: "a@example.test", OWNER_B: "b@example.test"})
        calls = []

        class Recorder(FakeMailer):
            def send(self, *args, **kw):
                calls.append(kw)
                super().send(*args, **kw)
        run(store, cfg=cfg, mailer=Recorder())
        assert calls[0]["cc"] == ["x@example.test", "b@example.test", "a@example.test"]      # oldest overdue request first
        assert calls[0]["display_only"] is True

    @pytest.mark.parametrize("cfg_over, ctx_kw, log_part", [
        ({"mail": {"to": []}}, {}, "no recipients"),
        ({"mail": {"to": [BOSS], "send": False}}, {}, "switched off"),
        ({}, {"no_email": True}, "does not send"),
    ])
    def test_no_mail_when_not_allowed_and_the_alerts_stay_pending(self, store, cfg_over, ctx_kw, log_part):
        result, ctx = run(store, cfg=sla_settings(**cfg_over), **ctx_kw)
        assert result.ok and result.emails_sent == 0 and ctx.mailer.sent == []
        assert any(log_part in line for line in ctx.logs)
        assert (store.report_dir("sla") / "SLA_latest.xlsx").is_file()                # the preview is still written
        assert store.read_state("sla")["alerted"] == {} and store.read_state("sla")["last_sent"] is None
        later, ctx2 = run(store, now=NOW + timedelta(minutes=1))                         # mail allowed again: all go out
        assert later.emails_sent == 1

    def test_dry_run_writes_and_sends_nothing(self, store):
        result, ctx = run(store, dry_run=True)
        assert result.ok and result.snapshot_data["totals"]["open"] == 12 and result.emails_sent == 0
        assert ctx.mailer.sent == []
        assert not (store.root / "state" / "sla.json").exists()
        assert list(store.report_dir("sla").iterdir()) == []

    def test_mail_failure_keeps_the_scan_and_retries_next_time(self, store):
        result, ctx = run(store, mailer=FakeMailer(fail=True))
        assert result.ok and result.emails_sent == 0 and result.snapshot_data["totals"]["open"] == 12
        assert "e-mail FAILED: simulated Outlook failure" in result.summary
        assert any("e-mail failed" in line for line in ctx.logs)
        assert store.read_state("sla")["last_sent"] is None and store.read_state("sla")["alerted"] == {}
        retry, ctx2 = run(store, now=NOW + timedelta(minutes=15))
        assert retry.emails_sent == 1 and len(ctx2.mailer.sent) == 1

    def test_a_live_ticket_is_aged_from_its_last_assignment_to_the_user_holding_it(self, store):
        """The owner's rule: the live SLA runs from the LAST «اسناد لمستخدم السرية» - the status history's assignment
        to the security user holding it now - reminder from warning_minutes, breached from breach_minutes; a ticket
        never re-assigned (no history yet) from «تاريخ الارسال للسرية» on the Re-Assign page."""
        assigned = (NOW - timedelta(minutes=30)).strftime("%Y-%m-%dT%H:%M:%S")
        store.update_sla_history_state(lambda d: d.update(open={"R1": {
            "id": "R1", "stage": "open", "user_no": "1000", "system": "نظام", "request_type": "نوع",
            "holder": OWNER_A, "assigned_at": assigned}}))
        result, ctx = run(store, (HEADERS, [row("a", 200, OWNER_A), row("b", 70, OWNER_B, user_no="1001"),
                                            row("c", 20, OWNER_B, user_no="1002")]))
        a, b, c = (by_name(result)[k] for k in "abc")
        assert (a["clock_from"], a["clock_at"], a["age_minutes"], a["status"]) == ("assignment", assigned, 30, schemas.SLA_WITHIN)
        assert a["sent_at"] == (NOW - timedelta(minutes=200)).isoformat(timespec="seconds")       # still shown
        assert (b["clock_from"], b["age_minutes"], b["status"]) == ("sent", 70, schemas.SLA_REMINDER)
        assert c["age_minutes"] == 20 and c["status"] == schemas.SLA_WITHIN
        assert [r["fields"]["اسم الموظف"] for r in result.snapshot_data["requests"]] == ["b", "a", "c"]   # oldest first
        totals = result.snapshot_data["totals"]
        assert (totals["within"], totals["reminder"], totals["breached"], totals["oldest_minutes"]) == (2, 1, 0, 70)
        assert result.history_extra["first_breach"] == {}
        assert any("3 request(s) aged from their last «اسناد لمستخدم السرية»: 1 from the status history, 0 re-assigned"
                   " since an earlier scan, 2 from «تاريخ الارسال للسرية»" in line for line in ctx.logs)

    def test_a_re_assignment_the_scans_see_restarts_the_clock_before_its_history_is_read(self, store):
        """No status history yet: when a later scan finds the request with another security user, its clock starts at
        that scan (it moved since the one before), and stays there at the next scans with the same user."""
        first, _ = run(store, (HEADERS, [row("x", 200, OWNER_A)]))
        x = by_name(first)["x"]
        assert (x["clock_from"], x["age_minutes"], x["status"]) == ("sent", 200, schemas.SLA_BREACHED)
        store.write_snapshot("sla", first.snapshot_data, generated_at=NOW.isoformat(timespec="seconds"))
        later = NOW + timedelta(minutes=5)
        moved, _ = run(store, (HEADERS, [row("x", 205, OWNER_B, now=later)]), now=later)
        x = by_name(moved)["x"]
        assert (x["security_user"], x["clock_from"], x["clock_at"], x["age_minutes"], x["status"]) == (
            OWNER_B, "seen", later.isoformat(timespec="seconds"), 0, schemas.SLA_WITHIN)
        store.write_snapshot("sla", moved.snapshot_data, generated_at=later.isoformat(timespec="seconds"))
        after = later + timedelta(minutes=70)
        still, _ = run(store, (HEADERS, [row("x", 275, OWNER_B, now=after)]), now=after)
        x = by_name(still)["x"]
        assert (x["clock_from"], x["age_minutes"], x["status"]) == ("seen", 70, schemas.SLA_REMINDER)

    def test_third_party_requests_are_left_out_of_the_live_sla(self, store):
        """As before: a request of a Third Party application is not in the live SLA - not in its list, totals, alerts
        or first breaches - however old it is."""
        store.update_auto_assign(lambda d: d.update(applications=[{"name": "TP Bank App", "team": "Third Party"}]))
        tp = row("tp", 400, OWNER_A, user_no="1009")
        tp[4] = "TP Bank App"
        result, ctx = run(store, (HEADERS, [tp, row("a", 130, OWNER_A), row("b", 20, OWNER_B, user_no="1001")]))
        assert set(by_name(result)) == {"a", "b"}
        totals = result.snapshot_data["totals"]
        assert (totals["open"], totals["breached"], totals["oldest_minutes"]) == (2, 1, 130)
        assert list(result.history_extra["first_breach"]) == [by_name(result)["a"]["key"]]
        assert ctx.mailer.sent and "TP Bank App" not in ctx.mailer.sent[0]["html"]
        assert any("1 request(s) of excluded applications left out of the SLA" in line for line in ctx.logs)

    def test_missing_mailer_falls_back_to_outlook_and_fails_cleanly_here(self, store):
        ctx = make_ctx(store)
        ctx.mailer = None
        result = sla.run(ctx, scraper=scraper_of(*sample_table()))
        assert result.ok and result.emails_sent == 0 and "e-mail FAILED" in result.summary        # no Outlook on Linux

    def test_timestamped_copies_are_limited_to_30(self, store):
        folder = store.report_dir("sla")
        for i in range(34):
            (folder / f"SLA_Report_202609{i % 28 + 1:02d}_{i:04d}.xlsx").write_bytes(b"old")
        (folder / "notes.xlsx").write_bytes(b"keep")
        (folder / "SLA_Report_final.xlsx").write_bytes(b"keep")
        run(store)
        copies = sorted(p.name for p in folder.iterdir() if sla.COPY_NAME.match(p.name))
        assert len(copies) == 30 and "SLA_Report_20261012_1000.xlsx" in copies
        assert (folder / "notes.xlsx").exists() and (folder / "SLA_Report_final.xlsx").exists()

    def test_an_excel_that_cannot_be_saved_does_not_stop_the_mail(self, store, monkeypatch):
        def broken(wb, path):
            raise PermissionError("file is open in Excel")
        monkeypatch.setattr(sla, "save_workbook", broken)
        result, ctx = run(store)
        assert result.ok and result.emails_sent == 1 and ctx.mailer.sent[0]["attachments"] == []
        assert any("could not be saved" in line for line in ctx.logs)

    def test_a_report_folder_that_cannot_be_used_does_not_stop_scan_or_mail(self, store, monkeypatch):
        def no_folder(module):
            raise OSError("share is down")
        monkeypatch.setattr(store, "report_dir", no_folder)
        result, ctx = run(store)
        assert result.ok and result.emails_sent == 1 and ctx.mailer.sent[0]["attachments"] == []
        assert sum("could not be saved" in line for line in ctx.logs) == 2

    def test_an_unexpected_error_after_the_scan_keeps_the_snapshot(self, store, monkeypatch):
        def broken(*args):
            raise ValueError("boom")
        monkeypatch.setattr(sla, "build_workbook", broken)
        result, ctx = run(store)
        assert result.ok and result.snapshot_data["totals"]["open"] == 12 and result.emails_sent == 0
        assert "report / e-mail step failed: boom" in result.summary

    def test_state_file_that_is_not_an_object_starts_fresh(self, store):
        store.write_state("sla", ["not", "an", "object"])
        result, _ = run(store)
        assert result.ok and result.emails_sent == 1

    def test_empty_queue_is_a_valid_scan(self, store):
        result, ctx = run(store, (HEADERS, []))
        assert result.ok and result.snapshot_data["totals"]["open"] == 0 and result.snapshot_data["requests"] == []
        assert result.emails_sent == 0 and result.summary == "0 open, 0 reminder, 0 breached"
        assert result.history_extra == {"first_breach": {}}

    def test_partial_settings_use_the_defaults(self, store):
        ctx = make_ctx(store, cfg={"username": "1001", "mail": {"to": [BOSS]}})
        result = sla.run(ctx, scraper=scraper_of(*sample_table()))
        assert result.ok and result.snapshot_data["thresholds"] == {"warning_minutes": 60, "breach_minutes": 120}

    def test_history_keeps_the_first_breach_time_of_the_day(self, store):
        first, _ = run(store)
        later, _ = run(store, now=NOW + timedelta(minutes=30))
        for res in (first, later):
            store.add_history_point("sla", "2026-10-12", res.history_point, res.history_extra)
        doc = store.read_history("sla", "2026-10-12")
        assert len(doc["points"]) == 2
        for key in first.history_extra["first_breach"]:                      # already breached at 10:00: time kept
            assert doc["first_breach"][key] == "2026-10-12T10:00:00"
        newly = set(later.history_extra["first_breach"]) - set(first.history_extra["first_breach"])
        assert newly and all(doc["first_breach"][k] == "2026-10-12T10:30:00" for k in newly)

    def test_fields_follow_the_configured_detail_headers(self, store):
        result, _ = run(store, cfg=sla_settings(detail_headers=["النظام", "اسم الموظف", "لا يوجد"]))
        assert list(result.snapshot_data["requests"][0]["fields"]) == ["النظام", "اسم الموظف"]


class TestRunFailures:
    def seeded(self, store):
        before = {"alerted": {"k": 2}, "last_sent": "2026-10-12T09:00:00"}
        store.write_state("sla", before)
        return before

    def test_scan_error_leaves_state_and_snapshot_alone(self, store):
        before = self.seeded(store)
        ctx = make_ctx(store)

        def broken(ctx):
            raise sla.PortalError("Package not found on the groups page: x")
        result = sla.run(ctx, scraper=broken)
        assert result.ok is False and result.snapshot_data is None and result.error_kind is None
        assert result.error == "Package not found on the groups page: x" and result.emails_sent == 0
        assert store.read_state("sla") == before and ctx.mailer.sent == []
        assert store.read_snapshot("sla") is None and list(store.report_dir("sla").iterdir()) == []

    def test_login_error_is_an_auth_error(self, store):
        ctx = make_ctx(store)

        def refused(ctx):
            raise sla.LoginError("Login failed - still on the login page")
        result = sla.run(ctx, scraper=refused)
        assert result.ok is False and result.error_kind == "auth"

    def test_missing_columns_fail_with_the_headers_seen(self, store):
        before = self.seeded(store)
        result, _ = run(store, (["رقم", "الحالة"], [["1", "x"]]))
        assert result.ok is False and "Headers seen on the page: رقم, الحالة" in result.error
        assert store.read_state("sla") == before

    def test_missing_password_tells_where_to_put_it(self, store):
        ctx = make_ctx(store)

        def no_password(kind, user):
            raise CredentialError(f"nothing for {user}")
        ctx.credentials = no_password
        called = []
        result = sla.run(ctx, scraper=lambda c: called.append(1))
        assert result.ok is False and result.error_kind == "auth" and called == []
        assert "collector_config.json" in result.error and "sla.password" in result.error and "1001" in result.error

    def test_empty_password_counts_as_missing(self, store):
        result, _ = run(store, password="")
        assert result.ok is False and result.error_kind == "auth" and "collector_config.json" in result.error

    def test_missing_username(self, store):
        result, ctx = run(store, cfg=sla_settings(username=""))
        assert result.ok is False and result.error_kind == "auth" and "collector_config.json" in result.error

    def test_the_password_is_asked_for_the_sla_account_only(self, store):
        asked = []
        ctx = make_ctx(store)
        ctx.credentials = lambda kind, user: asked.append((kind, user)) or PASSWORD
        sla.run(ctx, scraper=scraper_of(*sample_table()))
        assert asked == [("sla", "1001")]

    def test_the_password_never_reaches_errors_logs_or_files(self, store):
        ctx = make_ctx(store, debug=True)

        def leaky(ctx):
            raise RuntimeError(f"Locator.fill: Timeout 5000ms exceeded.\nCall log:\n  - fill(\"{PASSWORD}\")\n"
                               f"  - url=http://x/?p=secret%2Dpw")
        result = sla.run(ctx, scraper=leaky)
        assert result.ok is False
        assert PASSWORD not in result.error and "***" in result.error
        assert all(PASSWORD not in line for line in ctx.logs)
        ok_result, ok_ctx = run(store)                                      # and a good run leaves it out of every file
        for path in Path(store.root).rglob("*"):
            if path.is_file() and path.suffix != ".xlsx":
                assert PASSWORD.encode() not in path.read_bytes(), path
        assert all(PASSWORD not in line for line in ok_ctx.logs)

    def test_damaged_state_file_starts_fresh(self, store):
        (store.root / "state").mkdir(parents=True, exist_ok=True)
        (store.root / "state" / "sla.json").write_text("{not json", encoding="utf-8")
        result, ctx = run(store)
        assert result.ok and result.emails_sent == 1
        assert any("starting fresh" in line for line in ctx.logs)


# --------------------------------------------------------------------------------------------------- browser discovery (CDP)
# The SLA job drives the PC's own Edge / Chrome over the DevTools protocol (calllog.ops.engines.cdp) instead of
# Playwright: no download, no Node.js. These check how a browser is found - the configured path first, then the standard
# Windows install locations, then CHROME_PATH, then a Chromium that Playwright already put on the PC - without launching one.
@pytest.fixture
def no_windows_browsers(monkeypatch):
    # the hard-coded Windows install paths may really exist on the test PC; blank them so only the inputs below decide
    monkeypatch.setattr(cdp, "WINDOWS_BROWSERS", ())


def _touch(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("", encoding="utf-8")
    return str(path)


@pytest.mark.usefixtures("no_windows_browsers")
class TestFindBrowser:
    def _pw_build(self, root):
        leaf = root / "chromium-1200" / ("chrome-win" if os.name == "nt" else "chrome-linux")
        return _touch(leaf / ("chrome.exe" if os.name == "nt" else "chrome"))

    def test_configured_path_wins(self, tmp_path):
        want = _touch(tmp_path / "myedge.exe")
        other = _touch(tmp_path / "chrome.exe")
        assert cdp.find_browser(want, environ={"CHROME_PATH": other}) == want

    def test_chrome_path_env_is_used_when_nothing_is_configured(self, tmp_path):
        chrome = _touch(tmp_path / "chrome.exe")
        env = {"CHROME_PATH": chrome, "PLAYWRIGHT_BROWSERS_PATH": str(tmp_path / "none")}
        assert cdp.find_browser("", environ=env) == chrome

    def test_a_playwright_chromium_is_the_last_resort(self, tmp_path):
        exe = self._pw_build(tmp_path / "pw")
        assert cdp.find_browser("", environ={"PLAYWRIGHT_BROWSERS_PATH": str(tmp_path / "pw")}) == exe

    def test_configured_beats_chrome_path_beats_playwright(self, tmp_path):
        want = _touch(tmp_path / "cfg.exe")
        chrome = _touch(tmp_path / "chrome.exe")
        self._pw_build(tmp_path / "pw")
        env = {"CHROME_PATH": chrome, "PLAYWRIGHT_BROWSERS_PATH": str(tmp_path / "pw")}
        assert cdp.find_browser(want, environ=env) == want
        assert cdp.find_browser("", environ=env) == chrome

    def test_none_when_nothing_exists(self, tmp_path):
        missing = str(tmp_path / "nope.exe")
        assert cdp.find_browser(missing, environ={"PLAYWRIGHT_BROWSERS_PATH": str(tmp_path / "empty")}) is None

    def test_candidates_skip_empty_and_keep_the_order(self, tmp_path):
        chrome = _touch(tmp_path / "chrome.exe")
        cands = cdp.browser_candidates("cfg.exe", environ={"CHROME_PATH": chrome, "PLAYWRIGHT_BROWSERS_PATH": str(tmp_path / "none")})
        assert "" not in cands and cands[0] == "cfg.exe" and chrome in cands


class TestDebugDumps:
    class Browser:
        """The CDP surface sla.Debug.snap uses: one page, screenshot(path) and eval(expression)."""

        def __init__(self, html):
            self.html = html

        def screenshot(self, path):
            Path(path).write_bytes(b"\x89PNG\r\n\x1a\n")

        def eval(self, expression):
            return self.html

    def test_disabled_writes_nothing(self, tmp_path):
        sla.Debug(False, tmp_path / "d", [PASSWORD]).snap(self.Browser("x"), "login")
        assert not (tmp_path / "d").exists()

    def test_html_is_scrubbed_and_files_are_numbered(self, tmp_path):
        dbg = sla.Debug(True, tmp_path / "d", [PASSWORD])
        dbg.snap(self.Browser(f"<p>echo {PASSWORD}</p>"), "login")
        dbg.snap(self.Browser("the error page"), "error")
        names = sorted(p.name for p in (tmp_path / "d").iterdir())
        assert names == ["01_login_0.html", "01_login_0.png", "02_error_0.html", "02_error_0.png"]
        assert (tmp_path / "d" / "01_login_0.html").read_text(encoding="utf-8") == "<p>echo ***</p>"

    def test_folder_is_under_data_dir_ops_debug_sla(self, store):
        ctx = make_ctx(store)
        assert sla.debug_folder(ctx) == store.root / "debug" / "sla"
        ctx.data_dir = ""
        assert sla.debug_folder(ctx) == store.root / "debug" / "sla"


def test_module_imports_without_browser_or_windows_packages():
    code = ("import sys, calllog.ops.engines.sla_engine as e\n"
            "bad = [m for m in ('playwright', 'win32com', 'pythoncom', 'keyring', 'pandas', 'matplotlib') if m in sys.modules]\n"
            "assert not bad, bad")
    root = Path(__file__).resolve().parent.parent
    proc = subprocess.run([sys.executable, "-c", code], cwd=root, capture_output=True, text=True)
    assert proc.returncode == 0, proc.stderr


# ===================================================================================================
# End to end: a real Chromium driven over CDP (the engine's own CdpBrowser) against tests/fake_portal.py
# ===================================================================================================
from tests.fake_portal import GROUPS_PATH, HOME_PATH, LOGIN_PATH, REASSIGN_PATH, FakePortal  # noqa: E402


BROWSER = {"executable": ""}      # filled by chromium_ok: a Chromium the engine's own CdpBrowser can start here
PROFILES = {"root": None}


def _candidate_browsers():
    """CHROME_PATH, the Chromium builds already in PLAYWRIGHT_BROWSERS_PATH (newest first), then cdp.find_browser."""
    root = Path(os.environ.get("PLAYWRIGHT_BROWSERS_PATH") or "")
    found = [str(p) for p in sorted(root.glob("chromium-*/chrome-linux/chrome"), reverse=True)] if str(root) != "." else []
    out = [os.environ.get("CHROME_PATH") or ""] + found + [cdp.find_browser("") or ""]
    return [e for e in out if e]


@pytest.fixture(scope="module")
def chromium_ok(tmp_path_factory):
    errors = []
    PROFILES["root"] = tmp_path_factory.mktemp("sla_profiles")
    for executable in _candidate_browsers():
        try:
            cdp.CdpBrowser({"browser_executable": executable, "profile_dir": str(PROFILES["root"] / "probe")},
                           headless=True).close()
        except Exception as exc:                          # missing binary, no sandbox permission, ...
            errors.append(str(exc).strip().splitlines()[0] if str(exc).strip() else repr(exc))
            continue
        BROWSER["executable"] = executable
        return executable
    pytest.skip("Chromium cannot start here: " + (" | ".join(errors) or "none found"))


@pytest.fixture
def quick(monkeypatch):
    monkeypatch.setattr(sla, "SETTLE_MS", 300)
    monkeypatch.setattr(sla, "PAGE_MS", 250)


def e2e_settings(portal, timeout=8, **over):
    """Settings that point the engine at the fake portal, with the Chromium found by chromium_ok and a throw-away
    browser profile (never the shared temp one)."""
    cfg = sla_settings(base_url=portal.url, browser_channel="", browser_executable=BROWSER["executable"], **over)
    cfg["timeout_seconds"] = timeout
    cfg["profile_dir"] = tempfile.mkdtemp(dir=PROFILES["root"]) if PROFILES["root"] else ""
    return cfg


def portal_run(store, portal, password=PASSWORD, timeout=8, **over):
    ctx = make_ctx(store, e2e_settings(portal, timeout, **over), password=password)
    return sla.run(ctx), ctx


def spy(monkeypatch, name, record, owner=sla):
    """Wrap <owner>.<name>; ``record(args, result)`` gets every call (to prove WHICH code path the browser flow took)."""
    real = getattr(owner, name)

    def wrapper(*args, **kw):
        out = real(*args, **kw)
        record(args, out)
        return out
    monkeypatch.setattr(owner, name, wrapper)


def click_log(monkeypatch):
    """[(text, clicked)] for every click-by-text the engine tried, in order."""
    texts, out = {}, []
    real = sla.js_click_text

    def finder(text):
        expr = real(text)
        texts[expr] = text
        return expr
    monkeypatch.setattr(sla, "js_click_text", finder)
    spy(monkeypatch, "_click", lambda a, res: out.append((texts[a[1]], res)) if a[1] in texts else None)
    return out


def window_log(monkeypatch):
    """How many page targets (tabs / windows) the browser had when it attached to the package's app page."""
    out = []
    spy(monkeypatch, "attach", lambda a, res: out.append(len([t for t in a[0]._targets() if t.get("type") == "page"])),
        owner=cdp.CdpBrowser)
    return out


def expected_by_name(now=NOW):
    out = {}
    for r in sample_requests(now):
        age = r["age"]
        status = "no_date" if age is None else "breached" if age >= 120 else "reminder" if age >= 60 else "within"
        out[r["employee"]] = (age, status, r["owner"] or "(Unassigned)")
    return out


def assert_scanned_everything(result):
    got = {r["fields"]["اسم الموظف"]: (r["age_minutes"], r["status"], r["security_user"])
           for r in result.snapshot_data["requests"]}
    assert got == expected_by_name()
    assert result.snapshot_data["totals"] == {"open": 12, "within": 3, "reminder": 4, "breached": 4, "no_date": 1,
                                              "oldest_minutes": 400}


@pytest.mark.usefixtures("chromium_ok", "quick")
class TestEndToEnd:
    def test_new_window_plain_table(self, store, monkeypatch):
        sources = []
        windows, clicks = window_log(monkeypatch), click_log(monkeypatch)
        spy(monkeypatch, "pick_table", lambda a, out: sources.append(out["source"] if out else None))
        with FakePortal(sample_requests(NOW), open_mode="blank", variant="plain") as portal:
            result, ctx = portal_run(store, portal)
            assert windows == [2]                                  # the package opened in a second window
            assert clicks == [("دخول", True)] and set(sources) == {"dom"}       # login button found by its text
            assert result.ok, result.error
            assert_scanned_everything(result)
            assert result.emails_sent == 1 and ctx.mailer.sent[0]["to"] == [BOSS]
            assert portal.logins == [True]
            paths = [p for _, p in portal.hits]
            assert LOGIN_PATH in paths and GROUPS_PATH in paths and HOME_PATH in paths and REASSIGN_PATH in paths
        assert (store.report_dir("sla") / "SLA_latest.xlsx").is_file()
        assert not (store.root / "debug").exists()                       # no screenshots unless ctx.debug
        users = {u["user"]: u for u in result.snapshot_data["by_user"]}
        assert users["(Unassigned)"]["no_date"] == 1 and users[OWNER_A]["total"] == 3

    def test_the_package_opened_by_a_post_back_script_like_the_real_portal(self, store, monkeypatch):
        """The real «دخــــــول» link posts back and the RESPONSE opens the application with window.open - a pop-up
        blocker would stop it; the collector's browser must let it open (as the bank's proven bot's browser does)."""
        windows = window_log(monkeypatch)
        with FakePortal(sample_requests(NOW), open_mode="postback_script", variant="plain") as portal:
            result, _ = portal_run(store, portal)
            assert result.ok, result.error
            assert windows == [2]                                  # the application's own window
            assert_scanned_everything(result)

    def test_the_real_portals_login_ids_are_used_even_with_a_search_box_first(self, store):
        """The real login fields are #ContentPlaceHolder1_txtName / _txtPassword / _btnAdd; a guess ("the first text box")
        would type the user name into the site's search box. They work even when the saved settings name no selector."""
        with FakePortal(sample_requests(NOW), login_style="real") as portal:
            result, _ = portal_run(store, portal, selectors={})
            assert result.ok, result.error
            assert portal.logins == [True]

    def test_a_login_page_without_its_fields_is_not_reported_as_a_wrong_password(self, store):
        """A changed or broken login page must not pause the SLA job as 'wrong password'."""
        with FakePortal(sample_requests(NOW), login_style="no_fields") as portal:
            result, _ = portal_run(store, portal)
            assert not result.ok and "fields were not found" in result.error
            assert result.error_kind != "auth" and portal.logins == []

    def test_the_package_is_opened_by_its_btnsubmit_link_even_with_old_settings(self, store):
        """The real row's Enter link is a[id*=btnSubmit] «دخــــــول»; saved settings may still say 'تحويل', and the row
        has another link after it (the old 'last control' guess would click that one)."""
        with FakePortal(sample_requests(NOW), open_mode="postback_script") as portal:
            result, _ = portal_run(store, portal, package_link_text="تحويل")
            assert result.ok, result.error
            assert "/help" not in [p for _, p in portal.hits]

    def test_the_re_assign_page_is_opened_from_the_portals_menu(self, store):
        """The real menu link is relative to the package's folder; a guessed folder in the settings must not matter."""
        with FakePortal(sample_requests(NOW)) as portal:
            result, _ = portal_run(store, portal, reassign_path="/SecurityCodeRequest/Old_Folder/Request_ReAssign.aspx")
            assert result.ok, result.error
            assert_scanned_everything(result)

    def test_an_empty_queue_is_accepted_only_when_the_portal_says_no_data(self, store):
        with FakePortal(sample_requests(NOW), variant="empty") as portal:
            result, _ = portal_run(store, portal)
            assert result.ok, result.error
            assert result.snapshot_data["totals"]["open"] == 0

    def test_a_grid_that_has_not_loaded_is_not_taken_for_an_empty_queue(self, store):
        """Recording an empty scan closes every open request - so a grid with no rows AND no «no data» line fails."""
        with FakePortal(sample_requests(NOW), variant="not_loaded") as portal:
            result, _ = portal_run(store, portal, timeout=3)
            assert not result.ok and "not finish loading" in result.error and result.snapshot_data is None

    def test_a_grid_holding_fewer_rows_than_the_portal_counts_is_refused(self, store):
        with FakePortal(sample_requests(NOW), variant="short_info") as portal:
            result, _ = portal_run(store, portal)
            assert not result.ok and "of 13" in result.error and result.snapshot_data is None

    def test_same_tab_datatables_reads_every_page_through_the_api(self, store, monkeypatch):
        sources = []
        windows = window_log(monkeypatch)
        spy(monkeypatch, "pick_table", lambda a, out: sources.append(out["source"] if out else None))
        with FakePortal(sample_requests(NOW), open_mode="same", variant="datatables", page_size=5) as portal:
            result, _ = portal_run(store, portal)
            assert result.ok, result.error
            assert windows == [1] and sources == ["datatables"]    # same tab; read through the API, no paging clicks
            assert_scanned_everything(result)             # 12 rows although only 5 are in the DOM

    def test_paged_plain_table_is_followed_page_by_page_and_enter_logs_in(self, store, monkeypatch):
        reads, enters = [], []
        clicks = click_log(monkeypatch)
        spy(monkeypatch, "_press_enter", lambda a, out: enters.append(True))
        spy(monkeypatch, "pick_table", lambda a, out: reads.append(len(out["rows"]) if out else None))
        with FakePortal(sample_requests(NOW), variant="paged", page_size=5, login_style="english") as portal:
            result, _ = portal_run(store, portal)
            assert result.ok, result.error
            assert_scanned_everything(result)
            assert portal.logins == [True]
            assert clicks == [("دخول", False)] and enters == [True]   # no button text matched: Enter submitted the form
            assert reads == [5, 5, 2]                      # first page, then "next" twice (5 + 5 + 2 = 12 rows)

    def test_selectors_override_and_search_button(self, store, monkeypatch):
        clicks = click_log(monkeypatch)
        selectors = {"username": "#txtUserName", "password": "#txtPassword", "login_button": "#btnLogin"}
        with FakePortal(sample_requests(NOW), variant="search", login_style="english") as portal:
            result, _ = portal_run(store, portal, selectors=selectors)
            assert result.ok, result.error
            assert_scanned_everything(result)
            assert clicks == [("بحث", True)]               # the login used the selector; only the search was clicked by text

    def test_gridview_header_and_footer_rows_are_not_requests(self, store, monkeypatch):
        reads = []
        spy(monkeypatch, "pick_table", lambda a, out: reads.append(len(out["rows"]) if out else None))
        with FakePortal(sample_requests(NOW), variant="gridview") as portal:
            result, _ = portal_run(store, portal)
        assert result.ok, result.error
        assert reads[0] == 12                              # the <th> header and footer rows of the body are left out
        assert_scanned_everything(result)

    def test_wrong_password_is_a_clear_auth_error_and_debug_files_hide_it(self, store):
        before = {"alerted": {"k": 1}, "last_sent": "2026-10-12T09:00:00"}
        store.write_state("sla", before)
        with FakePortal(sample_requests(NOW), echo_password=True) as portal:
            ctx = make_ctx(store, e2e_settings(portal, timeout=3), password="wrong-pw", debug=True)
            result = sla.run(ctx)
            assert portal.logins == [False]
        assert result.ok is False and result.error_kind == "auth" and result.snapshot_data is None
        assert "Login failed" in result.error and "collector_config.json" in result.error
        assert "wrong-pw" not in result.error and all("wrong-pw" not in line for line in ctx.logs)
        assert store.read_state("sla") == before and ctx.mailer.sent == []
        dumps = list((store.root / "debug" / "sla").iterdir())
        names = sorted(p.name for p in dumps)
        assert "01_login_0.html" in names and "01_login_0.png" in names and any("_error_0.html" in n for n in names)
        error_html = next(p for p in dumps if p.name.endswith("_error_0.html")).read_text(encoding="utf-8")
        assert "***" in error_html                          # the portal echoed the password back; the dump hides it
        for path in dumps:
            data = path.read_bytes()
            assert data, path
            assert path.suffix != ".png" or data.startswith(b"\x89PNG")
            assert b"wrong-pw" not in data and PASSWORD.encode() not in data
            if path.suffix == ".html":
                text = path.read_text(encoding="utf-8")
                assert "value=" not in text.replace("value='دخول'", "").replace('value="دخول"', "")  # no typed / stored values

    def test_no_matching_table_lists_the_headers_seen(self, store):
        with FakePortal(sample_requests(NOW), variant="no_table") as portal:
            result, _ = portal_run(store, portal)
        assert result.ok is False and result.error_kind is None and result.snapshot_data is None
        assert "Requests table not found" in result.error and "Headers seen on the page: الرقم, الحالة" in result.error

    def test_package_not_on_the_groups_page(self, store):
        with FakePortal(sample_requests(NOW), package=False) as portal:
            result, _ = portal_run(store, portal)
        assert result.ok is False and "Package not found on the groups page" in result.error

    def test_reassign_page_that_redirects_away(self, store):
        with FakePortal(sample_requests(NOW), reassign_redirect=True) as portal:
            result, _ = portal_run(store, portal)
        assert result.ok is False and "Could not open the Re-Assign page" in result.error and GROUPS_PATH in result.error

    def test_an_empty_queue_is_a_good_scan(self, store):
        with FakePortal(sample_requests(NOW), variant="empty") as portal:
            result, ctx = portal_run(store, portal)
        assert result.ok, result.error
        assert result.snapshot_data["totals"]["open"] == 0 and result.emails_sent == 0

    def test_server_down_is_a_plain_error_without_the_password(self, store):
        with FakePortal(sample_requests(NOW)) as portal:
            url = portal.url
        class Gone:
            pass
        Gone.url = url
        ctx = make_ctx(store, e2e_settings(Gone, timeout=5))
        result = sla.run(ctx)
        assert result.ok is False and result.error_kind is None and PASSWORD not in result.error
        assert "ERR_CONNECTION_REFUSED" in result.error or "Page.goto" in result.error

    def test_js_normalisation_matches_python(self, tmp_path):
        samples = ["مُحَمَّد", "إأآ ا", "مدرسة", "على  مصطفى", "ـــ", "  Mixed CASE  ", "تحويل", "موظف اسناد - مستخدم السرية",
                   "٠١٢٣ ٤٥٦٧٨٩"]
        browser = cdp.CdpBrowser({"browser_executable": BROWSER["executable"], "profile_dir": str(tmp_path / "p")})
        try:
            got = [browser.eval("(() => {" + sla.JS_NORM + " return norm(" + json.dumps(s) + "); })()") for s in samples]
        finally:
            browser.close()
        assert got == [sla.norm(s) for s in samples]


def test_the_inquiry_follows_older_live_tickets_one_day_a_scan_never_a_month_at_once(store):
    """The analysis must know every open request, but a search a month back took 37 result pages and the whole scan:
    each scan searches the last ``history_days`` days plus ONE older day that still has a request in progress (a Live
    ticket's request date is a few days before it was sent), the oldest first and the next one at the next scan; never
    further back than HISTORY_REACH_DAYS."""
    ctx = make_ctx(store, None, NOW)
    plain = sla.history_plan(ctx, sla_settings(), NOW)
    assert plain["date_from"] == NOW.date() - timedelta(days=2) and plain["slices"] == [(plain["date_from"], NOW.date())]
    store.write_snapshot("sla", {"requests": [{"key": "old", "sent_at": (NOW - timedelta(days=9)).isoformat()},
                                              {"key": "ancient", "sent_at": (NOW - timedelta(days=90)).isoformat()}]},
                         generated_at=NOW.isoformat(timespec="seconds"))
    plan = sla.history_plan(ctx, sla_settings(), NOW)
    assert plan["date_from"] == NOW.date() - timedelta(days=2)                          # the recent window stays
    first = NOW.date() - timedelta(days=12)                                              # sent 9 days ago, 3 before
    assert plan["slices"][1] == (first, first) and plan["older_days"] == 4               # the 90-day one is too old
    store.update_sla_history_state(lambda d: d.update(history_older_day=first.isoformat()))
    after = sla.history_plan(ctx, sla_settings(), NOW)
    assert after["slices"][1] == (first + timedelta(days=1),) * 2                       # the next one, in turn


