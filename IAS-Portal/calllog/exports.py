"""Excel / CSV downloads. Everything here is read-only; each export is written to the audit log."""
import csv
import io
import json

from flask import Blueprint, g, request, send_file
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

from . import auth, overview as overview_data, stats
from .queries import ApiError, filtered_records, parse_range, reasons_by_key
from .storage import norm_uid

bp = Blueprint("exports", __name__)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
RECORD_HEAD = ["Date", "Time", "User ID", "Employee", "Department", "IP Phone", "Branch No.", "Branch",
               "Region", "Action Taken", "Action (English)", "Status", "Note", "Edited at", "IP phone location"]
BRANCH_HEAD = ["Branch No.", "Branch", "Region", "Total calls", "Resolved", "Transferred", "Ticket Inquiry"]
SUMMARY_HEAD = ["User ID", "Employee", "Department", "Total calls", "Resolved", "Transferred",
                "Ticket Inquiry", "First call", "Last call"]
AUDIT_HEAD = ["Time", "User ID", "User", "Action", "Target", "Details", "From address"]
_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


def _record_row(r):
    return [r["date"], r["time"], r["user_id"], r["name"], r["department"], r["ip_phone"],
            r["branch"], r["branch_name"], r["region"], r["reason_ar"], r["reason_en"],
            r["status_label"], r.get("note") or "", r.get("updated_at") or "", r.get("ip_unit") or ""]


def _summary_rows(d_from, d_to):
    rows = stats.employee_rows(auth.store().users(), overview_data.merged(d_from, d_to)["users"])
    return [[r["user_id"], r["name"], r["department"], r["total"], r["resolved"],
             r["transferred"], r["ticket_inquiry"], r["first_call"] or "", r["last_call"] or ""]
            for r in rows]


def _branch_rows(d_from, d_to):
    rows = stats.branch_rows(overview_data.merged(d_from, d_to)["branches"])
    return [[r["code"], r["name"], r["region"], r["total"], r["resolved"], r["transferred"],
             r["ticket_inquiry"]] for r in rows]


def _audit_rows(d_from, d_to):
    entries = auth.store().read_audit(d_from, d_to)
    entries.reverse()
    return [[e.get("ts", ""), e.get("actor") or "", e.get("actor_name") or "", e.get("action", ""),
             e.get("target") or "", json.dumps(e.get("details") or {}, ensure_ascii=False),
             e.get("ip") or ""] for e in entries]


# ------------------------------------------------------------------ writers
def _xlsx(sheets):
    """sheets: [(title, header, rows, widths)] -> BytesIO. Streams rows (write-only mode)."""
    wb = Workbook(write_only=True)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1C5CAB")
    for title, header, rows, widths in sheets:
        ws = wb.create_sheet(title)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[_col(i)].width = w
        ws.freeze_panes = "A2"
        cells = []
        for h in header:
            c = WriteOnlyCell(ws, value=h)
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(vertical="center")
            cells.append(c)
        ws.append(cells)
        for row in rows:
            out = []
            for v in row:
                if isinstance(v, str) and v.startswith("="):
                    c = WriteOnlyCell(ws, value=v)
                    c.data_type = "s"  # keep text that starts with '=' from becoming a formula
                    out.append(c)
                else:
                    out.append(v)
            ws.append(out)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _col(i):
    letters = ""
    while i:
        i, rem = divmod(i - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _csv(header, rows):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    for row in rows:
        # A leading ' stops Excel from running text such as =HYPERLINK(...) as a formula.
        writer.writerow(["'" + v if isinstance(v, str) and v.startswith(_FORMULA_STARTS) else v
                         for v in row])
    return io.BytesIO(buf.getvalue().encode("utf-8-sig"))  # BOM so Excel shows Arabic correctly


def _send(kind, fmt, d_from, d_to, header, rows, widths, extra=None):
    rows = list(rows)
    name = f"calllog-{kind}-{d_from}" + (f"_to_{d_to}" if d_to != d_from else "")
    auth.audit("export", target=f"{kind}.{fmt}", details={"from": d_from, "to": d_to, "rows": len(rows)})
    if fmt == "csv":
        return send_file(_csv(header, rows), mimetype="text/csv", as_attachment=True,
                         download_name=f"{name}.csv")
    sheets = [(SHEET_TITLES.get(kind, kind.capitalize()), header, rows, widths)]
    if extra:
        sheets = extra + sheets
    return send_file(_xlsx(sheets), mimetype=XLSX_MIME, as_attachment=True,
                     download_name=f"{name}.xlsx")


def _fmt(value):
    if value not in ("xlsx", "csv"):
        raise ApiError("Unknown format.", 404)
    return value


SHEET_TITLES = {"records": "Calls", "report": "Calls", "my-calls": "My calls", "audit": "Audit log"}
RECORD_WIDTHS = [12, 10, 12, 26, 24, 12, 11, 26, 30, 28, 30, 16, 36, 20, 30]
BRANCH_WIDTHS = [11, 30, 34, 12, 11, 13, 14]
SUMMARY_WIDTHS = [12, 26, 24, 12, 11, 13, 14, 20, 20]


# ------------------------------------------------------------------- routes
@bp.get("/records.<fmt>")
@auth.admin_required
def records(fmt):
    fmt = _fmt(fmt)
    d_from, d_to = parse_range()
    rows = [_record_row(r) for r in filtered_records(d_from, d_to, request.args)]
    return _send("records", fmt, d_from, d_to, RECORD_HEAD, rows, RECORD_WIDTHS)


@bp.get("/report.xlsx")
@auth.admin_required
def report():
    """Summary per employee + every record, in one workbook."""
    d_from, d_to = parse_range()
    rows = [_record_row(r) for r in filtered_records(d_from, d_to, {})]
    summary = ("Summary", SUMMARY_HEAD, _summary_rows(d_from, d_to), SUMMARY_WIDTHS)
    by_branch = ("By branch", BRANCH_HEAD, _branch_rows(d_from, d_to), BRANCH_WIDTHS)
    return _send("report", "xlsx", d_from, d_to, RECORD_HEAD, rows, RECORD_WIDTHS, extra=[summary, by_branch])


@bp.get("/audit.<fmt>")
@auth.admin_required
def audit_log(fmt):
    fmt = _fmt(fmt)
    d_from, d_to = parse_range()
    return _send("audit", fmt, d_from, d_to, AUDIT_HEAD, _audit_rows(d_from, d_to),
                 [20, 12, 24, 18, 30, 60, 16])


@bp.get("/my.<fmt>")
@auth.login_required
def my_records(fmt):
    """An employee's own records (the digital version of their Excel sheet)."""
    fmt = _fmt(fmt)
    d_from, d_to = parse_range()
    rows = [_record_row(r) for r in
            filtered_records(d_from, d_to, {"user_id": norm_uid(g.user["user_id"])})]
    return _send("my-calls", fmt, d_from, d_to, RECORD_HEAD, rows, RECORD_WIDTHS)
