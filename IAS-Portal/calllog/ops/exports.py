"""Downloads of the Operations Hub, mounted at ``/export/ops``.

* ``/export/ops/<module>.xlsx`` and ``.csv``: the latest snapshot as a workbook (NBE green header, like the Call Log
  exports) or as a plain CSV of the main table. E-mail rows have the overrides applied and are limited to the teams
  of a manager (an account can hold several teams).
* ``/export/ops/report/<module>/<name>``: a report file the collector produced (Excel / HTML), never outside its folder.

Every download is written to the audit log. Text that starts with ``=`` is written as text, never as a formula.
"""
import csv
import io

from flask import Blueprint, jsonify, send_file
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

from .. import auth
from ..exports import XLSX_MIME, _col
from . import api as ops_api, overlay, schemas

bp = Blueprint("ops_exports", __name__)

NBE_GREEN = "00573F"
_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")

EMAIL_HEAD = ["Received", "Subject", "Customer", "Customer e-mail", "Team", "Status", "SLA", "Escalated", "Waiting",
              "First response", "Response (min)", "First reply by", "First reply at", "Assigned to", "Quality",
              "Systems", "Carry-over", "Age (days)", "Reminders"]


# ------------------------------------------------------------------------------ writers
def _when(value):
    return str(value).replace("T", " ") if value else ""


def _width(header, rows, index):
    longest = max([len(str(header))] + [len(str(r[index])) for r in rows[:200] if index < len(r)])
    return min(max(longest + 2, 10), 60)


def _xlsx(sheets):
    """sheets: [(title, header, rows)] -> BytesIO (write-only workbook, green header row)."""
    wb = Workbook(write_only=True)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=NBE_GREEN)
    for title, header, rows in sheets:
        ws = wb.create_sheet(title[:31])
        for i, name in enumerate(header):
            ws.column_dimensions[_col(i + 1)].width = _width(name, rows, i)
        ws.freeze_panes = "A2"
        cells = []
        for name in header:
            cell = WriteOnlyCell(ws, value=name)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="center")
            cells.append(cell)
        ws.append(cells)
        for row in rows:
            out = []
            for value in row:
                if isinstance(value, str) and value.startswith("="):
                    cell = WriteOnlyCell(ws, value=value)
                    cell.data_type = "s"           # text that starts with '=' must not become a formula
                    out.append(cell)
                else:
                    out.append(value)
            ws.append(out)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _csv(header, rows):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    for row in rows:
        writer.writerow(["'" + v if isinstance(v, str) and v.startswith(_FORMULA_STARTS) else v for v in row])
    return io.BytesIO(buf.getvalue().encode("utf-8-sig"))      # BOM so Excel shows Arabic correctly


# ------------------------------------------------------------------------------ tables per module
def _email_tables(snap):
    scope = auth.team_scope()
    view = ops_api.email_view(snap, scope)
    stats, rows = view["stats"], view["rows"]
    summary = [["Date", snap.get("date")], ["Scan time", _when(snap.get("generated_at"))],
               ["Team", overlay.scope_label(scope) or "All teams"],
               ["Business e-mails", stats.get("business_emails", 0)], ["Replied", stats.get("replied", 0)],
               ["Pending reply", stats.get("pending", 0)], ["No action needed", stats.get("no_action_needed", 0)],
               ["Needed a reply", stats.get("needed_reply", 0)], ["Reply rate (%)", round(stats.get("reply_rate", 0), 1)],
               ["SLA target (min)", stats.get("sla_target_minutes", "")], ["SLA breached", stats.get("sla_breached", 0)],
               ["SLA compliance (%)", round(stats.get("sla_compliance_rate", 0), 1)],
               ["Escalated", stats.get("escalated_count", 0)], ["Carried over", stats.get("carryover_count", 0)],
               ["Assigned", stats.get("assigned_count", 0)],
               ["Average response (min)", round(stats.get("average_response", 0), 1)],
               ["Fastest response (min)", stats.get("fastest_response", 0)],
               ["Slowest response (min)", stats.get("slowest_response", 0)]]
    teams = []
    for name, t in sorted((stats.get("team_stats") or {}).items()):
        count = t.get("response_count") or 0
        teams.append([name, t.get("business_emails", 0), t.get("replied", 0), t.get("pending", 0),
                      t.get("no_action_needed", 0), t.get("escalated_count", 0), t.get("carryover_count", 0),
                      t.get("sla_breached_count", 0),
                      round(t.get("response_minutes_sum", 0) / count, 1) if count else ""])
    table = [[_when(r.get("received_time")), r.get("subject", ""), r.get("customer_name", ""),
              r.get("customer_email", ""), r.get("team", ""), r.get("status", ""), r.get("sla_status", ""),
              "Yes" if r.get("escalated") else "", r.get("current_waiting_time", ""), r.get("first_response_time", ""),
              r.get("response_minutes") if r.get("response_minutes") is not None else "", r.get("first_reply_by", ""),
              _when(r.get("first_reply_time")), r.get("assigned_to_name") or r.get("assigned_to_email", ""),
              r.get("quality_rating", ""), r.get("detected_systems", ""), "Yes" if r.get("is_carryover") else "",
              r.get("age_days", 0), r.get("reminder_count", 0)] for r in rows]
    return "E-mails", EMAIL_HEAD, table, [
        ("Summary", ["Measure", "Value"], summary),
        ("Teams", ["Team", "E-mails", "Replied", "Pending", "No action", "Escalated", "Carry-over", "SLA breached",
                   "Average response (min)"], teams)]


def _sla_tables(snap):
    data = snap.get("data") or {}
    limits, totals = data.get("thresholds") or {}, data.get("totals") or {}
    summary = [["Scan time", _when(snap.get("generated_at"))], ["Reminder after (min)", limits.get("warning_minutes", "")],
               ["Breach after (min)", limits.get("breach_minutes", "")], ["Open", totals.get("open", 0)],
               ["Within SLA", totals.get("within", 0)], ["Reminder", totals.get("reminder", 0)],
               ["Breached", totals.get("breached", 0)], ["No sent date", totals.get("no_date", 0)],
               ["Oldest (min)", totals.get("oldest_minutes", 0)]]
    by_user = [[u.get("user", ""), u.get("total", 0), u.get("within", 0), u.get("reminder", 0), u.get("breached", 0),
                u.get("no_date", 0), u.get("oldest_minutes", 0)] for u in data.get("by_user") or []]
    requests = [r for r in data.get("requests") or [] if isinstance(r, dict)]
    columns = []
    for r in requests:
        for key in (r.get("fields") or {}):
            if key not in columns:
                columns.append(key)
    head = ["Security user", "Sent at", "Age (min)", "Status"] + columns
    table = [[r.get("security_user", ""), _when(r.get("sent_at")), r.get("age_minutes") if r.get("age_minutes") is not None else "",
              schemas.SLA_LABELS.get(r.get("status"), r.get("status", ""))] + [(r.get("fields") or {}).get(c, "") for c in columns]
             for r in requests]
    return "Requests", head, table, [
        ("Summary", ["Measure", "Value"], summary),
        ("By user", ["Security user", "Total", "Within SLA", "Reminder", "Breached", "No sent date", "Oldest (min)"], by_user)]


def _itsm_tables(snap):
    data = snap.get("data") or {}
    categories = data.get("categories") or []
    summary = [["Scan time", _when(snap.get("generated_at"))]] + [[c, (data.get("totals") or {}).get(c, 0)] for c in categories]
    by_role, table = [], []
    for role in data.get("roles") or []:
        per = role.get("categories") or {}
        by_role.append([role.get("role", ""), role.get("error") or ""] + [(per.get(c) or {}).get("count", "") for c in categories])
        for c in categories:
            for item in (per.get(c) or {}).get("rows") or []:
                table.append([role.get("role", ""), c, item.get("number", ""), item.get("summary", "")])
    changes = [[kind, c.get("role", ""), c.get("category", ""), c.get("number", ""), c.get("summary", "")]
               for kind in ("new", "closed") for c in (data.get("changes") or {}).get(kind) or []]
    return "Requests", ["Role", "Category", "Number", "Summary"], table, [
        ("Summary", ["Measure", "Value"], summary),
        ("By role", ["Role", "Error"] + list(categories), by_role),
        ("Changes", ["Change", "Role", "Category", "Number", "Summary"], changes)]


_TABLES = {"email": _email_tables, "sla": _sla_tables, "itsm": _itsm_tables}


# ------------------------------------------------------------------------------ routes
@bp.get("/<any(email,sla,itsm):module>.<any(xlsx,csv):fmt>")
@auth.ops_view_required
def snapshot_export(module, fmt):
    snap, _ = ops_api.read_snapshot(module)
    if not snap:
        return jsonify(error="There is no data yet - the collector has not written a scan."), 404
    title, head, rows, extra = _TABLES[module](snap)
    auth.audit("ops_export", target=f"{module}.{fmt}", details={"rows": len(rows)})
    name = f"nbe-{module}-{snap.get('date') or auth.today()}"
    if fmt == "csv":
        return send_file(_csv(head, rows), mimetype="text/csv", as_attachment=True, download_name=f"{name}.csv")
    return send_file(_xlsx(extra[:1] + [(title, head, rows)] + extra[1:]), mimetype=XLSX_MIME, as_attachment=True,
                     download_name=f"{name}.xlsx")


@bp.get("/report/<module>/<name>")
@auth.ops_view_required
def report_file(module, name):
    path = ops_api.ops().report_path(module, name)
    if path is None:
        return jsonify(error="That report does not exist."), 404
    auth.audit("ops_report_download", target=f"{module}/{path.name}")
    return send_file(path, as_attachment=True, download_name=path.name)
