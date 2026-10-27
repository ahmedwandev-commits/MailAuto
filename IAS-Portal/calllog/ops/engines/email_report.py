"""The "email_report" job: the reporting half of the old Daily Team Reply Analytics script.

The scan and the analysis are email_engine.py; this module turns their result into what the old script produced once a
day: the Excel workbook, the HTML e-mails (main report, per-team pending lists, assignment notices, quality feedback,
end-of-day summary), the zipped ``.msg`` copies of the unanswered mails and the daily history. Entry point
(docs/HUB-DESIGN.md section 12)::

    run_report(ctx, namespace=None) -> JobResult        ``namespace`` lets tests inject the fake Outlook

Steps (the old ``main()``), each in its own try/except so one failure never skips the others; a failed step is listed
in ``JobResult.error`` and the run is then not ``ok`` (the mails that did go out are still counted in ``emails_sent``)::

    [1] scan                  email_engine.scan_and_analyze (keeps the Outlook items only when .msg files are wanted)
    [2] history               today's stats + quality history, trend and repeat-issue alerts
    [3] Excel                 <reports>/email/Team_Reply_Analytics_<YYYY-MM-DD>.xlsx, files older than 60 days removed
    [4] charts                matplotlib PNGs (in memory, sent as inline images) or a pure-table bar fallback
    [5] .msg files            COM ``item.SaveAs`` into an OS temp folder, one zip per team and chunk, deleted afterwards
    [6] report + team lists   main report to report.to/cc, one follow-up per team that has pending mail
    [7] assignment notices    only assignments not yet "notified" (the flag is written back to the assignments override)
    [7b] quality feedback     only reviews not yet sent (state "email_quality_sent")
    [8] end-of-day summary    to end_of_day.to (else the report recipients)

Switches. ``report.send`` is the switch of the daily report: it turns off the main report, the team follow-ups AND the
end-of-day summary (which is the wrap-up of that report, even when it has its own recipient list). Assignment notices
and quality feedback have their own flags (``assignment_notifications.send``, ``quality_feedback.send``) because a
manager acts on them all day, not only at report time. Every mail step also needs recipients and ``ctx.can_email()``:
``ctx.no_email`` builds everything and sends nothing (no state, no "notified" flag); ``ctx.dry_run`` additionally writes
nothing into the data folder. Nothing is sent to an address that is not exactly one well-formed address.

Idempotence: ``ctx.store`` state "email_report" remembers, per day and report slot (the latest ``report.send_times``
entry that has passed, or "manual"), which mails went out - the report, every part of every team follow-up, the
summary. A second call for the same slot sends no report again - and returns ``skipped`` without touching Outlook when
nothing else is waiting - unless ``ctx.options["force"]``. A step that failed is remembered and retried alone next
time (a team's parts that already went out are not sent twice); a failure that no longer applies (the step was
switched off meanwhile) is forgotten. Assignment notices and quality feedback are never repeated, they carry their own
flags, so a "Run now" still delivers an assignment made after the report; such a run does not save the ``.msg`` files or
draw the charts again. No weekend / holiday logic: the runner owns it.

Files. The workbook is written under a hidden temp name and moved into place; a copy open in Excel is not overwritten
(the new one gets ``_<HHMMSS>`` in its name). The ``.msg`` files, zips and charts live in one temp folder that is deleted
when the run ends, even after a failure; folders a killed run left behind are removed after a day.

HTML rules: tables, inline styles and bgcolor attributes only (Outlook's Word engine), no SVG / grid / flex, every
dynamic value HTML-escaped (subjects and names are hostile input), NBE colours from common.py, the logo as an inline
CID image (``assets/nbe_logo.png``; a text badge when the file is missing).
"""
import contextlib
import copy
import io
import os
import random
import re
import shutil
import tempfile
import time
import zipfile
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import email_engine as eng
from .common import (
    FONT_STACK, NBE_GOLD, NBE_GREEN, NBE_GREEN_LIGHT, OUTLOOK_LOCK, JobResult, OutlookMailer, esc, norm_email)
from .email_rules import UNCLASSIFIED, format_duration, to_int

STATE_NAME = "email_report"                     # day -> slot -> mails already sent
QUALITY_SENT_STATE = "email_quality_sent"       # message_id -> the review that was e-mailed
REPORT_PREFIX = "Team_Reply_Analytics_"
REPORT_KEEP_DAYS = 60
ASSETS = Path(__file__).resolve().parent.parent / "assets"
LOGO_PATH = ASSETS / "nbe_logo.png"
LOGO_CID = "nbe_logo"
STATUS_CHART_CID, TEAM_CHART_CID = "reply_status_chart", "team_chart"
OL_MSG = 9                                      # olMSGUnicode: keeps Arabic subjects/body (olMSG=3 is ANSI -> "?")
MAX_LISTED_ROWS = 60                            # rows one mail lists in a table
CHART_MEMBERS = 12                              # people in the "first replies" chart
XL_MAX_CHARS = 32000                            # Excel keeps 32767 characters per cell

ASSIGNMENT_SUBJECT = "Action Required - Emails Assigned to You"
QUALITY_SUBJECT = "Reply Quality Feedback"
END_OF_DAY_SUBJECT = "End of Day Summary - Team Reply Analytics"

# Status colours of the old script (kept: people read them the same way in Excel, charts and mails).
GREEN, RED, AMBER, BLUE, PURPLE, BROWN = "#007336", "#E4572E", "#FFA320", "#2E86AB", "#7B4FA6", "#B0472B"
INK, MUTED, LINE, PAGE_BG, CARD_BG, GREY = "#33404f", "#667488", "#eef1f6", "#eef1f6", "#F4F8FC", "#8a95a3"
VERY_GOOD = "#3E9E63"
RATING_COLOURS = {"Excellent": GREEN, "Very Good": VERY_GOOD, "Good": BLUE, "Needs Improvement": BROWN}


# ---------------------------------------------------------------------------------------------------------------------
# Settings
# ---------------------------------------------------------------------------------------------------------------------
def _section(settings, name):
    value = settings.get(name) if isinstance(settings, dict) else None
    return value if isinstance(value, dict) else {}


_ADDRESS = re.compile(r"^[^@\s;,<>]+@[^@\s;,<>]+$")


def single_address(value):
    """The lower-case address, or "" when ``value`` is not exactly one address (``a@x; b@y`` must never reach two
    people through a field that was meant for one)."""
    value = norm_email(value)
    return value if _ADDRESS.match(value) else ""


def addresses(values):
    """Lower-case, de-duplicated, sorted e-mail addresses (the settings already store them that way)."""
    if isinstance(values, str):
        values = re.split(r"[;,\s]+", values)
    return sorted({single_address(v) for v in (values or ()) if single_address(v)})


def one_line(text):
    """A subject / header value can never carry a line break."""
    return " ".join(str(text or "").split())


class ReportConfig:
    """The report-related e-mail settings with the defaults of settings.py, prepared once per run."""

    def __init__(self, settings):
        report = _section(settings, "report")
        assign = _section(settings, "assignment_notifications")
        quality = _section(settings, "quality_feedback")
        eod = _section(settings, "end_of_day")
        self.send = bool(report.get("send", True))
        self.to, self.cc = addresses(report.get("to")), addresses(report.get("cc"))
        self.subject_prefix = one_line(report.get("subject_prefix")) or "Daily Team Reply Analytics"
        self.attach_excel = bool(report.get("attach_excel", True))
        self.attach_msg = bool(report.get("attach_pending_msg", True))
        self.max_per_zip = max(1, to_int(report.get("max_pending_attachments"), 40))
        self.send_times = sorted(str(t) for t in (report.get("send_times") or []) if re.match(r"^\d\d:\d\d$", str(t)))
        self.team_recipients = {}
        for team, value in (report.get("team_recipients") or {}).items():
            value = value if isinstance(value, dict) else {}
            self.team_recipients[str(team)] = (addresses(value.get("to")), addresses(value.get("cc")))
        self.assign_send, self.assign_cc = bool(assign.get("send", True)), addresses(assign.get("cc"))
        self.quality_send, self.quality_cc = bool(quality.get("send", True)), addresses(quality.get("cc"))
        self.sample_size = max(0, to_int(quality.get("sample_size"), 5))
        self.eod_send = bool(eod.get("send", True))
        # the old script sent the end-of-day summary to the report recipients unless told otherwise
        self.eod_to = addresses(eod.get("to")) or self.to
        self.eod_cc = addresses(eod.get("cc")) or ([] if addresses(eod.get("to")) else self.cc)


# ---------------------------------------------------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------------------------------------------------
def _int(value):
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _float(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _duration(minutes):
    return format_duration(_float(minutes))


def _cut(text, limit):
    text = " ".join(str(text or "").split())
    return text if len(text) <= limit else text[:limit - 1].rstrip() + "\u2026"


def _counts(counter):
    """A Counter or plain dict as ``[(name, number)]``, biggest first."""
    return sorted(((str(k), _int(v)) for k, v in dict(counter or {}).items()), key=lambda kv: (-kv[1], kv[0]))


def _day_text(day):
    return f"{day:%d/%m/%Y}"


def chunk_list(items, size):
    size = max(1, size)
    return [items[i:i + size] for i in range(0, len(items), size)]


_RESERVED_NAMES = {"con", "prn", "aux", "nul", *(f"com{i}" for i in range(1, 10)), *(f"lpt{i}" for i in range(1, 10))}


def safe_name(text, max_len=50, default="email"):
    """A file or folder name Windows accepts: no separators / wildcards / control characters, no trailing dots."""
    text = re.sub(r'[\\/*?:"<>|\x00-\x1f]', "_", str(text or ""))
    text = text.strip(" ._")[:max_len].strip(" ._")
    if not text:
        return default
    return "_" + text if text.split(".")[0].lower() in _RESERVED_NAMES else text


def file_stem(text, max_len=50, default="email"):
    """``safe_name`` with underscores instead of spaces: attachment names travel through mail gateways and browsers."""
    return re.sub(r"\s+", "_", safe_name(text, max_len, default))


def pending_rows_by_team(rows):
    """{team: [rows]} of the mails each team still has to answer (a mail for two teams is listed under both)."""
    by_team = {}
    for row in rows:
        if row.get("status") == "No Action Needed":
            continue
        for team in row.get("teams_pending") or []:
            by_team.setdefault(team, []).append(row)
    return by_team


def review_sample(rows, size, day):
    """Replied mails nobody rated yet, picked at random for a spot check. Seeded with the day: the same list every run."""
    pool, seen = [], set()
    for row in rows:
        key = row.get("message_id") or id(row)
        if (row.get("status") == "Replied" and row.get("first_reply_by") and not row.get("quality_rating")
                and not row.get("reply_by_manager") and key not in seen):
            seen.add(key)
            pool.append(row)
    picked = random.Random(f"review-sample-{day}").sample(pool, min(max(0, size), len(pool)))
    return sorted(picked, key=lambda r: r["received_time"])


# ---------------------------------------------------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------------------------------------------------
XL_GREEN = NBE_GREEN.lstrip("#")
TEAM_EXCEL_COLORS = {                           # the tint of a team in the "Team" column
    "MFA Team": "FCE9CE", "Flex Cube": "DCEBFB", "Card Management": "EAE0F5", "Third Party": "FBE0D9",
    "Open System": "DFF5E7", "Unclassified": "EFEFEF"}
DETAIL_HEADERS = [
    "Subject", "Detected System(s)", "Team", "Customer / Sender", "Customer Email", "Received Time", "Mailbox / PST",
    "Folder", "Conversation ID", "Business Address in To", "Match Method", "First Team Reply By",
    "First Team Reply Email", "First Reply Time", "First Response Time", "Last Team Reply By",
    "Last Team Reply Email", "Last Reply Time", "Team Reply Count", "Current Waiting Time", "Status",
    "SLA Target (Min)", "SLA Status", "Escalated", "Reminders Sent", "All Matched Teams", "Teams Still Pending",
    "Diag: Nearest Team Email By", "Diag: Nearest Team Email", "Diag: Nearest Team Email Time",
    "Diag: Nearest Team Email Subject", "Diag: Gap (Min)", "Diag: Why Not Matched", "Diag: No Action Reason"]
DATE_FORMAT = "dd/mm/yyyy hh:mm AM/PM"
STATUS_COLUMN, ESCALATED_COLUMN = "U", "X"      # letters of "Status" and "Escalated" in DETAIL_HEADERS


def _xl_value(value):
    """Mail text is hostile input: control characters make openpyxl raise, a leading '=' would become a formula."""
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)[:XL_MAX_CHARS]
    return value


def _put(ws, row, col, value):
    cell = ws.cell(row, col)
    cell.value = _xl_value(value)
    if cell.data_type == "f":
        cell.data_type = "s"
    return cell


def _append(ws, values):
    ws.append([_xl_value(v) for v in values])
    for cell in ws[ws.max_row]:
        if cell.data_type == "f":
            cell.data_type = "s"


def _fill(color):
    return PatternFill(fill_type="solid", start_color=color, end_color=color)


def style_header(ws):
    for cell in ws[1]:
        cell.fill = _fill(XL_GREEN)
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 20


def auto_width(ws):
    for column_cells in ws.columns:
        longest = max((len(str(c.value or "")) for c in column_cells), default=0)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(longest + 2, 12), 55)


def apply_status_colors(ws, status_column):
    if ws.max_row < 2:
        return
    area = f"A2:{get_column_letter(ws.max_column)}{ws.max_row}"
    for text, color in (("Replied", "DDF3E8"), ("Pending Reply", "FBDCD3"), ("No Action Needed", "FCEFD0")):
        ws.conditional_formatting.add(area, FormulaRule(formula=[f'${status_column}2="{text}"'], fill=_fill(color)))


def apply_escalation_highlight(ws, escalated_column):
    """Bold dark-red text (the fill stays the status colour) for every escalated row."""
    if ws.max_row < 2:
        return
    area = f"A2:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.conditional_formatting.add(area, FormulaRule(formula=[f'${escalated_column}2="Yes"'],
                                                    font=Font(bold=True, color="C00000")))


def _finish_table(ws):
    ws.freeze_panes = "A2"
    if ws.max_row >= 2:
        ws.auto_filter.ref = ws.dimensions


def create_dashboard(workbook, rows, stats, today):
    ws = workbook.create_sheet("Dashboard", 0)
    _put(ws, 1, 1, "TEAM REPLY DAILY ANALYTICS").font = Font(name="Calibri", size=20, bold=True, color=XL_GREEN)
    _put(ws, 2, 1, "Analysis Date").font = Font(bold=True, color="667488")
    _put(ws, 2, 2, today.strftime("%d/%m/%Y")).font = Font(bold=True, color=XL_GREEN)

    kpis = [
        ("Business Emails", stats["business_emails"]), ("Team Replies", stats["team_replies"]),
        ("Replied", stats["replied"]), ("Pending", stats["pending"]),
        ("No Action Needed", stats["no_action_needed"]), ("Reply Rate", f'{stats["reply_rate"]:.1f}%'),
        ("Total Team Replies", stats["total_team_replies"]),
        ("Avg First Response", format_duration(stats["average_response"])),
        ("Fastest Response", format_duration(stats["fastest_response"])),
        ("Slowest Response", format_duration(stats["slowest_response"])),
        (f'SLA Breached (>{stats["sla_target_minutes"]}m)', stats["sla_breached"]),
        ("Escalated (Pending)", stats["escalated_count"]),
    ]
    accents = ["007336", "E4572E", "2E86AB"]
    for index, (label, value) in enumerate(kpis):
        row, col = 4 + (index // 3) * 3, 1 + (index % 3) * 3
        _put(ws, row, col, label)
        _put(ws, row + 1, col, value)
        for r in range(row, row + 2):
            for c in range(col, col + 2):
                cell = ws.cell(r, c)
                cell.fill = _fill("EEF3F9")
                if r == row:
                    cell.border = Border(top=Side(style="medium", color=accents[index % 3]))
        ws.cell(row, col).font = Font(bold=True, size=10, color="667488")
        ws.cell(row + 1, col).font = Font(bold=True, size=16, color=XL_GREEN)

    # Chart data sits beside the charts; the charts themselves start below the KPI cards (the old layout let the pie
    # cover three of them).
    status_row = 16
    for offset, (label, value) in enumerate([("Status", "Count"), ("Replied", stats["replied"]),
                                             ("Pending", stats["pending"]),
                                             ("No Action Needed", stats["no_action_needed"])]):
        _put(ws, status_row + offset, 1, label)
        _put(ws, status_row + offset, 2, value)
    pie = PieChart()
    pie.add_data(Reference(ws, min_col=2, min_row=status_row, max_row=status_row + 3), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=1, min_row=status_row + 1, max_row=status_row + 3))
    pie.title, pie.height, pie.width = "Reply Status", 8, 12
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    for point_index, color in enumerate(("007336", "E4572E", "FFA320")):
        point = DataPoint(idx=point_index)
        point.graphicalProperties.solidFill = color
        pie.series[0].data_points.append(point)
    ws.add_chart(pie, "A21")

    _put(ws, 2, 11, "Team Member")
    _put(ws, 2, 12, "First Replies")
    members = _counts(stats["first_reply_by"])
    for r, (member, count) in enumerate(members, 3):
        _put(ws, r, 11, member)
        _put(ws, r, 12, count)
    if members:
        bar = BarChart()
        bar.add_data(Reference(ws, min_col=12, min_row=2, max_row=2 + len(members)), titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=11, min_row=3, max_row=2 + len(members)))
        bar.title, bar.height, bar.width = "First Replies by Team Member", 8, 14
        bar.series[0].graphicalProperties.solidFill = "2E86AB"
        ws.add_chart(bar, "D21")

    team_row = 40
    headers = ["Team", "Business Emails", "Replied", "Pending", "No Action Needed", "Escalated", "SLA Breached",
               "Avg First Response"]
    for c, text in enumerate(headers, 1):
        cell = _put(ws, team_row, c, text)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), _fill(XL_GREEN)
    team_order = sorted(stats["team_stats"].items(), key=lambda kv: kv[1]["business_emails"], reverse=True)
    for r, (team, data) in enumerate(team_order, team_row + 1):
        average = data["response_minutes_sum"] / data["response_count"] if data["response_count"] else None
        for c, value in enumerate([team, data["business_emails"], data["replied"], data["pending"],
                                   data["no_action_needed"], data["escalated_count"], data["sla_breached_count"],
                                   format_duration(average) if average is not None else "-"], 1):
            _put(ws, r, c, value)

    address_row = team_row + len(team_order) + 3
    _put(ws, address_row, 1, "Business Address")
    _put(ws, address_row, 2, "Included Emails")
    for r, (address, count) in enumerate(_counts(stats["business_address_volume"]), address_row + 1):
        _put(ws, r, 1, address)
        _put(ws, r, 2, count)

    ws.freeze_panes = "A2"
    for letter, width in zip("ABCDEFGHIJKL", (36, 18, 14, 18, 18, 14, 18, 18, 4, 4, 30, 14)):
        ws.column_dimensions[letter].width = width


def create_detailed_report(workbook, rows):
    ws = workbook.create_sheet("Detailed Report")
    ws.append(DETAIL_HEADERS)
    for r in rows:
        _append(ws, [
            r["subject"], r["detected_systems"], r["team"], r["customer_name"], r["customer_email"],
            r["received_time"], r["mailbox"], r["folder"], r["conversation_id"], r["business_addresses_in_to"],
            r["match_method"], r["first_reply_by"], r["first_reply_email"], r["first_reply_time"],
            r["first_response_time"], r["last_reply_by"], r["last_reply_email"], r["last_reply_time"],
            r["reply_count"], r["current_waiting_time"], r["status"], r["sla_minutes"], r["sla_status"],
            "Yes" if r["escalated"] else "No", r.get("reminder_count", 0),
            "; ".join(r.get("all_teams") or [r["team"]]), "; ".join(r.get("teams_pending") or []) or "\u2014",
            r.get("diag_candidate_by", ""), r.get("diag_candidate_email", ""), r.get("diag_candidate_time"),
            r.get("diag_candidate_subject", ""),
            r["diag_gap_minutes"] if r.get("diag_gap_minutes") is not None else "",
            r.get("diag_reason", ""), r.get("diag_no_action_reason", "")])
    style_header(ws)
    for r in range(2, ws.max_row + 1):
        color = TEAM_EXCEL_COLORS.get(ws.cell(r, 3).value)
        if color:
            ws.cell(r, 3).fill = _fill(color)
        for c in (6, 14, 18, 30):
            ws.cell(r, c).number_format = DATE_FORMAT
    _finish_table(ws)
    apply_status_colors(ws, STATUS_COLUMN)
    apply_escalation_highlight(ws, ESCALATED_COLUMN)
    auto_width(ws)


def create_team_sheet(workbook, team_replies):
    ws = workbook.create_sheet("Team Replies")
    ws.append(["Team Member Name", "Team Member Email", "Reply Time", "Subject", "Mailbox / PST", "Folder",
               "Conversation ID"])
    for r in sorted(team_replies, key=lambda x: x["received_time"]):
        _append(ws, [r["sender_name"], r["sender_email"], r["received_time"], r["subject"], r["mailbox"], r["folder"],
                     r["conversation_id"]])
    style_header(ws)
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 3).number_format = DATE_FORMAT
    _finish_table(ws)
    auto_width(ws)


def create_team_analysis(workbook, rows):
    ws = workbook.create_sheet("Team Analysis")
    analysis = defaultdict(lambda: {"count": 0, "times": []})
    for r in rows:
        member = r["first_reply_by"]
        if not member:
            continue
        analysis[member]["count"] += 1
        if r["first_reply_time"]:
            analysis[member]["times"].append(int((r["first_reply_time"] - r["received_time"]).total_seconds() // 60))
    ws.append(["Team Member", "First Replies", "Avg First Response", "Fastest", "Slowest"])
    for member, data in sorted(analysis.items(), key=lambda x: x[1]["count"], reverse=True):
        times = data["times"]
        _append(ws, [member, data["count"], format_duration(sum(times) / len(times) if times else 0),
                     format_duration(min(times) if times else 0), format_duration(max(times) if times else 0)])
    style_header(ws)
    _finish_table(ws)
    auto_width(ws)


def create_workbook(rows, team_replies, stats, today):
    workbook = Workbook()
    workbook.remove(workbook.active)
    create_dashboard(workbook, rows, stats, today)
    create_detailed_report(workbook, rows)
    create_team_sheet(workbook, team_replies)
    create_team_analysis(workbook, rows)
    return workbook


def create_excel(rows, team_replies, stats, today, folder, now=None):
    """Write ``Team_Reply_Analytics_<YYYY-MM-DD>.xlsx`` into ``folder`` and return its path.

    The workbook is saved under a hidden temp name and moved into place, so a reader on the share never sees half a
    file. Windows refuses to replace a workbook somebody has open in Excel (PermissionError): then the new one is kept
    next to it as ``..._<HHMMSS>.xlsx`` instead of the run failing or the open copy being clobbered. The temp file
    never survives the call."""
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    final = folder / f"{REPORT_PREFIX}{today:%Y-%m-%d}.xlsx"
    temp = folder / f".tmp-{final.name}"
    try:
        create_workbook(rows, team_replies, stats, today).save(temp)
        try:
            os.replace(temp, final)
        except PermissionError:
            final = folder / f"{REPORT_PREFIX}{today:%Y-%m-%d}_{(now or datetime.now()):%H%M%S}.xlsx"
            os.replace(temp, final)
    finally:
        if temp.exists():
            temp.unlink()
    return final


_REPORT_NAME = re.compile(rf"^{REPORT_PREFIX}(\d{{4}}-\d{{2}}-\d{{2}})(_\d{{6}})?\.xlsx$")
_LEFTOVER_NAME = re.compile(rf"^\.tmp-{REPORT_PREFIX}\d{{4}}-\d{{2}}-\d{{2}}\.xlsx$")
STALE_TEMP_SECONDS = 24 * 3600


def prune_reports(folder, today, keep_days=REPORT_KEEP_DAYS):
    """Delete this job's workbooks older than ``keep_days`` and its own temp files that a killed run left behind
    (other files in the folder are never touched). Returns the number of workbooks removed."""
    removed = 0
    cutoff = today - timedelta(days=keep_days)
    try:
        names = [p for p in Path(folder).iterdir() if p.is_file()]
    except OSError:
        return 0
    for path in names:
        try:
            if _LEFTOVER_NAME.match(path.name):
                if time.time() - path.stat().st_mtime > STALE_TEMP_SECONDS:
                    path.unlink()
                continue
            match = _REPORT_NAME.match(path.name)
            if match and datetime.strptime(match.group(1), "%Y-%m-%d").date() < cutoff:
                path.unlink()
                removed += 1
        except (ValueError, OSError):
            continue
    return removed


# ---------------------------------------------------------------------------------------------------------------------
# Charts (matplotlib is optional: the e-mail falls back to table bars)
# ---------------------------------------------------------------------------------------------------------------------
def _style_axes(ax):
    from matplotlib.ticker import MaxNLocator
    ax.yaxis.set_major_locator(MaxNLocator(integer=True))      # mails are whole numbers
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.grid(axis="y", color="#E3E8EF", linewidth=0.9, zorder=0)
    ax.set_axisbelow(True)
    ax.tick_params(length=0)


def _png(fig):
    buffer = io.BytesIO()
    fig.savefig(buffer, format="png", dpi=110, bbox_inches="tight")
    return buffer.getvalue()


def _chart_label(text):
    return _cut(text, 24).replace("$", r"\$")           # "$...$" in a name would be read as math text


def render_charts(stats, today):
    """``{content_id: PNG bytes}`` for the status and team charts, drawn in memory with the Agg backend (no window, no
    pyplot state). Raises ImportError when matplotlib is not installed - the caller then uses the table fallback."""
    import matplotlib
    try:
        matplotlib.use("Agg")
    except Exception:
        pass
    from matplotlib.figure import Figure
    style = {"font.family": "sans-serif", "font.sans-serif": ["Segoe UI", "Calibri", "Arial", "DejaVu Sans"],
             "axes.edgecolor": "#E3E8EF", "axes.labelcolor": INK, "text.color": INK, "xtick.color": INK,
             "ytick.color": INK, "figure.facecolor": "white", "axes.facecolor": "white"}
    charts = {}
    with matplotlib.rc_context(style):
        fig = Figure(figsize=(6, 4))
        ax = fig.subplots()
        values = [_int(stats["replied"]), _int(stats["pending"]), _int(stats["no_action_needed"])]
        bars = ax.bar(["Replied", "Pending", "No Action\nNeeded"], values, color=[GREEN, RED, AMBER], width=0.55,
                      zorder=3)
        for bar, value in zip(bars, values):
            ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height(), f"{value}", ha="center", va="bottom",
                    fontsize=12, fontweight="bold", color=NBE_GREEN)
        ax.set_title(f"Reply Status  \u2022  {_day_text(today)}", fontsize=14, fontweight="bold", color=NBE_GREEN, pad=14)
        ax.set_ylabel("Emails", fontsize=10)
        _style_axes(ax)
        fig.tight_layout()
        charts[STATUS_CHART_CID] = _png(fig)

        fig = Figure(figsize=(8, 4))
        ax = fig.subplots()
        members = _counts(stats["first_reply_by"])[:CHART_MEMBERS]
        if members:
            positions = list(range(len(members)))              # positions, not names: two people can share a label
            bars = ax.bar(positions, [c for _, c in members], color=BLUE, width=0.6, zorder=3)
            for bar, (_, value) in zip(bars, members):
                ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height(), f"{value}", ha="center", va="bottom",
                        fontsize=10, fontweight="bold", color=NBE_GREEN)
            ax.set_xticks(positions)
            ax.set_xticklabels([_chart_label(m) for m, _ in members], rotation=35, ha="right")
        else:
            ax.text(0.5, 0.5, "No Team Replies Today", ha="center", va="center", fontsize=12, color=GREY,
                    transform=ax.transAxes)
            ax.set_xticks([])
            ax.set_yticks([])
        ax.set_title("First Replies by Team Member", fontsize=14, fontweight="bold", color=NBE_GREEN, pad=14)
        ax.set_ylabel("First Replies", fontsize=10)
        _style_axes(ax)
        fig.tight_layout()
        charts[TEAM_CHART_CID] = _png(fig)
    return charts


# ---------------------------------------------------------------------------------------------------------------------
# HTML building blocks (tables + inline styles only: this is what Outlook's Word engine renders reliably)
# ---------------------------------------------------------------------------------------------------------------------
def logo_available():
    return LOGO_PATH.is_file()


def mail_images(extra=None):
    """The ``inline_images`` of a mail: the logo (when the file exists) plus whatever else the body references."""
    images = {LOGO_CID: str(LOGO_PATH)} if logo_available() else {}
    images.update(extra or {})
    return images


def _logo_html(image=None):
    if logo_available() if image is None else image:
        return ('<table cellpadding="0" cellspacing="0" role="presentation" border="0"><tr>'
                '<td width="54" height="54" align="center" valign="middle" bgcolor="#ffffff" '
                'style="background:#ffffff;"><img src="cid:' + LOGO_CID + '" alt="NBE" width="42" height="42" '
                'style="display:block;border:0;"></td></tr></table>')
    return ('<table cellpadding="0" cellspacing="0" role="presentation" border="0"><tr>'
            f'<td width="54" height="54" align="center" valign="middle" bgcolor="{NBE_GOLD}" '
            f'style="background:{NBE_GOLD};color:{NBE_GREEN};font-weight:bold;font-size:15px;">NBE</td></tr></table>')


def _page(title, subtitle_html, inner, width=780):
    return (
        '<html><head><meta charset="UTF-8"></head>'
        f'<body style="margin:0;padding:0;background:{PAGE_BG};font-family:{FONT_STACK};color:{INK};">'
        f'<table width="100%" cellpadding="0" cellspacing="0" border="0" bgcolor="{PAGE_BG}" '
        f'style="background:{PAGE_BG};"><tr><td align="center" style="padding:20px 8px;">'
        f'<table width="{width}" cellpadding="0" cellspacing="0" border="0" bgcolor="#ffffff" '
        f'style="background:#ffffff;width:{width}px;max-width:100%;border:1px solid #dfe5ec;">'
        f'<tr><td bgcolor="{NBE_GREEN}" style="background:{NBE_GREEN};padding:22px 28px;border-bottom:4px solid {NBE_GOLD};">'
        '<table width="100%" cellpadding="0" cellspacing="0" border="0" role="presentation"><tr>'
        f'<td width="58" valign="middle">{_logo_html()}</td>'
        '<td valign="middle" style="padding-left:14px;">'
        f'<div style="color:#ffffff;font-size:23px;font-weight:700;font-family:{FONT_STACK};">{esc(title)}</div>'
        f'<div style="color:{NBE_GREEN_LIGHT};margin-top:5px;font-size:12.5px;font-family:{FONT_STACK};">{subtitle_html}</div>'
        '</td></tr></table></td></tr>'
        f'<tr><td style="padding:24px 28px;font-family:{FONT_STACK};">{inner}</td></tr>'
        '</table></td></tr></table></body></html>')


def without_inline_logo(html_body):
    """The mail with the NBE text badge instead of the inline logo image: a mail that goes without the logo attached."""
    return html_body.replace(_logo_html(True), _logo_html(False))


def _subtitle(day):
    return f"National Bank of Egypt &nbsp;&bull;&nbsp; {esc(_day_text(day))}"


def _heading(title, color=NBE_GREEN):
    return (f'<div style="font-size:15px;font-weight:700;color:{color};margin:26px 0 10px 0;'
            f'font-family:{FONT_STACK};">{esc(title)}</div>')


def _footer(text):
    return (f'<div style="margin-top:26px;padding-top:14px;border-top:1px solid {LINE};font-size:11.5px;color:{GREY};'
            f'font-family:{FONT_STACK};">{esc(text)}</div>')


def _note(inner_html, color=BLUE, background=CARD_BG):
    return (f'<table width="100%" cellpadding="0" cellspacing="0" border="0" style="margin-top:16px;"><tr>'
            f'<td bgcolor="{background}" style="background:{background};border-left:4px solid {color};padding:13px 16px;'
            f'font-size:12.5px;line-height:1.55;color:{INK};font-family:{FONT_STACK};">{inner_html}</td></tr></table>')


def _td(content, bold=False, color=None, align="left", size=12.5, raw=False):
    style = (f"padding:9px 12px;border-bottom:1px solid {LINE};font-size:{size}px;color:{color or INK};"
             f"font-family:{FONT_STACK};{'font-weight:700;' if bold else ''}")
    return f'<td align="{align}" valign="top" dir="auto" style="{style}">{content if raw else esc(content)}</td>'


def _tr(*cells):
    return "<tr>" + "".join(cells) + "</tr>"


def _table(headers, rows, empty_text="Nothing to show.", header_color=NBE_GREEN):
    head = "".join(
        f'<th align="left" bgcolor="{header_color}" style="background:{header_color};color:#ffffff;padding:10px 12px;'
        f'font-size:12px;text-align:left;font-family:{FONT_STACK};">{esc(h)}</th>' for h in headers)
    body = "".join(rows) or _tr(f'<td colspan="{len(headers)}" align="center" style="padding:12px;color:{GREY};'
                                f'font-size:12.5px;font-family:{FONT_STACK};">{esc(empty_text)}</td>')
    return (f'<table width="100%" cellpadding="0" cellspacing="0" border="0" style="border-collapse:collapse;">'
            f'<tr>{head}</tr>{body}</table>')


def _kpi(label, value, accent=BLUE, trend=""):
    # the accent border sits on the cell: Outlook's Word engine ignores borders on a <table>
    return (f'<td width="25%" valign="top" style="padding:5px;">'
            f'<table width="100%" cellpadding="0" cellspacing="0" border="0" bgcolor="{CARD_BG}" '
            f'style="background:{CARD_BG};"><tr>'
            f'<td align="center" bgcolor="{CARD_BG}" style="background:{CARD_BG};border-top:3px solid {accent};'
            f'padding:13px 6px;font-family:{FONT_STACK};">'
            f'<div style="font-size:11px;color:{MUTED};text-transform:uppercase;letter-spacing:0.4px;">{esc(label)}</div>'
            f'<div style="font-size:23px;font-weight:700;color:{NBE_GREEN};margin-top:4px;">{esc(value)}</div>'
            f'{trend}</td></tr></table></td>')


def _kpi_grid(cards, per_row=4):
    rows = []
    for start in range(0, len(cards), per_row):
        group = cards[start:start + per_row]
        group += ['<td width="25%" style="padding:5px;">&nbsp;</td>'] * (per_row - len(group))
        rows.append("<tr>" + "".join(group) + "</tr>")
    return ('<table width="100%" cellpadding="0" cellspacing="0" border="0" role="presentation">'
            + "".join(rows) + "</table>")


def _trend_badge(comparison, label="vs yesterday", invert=False, kind="count"):
    """"▲ 3 vs yesterday". ``invert``: a LOWER number is the good direction (pending, SLA breaches, waiting time)."""
    if not isinstance(comparison, dict) or comparison.get("direction") not in ("up", "down", "flat"):
        return ""
    direction, delta = comparison["direction"], abs(_float(comparison.get("delta")))
    if direction == "flat":
        color, arrow = MUTED, "&#8594;"
    else:
        color = GREEN if (direction == "up") != invert else RED
        arrow = "&#9650;" if direction == "up" else "&#9660;"
    amount = f"{delta:.1f} pt" if kind == "percent" else _duration(delta) if kind == "time" else f"{delta:.0f}"
    return f'<div style="font-size:11px;color:{color};margin-top:3px;">{arrow} {esc(amount)} {esc(label)}</div>'


def _bar(value, maximum, color):
    """A horizontal bar made of one coloured table cell (no CSS widths on divs, which Outlook ignores)."""
    percent = 0 if maximum <= 0 or value <= 0 else min(100, max(2, round(value / maximum * 100)))
    cell = f'bgcolor="{color}" style="background:{color};font-size:1px;line-height:12px;height:12px;"'
    if percent == 0:
        inner = '<td style="font-size:1px;line-height:12px;height:12px;">&nbsp;</td>'
    elif percent == 100:
        inner = f"<td {cell}>&nbsp;</td>"
    else:
        inner = f'<td width="{percent}%" {cell}>&nbsp;</td><td>&nbsp;</td>'
    return f'<table width="100%" cellpadding="0" cellspacing="0" border="0" role="presentation"><tr>{inner}</tr></table>'


def _bar_table(items, empty_text):
    """``items``: [(label, value, color)] drawn as a labelled bar list."""
    maximum = max((v for _, v, _ in items), default=0)
    rows = [_tr(_td(label), f'<td width="55%" style="padding:9px 12px;border-bottom:1px solid {LINE};">'
                            f'{_bar(value, maximum, color)}</td>', _td(value, bold=True, align="right"))
            for label, value, color in items]
    empty = _tr(f'<td align="center" style="padding:12px;color:{GREY};font-size:12.5px;font-family:{FONT_STACK};">'
                f'{esc(empty_text)}</td>')
    return ('<table width="100%" cellpadding="0" cellspacing="0" border="0" style="border-collapse:collapse;">'
            + ("".join(rows) or empty) + "</table>")


# ---------------------------------------------------------------------------------------------------------------------
# The main report e-mail
# ---------------------------------------------------------------------------------------------------------------------
def _kpi_cards(stats):
    trend = (stats.get("trend") or {}).get("vs_yesterday") or {}

    def badge(metric, **kw):
        return _trend_badge(trend.get(metric), **kw)
    target = _int(stats.get("sla_target_minutes"))
    return [
        _kpi("Business Emails", _int(stats["business_emails"]), BLUE, badge("business_emails")),
        _kpi("Replied", _int(stats["replied"]), GREEN, badge("replied")),
        _kpi("Pending", _int(stats["pending"]), RED, badge("pending", invert=True)),
        _kpi("No Action Needed", _int(stats["no_action_needed"]), AMBER, badge("no_action_needed")),
        _kpi("Reply Rate", f"{_float(stats['reply_rate']):.1f}%", BLUE, badge("reply_rate", kind="percent")),
        _kpi("Team Replies", _int(stats["team_replies"]), BLUE),
        _kpi("Total Team Replies", _int(stats["total_team_replies"]), BLUE),
        _kpi("Avg First Response", _duration(stats["average_response"]), BLUE,
             badge("average_response", invert=True, kind="time")),
        _kpi("Fastest Response", _duration(stats["fastest_response"]), GREEN),
        _kpi(f"SLA Breached (>{target}m)", _int(stats["sla_breached"]), RED, badge("sla_breached", invert=True)),
        _kpi("Escalated (Pending)", _int(stats["escalated_count"]), RED),
        _kpi("SLA Compliance", f"{_float(stats.get('sla_compliance_rate')):.1f}%", GREEN),
    ]


def _team_rows(stats):
    teams = [(t, d) for t, d in (stats.get("team_stats") or {}).items()
             if d["business_emails"] or d["pending"] or d.get("carryover_count")]
    return sorted(teams, key=lambda kv: (kv[0] == UNCLASSIFIED, -kv[1]["business_emails"], kv[0]))


def _team_section(stats):
    rows = []
    for team, d in _team_rows(stats):
        average = _duration(d["response_minutes_sum"] / d["response_count"]) if d["response_count"] else "-"
        carry = f"{d['carryover_count']} ({d['carryover_oldest_days']} d)" if d["carryover_count"] else "0"
        rows.append(_tr(_td(team, bold=True), _td(d["business_emails"]), _td(d["replied"], bold=True, color=GREEN),
                        _td(d["pending"], bold=True, color=RED), _td(d["no_action_needed"], color=AMBER, bold=True),
                        _td(d["escalated_count"], color=RED), _td(d["sla_breached_count"], color=RED),
                        _td(average), _td(carry)))
    return (_heading("Performance by Team")
            + _table(["Team", "Emails", "Replied", "Pending", "No Action", "Escalated", "SLA breached", "Avg reply",
                      "Carry-over"], rows, "No business e-mails today."))


def _systems_section(stats):
    items = _counts(stats.get("system_volume"))[:10]
    maximum = max((v for _, v in items), default=0)
    rows = [_tr(_td(name), _td(count, bold=True),
                f'<td width="45%" style="padding:9px 12px;border-bottom:1px solid {LINE};">{_bar(count, maximum, BLUE)}</td>')
            for name, count in items]
    return _heading("Systems") + _table(["System", "Emails", ""], rows, "No system was detected today.")


def _sla_section(stats):
    target = _int(stats.get("sla_target_minutes"))
    carry = _int(stats.get("carryover_count"))
    lines = [
        (f"SLA target (first reply within {target} min)", f"{_float(stats.get('sla_compliance_rate')):.1f}% met", GREEN),
        ("Replies slower than the target / pending longer than the target", _int(stats["sla_breached"]), RED),
        ("Escalated - still unanswered past the target", _int(stats["escalated_count"]), RED),
        ("Carried over from earlier days, still unanswered",
         f"{carry} (oldest {_int(stats.get('carryover_oldest_days'))} d)" if carry else 0, AMBER),
    ]
    rows = [_tr(_td(label), _td(value, bold=True, color=color)) for label, value, color in lines]
    return _heading("SLA") + _table(["Measure", "Value"], rows)


def _matching_section(stats):
    lines = [("ConversationID", stats["conversation_matches"]),
             ("ConversationIndex", stats.get("conversation_index_matches", 0)),
             ("Subject Fallback", stats["subject_matches"]), ("No Match", stats["no_matches"])]
    return (_heading("Reply Matching Quality")
            + _table(["Method", "Count"], [_tr(_td(name), _td(_int(count))) for name, count in lines]))


def _charts_section(stats, include_charts):
    if include_charts:
        return (_heading("Charts")
                + '<table width="100%" cellpadding="0" cellspacing="0" border="0"><tr>'
                + f'<td width="50%" align="center" style="padding:6px;"><img src="cid:{STATUS_CHART_CID}" width="330" '
                  'alt="Reply status chart" style="display:block;border:0;max-width:100%;"></td>'
                + f'<td width="50%" align="center" style="padding:6px;"><img src="cid:{TEAM_CHART_CID}" width="330" '
                  'alt="First replies chart" style="display:block;border:0;max-width:100%;"></td></tr></table>')
    status = [("Replied", _int(stats["replied"]), GREEN), ("Pending", _int(stats["pending"]), RED),
              ("No Action Needed", _int(stats["no_action_needed"]), AMBER)]
    members = [(name, count, BLUE) for name, count in _counts(stats["first_reply_by"])[:CHART_MEMBERS]]
    return (_heading("Reply Status") + _bar_table(status, "No e-mails today.")
            + _heading("First Replies by Team Member") + _bar_table(members, "No team replies today."))


def _escalated_section(stats):
    items = stats.get("escalated_items") or []
    if not items:
        return ""
    rows = [_tr(_td(_cut(i["subject"], 70)), _td(_cut(i["customer_name"], 40)), _td(i["team"]),
                _td(i["current_waiting_time"], bold=True, color=RED)) for i in items]
    return (_heading(f"Escalated - SLA breached (>{_int(stats.get('sla_target_minutes'))}m), oldest first", RED)
            + _table(["Subject", "Customer", "Team", "Waiting"], rows, header_color=RED))


def _slowest_section(stats):
    items = stats.get("top_slowest_replies") or []
    if not items:
        return ""
    rows = [_tr(_td(_cut(i["subject"], 70)), _td(_cut(i["customer_name"], 40)), _td(i["team"]),
                _td(i.get("first_reply_by") or "-"), _td(_duration(i["response_minutes"]), bold=True, color=RED))
            for i in items]
    return _heading("Top Slowest First Replies") + _table(["Subject", "Customer", "Team", "Replied by", "Took"], rows)


def _leaderboard_rows(stats):
    rows = []
    for tag, color, entries in (("Top", GREEN, stats.get("quality_leaderboard_top")),
                                ("Needs attention", BROWN, stats.get("quality_leaderboard_bottom"))):
        for e in entries or []:
            rows.append(_tr(_td(e["name"], bold=True), _td(e["team"]), _td(e["reviewed"]),
                            _td(f"{_float(e['score']):.0f}", bold=True, color=color), _td(tag, color=color)))
    return rows


def _quality_section(stats, sample):
    reviewed, repliable = _int(stats.get("quality_reviewed")), _int(stats.get("repliable_count"))
    cards = [_kpi("Replies Reviewed", f"{reviewed} / {repliable}", PURPLE),
             _kpi("Review Coverage", f"{_float(stats.get('quality_coverage_rate')):.0f}%", BLUE),
             _kpi("Excellent / Very Good / Good", f"{_int(stats.get('quality_excellent'))} / "
                  f"{_int(stats.get('quality_very_good'))} / {_int(stats.get('quality_good'))}", GREEN),
             _kpi("Needs Improvement", _int(stats.get("quality_needs_improvement")), RED)]
    html_out = _heading("Reply Quality") + _kpi_grid(cards)
    if _int(stats.get("quality_appeals")):
        html_out += _note(f"<b>{_int(stats['quality_appeals'])}</b> rating(s) have an appeal from the person who replied.",
                          AMBER, "#FFF8E8")
    index = [_tr(_td(t["team"], bold=True), _td(f"{_float(t['index']):.0f}", bold=True), _td(t["reviewed"]))
             for t in stats.get("team_quality_index") or []]
    if index:
        html_out += _heading("Team Quality Index") + _table(["Team", "Index /100", "Reviewed"], index)
    board = _leaderboard_rows(stats)
    if board:
        html_out += _heading("Quality Leaderboard") + _table(["Team member", "Team", "Reviewed", "Score", ""], board)
    alerts = [_tr(_td(a["name"], bold=True), _td(a["team"]), _td(f"{_int(a['needs'])} flags", bold=True, color=BROWN))
              for a in stats.get("repeat_issue_alerts") or []]
    if alerts:
        html_out += (_heading("Repeated Quality Issues", BROWN)
                     + _table(["Team member", "Team", "Recent issues"], alerts, header_color=BROWN))
    if sample:
        rows = [_tr(_td(_cut(r["subject"], 60)), _td(r["first_reply_by"]), _td(r["team"]),
                    _td(_duration(r["response_minutes"]) if r.get("response_minutes") is not None else "-"))
                for r in sample]
        html_out += (_heading(f"Suggested Review Sample ({len(sample)})")
                     + _table(["Subject", "Replied by", "Team", "First reply"], rows))
    return html_out


def build_email_html(stats, today, include_charts=False, pending_attached=0, pending_team_emails=0, review_sample_rows=(),
                     excel_attached=True):
    """The main report. ``include_charts``: the body points at the inline images ``cid:reply_status_chart`` /
    ``cid:team_chart`` (the caller attaches them); otherwise the same numbers are drawn as table bars.
    ``excel_attached``: the footer only claims the workbook is attached when it really is."""
    body = _kpi_grid(_kpi_cards(stats))
    body += (f'<div style="font-size:11px;color:{GREY};margin-top:6px;">'
             'Reply Rate = Replied &divide; (Business Emails &minus; No Action Needed)</div>')
    body += _note("<b>Business filter:</b> an e-mail is included only when one of the configured business addresses is in "
                  "its To line (CC does not count). <b>Reply rule:</b> a reply counts only when its sender is a team member.")
    body += _team_section(stats) + _systems_section(stats) + _sla_section(stats) + _matching_section(stats)
    body += _charts_section(stats, include_charts) + _escalated_section(stats) + _slowest_section(stats)
    body += _quality_section(stats, review_sample_rows)
    if _int(pending_attached):
        body += _note(f"<b>{_int(pending_attached)}</b> unanswered e-mail(s) are attached as zipped .msg files, sent in "
                      f"<b>{_int(pending_team_emails)}</b> follow-up e-mail(s) - one per team - so this e-mail carries the "
                      "summary, the charts and the Excel workbook.")
    body += _footer("Automated Outlook current-day Team Reply report."
                    + (" The detailed Excel workbook is attached to this message." if excel_attached else ""))
    return _page("Daily Team Reply Analytics", _subtitle(today) + " &nbsp;&bull;&nbsp; Current Day Analysis", body)


# ---------------------------------------------------------------------------------------------------------------------
# Team follow-up, assignment, quality feedback and end-of-day e-mails
# ---------------------------------------------------------------------------------------------------------------------
def build_team_pending_email_html(today, team, count_in_zip, part_index, total_parts, rows=None):
    """The follow-up for one team: how many .msg files are in the zip and, when ``rows`` are given, what is waiting."""
    part_text = f" (Part {part_index} of {total_parts})" if total_parts > 1 else ""
    rows = list(rows or [])
    if count_in_zip:
        text = (f"<b>{_int(count_in_zip)}</b> unanswered e-mail(s) for <b>{esc(team)}</b> are attached as a single .zip "
                "file (original .msg format inside) - ready to extract and action directly.")
    else:
        text = f"<b>{len(rows)}</b> e-mail(s) are still waiting for a reply from <b>{esc(team)}</b>."
    body = _note(text)
    if rows:
        shown = rows[:MAX_LISTED_ROWS]
        table = [_tr(_td(_cut(r["subject"], 70)), _td(_cut(r["customer_name"], 40)),
                     _td(r["received_time"].strftime("%d/%m %H:%M") if r.get("received_time") else "-"),
                     _td(r.get("current_waiting_time") or "-", bold=True, color=RED),
                     _td(_cut(r.get("detected_systems") or "-", 40)),
                     _td(r.get("assigned_to_name") or r.get("assigned_to_email") or "-")) for r in shown]
        body += _heading("Waiting for a reply") + _table(
            ["Subject", "Customer", "Received", "Waiting", "System(s)", "Assigned to"], table)
        if len(rows) > len(shown):
            body += _note(f"... and {len(rows) - len(shown)} more. The Excel workbook of the main report lists all of them.")
    body += _footer("Automated Outlook current-day Team Reply report. See the main report e-mail for the full KPI summary, "
                    "charts and the Excel workbook.")
    return _page(f"Pending Emails - {team}{part_text}", esc(_day_text(today)), body, width=820)


def build_assignment_email_html(person_name, items, today):
    """``items``: [(assignment, row)] - the assignment record and the analysed mail (``{}`` when it is not in today's rows)."""
    notes = list(dict.fromkeys(str(a.get("note") or "").strip() for a, _ in items if str(a.get("note") or "").strip()))
    names = list(dict.fromkeys(str(a.get("marked_by") or "").strip() for a, _ in items if str(a.get("marked_by") or "").strip()))
    who = ", ".join(names) or "the manager"
    table = []
    for assignment, row in items:
        table.append(_tr(_td(_cut(assignment.get("subject") or row.get("subject") or "(no subject)", 80)),
                         _td(_cut(row.get("customer_name") or "-", 40)), _td(row.get("team") or "-"),
                         _td(_cut(row.get("detected_systems") or "-", 40)), _td(row.get("status") or "-"),
                         _td(row.get("current_waiting_time") or "-", bold=True, color=RED)))
    body = (f'<div style="font-size:14px;">Dear {esc(person_name)},</div>'
            f'<div style="font-size:13.5px;margin-top:10px;line-height:1.6;"><b>{len(items)}</b> email(s) have been '
            f'assigned to you by <b>{esc(who)}</b> and are awaiting your action.</div>')
    if notes:
        note_html = "<br>".join(esc(n) for n in notes)
        body += _note(f'<div style="font-size:12px;font-weight:700;color:#8A5A00;margin-bottom:4px;">MANAGER\'S NOTE</div>'
                      f'<div style="font-size:13.5px;">{note_html}</div>', AMBER, "#FFF8E8")
    body += '<div style="height:12px;line-height:12px;">&nbsp;</div>'
    body += _table(["Subject", "Customer", "Team", "System(s)", "Status", "Waiting"], table)
    body += _footer("Automated message from Team Reply Daily Analytics. Please action these as soon as possible.")
    return _page("Emails Assigned to You", _subtitle(today), body)


def build_quality_feedback_html(person_name, items, today):
    """``items``: [(message_id, review)] - every review of this person that was not e-mailed yet."""
    positives = sum(1 for _, r in items if r.get("rating") in ("Excellent", "Very Good", "Good"))
    reviewers = sorted({str(r.get("reviewed_by") or "").strip() for _, r in items
                        if str(r.get("reviewed_by") or "").strip() and str(r.get("reviewed_by")).strip().lower() != "dashboard"})
    tone = ("Thank you for the quality of your replies - keep it up." if positives == len(items)
            else "Please review the points below so we can keep improving response quality.")
    table = [_tr(_td(_cut(r.get("subject") or "", 70)), _td(r.get("rating") or "", bold=True,
                                                          color=RATING_COLOURS.get(r.get("rating"), INK)),
                 _td(r.get("reasons") or "-"), _td(_cut(r.get("note") or "-", 300))) for _, r in items]
    body = (f'<div style="font-size:14px;">Dear {esc(person_name)},</div>'
            f'<div style="font-size:13.5px;margin-top:10px;line-height:1.6;"><b>{esc(", ".join(reviewers) or "your manager")}'
            f'</b> reviewed <b>{len(items)}</b> of your replies. {esc(tone)}</div><div style="height:12px;line-height:12px;">&nbsp;</div>')
    body += _table(["Subject", "Rating", "Reasons", "Reviewer note"], table)
    body += _note("<b>Disagree with a rating?</b> Open the Email page of the Operations Hub, find the e-mail and use the "
                  "<b>Appeal</b> option next to the rating. Your comment goes straight back to the reviewing manager.")
    body += _footer("Automated feedback from Team Reply Daily Analytics.")
    return _page("Reply Quality Feedback", _subtitle(today), body, width=820)


def _trend_arrow(direction):
    return {"improving": ("&#9650;", GREEN), "declining": ("&#9660;", RED)}.get(direction, ("&#8594;", MUTED))


def build_end_of_day_html(rows, stats, today, excel_attached=True):
    """The wrap-up of the day. Built from the same rows and stats as the other mails, so they can never disagree."""
    team_table = [_tr(_td(t, bold=True), _td(d["business_emails"]), _td(d["replied"], bold=True, color=GREEN),
                      _td(d["pending"], bold=True, color=RED), _td(d["no_action_needed"], bold=True, color=AMBER))
                  for t, d in _team_rows(stats) if d["business_emails"]]
    replied = [_tr(_td(rb["name"]), _td(rb["team"]), _td(rb["count"], bold=True)) for rb in stats.get("reply_breakdown") or []]
    assigned = [_tr(_td(_cut(r["subject"], 70)), _td(r.get("assigned_to_name") or r.get("assigned_to_email"), bold=True),
                    _td(r.get("team") or ""), _td(r.get("status") or "")) for r in rows if r.get("assigned_to_email")]
    reviews = [_tr(_td(_cut(r["subject"], 60)), _td(r.get("first_reply_by") or "-"),
                   _td(r["quality_rating"], bold=True, color=RATING_COLOURS.get(r["quality_rating"], INK)),
                   _td(_cut(r.get("quality_note") or "-", 200))) for r in rows if r.get("quality_rating")]
    directions = {t["name"]: t["direction"] for t in stats.get("quality_trends") or []}
    scorecard = []
    for sc in stats.get("quality_scorecard") or []:
        arrow, color = _trend_arrow(directions.get(sc["name"]))
        scorecard.append(_tr(_td(sc["name"], bold=True), _td(sc["team"]), _td(sc["reviewed"]),
                             _td(sc["excellent"], bold=True, color=GREEN),
                             _td(sc.get("very_good", 0), bold=True, color=VERY_GOOD), _td(sc["good"], bold=True, color=BLUE),
                             _td(sc["needs"], bold=True, color=BROWN),
                             _td(f'{_float(sc["score"]):.0f} <span style="color:{color};">{arrow}</span>', bold=True, raw=True)))
    index = [_tr(_td(t["team"], bold=True), _td(f"{_float(t['index']):.0f}", bold=True), _td(t["reviewed"]),
                 _td(t["excellent"], color=GREEN), _td(t.get("very_good", 0), color=VERY_GOOD), _td(t["good"], color=BLUE),
                 _td(t["needs"], color=BROWN))
             for t in stats.get("team_quality_index") or []]
    alerts = [_tr(_td(a["name"], bold=True), _td(a["team"]), _td(f"{_int(a['needs'])} flags", bold=True, color=BROWN))
              for a in stats.get("repeat_issue_alerts") or []]

    cards = [
        _kpi("Business Emails", _int(stats["business_emails"])), _kpi("Replied", _int(stats["replied"]), GREEN),
        _kpi("Pending", _int(stats["pending"]), RED), _kpi("No Action", _int(stats["no_action_needed"]), AMBER),
        _kpi("Reply Rate", f"{_float(stats['reply_rate']):.1f}%"), _kpi("SLA Breached", _int(stats["sla_breached"]), RED),
        _kpi("Avg Response", _duration(stats["average_response"])), _kpi("Assigned", _int(stats.get("assigned_count")), PURPLE),
        _kpi("Replies Reviewed", f"{_int(stats.get('quality_reviewed'))} / {_int(stats.get('repliable_count'))}", PURPLE),
        _kpi("Review Coverage", f"{_float(stats.get('quality_coverage_rate')):.0f}%"),
        _kpi("Excellent", _int(stats.get("quality_excellent")), GREEN), _kpi("Appeals", _int(stats.get("quality_appeals")), RED)]
    body = _kpi_grid(cards)
    body += _heading("Performance by Team") + _table(["Team", "Emails", "Replied", "Pending", "No Action"], team_table,
                                                     "No business e-mails today.")
    body += _heading("Who Replied Today") + _table(["Team member", "Team", "First replies"], replied, "No replies recorded today.")
    body += _heading("Assignments Made") + _table(["Subject", "Assigned to", "Team", "Status"], assigned, "No assignments made today.")
    body += _heading("Reply Quality Reviews") + _table(["Subject", "Replied by", "Rating", "Note"], reviews,
                                                       "No quality reviews recorded today.")
    body += _heading("Quality Scorecard by Person") + _table(
        ["Team member", "Team", "Reviewed", "Excellent", "Very good", "Good", "Needs improvement", "Score"], scorecard,
        "No quality reviews recorded today.")
    body += _heading("Team Quality Index") + _table(["Team", "Index /100", "Reviewed", "Excellent", "Very good", "Good",
                                                     "Needs improvement"],
                                                    index, "No reviews recorded today.")
    if alerts:
        body += _heading("Repeated Quality Issues", BROWN) + _table(["Team member", "Team", "Recent issues"], alerts,
                                                                    header_color=BROWN)
    body += _footer("Automated end-of-day wrap-up from Team Reply Daily Analytics."
                    + (" The detailed Excel workbook is attached." if excel_attached else ""))
    return _page("End of Day Summary", _subtitle(today), body, width=860)


# ---------------------------------------------------------------------------------------------------------------------
# Pending mails as .msg files (Outlook COM: item.SaveAs) and zips
# ---------------------------------------------------------------------------------------------------------------------
def _can_save(item):
    return item is not None and callable(getattr(item, "SaveAs", None))


def save_pending_msgs(rows, folder, log, fetch=None):
    """Save every unanswered mail as ``.msg`` under ``folder/<team>/``. Each mail is fetched again (``fetch(row)``:
    Outlook's GetItemFromID by its EntryID) only while it is saved, then let go: a live Outlook item kept for every mail
    of the day runs Outlook out of memory («Out of memory or system resources») for every other mail sent meanwhile.

    Returns ``({id(row): path}, failed_count)``. Rows without a usable item (no EntryID, a fake Outlook, a machine
    without Outlook) are skipped quietly with one log line - the follow-up e-mails then carry the list without a zip. A
    mail addressed to two teams is saved once and listed for both.
    """
    wanted = [r for r in rows if r.get("teams_pending") and r.get("status") != "No Action Needed"]
    files, saved, failed, number = {}, {}, 0, 0
    with OUTLOOK_LOCK:                                 # COM objects belong to the one Outlook conversation
        for row in wanted:
            key = row.get("message_id") or row.get("entry_id") or id(row)
            if key in saved:
                files[id(row)] = saved[key]
                continue
            item = row.get("mail_item")
            try:
                if item is None and fetch is not None and row.get("entry_id"):
                    item = fetch(row)
                if not _can_save(item):
                    continue
                number += 1
                team_folder = Path(folder) / safe_name(row.get("team") or UNCLASSIFIED, 30, UNCLASSIFIED)
                team_folder.mkdir(parents=True, exist_ok=True)
                name = f"{number:03d}_{safe_name(row.get('customer_name'), 25)}_{safe_name(row.get('subject'), 45)}.msg"
                path = team_folder / name
                item.SaveAs(str(path), OL_MSG)
                if path.is_file() and path.stat().st_size > 0:
                    saved[key] = files[id(row)] = path
                else:
                    failed += 1
            except Exception as exc:
                failed += 1
                log(f"Could not save a pending e-mail as .msg: {type(exc).__name__}: {exc}")
            finally:
                item = None                            # release the COM object before the next mail
    if wanted and not files and not failed:
        log("No Outlook items are available to save as .msg files - the follow-up e-mails carry the list only.")
    return files, failed


def zip_msg_files(paths, zip_path):
    """Compress the files into one zip. Returns ``zip_path``, or None when nothing could be zipped."""
    paths = list(dict.fromkeys(Path(p) for p in paths))
    if not paths:
        return None
    zip_path = Path(zip_path)
    zip_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for path in paths:
                try:
                    archive.write(path, arcname=path.name)
                except OSError:
                    continue
    except OSError:
        return None
    return zip_path if zip_path.is_file() and zip_path.stat().st_size > 0 else None


def build_team_parts(rows, files, max_per_part, zip_folder):
    """``{team: [{"rows": [...], "zip": path or None, "count": files in the zip}]}`` - one part per follow-up e-mail.

    A team's pending mails are cut into parts of at most ``max_per_part`` (the old MAX_PENDING_ATTACHMENTS), so nothing
    is ever dropped; each part gets its own zip when .msg files were saved for it.
    """
    parts_by_team, used = {}, set()
    for team, team_rows in pending_rows_by_team(rows).items():
        chunks = chunk_list(team_rows, max_per_part)
        parts = []
        for number, chunk in enumerate(chunks, 1):
            paths = [files[id(r)] for r in chunk if id(r) in files]
            zip_path = None
            if paths:
                stem = file_stem(team, 30, UNCLASSIFIED) + "_Pending" + (f"_Part{number}of{len(chunks)}" if len(chunks) > 1 else "")
                suffix, name = 1, f"{stem}.zip"
                while name.lower() in used:
                    suffix += 1
                    name = f"{stem}_{suffix}.zip"
                used.add(name.lower())
                zip_path = zip_msg_files(paths, Path(zip_folder) / name)
            parts.append({"rows": chunk, "zip": zip_path, "count": len(set(paths)) if zip_path else 0})
        parts_by_team[team] = parts
    return parts_by_team


# ---------------------------------------------------------------------------------------------------------------------
# State: what this day already sent
# ---------------------------------------------------------------------------------------------------------------------
def _load_state(store, day):
    try:
        doc = copy.deepcopy(store.read_state(STATE_NAME) or {})
    except Exception:
        doc = {}
    return doc if doc.get("day") == day and isinstance(doc.get("slots"), dict) else {"day": day, "slots": {}}


def slot_name(cfg, now, options=None):
    """The report slot of this run: an explicit ``options["slot"]``, else the latest ``send_times`` entry that has
    passed, else "manual" (a run before the first slot of the day, or with no times configured)."""
    explicit = str((options or {}).get("slot") or "").strip()
    if explicit:
        return explicit
    passed = [t for t in cfg.send_times if t <= f"{now:%H:%M}"]
    return passed[-1] if passed else "manual"


def part_id(team, number, total):
    """The state key of one follow-up e-mail of a team (a retry skips the parts that already went out)."""
    return f"team:{team}:{number}of{total}"


def _unsent_reviews(store):
    """[(message_id, review)] of ratings that were not e-mailed to the person who replied yet."""
    sent = store.read_state(QUALITY_SENT_STATE) or {}
    out = []
    for message_id, review in (store.overrides("quality") or {}).items():
        if not isinstance(review, dict) or not review.get("rating") or review.get("rating_withdrawn"):
            continue
        if review.get("feedback_sent"):                 # the old script's flag, kept for imported data
            continue
        done = sent.get(message_id) or {}
        if done.get("reviewed_at") == review.get("reviewed_at") and done.get("rating") == review.get("rating"):
            continue
        out.append((message_id, review))
    return out


def _notifications_waiting(ctx, cfg):
    """Is there anything the assignment / quality steps would still send? (Decided without Outlook, from the lists.)"""
    try:
        if cfg.assign_send and any(isinstance(a, dict) and not a.get("notified") and single_address(a.get("assigned_to_email"))
                                   for a in ctx.store.overrides("assignments").values()):
            return True
        return cfg.quality_send and any(single_address(r.get("replied_by_email")) for _, r in _unsent_reviews(ctx.store))
    except Exception:
        return True                                    # let the real run meet (and report) the problem


WORK_PREFIX = "nbe-email-report-"
STALE_WORK_SECONDS = 24 * 3600


def cleanup_stale_work_dirs():
    """A run the runner killed on its timeout cannot delete its temp folder (.msg files, zips): remove such leftovers
    once they are a day old. Returns how many folders went. Never raises."""
    removed = 0
    try:
        names = list(Path(tempfile.gettempdir()).iterdir())
    except OSError:
        return 0
    for path in names:
        try:
            if (path.name.startswith(WORK_PREFIX) and path.is_dir() and not path.is_symlink()
                    and time.time() - path.stat().st_mtime > STALE_WORK_SECONDS):
                shutil.rmtree(path, ignore_errors=True)
                removed += 1
        except OSError:
            continue
    return removed


# ---------------------------------------------------------------------------------------------------------------------
# One run
# ---------------------------------------------------------------------------------------------------------------------
class Run:
    """Everything one report run shares: the mailer, the temp folder, the day's state, mails sent, failed steps."""

    def __init__(self, ctx, cfg, outcome, work, slot, state, mailer):
        self.ctx, self.cfg, self.outcome, self.work = ctx, cfg, outcome, work
        self.stats, self.rows, self.today = outcome.stats, outcome.rows, outcome.today
        self.slot, self.state, self.mailer = slot, state, mailer
        self.record = state["slots"].setdefault(slot, {})
        self.record.setdefault("done", [])
        self.record.setdefault("failed", [])
        self.force = bool((ctx.options or {}).get("force"))
        self.persist = ctx.can_email()
        self.sent = 0
        self.errors = []                                # "step: reason" for the JobResult
        self.sent_parts = []                            # fragments of the summary
        self.main_state = None                          # "sent" now / "earlier" today / None
        self.excel = None
        self.charts = {}
        self.parts = {}
        self.namespace = None                           # Outlook's MAPI namespace (tests pass the fake one)

    def log(self, text):
        self.ctx.log(text)

    def step(self, name, fn, *args):
        """Run one step; a failure is logged and remembered, never raised."""
        try:
            return fn(*args)
        except Exception as exc:
            self.errors.append(f"{name}: {type(exc).__name__}: {exc}")
            self.log(f"WARNING: step '{name}' failed: {type(exc).__name__}: {exc}")
            return None

    def already_sent(self, step_id):
        return not self.force and step_id in self.record["done"]

    def _save_state(self):
        self.state["updated_at"] = self.ctx.now().isoformat(timespec="seconds")
        try:
            self.ctx.store.write_state(STATE_NAME, self.state)
        except Exception as exc:
            self.errors.append(f"state: {type(exc).__name__}: {exc}")
            self.log(f"WARNING: the sent-mail record could not be saved ({exc}); a retry may send a mail twice.")

    def mark_done(self, step_id):
        if not self.persist:
            return
        if step_id not in self.record["done"]:
            self.record["done"].append(step_id)
        if step_id in self.record["failed"]:
            self.record["failed"].remove(step_id)
        self._save_state()

    def mark_failed(self, step_id):
        if not self.persist:
            return
        if step_id not in self.record["failed"]:
            self.record["failed"].append(step_id)
        self._save_state()

    def clear_failed(self, step_ids):
        """These steps do not apply any more (switched off, no recipient, nothing pending): a failure remembered for
        them must not keep the day's record "open" and force a re-scan on every call."""
        if not self.persist:
            return
        gone = [i for i in self.record["failed"] if i in step_ids]
        if gone:
            self.record["failed"] = [i for i in self.record["failed"] if i not in step_ids]
            self._save_state()

    @contextlib.contextmanager
    def guarded(self, step_id):
        """A send that raises marks its step failed - the next run retries exactly that step - and re-raises."""
        try:
            yield
        except Exception:
            self.mark_failed(step_id)
            raise

    def send(self, to, cc, subject, body, attachments=(), images=None):
        self.mailer.send(to=list(to), cc=list(cc), subject=one_line(subject), html_body=body,
                         attachments=[str(a) for a in attachments if a], inline_images=images or mail_images())
        self.sent += 1


def _history_step(run):
    """[2] today's headline numbers into the day's history file, then the trend against the days before."""
    if run.ctx.dry_run:
        return
    stats, store, today = run.stats, run.ctx.store, run.today
    history = eng.save_today_stats_history(store, today, stats)
    stats["trend"] = eng.compute_trend(today, history, stats)
    eng.save_today_reply_history(store, today, run.rows)


def _quality_history_step(run):
    if run.ctx.dry_run:
        return
    stats, cfg = run.stats, run.outcome.config
    history = eng.save_today_quality_history(run.ctx.store, run.today, stats.get("quality_scorecard") or [])
    stats["quality_trends"] = eng.build_person_quality_trends(history)
    stats["repeat_issue_alerts"] = eng.build_repeat_issue_alerts(history, run.today, cfg.repeat_threshold,
                                                                  cfg.repeat_lookback_days)


def _excel_step(run):
    """[3] The workbook goes to the shared report folder (a dry run builds it in the temp folder instead)."""
    folder = Path(run.work) / "excel" if run.ctx.dry_run else run.ctx.store.report_dir("email")
    path = create_excel(run.rows, run.outcome.team_replies, run.stats, run.today, folder, now=run.ctx.now())
    run.excel = path
    run.log(f"Excel created: {path.name}")
    if not run.ctx.dry_run:
        removed = prune_reports(folder, run.today)
        if removed:
            run.log(f"{removed} workbook(s) older than {REPORT_KEEP_DAYS} days removed.")


def _main_report_due(run):
    """Will this run send the main report? (It decides whether charts are worth drawing.)"""
    return run.cfg.send and bool(run.cfg.to) and not run.already_sent("main")


def _charts_step(run):
    """[4] Charts are an extra: without matplotlib (or when drawing fails) the mail uses table bars instead."""
    if not _main_report_due(run):
        return
    try:
        images = render_charts(run.stats, run.today)
    except ImportError:
        run.log("Charts skipped: matplotlib is not installed, the e-mail uses table bars.")
        return
    except Exception as exc:
        run.log(f"Charts skipped: {type(exc).__name__}: {exc}; the e-mail uses table bars.")
        return
    folder = Path(run.work) / "charts"
    folder.mkdir(parents=True, exist_ok=True)
    for cid, data in images.items():
        path = folder / f"{cid}.png"
        path.write_bytes(data)
        run.charts[cid] = str(path)


def _msgs_wanted(run):
    """Saving the .msg files is the slow, COM-heavy part: do it only when a mail that carries them will go out now (a
    re-run that only delivers a new assignment must not save forty mails again)."""
    cfg = run.cfg
    if not (cfg.send and cfg.attach_msg and run.ctx.can_email()):
        return False
    if run.force or _main_report_due(run):
        return True
    for team, rows in pending_rows_by_team(run.rows).items():
        own_to = cfg.team_recipients.get(team, ([], []))[0]
        if not (own_to or cfg.to):
            continue                                    # nobody could receive this team's list anyway
        total = len(chunk_list(rows, cfg.max_per_zip))
        if not all(run.already_sent(part_id(team, n, total)) for n in range(1, total + 1)):
            return True
    return False


def _pending_step(run):
    """[5] .msg files + zips per team. Without Outlook items it still builds the (zip-less) parts for the follow-ups."""
    files, failed = {}, 0
    if _msgs_wanted(run):
        namespace = run.namespace
        if namespace is None:
            try:
                namespace = eng.connect_outlook(run.ctx)[1]
            except Exception as exc:                   # noqa: BLE001 - the follow-ups still go, with the list only
                run.log(f"Outlook could not be opened to save the pending e-mails as .msg: {exc}")
        if namespace is not None:
            files, failed = save_pending_msgs(run.rows, Path(run.work) / "msg", run.log,
                                              lambda row: eng.fetch_item(namespace, row))
    run.parts = build_team_parts(run.rows, files, run.cfg.max_per_zip, Path(run.work) / "zips")
    attached = sum(p["count"] for parts in run.parts.values() for p in parts)
    if attached or failed:
        run.log(f"{attached} pending e-mail(s) saved as .msg and zipped by team" + (f" ({failed} could not be saved)" if failed else ""))


def _main_report_step(run):
    cfg = run.cfg
    if not cfg.send:
        run.clear_failed({"main"})
        return run.log("Report e-mail is switched off (report.send).")
    if not cfg.to:
        run.clear_failed({"main"})
        return run.log("Report e-mail skipped: no recipient (email.report.to).")
    if run.already_sent("main"):
        run.main_state = "earlier"
        return run.log(f"Report e-mail already sent today ({run.slot} slot).")
    with run.guarded("main"):
        attached = sum(p["count"] for parts in run.parts.values() for p in parts)
        team_mails = sum(len(parts) for team, parts in run.parts.items() if _team_recipients(cfg, team, parts))
        sample = review_sample(run.rows, cfg.sample_size, run.today)
        files = [run.excel] if cfg.attach_excel and run.excel and Path(run.excel).is_file() else []
        body = build_email_html(run.stats, run.today, include_charts=len(run.charts) == 2, pending_attached=attached,
                                pending_team_emails=team_mails, review_sample_rows=sample, excel_attached=bool(files))
        run.send(cfg.to, cfg.cc, f"{cfg.subject_prefix} - {_day_text(run.today)}", body, files, mail_images(run.charts))
    run.mark_done("main")
    run.main_state = "sent"
    run.sent_parts.append("report sent")


def _team_recipients(cfg, team, parts):
    """(to, cc) for a team's follow-up, or None. A team with its own recipients always gets its list; the report
    recipients get a team's follow-up only when it carries the .msg zip (the old fallback), never as plain noise."""
    own_to, own_cc = cfg.team_recipients.get(team, ([], []))
    if own_to:
        return own_to, own_cc
    if cfg.to and any(p["zip"] for p in parts):
        return cfg.to, cfg.cc
    return None


def _team_step(run):
    cfg = run.cfg
    if not cfg.send:
        run.clear_failed({i for i in run.record["failed"] if i.startswith("team:")})
        return
    sent_teams, skipped, applicable = 0, [], set()
    for team, parts in run.parts.items():
        recipients = _team_recipients(cfg, team, parts)
        if recipients is None:
            skipped.append(team)
            continue
        step_id = f"team:{team}"
        applicable.add(step_id)
        if run.already_sent(step_id):
            continue
        to, cc = recipients
        sent_now = False
        try:
            for number, part in enumerate(parts, 1):
                pid = part_id(team, number, len(parts))
                if run.already_sent(pid):                # an earlier try of this slot delivered this part
                    continue
                suffix = f" - Part {number} of {len(parts)}" if len(parts) > 1 else ""
                body = build_team_pending_email_html(run.today, team, part["count"], number, len(parts), part["rows"])
                run.send(to, cc, f"{cfg.subject_prefix} - {_day_text(run.today)} - Pending: {team}{suffix}", body,
                         [part["zip"]] if part["zip"] else [])
                run.mark_done(pid)
                sent_now = True
            run.mark_done(step_id)
            sent_teams += sent_now
        except Exception as exc:
            run.mark_failed(step_id)
            run.errors.append(f"team follow-up {team}: {type(exc).__name__}: {exc}")
            run.log(f"WARNING: could not send the pending e-mail for {team}: {type(exc).__name__}: {exc}")
    run.clear_failed({i for i in run.record["failed"] if i.startswith("team:") and i not in applicable})
    if skipped:
        run.log("No recipients for the pending list of: " + ", ".join(skipped) + " (email.report.team_recipients).")
    if sent_teams:
        run.sent_parts.append(f"{sent_teams} team follow-up(s)")


def _rows_by_message(rows):
    by_id = {}
    for row in rows:
        if row.get("message_id"):
            by_id.setdefault(row["message_id"], row)
    return by_id


def _assignment_step(run):
    """[7] One mail per person with all their new assignments; the "notified" flag is written right after each send."""
    cfg, store = run.cfg, run.ctx.store
    if not cfg.assign_send:
        return run.log("Assignment notices are switched off.")
    by_person = defaultdict(list)
    for message_id, info in (store.overrides("assignments") or {}).items():
        if not isinstance(info, dict) or info.get("notified") or not norm_email(info.get("assigned_to_email")):
            continue
        email = single_address(info.get("assigned_to_email"))
        if email:
            by_person[email].append((message_id, info))
        else:
            run.log("An assignment was skipped: its address is not a single valid e-mail address.")
    if not by_person:
        return run.log("No new assignments to notify.")
    rows_by_id = _rows_by_message(run.rows)
    sent = 0
    for email, items in sorted(by_person.items()):
        try:
            name = next((i.get("assigned_to_name") for _, i in items if i.get("assigned_to_name")), "") or email
            body = build_assignment_email_html(name, [(i, rows_by_id.get(m, {})) for m, i in items], run.today)
            run.send([email], cfg.assign_cc, f"{ASSIGNMENT_SUBJECT} - {len(items)} email(s) - {_day_text(run.today)}", body)
        except Exception as exc:
            run.errors.append(f"assignment notice: {type(exc).__name__}: {exc}")
            run.log(f"WARNING: could not notify {email}: {type(exc).__name__}: {exc}")
            continue
        sent += 1
        for message_id, info in items:
            _mark_notified(run, message_id, info)
    run.log(f"{sent} assignment notification e-mail(s) sent.")
    if sent:
        run.sent_parts.append(f"{sent} assignment notice(s)")


def _mark_notified(run, message_id, sent_info):
    """Flag the assignment as notified - but only if it is still the one we mailed (it may have been re-assigned meanwhile).
    The web API clears the flag the same way when the person or the note changes."""
    def change(old):
        if not isinstance(old, dict):
            return old                                  # deleted meanwhile: never re-create it
        same = (norm_email(old.get("assigned_to_email")) == norm_email(sent_info.get("assigned_to_email"))
                and (old.get("note") or "") == (sent_info.get("note") or ""))
        return {**old, "notified": True} if same else old
    try:
        run.ctx.store.update_override("assignments", message_id, change)
    except Exception as exc:
        run.errors.append(f"notified flag: {type(exc).__name__}: {exc}")
        run.log(f"WARNING: could not mark an assignment as notified ({exc}); the person may be e-mailed again.")


def _quality_step(run):
    """[7b] One mail per person with every rating of their replies that was not e-mailed yet."""
    cfg, store = run.cfg, run.ctx.store
    if not cfg.quality_send:
        return run.log("Quality feedback is switched off.")
    rows_by_id = _rows_by_message(run.rows)
    by_person = defaultdict(list)
    for message_id, review in _unsent_reviews(store):
        email = (single_address(review.get("replied_by_email"))
                 or single_address(rows_by_id.get(message_id, {}).get("first_reply_email")))
        if email:
            by_person[email].append((message_id, review))
    if not by_person:
        return run.log("No new quality reviews to send.")
    sent_state = copy.deepcopy(store.read_state(QUALITY_SENT_STATE) or {})
    sent = 0
    for email, items in sorted(by_person.items()):
        items.sort(key=lambda mr: str(mr[1].get("reviewed_at") or ""))
        try:
            name = next((r.get("replied_by") for _, r in items if r.get("replied_by")), "") or email
            body = build_quality_feedback_html(name, items, run.today)
            run.send([email], cfg.quality_cc, f"{QUALITY_SUBJECT} - {len(items)} reply review(s) - {_day_text(run.today)}", body)
        except Exception as exc:
            run.errors.append(f"quality feedback: {type(exc).__name__}: {exc}")
            run.log(f"WARNING: could not send quality feedback to {email}: {type(exc).__name__}: {exc}")
            continue
        sent += 1
        stamp = run.ctx.now().isoformat(timespec="seconds")
        for message_id, review in items:
            sent_state[message_id] = {"reviewed_at": review.get("reviewed_at"), "rating": review.get("rating"), "sent_at": stamp}
        _save_quality_sent(run, sent_state)              # after every person, like the old flag: a crash loses nothing
    run.log(f"{sent} quality feedback e-mail(s) sent.")
    if sent:
        run.sent_parts.append(f"{sent} quality feedback")


def _save_quality_sent(run, sent_state):
    try:
        current = set((run.ctx.store.overrides("quality") or {}))
        run.ctx.store.write_state(QUALITY_SENT_STATE, {m: v for m, v in sent_state.items() if m in current})
    except Exception as exc:
        run.errors.append(f"quality sent record: {type(exc).__name__}: {exc}")
        run.log(f"WARNING: the quality feedback record could not be saved ({exc}); the people may be e-mailed again.")


def _end_of_day_step(run):
    """The wrap-up goes to ``end_of_day.to`` (else the report recipients). It belongs to the daily report: with
    ``report.send`` off it is not sent either."""
    cfg = run.cfg
    if not cfg.eod_send:
        run.clear_failed({"end_of_day"})
        return run.log("End-of-day summary is switched off.")
    if not cfg.send:
        run.clear_failed({"end_of_day"})
        return run.log("End-of-day summary skipped: the daily report is switched off (report.send).")
    if not cfg.eod_to:
        run.clear_failed({"end_of_day"})
        return run.log("End-of-day summary skipped: no recipient.")
    if run.already_sent("end_of_day"):
        return run.log(f"End-of-day summary already sent today ({run.slot} slot).")
    with run.guarded("end_of_day"):
        files = [run.excel] if cfg.attach_excel and run.excel and Path(run.excel).is_file() else []
        body = build_end_of_day_html(run.rows, run.stats, run.today, excel_attached=bool(files))
        run.send(cfg.eod_to, cfg.eod_cc, f"{END_OF_DAY_SUBJECT} - {_day_text(run.today)}", body, files)
    run.mark_done("end_of_day")
    run.sent_parts.append("end-of-day summary")


def _summary(run, outcome):
    s = outcome.stats
    head = f"{s['business_emails']} mails, {s['pending']} pending"
    others = [p for p in run.sent_parts if p != "report sent"]
    if run.main_state == "sent":
        parts = [f"report sent: {head}"] + others
    elif run.main_state == "earlier":
        parts = [f"report already sent earlier today: {head}"] + others
    else:
        reason = ("e-mail is off for this run" if not run.ctx.can_email()
                  else "no report e-mail" if run.sent else "nothing sent")
        parts = [f"report built ({reason}): {head}"] + others
    text = ", ".join(parts)
    if run.errors:
        text += f" - {len(run.errors)} step(s) failed"
    return text


def _run_log_line(run):
    s, today = run.stats, run.today
    return (f"Date={_day_text(today)} Business={s['business_emails']} Replied={s['replied']} Pending={s['pending']} "
            f"NoAction={s['no_action_needed']} ReplyRate={s['reply_rate']:.1f}% SLABreached={s['sla_breached']} "
            f"Escalated={s['escalated_count']} Excel={run.excel.name if run.excel else '-'} "
            f"EmailsSent={run.sent}")


def _mailer(ctx):
    return ctx.mailer if ctx.mailer is not None else OutlookMailer()


def _run(ctx, namespace):
    cfg = ReportConfig(ctx.settings)
    now = ctx.now()
    slot = slot_name(cfg, now, ctx.options)
    state = _load_state(ctx.store, now.date().isoformat())
    record = state["slots"].get(slot) or {}
    force = bool((ctx.options or {}).get("force"))
    if (ctx.can_email() and not force and record.get("done") and not record.get("failed")
            and not _notifications_waiting(ctx, cfg)):
        ctx.log(f"The report of the {slot} slot was already sent today and nothing new is waiting - nothing to do "
                "(use Run now with force to send it again).")
        return JobResult(ok=True, skipped=True,
                         summary=f"report already sent today ({slot} slot); nothing new to send")

    ctx.log("[1] Scanning Outlook and analysing the day...")
    try:
        outcome = eng.scan_and_analyze(ctx, namespace)  # no live Outlook item kept: the .msg step fetches each one
    except (eng.ConfigProblem, eng.OutlookUnavailable) as exc:
        ctx.log(f"E-mail report cannot run: {exc}")
        return JobResult(ok=False, error=str(exc))
    except Exception as exc:
        ctx.log(f"E-mail report failed: {type(exc).__name__}: {exc}")
        return JobResult(ok=False, error=f"{type(exc).__name__}: {exc}")

    cleanup_stale_work_dirs()
    work = tempfile.mkdtemp(prefix=WORK_PREFIX)
    try:
        run = Run(ctx, cfg, outcome, work, slot, state, _mailer(ctx) if ctx.can_email() else None)
        run.namespace = namespace                      # None: the .msg step connects to Outlook itself
        if not ctx.can_email():
            ctx.log("E-mail is switched off for this run: the report is built, nothing is sent.")
        ctx.log("[2] Updating the trend and quality history...")
        run.step("history", _history_step, run)
        run.step("quality history", _quality_history_step, run)
        snapshot = run.step("snapshot", eng.to_snapshot_data, outcome, ctx)
        ctx.log("[3] Creating the Excel workbook...")
        run.step("excel", _excel_step, run)
        if ctx.can_email():
            ctx.log("[4] Creating the e-mail charts...")
            run.step("charts", _charts_step, run)
            ctx.log("[5] Preparing the pending lists per team...")
            run.step("pending lists", _pending_step, run)
            ctx.log("[6] Sending the report e-mail and the team follow-ups...")
            run.step("report e-mail", _main_report_step, run)
            run.step("team follow-ups", _team_step, run)
            ctx.log("[7] Notifying assigned team members...")
            run.step("assignment notices", _assignment_step, run)
            ctx.log("[7b] Sending reply-quality feedback...")
            run.step("quality feedback", _quality_step, run)
            ctx.log("[8] Sending the end-of-day summary...")
            run.step("end-of-day summary", _end_of_day_step, run)
        history = run.step("history point", eng.history_point, outcome, ctx)
        summary = _summary(run, outcome)
        ctx.log(_run_log_line(run))
        return JobResult(ok=not run.errors, snapshot_data=snapshot, snapshot_date=outcome.today.isoformat(),
                         summary=summary, error="; ".join(run.errors) or None, emails_sent=run.sent,
                         history_point=history)
    finally:
        shutil.rmtree(work, ignore_errors=True)


def run_report(ctx, namespace=None):
    """Job "email_report". Never raises: a failure becomes ``JobResult(ok=False, error=...)``."""
    try:
        return _run(ctx, namespace)
    except Exception as exc:
        ctx.log(f"E-mail report failed: {type(exc).__name__}: {exc}")
        return JobResult(ok=False, error=f"{type(exc).__name__}: {exc}")
