"""The "sla" job: how long has every security request been with its security user (port of the old SLA.PY monitor).

It logs in to the SSO portal with the PC's Edge / Chrome, opens the "User Provisioning Portal - موظف اسناد - مستخدم السرية" package,
jumps to the Re-Assign page, reads the requests table (plain tables with paging and DataTables through their API), and
measures the age of every request since its LAST assignment to the security user holding it (``clock_requests``: a
re-assignment restarts the clock). Entry point (docs/HUB-DESIGN.md sections 5.1, 6, 12):

    run(ctx, scraper=None) -> JobResult     what the runner calls; ``scraper(ctx) -> (headers, rows)`` replaces the browser in tests

SLA is two hours: from ``warning_minutes`` a request is a *reminder*, from ``breach_minutes`` it is *breached*. The result
is a snapshot (``requests`` with a stable ``key`` hashed from the identifying columns, ``by_user``, ``totals``), a history
point and ``first_breach`` for the day file. The runner writes snapshot, history and heartbeat; this module writes only the
Excel report (``SLA_latest.xlsx`` always, a timestamped copy with every e-mail, newest 30 kept) and its private alert state
(``ctx.store`` state "sla": ``{"alerted": {key: level}, "last_sent": iso}``).

E-mail rule (unchanged from the old script): send when a request reaches a higher level than the one already alerted
(first reminder, then breach), or every ``repeat_minutes`` while anything is still overdue (0 = only on escalation).
"Run now" with ``options["send_mail"]`` sends regardless. Nothing is sent when ``ctx.can_email()`` is false, ``mail.send`` is
off or ``mail.to`` is empty; in that case the alerts stay pending, so they go out as soon as mail is possible. Mail goes
through ``ctx.mailer`` (an Outlook sender is created lazily when the runner gave none). A failed scan, or a failed mail,
never changes the snapshot; a failed scan never changes the alert state either, and a failed mail leaves the alerts pending
so the next scan retries it.

The account comes from ``ctx.settings["username"]`` and the password only from ``ctx.password("sla", username)``. The
password is never logged, put in an exception text, a file name or a saved page: every error text and debug HTML is
scrubbed of it, and debug HTML is written without any input values.

Only the standard library and openpyxl are needed to import this module; the browser is driven over the DevTools
protocol (``calllog.ops.engines.cdp``) - no Playwright, no Node.js and no browser download.
"""
import hashlib
import os
import re
import subprocess
import sys
import time
from collections import Counter
from contextlib import suppress
from datetime import datetime, timedelta
from pathlib import Path
from urllib.parse import quote, urlparse

import json

from .. import schemas, teams
from . import cdp
from ..settings import DEFAULTS, deep_merge
from .common import (
    FONT_STACK, NBE_GOLD, NBE_GREEN, NBE_GREEN_LIGHT, CredentialError, JobResult, OutlookMailer, esc, norm_arabic)

STATE_NAME = "sla"
LATEST_REPORT = "SLA_latest.xlsx"
REPORT_PREFIX = "SLA_Report_"
KEEP_REPORT_COPIES = 30
COPY_NAME = re.compile(r"^SLA_Report_\d{8}_\d{4}\.xlsx$")
MAX_EMAIL_ROWS = 200

ASSETS_DIR = Path(__file__).resolve().parent.parent / "assets"
LOGO_FILE = ASSETS_DIR / "nbe_logo.png"
LOGO_CID = "nbe_logo"

SETTLE_MS = 1500            # wait after a page opens / a search runs, before the tables are read
PAGE_MS = 700               # wait after clicking "next" on a plain paged table
MAX_PAGES = 500
NO_WINDOW = 0x08000000 if os.name == "nt" else 0
UNASSIGNED = "(Unassigned)"

LEVEL_REMINDER, LEVEL_BREACHED = 1, 2
OVERDUE = (schemas.SLA_REMINDER, schemas.SLA_BREACHED)

# The old script's English labels (Excel and e-mail); the snapshot carries the stable codes instead.
LABELS = {schemas.SLA_WITHIN: "WITHIN SLA", schemas.SLA_REMINDER: "REMINDER",
          schemas.SLA_BREACHED: "SLA BREACHED", schemas.SLA_NO_DATE: "NO SENT DATE"}
COLORS = {  # (text, background)
    schemas.SLA_BREACHED: ("#B71C1C", "#FDECEA"),
    schemas.SLA_REMINDER: ("#8A4B00", "#FFF3E0"),
    schemas.SLA_WITHIN: ("#1B5E20", "#E8F5E9"),
    schemas.SLA_NO_DATE: ("#455A64", "#ECEFF1"),
}


class PortalError(RuntimeError):
    pass


class LoginError(PortalError):
    """Wrong / expired account: retrying only risks locking it, so the job result says error_kind "auth"."""


# --------------------------------------------------------------------------- text and dates
ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")
DATE_DMY = re.compile(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})")
DATE_YMD = re.compile(r"(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})")
TIME_RE = re.compile(r"(\d{1,2}):(\d{2})(?::(\d{2}))?\s*(AM|PM|ص|م)?", re.I)
NOT_IDENTIFYING = re.compile("اسناد|ملاحظات|تفاصيل")      # reassigning / adding a note must not make a "new" request

norm = norm_arabic


def scrub(text, secrets=()):
    """``text`` with every secret (and its URL-encoded form) replaced by ***."""
    text = str(text)
    for secret in secrets:
        if secret:
            for form in {secret, quote(secret, safe=""), quote(secret)}:
                text = text.replace(form, "***")
    return text


def clean_error(exc, secrets=()):
    """One short line for the heartbeat / System page: the first lines of the exception, scrubbed of secrets."""
    text = str(exc).strip() or type(exc).__name__
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()][:6]
    return scrub(" | ".join(lines)[:600], secrets)


def fmt_age(minutes):
    if minutes is None:
        return "-"
    minutes = int(minutes)
    return f"{minutes // 60}h {minutes % 60:02d}m"


def fmt_hours(minutes):
    return f"{minutes / 60:g}h"


def parse_datetime(text, date_order="dmy"):
    """dd/mm/yyyy (or mm/dd with date_order "mdy", or yyyy-mm-dd) with optional time; Arabic digits; AM/PM or ص/م."""
    s = str(text or "").translate(ARABIC_DIGITS)
    m = DATE_YMD.search(s)
    if m:
        year, month, day = map(int, m.groups())
    else:
        m = DATE_DMY.search(s)
        if not m:
            return None
        a, b, year = map(int, m.groups())
        day, month = (b, a) if date_order == "mdy" else (a, b)
    rest = s[:m.start()] + " " + s[m.end():]
    hh = mm = ss = 0
    t = TIME_RE.search(rest)
    if t:
        hh, mm = int(t.group(1)), int(t.group(2))
        ss = int(t.group(3) or 0)
        ampm = (t.group(4) or "").upper()
        if ampm in ("PM", "م") and hh < 12:
            hh += 12
        elif ampm in ("AM", "ص") and hh == 12:
            hh = 0
    try:
        return datetime(year, month, day, hh, mm, ss)
    except ValueError:
        return None


def find_col(columns, wanted):
    """The column whose normalised name equals ``wanted``, else the first that contains it."""
    w = norm(wanted)
    for c in columns:
        if norm(c) == w:
            return c
    for c in columns:
        if w in norm(c):
            return c
    return None


# --------------------------------------------------------------------------- browser (DevTools protocol)
# Arabic-Indic digits, tatweel, the diacritics range, alef / taa marbuta / alef maqsura variants (same as common.norm_arabic)
JS_NORM = r"""
  const _AR = {"٠":"0","١":"1","٢":"2","٣":"3","٤":"4","٥":"5","٦":"6","٧":"7","٨":"8","٩":"9"};
  const norm = s => (s || "").replace(/[٠-٩]/g, d => _AR[d]).replace(/ـ/g, "").replace(/[ً-ْ]/g, "")
      .replace(/[إأآ]/g, "ا").replace(/ة/g, "ه").replace(/ى/g, "ي")
      .replace(/\s+/g, " ").trim().toLowerCase();
  const visible = el => !!(el.offsetWidth || el.offsetHeight || el.getClientRects().length);
  const textOf = el => norm(el.tagName === "INPUT" ? (el.value || "") : (el.innerText || el.textContent));
"""

# Reads every table on the page. DataTables are read through their API so all pages are included.
JS_READ_TABLES = "() => {" + r"""
  const cellText = n => {
    if (!n) return "";
    const c = n.cloneNode(true);
    c.querySelectorAll("br").forEach(b => b.replaceWith(" "));
    c.querySelectorAll("script, style").forEach(s => s.remove());
    return (c.textContent || "").replace(/\s+/g, " ").trim();
  };
  const strip = v => {
    if (v === null || v === undefined || typeof v === "object") return "";
    const d = document.createElement("div");
    d.innerHTML = String(v).replace(/<br\s*\/?>/gi, " ");
    return (d.textContent || "").replace(/\s+/g, " ").trim();
  };
  // the row's own «التفاصيل» post-back (an ASP.NET grid link: __doPostBack('grid', 'Select$N')), or null
  const PB = /__doPostBack\(\s*['"]([^'"]*)['"]\s*,\s*['"]([^'"]*)['"]\s*\)/, PBO = /WebForm_PostBackOptions\(\s*["']([^"']*)["']\s*,\s*["']([^"']*)["']/;
  const pbOf = tr => {
    if (!tr || !tr.querySelectorAll) return null;
    const ctls = [...tr.querySelectorAll('a[href*="PostBack"], [onclick*="PostBack"]')];
    if (!ctls.length) return null;
    const det = ctls.find(c => (c.textContent || c.value || '').includes('التفاصيل')) || ctls[ctls.length - 1];
    const src = (det.getAttribute('href') || '') + ' ' + (det.getAttribute('onclick') || '');
    const m = PB.exec(src) || PBO.exec(src);
    return m ? {t: m[1], a: m[2]} : null;
  };
  const out = [], seen = new Set();
  const DT = (window.jQuery && window.jQuery.fn && window.jQuery.fn.dataTable) || window.DataTable;
  if (DT && DT.tables && DT.Api) {
    for (const node of DT.tables()) {
      const api = new DT.Api(node);
      const cols = api.columns().indexes().toArray();
      const headers = cols.map(c => cellText(api.column(c).header()));
      const rows = api.rows().indexes().toArray().map(r => cols.map(c => {
        const n = api.cell(r, c).node();
        return n ? cellText(n) : strip(api.cell(r, c).data());
      }));
      let pbs = [];
      try { pbs = api.rows().nodes().toArray().map(pbOf); } catch (e) { pbs = rows.map(() => null); }
      out.push({source: "datatables", headers, rows, pbs});
      seen.add(node);
    }
  }
  for (const t of document.querySelectorAll("table")) {
    if (seen.has(t)) continue;
    const headRows = t.tHead ? [...t.tHead.rows] : [];
    let head = headRows.length ? [...headRows[headRows.length - 1].cells] : [];
    const allTh = tr => [...tr.cells].every(c => c.tagName === "TH");
    if (!head.length && t.rows.length && allTh(t.rows[0]))
      head = [...t.rows[0].cells];
    // no <thead> (an ASP.NET GridView): the header row, and a <th> footer row, sit in the body; neither is a request
    const body = [...t.tBodies].flatMap(b => [...b.rows]).filter(tr => tr.cells.length > 1 && !allTh(tr));
    out.push({source: "dom", headers: head.map(cellText), rows: body.map(tr => [...tr.cells].map(cellText)),
              pbs: body.map(pbOf)});
  }
  return out;
}"""

# The page without any typed or stored input value (buttons keep theirs: their text helps to find the right selector).
JS_REDACTED_HTML = "() => {" + r"""
  const root = document.documentElement.cloneNode(true);
  root.querySelectorAll("input, textarea").forEach(el => {
    const type = (el.getAttribute("type") || "text").toLowerCase();
    if (["submit", "button", "image", "reset"].includes(type)) return;
    el.removeAttribute("value");
    if (el.tagName === "TEXTAREA") el.textContent = "";
  });
  return "<!DOCTYPE html>\n" + root.outerHTML;
}"""


# The table-reading and page-dump scripts are arrow functions; CDP evaluates a complete expression, so they are called.
READ_TABLES = "(" + JS_READ_TABLES + ")()"
REDACTED_HTML = "(" + JS_REDACTED_HTML + ")()"



def js_click_text(text):
    """Find the best clickable element whose text contains ``text``; return its centre {x, y}, or null."""
    return ("(() => {" + JS_NORM
            + "const R = el => { try{el.scrollIntoView({block:'center'});}catch(e){} const r = el.getBoundingClientRect(); return {x: r.left + r.width/2, y: r.top + r.height/2}; };"
            + "const t = norm(" + json.dumps(text) + ");"
            + "const sel = \"a, button, input[type=submit], input[type=button], input[type=image], [onclick], [role=button], li, span\";"
            + "const cands = [...document.querySelectorAll(sel)].filter(visible).filter(el => textOf(el).includes(t));"
            + "if (!cands.length) return null;"
            + "const score = el => (/^(A|BUTTON|INPUT)$/.test(el.tagName) ? 0 : 10) + (textOf(el) === t ? 0 : 5) + Math.min(textOf(el).length/100, 4);"
            + "cands.sort((a,b) => score(a)-score(b));"
            + "return R(cands[0]); })()")


def js_click_selector(selector):
    """Centre {x, y} of the first visible element matching ``selector``, or null."""
    return ("(() => {" + JS_NORM
            + "const R = el => { try{el.scrollIntoView({block:'center'});}catch(e){} const r = el.getBoundingClientRect(); return {x: r.left + r.width/2, y: r.top + r.height/2}; };"
            + "const el = [...document.querySelectorAll(" + json.dumps(selector) + ")].filter(visible)[0];"
            + "return el ? R(el) : null; })()")


def js_open_package(target, link_text):
    """Centre {x, y} of the package row's open control (prefer one whose text has ``link_text``). 'notfound' if no row."""
    return ("(() => {" + JS_NORM
            + "const R = el => { try{el.scrollIntoView({block:'center'});}catch(e){} const r = el.getBoundingClientRect(); return {x: r.left + r.width/2, y: r.top + r.height/2}; };"
            + "const loose = s => s.replace(/[\\s\\-\u2013_]/g,'');"
            + "const t = loose(norm(" + json.dumps(target) + ")), lt = norm(" + json.dumps(link_text or "") + ");"
            + "const rows = [...document.querySelectorAll('tr')].filter(visible).filter(r => loose(textOf(r)).includes(t));"
            + "if (!rows.length) return 'notfound';"
            + "rows.sort((a,b) => textOf(a).length - textOf(b).length);"
            + "const row = rows[0];"
            + "const ctls = [...row.querySelectorAll(\"a, button, input[type=submit], input[type=button], input[type=image], [onclick]\")].filter(visible);"
            + "const enter = [...row.querySelectorAll(\"a[id*='btnSubmit'], input[id*='btnSubmit']\")].filter(visible)[0];"
            + "const ctl = enter || ctls.find(c => lt && textOf(c).includes(lt)) || ctls[ctls.length-1] || row;"
            + "return R(ctl); })()")


def js_next_page():
    """Centre {x, y} of the DataTables 'next' button when it is enabled, else null."""
    return ("(() => {" + 
            "const el = el => el; const R = e => { try{e.scrollIntoView({block:'center'});}catch(x){} const r = e.getBoundingClientRect(); return {x: r.left + r.width/2, y: r.top + r.height/2}; };"
            + "const n = document.querySelector('.dataTables_paginate .next:not(.disabled), .paginate_button.next:not(.disabled)');"
            + "return n ? R(n) : null; })()")


# the real portal's login and landing ids (the bank's proven bot uses exactly these); tried before any guess
PROVEN_LOGIN = {"username": "#ContentPlaceHolder1_txtName", "password": "#ContentPlaceHolder1_txtPassword",
                "login_button": "#ContentPlaceHolder1_btnAdd"}
LANDING_GRID = "#ContentPlaceHolder1_GridView_Application"


def js_fill_login(user_sel, pass_sel, username, password):
    """Set the user and password inputs - the configured selector, else the real portal's id, else the first visible
    text / password input - and fire input/change. Returns {user, pass}: whether each field was found."""
    return ("(() => {" + JS_NORM
            + "const one = (sels, guess) => { for (const s of sels) { if (!s) continue; const e = document.querySelector(s); if (e) return e; } return guess(); };"
            + "const u = one([" + json.dumps(user_sel or "") + ", " + json.dumps(PROVEN_LOGIN["username"]) + "], () => [...document.querySelectorAll(\"input[type=text], input:not([type])\")].filter(visible)[0]);"
            + "const p = one([" + json.dumps(pass_sel or "") + ", " + json.dumps(PROVEN_LOGIN["password"]) + "], () => [...document.querySelectorAll(\"input[type=password]\")].filter(visible)[0]);"
            + "const set=(el,val)=>{if(!el)return false; try{el.focus();}catch(e){} el.value=val; el.dispatchEvent(new Event('input',{bubbles:true})); el.dispatchEvent(new Event('change',{bubbles:true})); return true;};"
            + "return {user:set(u," + json.dumps(str(username)) + "), pass:set(p," + json.dumps(password) + ")}; })()")


def _click(browser, finder_expr):
    """Run ``finder_expr`` (returns a point {x, y} or null); do a trusted mouse click there. Returns True if clicked."""
    point = browser.eval(finder_expr)
    if not isinstance(point, dict):
        return False
    browser.click_at(point["x"], point["y"])
    return True


class Debug:
    """Screenshot + page HTML of every step, only when ``ctx.debug``. The HTML has no input values and no password."""

    def __init__(self, enabled, folder, secrets=()):
        self.enabled = enabled
        self.folder = Path(folder)
        self.secrets = list(secrets)
        self.n = 0

    def snap(self, browser, name):
        if not self.enabled:
            return
        with suppress(OSError):
            self.folder.mkdir(parents=True, exist_ok=True)
        self.n += 1
        base = self.folder / f"{self.n:02d}_{name}_0"
        try:
            browser.screenshot(f"{base}.png")
        except Exception:
            pass
        try:
            page_html = scrub(browser.eval(REDACTED_HTML), self.secrets)
            Path(f"{base}.html").write_text(page_html, encoding="utf-8")
        except Exception:
            pass


def debug_folder(ctx):
    root = Path(ctx.data_dir) / "ops" if ctx.data_dir else Path(ctx.store.root)
    return root / "debug" / "sla"


def open_browser(ctx, cfg, headless):
    """The real browser for a scan: the PC's Edge / Chrome driven over the DevTools protocol (no Playwright). Tests
    monkeypatch this to inject a scripted browser."""
    bcfg = {"browser_executable": cfg.get("browser_executable") or "",
            "cdp_port": cfg.get("cdp_port") or 0,
            "profile_dir": cfg.get("profile_dir") or ""}
    return cdp.CdpBrowser(bcfg, log=ctx.log, headless=headless, profile_leaf="sla_profile")


def _press_enter(browser):
    """Press Enter in the focused password field (a real, trusted key event), to submit a login form that has no button
    whose text we recognise."""
    try:
        browser.eval("(() => { const p=[...document.querySelectorAll('input[type=password]')]"
                     ".filter(e => e.offsetParent !== null)[0]; if (p) p.focus(); return !!p; })()")
    except Exception:
        pass
    for kind in ("rawKeyDown", "char", "keyUp"):
        try:
            browser.call("Input.dispatchKeyEvent", type=kind, key="Enter", code="Enter", text="\r",
                         windowsVirtualKeyCode=13, nativeVirtualKeyCode=13)
        except Exception:
            pass


def _sleep_ms(ms):
    time.sleep(max(0, int(ms)) / 1000.0)


def pick_table(tables, cfg):
    """The table that has both the sent-date and the owner column; a DataTable (all pages) beats a plain one."""
    best = None
    for t in tables:
        if find_col(t["headers"], cfg["sent_date_header"]) and find_col(t["headers"], cfg["owner_header"]):
            key = (t["source"] == "datatables", len(t["rows"]))
            if best is None or key > best[0]:
                best = (key, t)
    return best[1] if best else None


def js_landing_packages():
    """The package names on the after-login packages page (rows that carry an Enter link) - only to tell the user which
    packages the account holds when the wanted one is missing."""
    return ("(() => { const out = [];"
            + "for (const tr of document.querySelectorAll('table tr')) {"
            + "  if (!tr.querySelector(\"a[id*='btnSubmit'], a, input[type=submit], button\")) continue;"
            + "  const t = (tr.innerText || tr.textContent || '').replace(/\\s+/g, ' ').trim();"
            + "  if (t && t.length < 200 && !out.includes(t)) out.push(t); }"
            + "return out.slice(0, 15); })()")


def js_link_href(file_name):
    """The absolute address of the page's link to ``file_name`` (e.g. request_reassign.aspx), or null."""
    return ("(() => { const want = " + json.dumps(str(file_name or "").lower()) + ";"
            + "const a = [...document.querySelectorAll('a[href]')].find(x => (x.getAttribute('href') || '')"
            + ".toLowerCase().split(/[?#]/)[0].endsWith(want));"
            + "return a ? a.href : null; })()")


def login(browser, cfg, username, password, wrong_text):
    """Log in on the portal's login page (already open): the configured ids, then the real portal's ids
    (``PROVEN_LOGIN``), then a guess. Missing fields are a PortalError (the page changed), NOT a wrong password: only a
    login page that stays after a filled-in submit is a LoginError (``wrong_text``). Done when the URL left the login
    page or the real packages grid is shown."""
    sel = cfg.get("selectors") or {}
    timeout_s = int(cfg["timeout_seconds"])
    login_marker = url_marker(cfg["login_path"], "login.aspx", keep_extension=True)
    filled = browser.eval(js_fill_login(sel.get("username"), sel.get("password"), username, password))
    if not (isinstance(filled, dict) and filled.get("user") and filled.get("pass")):
        raise PortalError("The login page's user name / password fields were not found - the portal page may have "
                          "changed or not loaded. Nothing was typed.")
    buttons = [b for b in (sel.get("login_button"), PROVEN_LOGIN["login_button"]) if b]
    if not any(_click(browser, js_click_selector(b)) for b in buttons):
        if not _click(browser, js_click_text(cfg["login_button_text"])):
            _press_enter(browser)
    landing = "!!document.querySelector(" + json.dumps(LANDING_GRID) + ")"
    deadline = time.monotonic() + timeout_s
    while time.monotonic() < deadline:
        with suppress(Exception):
            if login_marker not in (browser.current_url() or "").lower() or browser.eval(landing):
                browser.wait_ready(timeout_s)
                return
        _sleep_ms(POLL_MS)
    raise LoginError(wrong_text)


JS_GRID_STATE = ("(() => { const i = document.querySelector('.dataTables_info, [id$=\"_info\"]');"
                 " return {empty: !!document.querySelector('td.dataTables_empty'),"
                 " info: i ? (i.textContent || '').trim() : ''}; })()")
_INFO_TOTAL = re.compile(r"\bof\s+([\d,٠-٩]+)|من\s+(?:أصل\s+)?([\d,٠-٩]+)")


def info_total(text):
    """The total a DataTables info line states ('Showing 1 to 10 of 130 records' -> 130), or None."""
    m = _INFO_TOTAL.search(str(text or ""))
    if not m:
        return None
    digits = (m.group(1) or m.group(2) or "").replace(",", "").translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789"))
    return int(digits) if digits.isdigit() else None


def _wait_for_grid(browser, cfg, seconds=None):
    """Read the requests table until it has rows or the portal's «no data» line (or the timeout). Returns the table."""
    deadline = time.monotonic() + (int(cfg["timeout_seconds"]) if seconds is None else seconds)
    table = None
    while True:
        with suppress(Exception):
            table = pick_table(browser.eval(READ_TABLES), cfg)
            if table is not None and (table["rows"] or (browser.eval(JS_GRID_STATE) or {}).get("empty")):
                return table
        if time.monotonic() >= deadline:
            return table
        _sleep_ms(POLL_MS)


def read_requests(browser, cfg, debug):
    """(headers, rows) of the requests table on the Re-Assign page."""
    table = _wait_for_grid(browser, cfg, SETTLE_MS * 2 / 1000.0)     # a grid drawn by the page itself: seconds at most
    said_empty = (browser.eval(JS_GRID_STATE) or {}).get("empty")
    if table is None or not (table["rows"] or said_empty):
        # results not loaded yet -> press the search button, then wait for rows or the portal's «no data» line
        mark_page(browser)
        if _click(browser, js_click_text(cfg["search_button_text"])):
            wait_postback(browser, int(cfg["timeout_seconds"]))
        table = _wait_for_grid(browser, cfg)
    debug.snap(browser, "results")
    if table is None:
        seen = [t["headers"] for t in browser.eval(READ_TABLES) if t["headers"]]
        raise PortalError("Requests table not found. Headers seen on the page: "
                          + " | ".join(", ".join(h) for h in seen))

    rows = list(table["rows"])
    if table["source"] == "dom":
        # plain paging: click "next" until it is disabled or nothing new arrives
        for _ in range(MAX_PAGES):
            if not _click(browser, js_next_page()):
                break
            t = _poll(lambda: pick_table(browser.eval(READ_TABLES), cfg),
                      lambda t: t is not None and any(r not in rows for r in t["rows"]), PAGE_MS * 4 / 1000.0)
            new = [r for r in (t["rows"] if t else []) if r not in rows]
            if not new:
                break
            rows.extend(new)
    state = browser.eval(JS_GRID_STATE) or {}
    # an empty or short read would be recorded as requests leaving the queue: refuse it rather than record it
    if not rows and not state.get("empty"):
        raise PortalError("The requests table did not finish loading (no rows and no «no data» line from the portal) - "
                          "this scan was not used, nothing was recorded.")
    total = info_total(state.get("info"))
    if total is not None and len(rows) < total:
        raise PortalError(f"Read {len(rows)} request(s) but the portal counts {total} (\"{state.get('info')}\") - this "
                          f"scan was not used, nothing was recorded.")
    return table["headers"], rows


def url_marker(path, default, keep_extension=False):
    """The file name of a portal path, lower-case: what must (not) appear in the URL once we are there."""
    name = urlparse(path or "").path.rstrip("/").rsplit("/", 1)[-1]
    return (name if keep_extension else os.path.splitext(name)[0]).lower() or default


# --------------------------------------------------------------------------- SLA Analysis: the live status history
# «الاستفسار عن الطلبات» lists the requests of a period; each row's «التفاصيل» opens «تفاصيل الطلب» with the request
# number and «الحالة التاريخية للطلب» (who did each step and when). Read in the same session as the Re-Assign page.
INQUIRY_DATE_FORMAT = "%d-%m-%Y"           # how the portal's date picker writes «21-10-2026»
DETAILS_WAIT_S = 12                        # one request's details: the portal's own time, checked every POLL_MS
CLOSE_WAIT_S = 3                           # a details box that closes on the client
POLL_MS = 150                              # how often a waiting step looks again (a local DevTools call: no portal load)
QUIET_S = 1.0                              # an action that posts nothing back is known after this
# The «تفاصيل الطلب» status-history grid of the bank's portal (the id its own lifecycle tool reads); found by its
# headers when the id is not there
DETAIL_GRID_ID = "ContentPlaceHolder1_UC_RequestDeatils_Inquiry1_UC_RequestDetails1_GridView_Req"


def js_fill_inquiry(date_from, date_to):
    """Set «تاريخ الطلب من / الى» and «مجموعة العمل» = All on the inquiry page. Returns {from, to, group, changed}:
    each found, and whether the work group had to be changed (on the real page that posts back)."""
    return ("(() => {" + JS_NORM
            + "const wantFrom = norm('تاريخ الطلب من'), wantTo = norm('تاريخ الطلب الى'), wantGroup = norm('مجموعة العمل');"
            + "const near = el => { let t = [el.getAttribute('placeholder'), el.getAttribute('aria-label'), el.id, el.name].join(' ');"
            + "  if (el.id) { const l = document.querySelector('label[for=\"' + el.id + '\"]'); if (l) t += ' ' + l.textContent; }"
            + "  let p = el.parentElement; for (let i = 0; i < 3 && p; i++, p = p.parentElement) {"
            + "    const l = p.querySelector('label'); if (l) { t += ' ' + l.textContent; break; } }"
            + "  return norm(t); };"
            + "const set = (el, v) => { if (!el) return false; try { el.removeAttribute('readonly'); el.focus(); } catch (e) {}"
            + "  el.value = v; for (const ev of ['input', 'change', 'blur']) el.dispatchEvent(new Event(ev, {bubbles: true}));"
            + "  return true; };"
            + "const ins = [...document.querySelectorAll('input')].filter(visible)"
            + "  .filter(i => !/^(hidden|submit|button|checkbox|radio|image|password)$/i.test(i.type || ''));"
            + "let f = ins.find(i => near(i).includes(wantFrom)), t = ins.find(i => near(i).includes(wantTo) && i !== f);"
            + "if (!f) f = ins.find(i => /(fromdate|datefrom|from$|_from)/i.test((i.id || '') + ' ' + (i.name || '')));"
            + "if (!t) t = ins.find(i => i !== f && /(todate|dateto|to$|_to)/i.test((i.id || '') + ' ' + (i.name || '')));"
            + "const sels = [...document.querySelectorAll('select')].filter(visible);"
            + "const allAt = s => [...s.options].findIndex(o => norm(o.text) === 'all' || norm(o.text) === norm('الكل'));"
            + "const g = sels.find(s => near(s).includes(wantGroup) && allAt(s) >= 0) || sels.find(s => allAt(s) >= 0);"
            + "let group = false, changed = false;"
            + "if (g) { const i = allAt(g); if (g.selectedIndex !== i) { g.selectedIndex = i; changed = true;"
            + "  g.dispatchEvent(new Event('change', {bubbles: true})); } group = true; }"
            + "return {from: set(f, " + json.dumps(date_from) + "), to: set(t, " + json.dumps(date_to) + "), group, changed}; })()")


JS_INQUIRY_FIELDS = ("(() => {" + JS_NORM
                     + "const want = [norm('تاريخ الطلب من'), norm('تاريخ الطلب الى')];"
                     + "const labels = [...document.querySelectorAll('label, span, td, th')].filter(e => !e.children.length)"
                     + "  .map(e => norm(e.textContent)); return want.every(w => labels.some(l => l.includes(w))); })()")

# DataTables shows 20 rows a page: show them all, so every row's «التفاصيل» is on the page (the click path only)
JS_SHOW_ALL_ROWS = ("(() => { const DT = (window.jQuery && jQuery.fn && jQuery.fn.dataTable) || window.DataTable; let n = 0;"
                    " if (DT && DT.tables && DT.Api) for (const t of DT.tables()) { try { const a = new DT.Api(t);"
                    " if (a.page.len() !== -1) { a.page.len(-1).draw(false); n++; } } catch (e) {} } return n; })()")

# Before an action that may post the page back: a mark on the window (gone when the page is replaced) and a flag set
# the moment the browser starts leaving the page. ASP.NET's partial post-backs (UpdatePanel) are seen through Sys.
JS_MARK = ("(() => { const m = {leaving: false}; window.__nbeMark = m;"
           " addEventListener('beforeunload', () => { m.leaving = true; }, {once: true});"
           " addEventListener('pagehide', () => { m.leaving = true; }, {once: true}); return true; })()")
JS_MARK_STATE = ("(() => { const m = window.__nbeMark; let busy = false;"
                 " try { busy = !!(window.Sys && Sys.WebForms && Sys.WebForms.PageRequestManager"
                 "   && Sys.WebForms.PageRequestManager.getInstance().get_isInAsyncPostBack()); } catch (e) {}"
                 " return {mark: !!m, leaving: !!(m && m.leaving), busy, ready: document.readyState}; })()")


def mark_page(browser):
    with suppress(Exception):
        browser.eval(JS_MARK)


def wait_postback(browser, timeout_s, quiet_s=QUIET_S):
    """After an action that may post the page back (call ``mark_page`` first): wait until the new page is ready, or a
    partial post-back has finished. Returns True when a post-back happened, False when nothing was posted within
    ``quiet_s`` (or at the timeout). Replaces fixed sleeps: a fast portal costs only its own time."""
    start, seen = time.monotonic(), False
    while True:
        try:
            s = browser.eval(JS_MARK_STATE) or {}
        except Exception:                                  # the old page is being replaced right now
            s, seen = None, True
        if s is not None:
            if not s.get("mark") and s.get("ready") == "complete":
                return True
            if s.get("busy") or s.get("leaving") or not s.get("mark"):
                seen = True
            elif seen:
                return True                                # a partial post-back finished (the page stayed)
            elif time.monotonic() - start >= quiet_s:
                return False                               # nothing was posted back
        if time.monotonic() - start >= timeout_s:
            return seen
        _sleep_ms(POLL_MS)


def js_details_point(target):
    """Centre {x, y} of the «التفاصيل» control in the inquiry row whose cells match ``target`` ({header: text});
    'notfound' when no row matches, 'nolink' when the row has no control."""
    return ("(() => {" + JS_NORM
            + "const R = el => { try{el.scrollIntoView({block:'center'});}catch(e){} const r = el.getBoundingClientRect();"
            + "  return {x: r.left + r.width/2, y: r.top + r.height/2}; };"
            + "const want = " + json.dumps(target, ensure_ascii=False) + ";"
            + "const cell = n => n ? norm(n.innerText || n.textContent || '') : '';"
            + "for (const t of document.querySelectorAll('table')) {"
            + "  const hr = t.tHead ? t.tHead.rows[t.tHead.rows.length - 1] : t.rows[0]; if (!hr) continue;"
            + "  const heads = [...hr.cells].map(c => cell(c)); const idx = {}; let ok = true;"
            + "  for (const h of Object.keys(want)) { const nh = norm(h); let i = heads.indexOf(nh);"
            + "    if (i < 0) i = heads.findIndex(x => x.includes(nh)); if (i < 0) { ok = false; break; } idx[h] = i; }"
            + "  if (!ok) continue;"
            + "  for (const tr of (t.tBodies[0] ? t.tBodies[0].rows : [])) {"
            + "    if (!visible(tr)) continue;"
            + "    if (Object.keys(want).every(h => cell(tr.cells[idx[h]]) === norm(want[h]))) {"
            + "      const ctls = [...tr.querySelectorAll('a, button, input[type=button], input[type=submit], [onclick]')].filter(visible);"
            + "      const ctl = ctls.find(c => textOf(c).includes(norm('التفاصيل'))) || ctls[ctls.length - 1];"
            + "      return ctl ? R(ctl) : 'nolink'; } } }"
            + "return 'notfound'; })()")


def js_postback(target, argument):
    """Open one row's «التفاصيل» the way its own link does - ``__doPostBack(target, argument)`` - with the page marked
    first (see ``wait_postback``). False when the page has no __doPostBack."""
    return ("(() => {" + "const m = {leaving: false}; window.__nbeMark = m;"
            + " addEventListener('beforeunload', () => { m.leaving = true; }, {once: true});"
            + " if (typeof window.__doPostBack !== 'function') return false;"
            + " window.__doPostBack(" + json.dumps(str(target)) + ", " + json.dumps(str(argument)) + "); return true; })()")


# The «الحالة التاريخية للطلب» table of the open details (by its id, else by its five headers), and its DataTables
# wrapper: the portal shows it 5 rows a page, sorted by name - every page must be read (see read_history)
JS_FIND_HISTORY = r"""
  const want = ['رقم المستخدم', 'اسم المستخدم', 'نوع المستخدم', 'التاريخ', 'الحالة'].map(norm);
  const strip = v => { if (v === null || v === undefined || typeof v === 'object') return '';
    const d = document.createElement('div'); d.innerHTML = String(v).replace(/<br\s*\/?>/gi, ' ');
    return (d.textContent || '').replace(/\s+/g, ' ').trim(); };
  const txt = n => (n ? strip(n.innerHTML) : '');               // «21/10/2026<br>08:54:23» -> «21/10/2026 08:54:23»
  const head = t => t.tHead ? t.tHead.rows[t.tHead.rows.length - 1] : t.rows[0];
  const fits = t => { const hr = head(t); if (!hr) return false; const nh = [...hr.cells].map(c => norm(txt(c)));
    return want.every(w => nh.some(h => h === w || h.includes(w))); };
  let table = document.getElementById(""" + json.dumps(DETAIL_GRID_ID) + r""");
  if (!(table && visible(table) && fits(table))) table = [...document.querySelectorAll('table')].find(t => visible(t) && fits(t));
  const wrap = table ? (table.closest('.dataTables_wrapper') || table.parentElement) : null;
  const DT = (window.jQuery && jQuery.fn && jQuery.fn.dataTable) || window.DataTable;
  const isDT = !!(table && DT && DT.isDataTable && DT.isDataTable(table));
"""

# Show every row of the history table at once: the DataTables API, else its "Show N" list set to the largest choice
JS_HISTORY_EXPAND = "(() => {" + JS_NORM + JS_FIND_HISTORY + r"""
  if (!table) return '';
  if (isDT) { try { const a = new DT.Api(table); if (a.page.len() !== -1) { a.page.len(-1).draw(false); return 'api'; } return ''; } catch (e) {} }
  const sel = wrap && wrap.querySelector('.dataTables_length select, select[name$="_length"]');
  if (sel && sel.options.length) {
    const opts = [...sel.options];
    const best = opts.find(o => o.value === '-1') || opts.reduce((a, b) => ((+b.value || 0) > (+a.value || 0) ? b : a), opts[0]);
    if (best && sel.value !== best.value) { sel.value = best.value; sel.dispatchEvent(new Event('change', {bubbles: true})); return 'length'; }
  }
  return '';
})()"""

# Next page of the history table (its own pager only); false when there is none
JS_HISTORY_NEXT = "(() => {" + JS_NORM + JS_FIND_HISTORY + r"""
  if (!wrap) return false;
  const next = wrap.querySelector('.paginate_button.next:not(.disabled), li.next:not(.disabled) > a, a.next:not(.disabled)');
  if (!next || next.closest('.disabled')) return false;
  next.click();
  return true;
})()"""


def js_detail_snapshot(tokens=()):
    """ONE cheap look at the details: {open, number (the «رقم الطلب» value), headers, rows, total (the «Showing 1 to 5
    of 6 records» count of the history table; null when it shows none), in_page (no box of its own: the panel is filled
    in place), verified (the box shows one of ``tokens`` - the row's employee / user number; null when it cannot tell),
    mark/leaving/busy/ready (see ``wait_postback``)}. ``rows`` = every row the DataTables API holds, else the rows on
    the page (read_history then expands / pages the table until ``total`` rows are read). The request number is looked
    for in labels first and only then in cells, inside the details box - never by reading every element's text."""
    return ("(() => {" + JS_NORM + JS_FIND_HISTORY + r"""
  const m = window.__nbeMark; let busy = false;
  try { busy = !!(window.Sys && Sys.WebForms && Sys.WebForms.PageRequestManager
    && Sys.WebForms.PageRequestManager.getInstance().get_isInAsyncPostBack()); } catch (e) {}
  const out = {open: false, number: '', headers: [], rows: [], total: null, in_page: false, verified: null,
               mark: !!m, leaving: !!(m && m.leaving), busy, ready: document.readyState};
  if (!table) return out;
  let rows;
  if (isDT) {
    const api = new DT.Api(table), nodes = api.rows().nodes().toArray();
    rows = nodes.every(Boolean) ? nodes.map(tr => [...tr.cells].map(txt))
      : api.rows().data().toArray().map(d => (Array.isArray(d) ? d : Object.values(d)).map(strip));
  } else {
    rows = [...table.querySelectorAll('tr')].filter(tr => tr !== head(table) && tr.cells.length > 1)
      .filter(tr => ![...tr.cells].every(c => c.tagName === 'TH')).filter(tr => !tr.querySelector('td.dataTables_empty'))
      .map(tr => [...tr.cells].map(txt));
  }
  const info = wrap && wrap.querySelector('.dataTables_info, [id$="_info"]');
  const mt = info ? /(?:of|من)\s+([\d,٠-٩]+)/.exec(info.textContent || '') : null;
  if (mt) out.total = parseInt(mt[1].replace(/,/g, '').replace(/[٠-٩]/g, d => '٠١٢٣٤٥٦٧٨٩'.indexOf(d)), 10);
  out.headers = [...head(table).cells].map(txt);
  out.rows = rows;
  out.open = rows.length > 0;
  const box = table.closest('.modal, [role=dialog], .ui-dialog, .modal-content, fieldset, .panel');
  out.in_page = !box;
  let scope = box;
  if (!scope) { scope = table.parentElement; for (let i = 0; i < 4 && scope && scope.parentElement; i++) scope = scope.parentElement; }
  const num = norm('رقم الطلب');
  const after = (a, b) => !!(a.compareDocumentPosition(b) & Node.DOCUMENT_POSITION_FOLLOWING);
  const findNo = sel => { for (const el of scope.querySelectorAll(sel)) {
      if (el.children.length || norm(el.textContent).replace(/[:：]/g, '').trim() !== num) continue;
      let p = el.parentElement;
      for (let i = 0; i < 3 && p; i++, p = p.parentElement) {
        const inp = [...p.querySelectorAll('input, textarea')].find(x => after(el, x));
        if (inp && (inp.value || '').trim()) return inp.value.trim();
        const sp = [...p.querySelectorAll('span, td, div')].find(x => after(el, x) && !x.children.length && (x.textContent || '').trim());
        if (sp && /\d/.test(sp.textContent)) return sp.textContent.trim();
      }
    } return ''; };
  out.number = findNo('label, span, b, strong, th, legend') || findNo('td, div');
  // «الفترة من / الى» (a request with a period, «نموذج (7)»): the details' text and field values in page order (the
  // history table left out), then the date (and time) right after the label - never a date further on
  out.period = null;
  if (scope) {
    const parts = [], walk = document.createTreeWalker(scope, NodeFilter.SHOW_ELEMENT | NodeFilter.SHOW_TEXT);
    for (let n = walk.currentNode; n; n = walk.nextNode()) {
      if (n.nodeType === 3) { if (!table.contains(n)) parts.push(n.nodeValue); }
      else if ((n.tagName === 'INPUT' && n.type !== 'hidden') || n.tagName === 'TEXTAREA') parts.push(n.value || '');
      else if (n.tagName === 'SELECT') { const o = n.options[n.selectedIndex]; parts.push(o ? o.text : ''); }
    }
    const flat = norm(parts.join(' | ')), D = '(\\d{1,2}[/.-]\\d{1,2}[/.-]\\d{4}|\\d{4}[/.-]\\d{1,2}[/.-]\\d{1,2})'
      + '(?:[\\s|]*(\\d{1,2}:\\d{2}(?::\\d{2})?(?:\\s*(?:am|pm|ص|م))?))?';
    const pf = new RegExp('الفتره\\s*(?:من)?\\s*(?:[/-]\\s*الي)?\\s*[:：]?[\\s|]*' + D).exec(flat);
    if (pf) {
      const rest = flat.slice(pf.index + pf[0].length);
      const pt = new RegExp('^[\\s|:：-]*(?:(?:الفتره\\s*)?الي\\s*[:：]?[\\s|]*)?' + D).exec(rest);
      out.period = {from: [pf[1], pf[2]].filter(Boolean).join(' '), to: pt ? [pt[1], pt[2]].filter(Boolean).join(' ') : ''};
    }
  }
  const tokens = """ + json.dumps([str(t) for t in tokens if str(t or "").strip()]) + r""";
  if (box && tokens.length) {
    const seen = (box.textContent || '') + ' ' + [...box.querySelectorAll('input, textarea')].map(i => i.value || '').join(' ');
    out.verified = tokens.some(t => seen.includes(t));
  }
  return out;
})()""")


# The inquiry results grid (its headers) and its DataTables wrapper: what the grid helpers below act on - never the
# status-history table of an open «تفاصيل الطلب»
JS_FIND_GRID = r"""
  const gwant = ['تاريخ الطلب', 'مستخدم السرية', 'الحالة'].map(norm);
  const ghead = t => t.tHead ? t.tHead.rows[t.tHead.rows.length - 1] : t.rows[0];
  const grid = [...document.querySelectorAll('table')].find(t => { const hr = ghead(t); if (!hr) return false;
    const nh = [...hr.cells].map(c => norm(c.textContent)); return gwant.every(w => nh.some(h => h === w || h.includes(w))); });
  const gwrap = grid ? (grid.closest('.dataTables_wrapper') || grid.parentElement) : null;
"""

# The grid's own «Showing 1 to 20 of N records» and «no data» line
JS_INQUIRY_STATE = ("(() => {" + JS_NORM + JS_FIND_GRID
                    + "const i = gwrap && gwrap.querySelector('.dataTables_info, [id$=\"_info\"]');"
                    + "return {found: !!grid, empty: !!(gwrap && gwrap.querySelector('td.dataTables_empty')),"
                    + " info: i ? (i.textContent || '').trim() : ''}; })()")

# «Show» of the grid set to its largest choice («All»): what the user picks when the grid holds more than one page
JS_GRID_EXPAND = ("(() => {" + JS_NORM + JS_FIND_GRID + r"""
  const sel = gwrap && gwrap.querySelector('.dataTables_length select, select[name$="_length"]');
  if (!sel || !sel.options.length) return false;
  const opts = [...sel.options];
  const best = opts.find(o => o.value === '-1') || opts.reduce((a, b) => ((+b.value || 0) > (+a.value || 0) ? b : a), opts[0]);
  if (!best || sel.value === best.value) return false;
  sel.value = best.value; sel.dispatchEvent(new Event('change', {bubbles: true}));
  return true; })()""")

# The grid's DataTables "next" button (its own pager, never the status history's), when enabled: its centre or null
JS_GRID_NEXT = ("(() => {" + JS_NORM + JS_FIND_GRID + r"""
  const n = gwrap && gwrap.querySelector('.dataTables_paginate .next:not(.disabled), .paginate_button.next:not(.disabled)');
  if (!n || n.closest('.disabled')) return null;
  try { n.scrollIntoView({block: 'center'}); } catch (e) {}
  const r = n.getBoundingClientRect(); return {x: r.left + r.width / 2, y: r.top + r.height / 2}; })()""")

# The portal's OWN next page of results (the server sends at most a page of rows, «All» = 1000 on the bank's portal):
# the button «عرض التالي» under the grid (next to «عرض السابق»), or a pager link «التالي» / Page$Next. An ASP.NET
# post-back {t, a}, else the control's centre {x, y}; null on the last page. The DataTables pager and anything in a
# details box are not it.
JS_SERVER_NEXT = ("(() => {" + JS_NORM + JS_FIND_GRID + r"""
  if (!grid) return null;
  const PB = /__doPostBack\(\s*['"]([^'"]*)['"]\s*,\s*['"]([^'"]*)['"]\s*\)/;
  const src = el => (el.getAttribute('href') || '') + ' ' + (el.getAttribute('onclick') || '');
  const after = el => !!(grid.compareDocumentPosition(el) & Node.DOCUMENT_POSITION_FOLLOWING);
  const ctls = [...document.querySelectorAll('a, input[type=submit], input[type=button], button')]
    .filter(visible).filter(after)
    .filter(el => !el.closest('.dataTables_paginate, .dataTables_length, .modal, [role=dialog], .ui-dialog'));
  const NEXT = norm('التالي'), BACK = norm('السابق');
  const isNext = el => { const t = textOf(el);       // «عرض التالي», «التالي», «الصفحة التالية»..., never «عرض السابق»
    return !!t && t.length <= 24 && !t.includes(BACK) && (t === NEXT || t.endsWith(NEXT) || t === 'next' || t.includes(norm('التاليه'))); };
  const n = ctls.find(el => /Page\$Next/i.test(src(el))) || ctls.find(isNext);
  if (!n || n.disabled || n.hasAttribute('disabled') || /aspNetDisabled|disabled/.test(n.className || '')) return null;
  const m = PB.exec(src(n));
  if (m) return {t: m[1], a: m[2]};
  try { n.scrollIntoView({block: 'center'}); } catch (e) {}
  const r = n.getBoundingClientRect(); return {x: r.left + r.width / 2, y: r.top + r.height / 2}; })()""")

MAX_SERVER_PAGES = 50                      # the portal's own result pages walked per search (1000 rows each)
MAX_REPEAT_ROWS = 50                       # rows a result page can repeat from the one before (requests came in meanwhile)
DRILL_GUESS_S = 3.0                        # one request's details, until the scan has timed its own
LIVE_HOLD_MAX_S = 30                       # at most this long kept for the Live tickets the pages to come may list
WALK_MARGIN, WALK_SLACK_S = 1.5, 3.0       # the pages still to come x their measured time x this, plus this

# The portal's own grand total under the grid: «اجمالي عدد الطلبات : - 4351» (every result page together), or null
JS_PORTAL_TOTAL = ("(() => {" + JS_NORM + r"""
  const t = norm(document.body ? document.body.innerText : '');
  const m = t.match(/اجمالي\s*عدد\s*الطلبات[^0-9]{0,12}([0-9][0-9,]*)/);
  return m ? parseInt(m[1].replace(/,/g, ''), 10) : null; })()""")


def _rows_sig(rows):
    return "\x1f".join("\x1e".join(map(str, r)) for r in rows or [])


def _cells(row):
    """One grid row's cells, to know the row again on a page shown anew."""
    return tuple(str(c).strip() for c in row or ())


JS_GRID_PRESENT = ("(() => {" + JS_NORM
                   + "const want = [norm('تاريخ الطلب'), norm('مستخدم السرية'), norm('الحالة')];"
                   + "return [...document.querySelectorAll('table')].some(t => { const hr = t.tHead ? t.tHead.rows[t.tHead.rows.length - 1] : t.rows[0];"
                   + " if (!hr) return false; const nh = [...hr.cells].map(c => norm(c.textContent));"
                   + " return want.every(w => nh.some(h => h === w || h.includes(w))); }); })()")


def _poll(read, ok, seconds, every_ms=None):
    """Call ``read()`` until ``ok(value)`` or ``seconds`` pass; the last value (or None)."""
    deadline = time.monotonic() + seconds
    value = None
    while True:
        with suppress(Exception):
            value = read()
            if ok(value):
                return value
        if time.monotonic() >= deadline:
            return value
        _sleep_ms(POLL_MS if every_ms is None else every_ms)


def pick_inquiry_table(tables):
    """The inquiry results table: it has the request date, the security user and the status columns."""
    from .. import sla_history
    best = None
    for t in tables or []:
        h = t.get("headers") or []
        if all(sla_history.find_col(h, sla_history.GRID_COLUMNS[k]) is not None for k in ("date", "executor", "status")):
            key = (t.get("source") == "datatables", len(t.get("rows") or []))
            if best is None or key > best[0]:
                best = (key, t)
    return best[1] if best else None


def _close_details(browser):
    """Close «تفاصيل الطلب» when it is a box of its own: its «اغلاق» button, else a modal close control, else Escape."""
    if not _click(browser, js_click_text("اغلاق")):
        if not _click(browser, js_click_selector(".modal .close, [data-dismiss=modal], .ui-dialog-titlebar-close")):
            with suppress(Exception):
                for kind in ("rawKeyDown", "keyUp"):
                    browser.call("Input.dispatchKeyEvent", type=kind, key="Escape", code="Escape",
                                 windowsVirtualKeyCode=27, nativeVirtualKeyCode=27)
    _poll(lambda: browser.eval(js_detail_snapshot()) or {}, lambda s: not s.get("open"), CLOSE_WAIT_S)


def _signature(snap):
    return "\x1f".join("\x1e".join(map(str, r)) for r in (snap or {}).get("rows") or [])


def _detail_sig(snap):
    """What tells one open «تفاصيل الطلب» from the previous one: its request number with its history lines - the rows
    of one form (one per application, each its own number) can have the very same history lines."""
    return str((snap or {}).get("number") or "") + "\x1d" + _signature(snap)


def _complete_history(browser, snap, tokens=()):
    """``snap`` with EVERY row of «الحالة التاريخية للطلب», or None when they cannot all be read. The portal shows that
    table 5 rows a page sorted by name: one page alone can miss the assignment or the registration and give a wrong age
    or a request in the wrong stage - so a partial read is never recorded (the request is read again next scan).
    The table's own «of N records» count is the yardstick: the DataTables rows, else every row shown at once
    ("Show" set to the largest choice), else the pages one by one."""
    total = snap.get("total")
    if not total or len(snap.get("rows") or []) >= total:
        return snap
    read = lambda: browser.eval(js_detail_snapshot(tokens)) or {}          # noqa: E731
    if browser.eval(JS_HISTORY_EXPAND):
        s = _poll(read, lambda s: len(s.get("rows") or []) >= total, 3)
        if s and len(s.get("rows") or []) >= total:
            return {**snap, "rows": s["rows"]}
    rows, pages = list(snap.get("rows") or []), {_signature(snap)}
    for _ in range(40):
        if len(rows) >= total or not browser.eval(JS_HISTORY_NEXT):
            break
        s = _poll(read, lambda s: _signature(s) not in pages and bool(s.get("rows")), 3) or {}
        if _signature(s) in pages or not s.get("rows"):
            break
        pages.add(_signature(s))
        rows += list(s["rows"])
    return {**snap, "rows": rows} if len(rows) >= total else None


def history_seconds(cfg, general=None):
    """The history read's time budget (seconds): the setting, kept two minutes inside the scan's own time limit
    (``general.job_timeout_minutes.sla``, 5 minutes by default) so a scan is never killed with its work unsaved."""
    seconds = int(cfg.get("history_seconds") or DEFAULTS["sla"]["history_seconds"])
    job_minutes = 5
    with suppress(Exception):
        job_minutes = int(((general or {}).get("job_timeout_minutes") or {}).get("sla") or 5)
    return max(30, min(seconds, job_minutes * 60 - 120))


def read_history(browser, cfg, debug, plan, log=lambda m: None):
    """Read «الاستفسار عن الطلبات» for the plan's period, then the status history of the requests that are new or whose
    status moved (``plan['index']``): at most ``history_max_details`` per scan and only within the plan's time budget
    (``plan['seconds']``) - what was read is always returned, the rest follows at the next scan. Returns {"items",
    "rows", "opened", "left", "failed", "timing"}: ``items`` = [{"grid", "id", "history"}] for
    ``sla_history.fold_history`` (``failed`` marks a row whose details could not be read). Raises PortalError only when
    the page itself cannot be used; one request that cannot be read is skipped.

    The rows of each result page are opened WHILE that page is shown: the portal's next page is only reached by
    «التالي» from the first one, and on a busy portal the first page changes between two searches, so a page left
    to be found again later is not found at all. Time is kept for the pages still to come (the walk always finishes:
    every row is counted) and the details budget for the Live tickets they may list; the pages before the one where
    the last scan stopped (``plan['resume_page']``) are opened after the others - with the rows held back so the walk
    could finish - by a second search, while time is left.

    Fast path (the bank's portal): a row's «التفاصيل» is a ``__doPostBack`` link, so each drill is ONE call of that
    post-back and ONE cheap look per ``POLL_MS`` until a NEW status history with a fresh request number shows - no
    mouse, no re-drawing the grid, no closing, no fixed sleep. Rows without such a link are clicked (and their box
    closed) the old way."""
    from .. import sla_history
    started = time.monotonic()
    deadline = plan.get("deadline") or started + float(plan.get("seconds") or history_seconds(cfg))
    timeout_s = int(cfg["timeout_seconds"])
    base = cfg["base_url"].rstrip("/")
    path = cfg.get("inquiry_path") or DEFAULTS["sla"]["inquiry_path"]
    marker = url_marker(path, "requestdetails_inquiry")
    href = browser.eval(js_link_href(url_marker(path, "requestdetails_inquiry.aspx", keep_extension=True)))
    browser.navigate(href if isinstance(href, str) and href.startswith("http") else base + path, timeout=timeout_s)
    if marker not in (browser.current_url() or "").lower():
        raise PortalError(f"Could not open the inquiry page «الاستفسار عن الطلبات» (ended on {browser.current_url()})")
    _poll(lambda: browser.eval(JS_INQUIRY_FIELDS), bool, min(timeout_s, 15))
    fmt = cfg.get("inquiry_date_format") or INQUIRY_DATE_FORMAT
    dfrom, dto = plan["date_from"].strftime(fmt), plan["date_to"].strftime(fmt)

    def search():
        mark_page(browser)
        filled = browser.eval(js_fill_inquiry(dfrom, dto)) or {}
        if filled.get("changed") and wait_postback(browser, timeout_s):   # the work group posted back: dates again
            filled = browser.eval(js_fill_inquiry(dfrom, dto)) or {}
        if not (filled.get("from") and filled.get("to")):
            raise PortalError("The inquiry page's date fields («تاريخ الطلب من / الى») were not found - the page may "
                              "have changed. The SLA history was not read this scan.")
        mark_page(browser)
        if not _click(browser, js_click_text(cfg["search_button_text"])):
            raise PortalError("The inquiry page's search button was not found.")
        wait_postback(browser, timeout_s)
        return _poll(lambda: pick_inquiry_table(browser.eval(READ_TABLES)),
                     lambda t: t is not None and (t["rows"] or (browser.eval(JS_GRID_STATE) or {}).get("empty")),
                     timeout_s)

    read_table = lambda: pick_inquiry_table(browser.eval(READ_TABLES))        # noqa: E731

    def grid_total_now():
        with suppress(Exception):
            return info_total((browser.eval(JS_INQUIRY_STATE) or {}).get("info"))
        return None

    def whole_page(table):
        """Every row of the result page shown now, with the page's own «of N records» (None when it shows none): the
        DataTables rows, else - when the grid holds only the rows of its current page - «Show» set to its largest
        choice and then its own pages one by one. Identical rows are kept (two requests can look the same)."""
        rows, pbs = list(table["rows"]), list(table.get("pbs") or [None] * len(table["rows"]))
        shown = table["rows"]                             # the rows the grid shows right now
        total = grid_total_now()
        wait = min(timeout_s, 15)                         # a DataTable that fetches each page from the portal
        if total is not None and len(rows) < total and browser.eval(JS_GRID_EXPAND):
            t = _poll(read_table, lambda t: t is not None and len(t["rows"]) >= total, wait)
            if t is not None and len(t["rows"]) > len(rows):
                rows, pbs = list(t["rows"]), list(t.get("pbs") or [None] * len(t["rows"]))
                shown = t["rows"]
        if table["source"] == "dom" or (total is not None and len(rows) < total):
            sig = _rows_sig(shown)
            for _ in range(MAX_PAGES):
                if total is not None and len(rows) >= total:
                    break
                if not _click(browser, JS_GRID_NEXT):
                    break
                t = _poll(read_table, lambda t: t is not None and t["rows"] and _rows_sig(t["rows"]) != sig, wait)
                if t is None or not t["rows"] or _rows_sig(t["rows"]) == sig:
                    break
                sig = _rows_sig(t["rows"])
                rows.extend(t["rows"])
                pbs.extend(t.get("pbs") or [None] * len(t["rows"]))
        return rows, pbs, total

    def server_next(sig):
        """Go to the portal's next result page; the new page's table, or None (last page / it did not change)."""
        nxt = browser.eval(JS_SERVER_NEXT)
        if not nxt:
            return None
        mark_page(browser)
        if "t" in nxt:
            if not browser.eval(js_postback(nxt["t"], nxt["a"])):
                return None
        else:
            browser.click_at(nxt["x"], nxt["y"])
        wait_postback(browser, timeout_s)
        t = _poll(read_table, lambda t: t is not None and t["rows"] and _rows_sig(t["rows"]) != sig, timeout_s)
        return t if t is not None and t["rows"] and _rows_sig(t["rows"]) != sig else None

    t0 = time.monotonic()
    table = search()
    searches, search_s = 1, time.monotonic() - t0
    debug.snap(browser, "inquiry")
    if table is None:
        raise PortalError("The inquiry results table did not appear after the search.")
    headers = table["headers"]
    # the portal's own grand total «اجمالي عدد الطلبات» says how many requests there are in all: the last result page
    # is known without one click too many, and the pages still to come are known while the details are opened
    portal_total = None
    with suppress(Exception):
        portal_total = browser.eval(JS_PORTAL_TOTAL)
    if not isinstance(portal_total, int):
        portal_total = None
    index = plan.get("index") or {}
    budget = int(cfg.get("history_max_details") or DEFAULTS["sla"]["history_max_details"])
    # a request still with the branch (registered / confirmed by its supervisor) has no SLA step yet: opened once it moves
    # and a request of an excluded application (the Third Party team's by default) is never opened at all
    excluded = plan.get("excluded") or (lambda system: False)
    live_keys = plan.get("live_keys") or set()
    # the Live tickets this search can list (their request date is a few days before they were sent): only those are
    # waited for while the pages before them are opened
    spans = plan.get("live_spans") or {}
    lo_day, hi_day = plan.get("date_from"), plan.get("date_to")
    live_here = {k for k in live_keys if k not in spans or not hasattr(lo_day, "isoformat")
                 or (spans[k][0] <= hi_day.isoformat() and spans[k][1] >= lo_day.isoformat())}
    resume = max(0, int(plan.get("resume_page") or 0))
    after_s = float(plan.get("after_s") or 0)            # kept for the slices read after this one (their pages)

    def wanted(it):
        return (sla_history.needs_details(it, index, plan["cfg"])
                and sla_history.detail_priority(it["status"], plan["cfg"], it.get("executor") or "") is not None
                and not excluded(it["system"]))

    def live_wanted(it):
        return (bool(live_keys) and sla_history.grid_live_key(it) in live_keys
                and sla_history.grid_stage(it["status"], plan["cfg"], it.get("executor") or "") == "open")

    # On the page shown: a Live ticket whose last assignment is not known yet first (the Live tab's clock), then the
    # FINISHED requests (confirmed, then rejected: one read gives every stamp of a request), then the ones awaiting the
    # reviewer, then the ones with a security user, the newest request first; a «نموذج (7)» row read before its period
    # was taken from the details comes last (``late``)
    def first(it):
        return (it["key"] in late, not live_wanted(it),
                sla_history.detail_priority(it["status"], plan["cfg"], it.get("executor") or "") or 0,
                -((sla_history.parse_time(it["date"]) or datetime.min).toordinal()))

    pages = []                       # the result pages walked: {"rows", "total", "keys", "repeats", "postbacks", "listed"}
    grid, item_of, page_of, by_cells = [], {}, {}, {}
    seen = {}                        # grid key occurrences so far: two identical rows stay two requests, across pages
    todo, late, repeats, live_found = set(), set(), [], set()
    done = {}                        # grid key -> its item (read, or could not be read)
    read_numbers = {}                # request number -> grid key, this scan
    drill_s = []
    opened = failed = streak = unverified = partial = 0
    prev_sig, stopped, held, resume_next = "", "", "", None
    opened_on = set()                # the pages whose requests this pass opened (where the next scan does not start)
    page_s = search_s                # what one more result page costs (shown and read): the search's own time at first
    pages_s = search_s               # the time spent showing and reading the result pages (the details' not counted)

    def rows_read():
        return sum(len(p["rows"]) - p["repeats"] for p in pages)

    def take_page(table):
        """Read the result page shown now into the grid. Returns its rows' «التفاصيل» post-backs by grid key (None
        for a row without one). Requests registered while the pages are walked push the last rows of a page onto the
        next one: when a page starts with the rows the page before ended with AND the portal's total grew in between,
        those rows are kept apart (``repeats``) - never opened, and left out when the walk counts more rows than the
        portal's total. Two real requests that look the same (the total did not grow) stay two."""
        rows, pbs, total = whole_page(table)
        skip, listed = 0, (pages[-1]["listed"] if pages else portal_total)
        if pages and portal_total is not None:
            prev = pages[-1]["rows"]
            same = next((k for k in range(min(len(prev), len(rows), MAX_REPEAT_ROWS), 0, -1)
                         if prev[-k:] == rows[:k]), 0)
            if same:
                now_listed = None
                with suppress(Exception):
                    now_listed = browser.eval(JS_PORTAL_TOTAL)
                if isinstance(now_listed, int) and isinstance(listed, int) and now_listed > listed:
                    skip, listed = min(same, now_listed - listed), now_listed
        n, keys, shown = len(pages), [], {}
        for i, it in sla_history.grid_rows_indexed(headers, rows, seen):
            key = it["key"]
            grid.append(it)
            keys.append(key)
            item_of[key], page_of[key] = it, n
            by_cells.setdefault(_cells(rows[i]), []).append(key)
            shown[key] = pbs[i] if i < len(pbs) else None
            if i < skip:
                repeats.append(key)
            elif wanted(it):
                todo.add(key)
                if not sla_history.needs_details(it, index, plan["cfg"], period=False):
                    late.add(key)
                if live_wanted(it):
                    live_found.add(sla_history.grid_live_key(it))
        pages.append({"rows": rows, "total": total, "keys": keys, "repeats": skip, "listed": listed,
                      "postbacks": sum(1 for pb in shown.values() if pb)})
        return shown

    def pages_left():
        """How many result pages are still to come (the portal's total says it; else whether «التالي» is there)."""
        if len(pages) >= MAX_SERVER_PAGES:
            return 0
        if portal_total is not None:
            left = portal_total - rows_read()
            return -(-left // max(1, len(pages[0]["rows"]))) if left > 0 else 0
        try:
            return 1 if browser.eval(JS_SERVER_NEXT) else 0
        except Exception:
            return 1

    def reshow(n):
        """Show result page ``n`` again (a new search starts on the first page, then «التالي» forward) and know its
        rows again by their cells: a busy portal moves rows between pages and changes statuses, so a row is never
        taken by its place. Its rows' post-backs by grid key, or None when it cannot be shown in the time left."""
        nonlocal searches, prev_sig
        if time.monotonic() + page_s * (n + 1) >= deadline:
            return None
        t = search()
        searches += 1
        prev_sig = ""
        for _ in range(n):
            if t is None or time.monotonic() >= deadline:
                return None
            t = server_next(_rows_sig(t["rows"]))
        if t is None:
            return None
        return known_again(t)

    def known_again(t):
        """The rows of the page shown now (``t``) known again by their cells: their post-backs by grid key."""
        rows, pbs, _ = whole_page(t)
        shown = {}
        for i, row in enumerate(rows):
            key = next((k for k in by_cells.get(_cells(row), ()) if k not in shown), None)
            if key is not None:
                shown[key] = pbs[i] if i < len(pbs) else None
        return shown

    def open_one(it, pb, page):
        """Open one row's «التفاصيل» (on result page ``page``) and read its whole status history into ``done``."""
        nonlocal prev_sig, opened, failed, streak, unverified, partial
        t1 = time.monotonic()
        tokens = [it.get("employee_no"), it.get("user_no")]
        fresh = lambda s: bool(s.get("open") and s.get("number") and s.get("ready") == "complete"   # noqa: E731
                               and not s.get("busy") and not s.get("leaving") and _detail_sig(s) != prev_sig
                               and read_numbers.get(s["number"], it["key"]) == it["key"])
        got, snap, in_page = None, None, False
        try:
            if pb and browser.eval(js_postback(pb["t"], pb["a"])):
                snap = _poll(lambda: browser.eval(js_detail_snapshot(tokens)) or {}, fresh, DETAILS_WAIT_S)
                in_page = True                            # never closed: the next post-back replaces it
            else:
                with suppress(Exception):
                    browser.eval(JS_SHOW_ALL_ROWS)        # a post-back may have put the grid back to one page
                target = {sla_history.GRID_COLUMNS[k]: it[k]
                          for k in ("date", "employee_no", "user_no", "system", "request_type", "form", "branch") if it[k]}
                point = browser.eval(js_details_point(target))
                if isinstance(point, dict):
                    mark_page(browser)
                    browser.click_at(point["x"], point["y"])
                    snap = _poll(lambda: browser.eval(js_detail_snapshot(tokens)) or {}, fresh, DETAILS_WAIT_S)
                    in_page = bool((snap or {}).get("in_page"))
            if snap and fresh(snap):
                number = str(snap["number"]).strip()
                prev_sig = _detail_sig(snap)
                whole = _complete_history(browser, snap, tokens)
                if whole is None:
                    partial += 1
                    log(f"SLA history: request {number} - only {len(snap['rows'])} of its {snap.get('total')} history "
                        "lines could be read; it is read again next scan (nothing partial is recorded)")
                else:
                    read_numbers[number] = it["key"]
                    unverified += snap.get("verified") is False
                    got = {"grid": it, "id": number,
                           "history": sla_history.events_from_table(whole["headers"], whole["rows"], plan["cfg"])}
                    span = snap.get("period") or {}
                    if span.get("from"):                  # «الفترة من / الى»: its clocks start there
                        got["period_from"] = sla_history.period_start(span["from"], plan.get("day_start") or "08:00")
                        got["period_to"] = sla_history.period_start(span.get("to"), "00:00") if span.get("to") else ""
            if snap and snap.get("open") and not in_page:
                _close_details(browser)
        except Exception as exc:                          # one request must not stop the others
            log(f"SLA history: one request could not be read ({type(exc).__name__})")
        drill_s.append(time.monotonic() - t1)
        if got:
            done[it["key"]] = got
            opened += 1
            streak = 0
            opened_on.add(page)
        else:
            done[it["key"]] = {"grid": it, "id": "", "history": None, "failed": True}
            failed += 1
            streak += 1

    def drill(shown, n, reserve, hold=0, rotate=True):
        """Open the details of the wanted rows of the page shown now (``shown``: its rows' post-backs by grid key), the
        most wanted first, while the detail budget and the time allow - ``reserve`` seconds are kept for the result pages
        and the slices still to come, and ``hold`` details (with their time) for the Live tickets the pages to come may
        still list. Returns the page's post-backs as shown now; None when a post-back lost the grid and the page could
        not be shown again. A row left for lack of budget or time marks where the next scan starts (``rotate``): the
        page after this one when this one had its turn (some of its requests opened), so no page keeps the others
        waiting scan after scan; this one when none of its requests could be opened."""
        nonlocal stopped, held, resume_next
        for it in sorted((item_of[k] for k in shown if k in todo and k not in done), key=first):
            if it["key"] not in shown:                    # not on the page shown again after a lost grid
                continue
            keep = 0 if live_wanted(it) else hold
            per = sum(drill_s) / len(drill_s) if drill_s else DRILL_GUESS_S
            why = stopped or ("detail budget" if opened + failed + keep >= budget else "3 failures in a row"
                              if streak >= 3 else "time budget"
                              if time.monotonic() >= deadline - reserve - min(keep * per, LIVE_HOLD_MAX_S) else "")
            if why:
                if why == "3 failures in a row" or (not keep and (why != "time budget" or reserve <= after_s)):
                    stopped = why                         # final: no later page can do better
                held = why
                if rotate and resume_next is None and it["key"] not in late:
                    resume_next = n + 1 if n in opened_on else n      # the next scan starts here
                return shown
            open_one(it, shown[it["key"]], n)
            try:                                          # a post-back that lost the grid: search again, back to page n
                if not browser.eval(JS_GRID_PRESENT):
                    shown = reshow(n)
                    if shown is None:
                        stopped = "the inquiry grid did not come back"
                        if rotate and resume_next is None:
                            resume_next = n
                        return None
            except Exception as exc:                      # keep what was read: the rest waits for the next scan
                stopped = f"the inquiry page stopped answering ({type(exc).__name__})"
                return None
        return shown

    # 1) Walk the result pages, opening the wanted rows of each page WHILE it is shown: a busy portal changes its first
    #    page between two searches (new requests, statuses moving), so a page is never left to be found again later.
    #    The pages before the one where the last scan stopped (``resume_page``) only get their Live tickets opened here.
    walk_cut, walked_all = False, False
    shown = take_page(table)
    if portal_total and pages[0]["rows"] and resume * len(pages[0]["rows"]) >= portal_total:
        resume = 0                                        # fewer pages than when the last scan stopped
    while True:
        n = len(pages) - 1
        more = pages_left()
        # generous: what the early pages leave is opened on the last one, while a walk cut short is a day not counted
        reserve = after_s + (more * page_s * WALK_MARGIN + PAGE_RESERVE_S + WALK_SLACK_S if more else 0.0)
        hold = len(live_here - live_found) if more else 0  # Live tickets the pages to come may list
        mine = {k: pb for k, pb in shown.items() if n >= resume or live_wanted(item_of[k])}
        drilled = not stopped and any(k in todo and k not in done for k in mine)
        if drilled and drill(mine, n, reserve, hold, rotate=n >= resume) is None:
            walked_all = not more
            break
        if not more:
            walked_all = True
            break
        if time.monotonic() >= deadline - PAGE_RESERVE_S:   # the run must never outlive its time limit
            walk_cut = True
            break
        t1 = time.monotonic()
        nxt = server_next(_rows_sig(((read_table() if drilled else None) or table)["rows"]))
        if nxt is None:
            walked_all = True                             # the last page (or «التالي» did not open: reported below)
            break
        table = nxt
        shown = take_page(table)
        page_s = time.monotonic() - t1
        pages_s += page_s
    walk_s = time.monotonic() - t0

    # 2) The rows the walk counted twice (pushed onto the next page while it was walked) are left out when the walk
    #    counted more rows than the portal's total; else they are real requests that look the same (read next scan)
    raw = sum(len(p["rows"]) for p in pages)
    drop = set(repeats[:raw - portal_total]) if portal_total is not None and raw > portal_total else set()
    if drop:
        grid = [it for it in grid if it["key"] not in drop]
    todo |= {k for k in repeats if k not in drop and wanted(item_of[k])}

    totals = [p["total"] for p in pages]
    grid_total = sum(totals) if totals and all(t is not None for t in totals) else None
    short = [(n + 1, len(p["rows"]), p["total"]) for n, p in enumerate(pages)
             if p["total"] is not None and len(p["rows"]) < p["total"]]
    days = [d for d in (sla_history.parse_time(it["date"]) for it in grid) if d]
    log(f"SLA history: inquiry {dfrom} – {dto}: {raw - len(drop)} row(s)"
        + (f" over {len(pages)} result pages" if len(pages) > 1 else "")
        + (f", the portal lists {portal_total} in all («اجمالي عدد الطلبات»)" if portal_total is not None
           else f", the portal lists {grid_total}" if grid_total is not None else "")
        + (f"; request dates {min(days):%d/%m/%Y} – {max(days):%d/%m/%Y}" if days else "")
        + (f"; {len(drop)} row(s) seen twice left out (requests came in while the pages were read)" if drop else "")
        + (f"; the pages took {pages_s:.0f} s ({pages_s / len(pages):.1f} s a page)" if len(pages) > 1 else "")
        + "".join(f"; page {n}: only {got} of its {want} rows could be read" for n, got, want in short))
    d_from, d_to = plan["date_from"], plan["date_to"]
    if days and hasattr(d_from, "year") and (min(days).date() < d_from or max(days).date() > d_to):
        log(f"SLA history: WARNING the portal listed requests dated {min(days):%d/%m/%Y} – {max(days):%d/%m/%Y} for a "
            f"search of {dfrom} – {dto}: the dates typed may not have been taken (System › Settings › Ticketing › the "
            "inquiry's date format)")
    if walk_cut:
        log(f"SLA history: the time budget ran out after {len(pages)} result page(s) ({raw} rows) - the rows of the "
            "other pages keep what the earlier scans read")
    elif portal_total is not None and rows_read() < portal_total:
        log(f"SLA history: WARNING only {rows_read()} of the {portal_total} requests could be read - the next result "
            "page («عرض التالي») did not open; the rest are read at the next scan")

    # 3) A second pass while time is left: the rows still unread - on the pages before the resume page, and the ones
    #    held back so the walk could finish (what the early pages leave is used here). A new search (it starts on the
    #    first page), then «التالي» up to the last page with such a row, each page known again by its rows' cells
    waiting = lambda: [k for k in todo if k not in done and k not in drop]       # noqa: E731
    if (walked_all and not stopped and waiting()
            and deadline - after_s - time.monotonic() > search_s + 2 * page_s):
        last_page = max(page_of.get(k, 0) for k in waiting())
        opened_on.clear()
        resume_next = None                                # this pass decides where the next scan starts
        m, shown = 0, reshow(0)
        while shown is not None and not stopped:
            shown = drill(shown, m, after_s)
            if shown is None or stopped or not waiting() or m + 1 > last_page:
                break
            if time.monotonic() + page_s >= deadline - after_s:
                stopped = "time budget"
                if resume_next is None:
                    resume_next = m + 1
                break
            nxt = server_next(_rows_sig((read_table() or {}).get("rows") or []))
            if nxt is None:
                break
            m += 1
            shown = known_again(nxt)

    items = [done.get(it["key"]) or {"grid": it, "id": "", "history": None} for it in grid]
    left = sum(1 for k in todo if k not in done and k not in drop)
    total_s = time.monotonic() - started
    timing = {"seconds": round(total_s, 1), "search_seconds": round(search_s, 1), "searches": searches,
              "walk_seconds": round(walk_s, 1), "page_seconds": round(pages_s / len(pages), 1) if pages else None,
              "drills": len(drill_s), "drill_avg_seconds": round(sum(drill_s) / len(drill_s), 2) if drill_s else None,
              "drill_max_seconds": round(max(drill_s), 2) if drill_s else None,
              "postback_rows": sum(p["postbacks"] for p in pages), "stopped": (stopped or held) if left else "",
              "partial": partial, "grid_total": grid_total, "grid_pages": len(pages), "grid_short": bool(short),
              "repeats": len(drop)}
    log(f"SLA history: {len(grid)} request(s) on the inquiry page, {opened} status histor(ies) read"
        + (f", {failed} could not be read" if failed else "") + (f", {left} left for the next scan" if left else "")
        + f" - {timing['seconds']} s (search {timing['search_seconds']} s"
        + (f", {timing['drill_avg_seconds']} s a request, slowest {timing['drill_max_seconds']} s" if drill_s else "")
        + ")" + (f"; stopped: {stopped or held}" if (stopped or held) and left else "")
        + (f"; {unverified} detail(s) could not be matched to their row by number" if unverified else ""))
    return {"items": items, "rows": len(grid), "opened": opened, "left": left, "failed": failed, "timing": timing,
            "portal_total": portal_total,
            "complete": walked_all and not short and not (portal_total is not None and rows_read() < portal_total),
            "resume_page": resume_next if resume_next is not None and resume_next < len(pages) else 0,
            "date_from": plan["date_from"].isoformat() if hasattr(plan.get("date_from"), "isoformat") else ""}


def scrape(ctx, username, password, history=None, history_sink=None):
    """Log in, open the package, open the Re-Assign page, read the table. Returns (headers, rows). With a ``history``
    plan (see ``history_plan``), the same session then reads «الاستفسار عن الطلبات» and the status history of the
    requests that moved into ``history_sink`` - a failure there is logged and never costs the Re-Assign scan.

    Driven over the DevTools protocol (no Playwright): the PC's Edge / Chrome is launched with a debugging port and
    walked through the pages with small JavaScript snippets. The debug dumps hold no input value and no password."""
    cfg = _config(ctx)
    timeout_s = int(cfg["timeout_seconds"])
    base = cfg["base_url"].rstrip("/")
    reassign_marker = url_marker(cfg["reassign_path"], "request_reassign")
    debug = Debug(ctx.debug, debug_folder(ctx), [password])

    browser = open_browser(ctx, cfg, headless=not ctx.headed)
    try:
        # 1) Login. The page is dumped BEFORE the password is typed.
        browser.navigate(base + cfg["login_path"], timeout=timeout_s)
        debug.snap(browser, "login")
        login(browser, cfg, username, password, "Login failed - still on the login page. Check the SLA user and "
                                                "password in collector_config.json (sla.username / sla.password).")
        debug.snap(browser, "groups")

        # 2) Open the package (may open in a new window, or load in the same tab)
        pkg = browser.eval(js_open_package(cfg["package_text"], cfg["package_link_text"]))
        if pkg == "notfound" or not isinstance(pkg, dict):
            held = browser.eval(js_landing_packages()) or []
            raise PortalError(f"Package not found on the groups page: {cfg['package_text']}"
                              + (f". The account holds: {' | '.join(held)}" if held else ""))
        browser.click_at(pkg["x"], pkg["y"])
        browser.attach(cfg["app_url_fragment"].lower(), timeout=timeout_s)
        debug.snap(browser, "app_home")

        # 3) Re-Assign page: through the portal's own menu link (relative to the package's folder), else the path
        link = js_link_href(url_marker(cfg["reassign_path"], "request_reassign.aspx", keep_extension=True))
        href = _poll(lambda: browser.eval(link), lambda h: isinstance(h, str) and h.startswith("http"),
                     SETTLE_MS * 4 / 1000.0)
        browser.navigate(href if isinstance(href, str) and href.startswith("http") else base + cfg["reassign_path"],
                         timeout=timeout_s)
        if reassign_marker not in (browser.current_url() or "").lower():
            raise PortalError(f"Could not open the Re-Assign page (ended on {scrub(browser.current_url(), [password])})")
        result = read_requests(browser, cfg, debug)

        # 4) SLA Analysis: the inquiry page and the status history of what moved (never costs the scan above)
        if history is not None and history_sink is not None:
            try:
                history_sink.update(read_history_slices(browser, cfg, debug, history, ctx.log))
            except Exception as exc:
                history_sink["error"] = clean_error(exc, [password])
                ctx.log(f"SLA history not read this scan: {history_sink['error']}")
                with suppress(Exception):
                    debug.snap(browser, "history_error")
        return result
    except Exception:
        try:
            debug.snap(browser, "error")
        except Exception:
            pass
        raise
    finally:
        browser.close()


# --------------------------------------------------------------------------- analysis
def _text(value):
    return "" if value is None else str(value)


def build_table(headers, rows):
    """Page table -> (columns, rows): unique non-empty column names, rows padded / cut to the width, and empty helper
    columns (checkboxes, icons: nameless and blank in every row) dropped."""
    cols, used = [], set()
    for i, h in enumerate(headers):
        name = _text(h).strip() or f"col_{i + 1}"
        while name in used:
            name += "_"
        used.add(name)
        cols.append(name)
    width = len(cols)
    body = [(list(map(_text, r)) + [""] * width)[:width] for r in rows]
    keep = [i for i, c in enumerate(cols)
            if not (c.startswith("col_") and all(not row[i].strip() for row in body))]
    return [cols[i] for i in keep], [[row[i] for i in keep] for row in body]


def note_assignments(ctx, cfg, requests, secrets=()):
    """Note on each live request its assignments from the SLA history (``sla_history.note_assignments``:
    ``assigned_at`` the first, ``holder_since`` / ``last_assigned_at`` the last one, to the user holding it now - the
    clock of ``clock_requests``). A missing or broken recording notes nothing."""
    from .. import sla_history, sla_records
    try:
        rules = sla_records.clean_config((ctx.store.read_sla_records() or {}).get("config"))
        open_records = {k: sla_history.rederive(r, rules)
                        for k, r in ((ctx.store.read_sla_history_state() or {}).get("open") or {}).items()}
    except Exception as exc:
        open_records = {}
        ctx.log(f"SLA history could not be read for the assignment times: {type(exc).__name__}: "
                f"{scrub(str(exc), secrets)}")
    return sla_history.note_assignments(requests, open_records)


def clock_requests(requests, previous, now, cfg, previous_at=None):
    """Start each live request's SLA clock at its LAST «اسناد لمستخدم السرية» - when the security user holding it
    now got it (the owner's rule: a re-assignment restarts the clock). «تاريخ الارسال للسرية» is NOT that moment (a
    request can wait days before it is assigned), so it is never used for an assigned request:

    * ``assignment`` - the status history's assignment to that user (``last_assigned_at``, exact);
    * ``seen`` - its assignment is not read yet: since the scans first saw it with that user (``holder_seen_at``,
      carried from scan to scan; a new holder at this scan, or a request seen for the first time: this scan; a request
      the last snapshot already had without one: that snapshot's time ``previous_at``) - at least that old, never a
      guess beyond it; the history read gives the exact time at the next scans;
    * ``sent`` - a request with no security user yet (waiting to be assigned): «تاريخ الارسال للسرية».

    Sets ``clock_at``, ``clock_from``, ``holder_seen_at``, ``age_minutes`` and ``status``, oldest first again.
    ``previous`` = the requests of the last snapshot. Returns {source: count}."""
    warn, breach = int(cfg["warning_minutes"]), int(cfg["breach_minutes"])
    before = {r.get("key"): r for r in previous or [] if isinstance(r, dict)}
    stamp = now.strftime("%Y-%m-%dT%H:%M:%S")
    counts = {"assignment": 0, "seen": 0, "sent": 0}
    for r in requests:
        prev = before.get(r["key"])
        assigned = r["security_user"] != UNASSIGNED
        if not assigned:
            seen = None
        elif prev is None:
            seen = stamp                                 # first sighting: with this user since this scan, at least
        elif norm(prev.get("security_user")) == norm(r["security_user"]):
            seen = prev.get("holder_seen_at") or (previous_at or stamp)[:19]   # the same user: since then still
        else:
            seen = stamp                                 # moved to this user since the last scan
        r["holder_seen_at"] = seen
        start, source = (parse_time_iso(r.get("last_assigned_at")), "assignment") if assigned else (None, "sent")
        if start is None and seen:
            start, source = parse_time_iso(seen), "seen"
        if start is None:
            start, source = r.get("_sent"), "sent"
        age = None if start is None else max(0, round((now - start).total_seconds() / 60))
        r.update(clock_at=start.isoformat(timespec="seconds") if start else None, clock_from=source,
                 age_minutes=age, status=status_of(age, warn, breach))
        counts[source] += 1
    requests.sort(key=lambda r: (r["age_minutes"] is None, -(r["age_minutes"] or 0)))
    return counts


def parse_time_iso(stamp):
    try:
        return datetime.fromisoformat(stamp) if stamp else None
    except (TypeError, ValueError):
        return None


def status_of(age, warn, breach):
    if age is None:
        return schemas.SLA_NO_DATE
    return schemas.SLA_BREACHED if age >= breach else schemas.SLA_REMINDER if age >= warn else schemas.SLA_WITHIN


def request_key(columns, cells, owner_col):
    """Stable id of a request: a hash of the identifying columns (not the owner, notes or assignee, so reassigning a
    request does not make it a "new" one)."""
    idx = [i for i, c in enumerate(columns) if c != owner_col and not NOT_IDENTIFYING.search(norm(c))
           and not c.startswith("col_")]
    if not idx:
        idx = [i for i, c in enumerate(columns) if not c.startswith("col_")] or list(range(len(columns)))
    raw = "|".join(cells[i] for i in idx)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:16]


def analyze(columns, rows, cfg, now):
    """One record per request, oldest first (no date last; equal ages keep the page order).

    Record: key, security_user, sent_at (iso | None), age_minutes (| None), status, fields {detail header: text},
    plus private ``_sent`` (datetime) and ``_cells`` (the page row) used by the Excel report."""
    sent_col = find_col(columns, cfg["sent_date_header"])
    owner_col = find_col(columns, cfg["owner_header"])
    if sent_col is None or owner_col is None:
        raise PortalError("Requests table not found. Headers seen on the page: " + ", ".join(columns))
    sent_i, owner_i = columns.index(sent_col), columns.index(owner_col)
    warn, breach = int(cfg["warning_minutes"]), int(cfg["breach_minutes"])

    details, taken = [], set()
    for wanted in cfg["detail_headers"]:
        col = find_col(columns, wanted)
        if col is not None and col not in taken:
            taken.add(col)
            details.append((wanted, columns.index(col)))

    out = []
    for cells in rows:
        sent = parse_datetime(cells[sent_i], cfg["date_order"])
        age = None if sent is None else max(0, round((now - sent).total_seconds() / 60))
        out.append({
            "key": request_key(columns, cells, owner_col),
            "security_user": cells[owner_i].strip() or UNASSIGNED,
            "sent_at": sent.isoformat(timespec="seconds") if sent else None,
            "age_minutes": age,
            "status": status_of(age, warn, breach),
            "fields": {name: cells[i] for name, i in details},
            "_sent": sent, "_cells": cells,
        })
    out.sort(key=lambda r: (r["age_minutes"] is None, -(r["age_minutes"] or 0)))
    return out


def _counts(requests):
    by_status = Counter(r["status"] for r in requests)
    ages = [r["age_minutes"] for r in requests if r["age_minutes"] is not None]
    return {"within": by_status[schemas.SLA_WITHIN], "reminder": by_status[schemas.SLA_REMINDER],
            "breached": by_status[schemas.SLA_BREACHED], "no_date": by_status[schemas.SLA_NO_DATE],
            "oldest_minutes": max(ages, default=0)}


def summarize(requests):
    """Per security user, worst first: breached desc, reminder desc, total desc (then name)."""
    groups = {}
    for r in requests:
        groups.setdefault(r["security_user"], []).append(r)
    rows = [{"user": user, "total": len(items), **_counts(items)} for user, items in sorted(groups.items())]
    rows.sort(key=lambda s: (-s["breached"], -s["reminder"], -s["total"]))
    return rows


def totals_of(requests):
    return {"open": len(requests), **_counts(requests)}


def snapshot_data(requests, by_user, cfg):
    """The ``data`` of the sla snapshot (HUB-DESIGN 5.1)."""
    return {
        "thresholds": {"warning_minutes": int(cfg["warning_minutes"]), "breach_minutes": int(cfg["breach_minutes"])},
        "totals": totals_of(requests),
        "by_user": by_user,
        "requests": [{k: v for k, v in r.items() if not k.startswith("_")} for r in requests],
    }


# --------------------------------------------------------------------------- alert state and decision
def load_state(ctx):
    """{"alerted": {key: level}, "last_sent": datetime | None}; a missing or damaged file is a fresh state."""
    try:
        raw = ctx.store.read_state(STATE_NAME) or {}
    except Exception as exc:
        ctx.log(f"SLA alert state could not be read ({type(exc).__name__}); starting fresh")
        raw = {}
    if not isinstance(raw, dict):
        raw = {}
    alerted = {str(k): int(v) for k, v in (raw.get("alerted") or {}).items()
               if isinstance(v, int) and not isinstance(v, bool)} if isinstance(raw.get("alerted"), dict) else {}
    try:
        last_sent = datetime.fromisoformat(raw["last_sent"]) if raw.get("last_sent") else None
    except (TypeError, ValueError):
        last_sent = None
    return {"alerted": alerted, "last_sent": last_sent}


def alert_levels(requests):
    return {r["key"]: LEVEL_BREACHED if r["status"] == schemas.SLA_BREACHED else LEVEL_REMINDER
            for r in requests if r["status"] in OVERDUE}


def decide_alert(state, levels, now, repeat_minutes, force=False):
    """Should an e-mail go out now? Returns {"send", "escalated": [keys], "repeat": bool}."""
    escalated = [k for k, lv in levels.items() if lv > state["alerted"].get(k, 0)]
    since = None
    if state["last_sent"]:
        since = (now - state["last_sent"]).total_seconds() / 60
    repeat = int(repeat_minutes or 0)
    # a negative gap means the clock moved back (daylight saving, another PC): do not wait for it
    due_repeat = bool(levels) and repeat > 0 and (since is None or since < 0 or since >= repeat)
    return {"send": bool(force or escalated or due_repeat), "escalated": escalated, "repeat": due_repeat}


def next_state(state, levels, sent, now):
    """After a sent mail the levels are raised (never lowered); otherwise only closed / reassigned requests are forgotten."""
    if sent:
        alerted = {k: max(lv, state["alerted"].get(k, 0)) for k, lv in levels.items()}
        last_sent = now
    else:
        alerted = {k: v for k, v in state["alerted"].items() if k in levels}
        last_sent = state["last_sent"]
    return {"alerted": alerted, "last_sent": last_sent.isoformat(timespec="seconds") if last_sent else None}


# --------------------------------------------------------------------------- Excel
def _excel_text(value):
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    return ILLEGAL_CHARACTERS_RE.sub("", value)


def _append(ws, values):
    """Append a row; text is always stored as text, so a scraped '=HYPERLINK(...)' can never run as a formula."""
    ws.append([None] * len(values))
    for cell, value in zip(ws[ws.max_row], values):
        if isinstance(value, str):
            cell.value = _excel_text(value)
            cell.data_type = "s"
        else:
            cell.value = value


def build_workbook(columns, requests, by_user):
    """Workbook with the three sheets of the old report: By Security User / Over 1 Hour / All Requests."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    first = ["SECURITY_USER", "SENT_AT", "ASSIGNED_AT", "AGE_MIN", "AGE", "SLA_STATUS"]

    def request_row(r):                               # ASSIGNED_AT: the last assignment, the age's clock (when known)
        sent = r["_sent"].strftime("%d/%m/%Y %H:%M") if r["_sent"] else ""
        last = parse_time_iso(r.get("clock_at")) if r.get("clock_from") == "assignment" else None
        assigned = last.strftime("%d/%m/%Y %H:%M") if last else ""
        return [r["security_user"], sent, assigned, r["age_minutes"], fmt_age(r["age_minutes"]), LABELS[r["status"]],
                *r["_cells"]]

    wb = Workbook()
    sheets = [
        ("By Security User", ["SECURITY_USER", "TOTAL", "WITHIN_SLA", "REMINDER", "BREACHED", "NO_DATE", "OLDEST"],
         [[s["user"], s["total"], s["within"], s["reminder"], s["breached"], s["no_date"],
           fmt_age(s["oldest_minutes"]) if s["total"] > s["no_date"] else "-"] for s in by_user]),
        ("Over 1 Hour", first + columns, [request_row(r) for r in requests if r["status"] in OVERDUE]),
        ("All Requests", first + columns, [request_row(r) for r in requests]),
    ]
    thin = Side(style="thin", color="D0D7D4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    label_status = {label: status for status, label in LABELS.items()}
    for n, (title, headers, rows) in enumerate(sheets):
        ws = wb.active if n == 0 else wb.create_sheet()
        ws.title = title
        _append(ws, headers)
        for row in rows:
            _append(ws, row)
        ws.sheet_view.rightToLeft = False
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.fill = PatternFill("solid", fgColor=NBE_GREEN[1:])
            c.font = Font(bold=True, color="FFFFFF")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        names = [c.value for c in ws[1]]
        st = names.index("SLA_STATUS") if "SLA_STATUS" in names else None
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.border = border
            if st is not None and row[st].value in label_status:
                fg, bg = COLORS[label_status[row[st].value]]
                row[st].fill = PatternFill("solid", fgColor=bg[1:])
                row[st].font = Font(bold=True, color=fg[1:])
            for name, status in (("BREACHED", schemas.SLA_BREACHED), ("REMINDER", schemas.SLA_REMINDER)):
                if name in names and (row[names.index(name)].value or 0) > 0:
                    fg, bg = COLORS[status]
                    row[names.index(name)].fill = PatternFill("solid", fgColor=bg[1:])
                    row[names.index(name)].font = Font(bold=True, color=fg[1:])
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 3, 12), 45)
    return wb


def save_workbook(wb, path):
    """Write next to the target and rename, so a reader never sees a half-written file."""
    path = Path(path)
    tmp = path.with_name(f".{path.name}.tmp")
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        with suppress(OSError):
            tmp.unlink()
    return path


def prune_report_copies(folder, keep=KEEP_REPORT_COPIES):
    try:
        names = sorted(p.name for p in Path(folder).iterdir() if COPY_NAME.match(p.name))
    except OSError:
        return
    for name in names[:-keep] if keep else names:
        with suppress(OSError):
            (Path(folder) / name).unlink()


# --------------------------------------------------------------------------- e-mail
def _badge(status):
    fg, bg = COLORS.get(status, ("#333333", "#EEEEEE"))
    return (f'<span style="background:{bg};color:{fg};font-weight:bold;padding:3px 8px;'
            f'border-radius:10px;font-size:9pt;white-space:nowrap">{esc(LABELS.get(status, status))}</span>')


def _kpi(label, value, color):
    return (f'<td width="25%" style="padding:6px">'
            f'<table width="100%" cellpadding="0" cellspacing="0" style="border-collapse:collapse;'
            f'background:#FFFFFF;border:1px solid #DDE3E0;border-top:4px solid {color}">'
            f'<tr><td style="padding:12px 14px">'
            f'<div style="font-size:9pt;color:#5F6B66;text-transform:uppercase">{esc(label)}</div>'
            f'<div style="font-size:22pt;font-weight:bold;color:{color}">{esc(value)}</div>'
            f'</td></tr></table></td>')


def html_body(requests, by_user, cfg, now, has_logo=False, max_rows=MAX_EMAIL_ROWS):
    """The reminder e-mail. Every scraped value goes through ``esc``."""
    warn, breach = int(cfg["warning_minutes"]), int(cfg["breach_minutes"])
    counts = _counts(requests)
    n_total, n_breach, n_warn, n_ok = len(requests), counts["breached"], counts["reminder"], counts["within"]
    font = f"font-family:{FONT_STACK}"
    th = (f'style="background:{NBE_GREEN};color:#FFFFFF;padding:8px 10px;text-align:center;'
          f'font-size:9.5pt;border:1px solid {NBE_GREEN}"')
    td = 'style="padding:7px 10px;border:1px solid #E1E6E4;font-size:9.5pt;vertical-align:top"'
    tdc = 'style="padding:7px 10px;border:1px solid #E1E6E4;font-size:9.5pt;text-align:center;white-space:nowrap"'
    logo = (f'<img src="cid:{LOGO_CID}" height="56" alt="NBE" style="display:block">' if has_logo else
            '<div style="color:#FFFFFF;font-size:20pt;font-weight:bold">NBE</div>')

    def num(v, status):
        if not v:
            return '<span style="color:#9AA5A0">0</span>'
        fg, bg = COLORS[status]
        return (f'<span style="background:{bg};color:{fg};font-weight:bold;padding:2px 10px;'
                f'border-radius:10px">{int(v)}</span>')

    # --- per security user
    user_rows = []
    for i, s in enumerate(by_user):
        zebra = "#FFFFFF" if i % 2 == 0 else "#F7FAF9"
        oldest = fmt_age(s["oldest_minutes"]) if s["total"] > s["no_date"] else "-"
        user_rows.append(
            f'<tr style="background:{zebra}"><td {td} dir="auto"><b>{esc(s["user"])}</b></td>'
            f'<td {tdc}><b>{int(s["total"])}</b></td><td {tdc}>{num(s["within"], schemas.SLA_WITHIN)}</td>'
            f'<td {tdc}>{num(s["reminder"], schemas.SLA_REMINDER)}</td>'
            f'<td {tdc}>{num(s["breached"], schemas.SLA_BREACHED)}</td>'
            f'<td {tdc}>{esc(oldest)}</td></tr>')

    # --- overdue requests grouped by security user
    over = [r for r in requests if r["status"] in OVERDUE]
    detail_names = list(over[0]["fields"]) if over else []
    ncols = len(detail_names) + 3
    det_rows, shown = [], 0
    for s in by_user:
        grp = [r for r in over if r["security_user"] == s["user"]]
        if not grp or shown >= max_rows:
            continue
        nb = sum(1 for r in grp if r["status"] == schemas.SLA_BREACHED)
        grp_style = (f'background:{NBE_GREEN_LIGHT};padding:8px 10px;border:1px solid #D5E2DC;'
                     f'font-size:10pt;color:{NBE_GREEN}')
        breached_txt = f' &middot; <b style="color:#B71C1C">{nb} breached</b>' if nb else ""
        det_rows.append(
            f'<tr><td colspan="3" style="{grp_style}" dir="ltr">{len(grp)} request(s) over '
            f'{fmt_hours(warn)}{breached_txt}</td>'
            f'<td colspan="{ncols - 3}" style="{grp_style};text-align:right" dir="rtl">'
            f'<b>{esc(s["user"])}</b></td></tr>')
        for r in grp[:max_rows - shown]:
            shown += 1
            cells = "".join(f'<td {td} dir="auto">{esc(r["fields"].get(n))}</td>' for n in detail_names)
            sent = r["_sent"].strftime("%d/%m %H:%M") if r["_sent"] else "-"
            det_rows.append(
                f'<tr>{cells}<td {tdc}>{esc(sent)}</td>'
                f'<td {tdc}><b>{esc(fmt_age(r["age_minutes"]))}</b></td><td {tdc}>{_badge(r["status"])}</td></tr>')
    more = (f'<p style="font-size:9pt;color:#5F6B66">Showing {shown} of {len(over)} requests. '
            f'Full list in the attached Excel.</p>' if len(over) > shown else "")
    head = "".join(f"<th {th}>{esc(n)}</th>" for n in detail_names)
    details = (
        f'<table width="100%" cellpadding="0" cellspacing="0" style="border-collapse:collapse">'
        f'<tr>{head}<th {th}>Sent to Security</th><th {th}>Age</th><th {th}>Status</th></tr>'
        f'{"".join(det_rows)}</table>{more}'
        if det_rows else
        f'<p style="color:#2E7D32;font-weight:bold">&#10004; No request is older than {fmt_hours(warn)}.</p>')

    def section(title):
        return (f'<tr><td style="padding:22px 24px 8px 24px"><div style="font-size:13pt;font-weight:bold;'
                f'color:{NBE_GREEN};border-left:4px solid {NBE_GOLD};padding-left:10px">{title}</div></td></tr>')

    return f"""<html><body style="margin:0;padding:0;background:#EEF2F0;{font}">
<table width="100%" cellpadding="0" cellspacing="0" style="background:#EEF2F0"><tr><td align="center" style="padding:20px 10px">
<table width="900" cellpadding="0" cellspacing="0" style="border-collapse:collapse;background:#FFFFFF;{font};color:#1F2A26">
<tr><td style="background:{NBE_GREEN};padding:16px 24px">
  <table width="100%" cellpadding="0" cellspacing="0"><tr>
    <td valign="middle">{logo}</td>
    <td valign="middle" align="right" style="color:#FFFFFF">
      <div style="font-size:16pt;font-weight:bold">Security Requests SLA Reminder</div>
      <div style="font-size:9.5pt;color:#DCE9E4">Re-Assign queue &middot; {now:%d %b %Y, %H:%M}</div>
    </td></tr></table>
</td></tr>
<tr><td style="background:{NBE_GOLD};height:4px;font-size:1px">&nbsp;</td></tr>
<tr><td style="padding:20px 24px 4px 24px;font-size:10.5pt">
  Dear Team,<br><br>
  Security requests must be completed within <b>{fmt_hours(breach)}</b> of being sent to security.
  Below are the open requests per security user; requests over <b>{fmt_hours(warn)}</b> need action now.
</td></tr>
<tr><td style="padding:10px 18px"><table width="100%" cellpadding="0" cellspacing="0"><tr>
  {_kpi("Open Requests", n_total, NBE_GREEN)}
  {_kpi(f"Within {fmt_hours(warn)}", n_ok, "#2E7D32")}
  {_kpi(f"{fmt_hours(warn)} - {fmt_hours(breach)}", n_warn, "#EF6C00")}
  {_kpi(f"Over {fmt_hours(breach)} (Breached)", n_breach, "#C62828")}
</tr></table></td></tr>
{section("Requests per Security User")}
<tr><td style="padding:0 24px">
  <table width="100%" cellpadding="0" cellspacing="0" style="border-collapse:collapse">
  <tr><th {th}>Security User</th><th {th}>Total</th>
      <th {th}>Under {fmt_hours(warn)}</th><th {th}>{fmt_hours(warn)} - {fmt_hours(breach)}</th>
      <th {th}>Over {fmt_hours(breach)}</th><th {th}>Oldest</th></tr>
  {"".join(user_rows)}
  </table>
</td></tr>
{section(f"Requests Over {fmt_hours(warn)} ({len(over)})")}
<tr><td style="padding:0 24px">{details}</td></tr>
<tr><td style="padding:22px 24px;font-size:10.5pt">Please action the requests above. Full details are attached.<br><br>Regards,</td></tr>
<tr><td style="background:{NBE_GREEN_LIGHT};padding:10px 24px;font-size:8.5pt;color:#5F6B66;border-top:1px solid #D5E2DC">
  National Bank of Egypt &middot; Automated SLA monitor &middot; Internal use only
</td></tr>
</table></td></tr></table>
</body></html>"""


def mail_subject(mail_cfg, cfg, counts, now):
    subject = (f"{mail_cfg.get('subject') or 'Security Requests SLA'} - {counts['breached']} breached, "
               f"{counts['reminder']} over {fmt_hours(int(cfg['warning_minutes']))} - {now:%d-%b %H:%M}")
    return re.sub(r"[\r\n]+", " ", subject)


def mail_cc(mail_cfg, owner_emails, requests):
    """The configured cc list, plus (cc_security_users) the owners of overdue requests that have an address."""
    to = set(mail_cfg.get("to") or [])
    cc = [a for a in (mail_cfg.get("cc") or []) if a]
    if mail_cfg.get("cc_security_users"):
        emails = {norm(k): v for k, v in (owner_emails or {}).items()}
        for r in requests:
            addr = emails.get(norm(r["security_user"])) if r["status"] in OVERDUE else None
            if addr and addr not in cc and addr not in to:
                cc.append(addr)
    return cc


# --------------------------------------------------------------------------- the job
def _config(ctx):
    return deep_merge(DEFAULTS["sla"], ctx.settings or {})


def _fail(ctx, message, kind=None):
    ctx.log(f"SLA scan failed: {message}")
    return JobResult(ok=False, error=message, error_kind=kind)


def _mailer(ctx):
    return ctx.mailer or OutlookMailer()


def _write_reports(ctx, wb, now, send):
    """SLA_latest.xlsx always, a timestamped copy when a mail goes out. Returns the copy's path (or None).
    A report that cannot be written (share down, file open in Excel) is logged, never fatal."""
    copy = None
    try:
        folder = ctx.store.report_dir("sla")
        save_workbook(wb, folder / LATEST_REPORT)
    except OSError as exc:
        ctx.log(f"Excel {LATEST_REPORT} could not be saved ({type(exc).__name__}: {exc})")
    if send:
        try:
            copy = save_workbook(wb, ctx.store.report_dir("sla") / f"{REPORT_PREFIX}{now:%Y%m%d_%H%M}.xlsx")
            prune_report_copies(ctx.store.report_dir("sla"))
        except OSError as exc:
            ctx.log(f"Excel copy could not be saved, the e-mail goes without it ({type(exc).__name__}: {exc})")
    return copy


def _send_mail(ctx, cfg, requests, by_user, now, attachment):
    mail_cfg = cfg["mail"]
    cc = mail_cc(mail_cfg, cfg.get("owner_emails"), requests)
    has_logo = LOGO_FILE.is_file()
    body = html_body(requests, by_user, cfg, now, has_logo=has_logo)
    _mailer(ctx).send(
        to=list(mail_cfg["to"]), cc=cc, subject=mail_subject(mail_cfg, cfg, _counts(requests), now), html_body=body,
        attachments=[str(attachment)] if attachment else [], inline_images={LOGO_CID: str(LOGO_FILE)} if has_logo else None,
        display_only=bool(mail_cfg.get("display_only")))
    return cc


def _notify(ctx, cfg, requests, by_user, columns, now):
    """Decide, write the Excel files, send, update the alert state. Returns (sent: bool, note: str)."""
    state = load_state(ctx)
    levels = alert_levels(requests)
    mail_cfg = cfg.get("mail") or {}
    force = bool((ctx.options or {}).get("send_mail"))
    decision = decide_alert(state, levels, now, cfg.get("repeat_minutes"), force)

    if not ctx.can_email():
        ctx.log("E-mail skipped (this run does not send e-mail)")
        can_mail = False
    elif not mail_cfg.get("send", True):
        ctx.log("E-mail skipped (switched off in the SLA settings)")
        can_mail = False
    elif not mail_cfg.get("to"):
        ctx.log("E-mail skipped (no recipients in the SLA settings)")
        can_mail = False
    else:
        can_mail = True
        if not decision["send"]:
            ctx.log("No new request crossed a threshold - no e-mail")
    send = can_mail and decision["send"]

    sent, note = False, ""
    attachment = None if ctx.dry_run else _write_reports(ctx, build_workbook(columns, requests, by_user), now, send)
    if send:
        try:
            cc = _send_mail(ctx, cfg, requests, by_user, now, attachment)
            sent = True
            ctx.log(f"E-mail sent to {', '.join(mail_cfg['to'])}" + (f" (cc {', '.join(cc)})" if cc else ""))
        except Exception as exc:
            note = f"e-mail FAILED: {clean_error(exc)}"
            ctx.log(f"SLA e-mail failed, it will be tried again on the next scan: {clean_error(exc)}")
    if not ctx.dry_run:
        try:
            ctx.store.write_state(STATE_NAME, next_state(state, levels, sent, now))
        except Exception as exc:
            ctx.log(f"SLA alert state could not be saved ({type(exc).__name__}: {exc})")
    return sent, note


HISTORY_REACH_DAYS = 31        # how far back a request still in progress is followed (one older day a scan)
PAGE_RESERVE_S = 20            # seconds kept for the details and the fold when the result pages are walked
SLICE_RESERVE_S = 20           # seconds kept, while one slice's details are opened, for each slice still to search
LIVE_REACH_MARGIN_DAYS = 3     # a live ticket's request date can be this many days before it was sent to the team


def sla_excluder(ctx, cfg):
    """``excluded(system)`` for this scan (the Third Party team's applications by default - Settings › Ticketing), or
    None. Never fails the scan: a broken application list only means nothing is excluded by it."""
    applications = []
    with suppress(Exception):
        applications = (ctx.store.read_auto_assign() or {}).get("applications") or []
    return teams.sla_exclusion(cfg, teams.Resolver(ctx.hub or {}, None, applications))


def left_out_count(ctx, cfg, requests):
    """The Live tickets left out of the SLA (Third Party by default), counted only: {open, by_team: {team: n}, teams
    (the teams left out, as the settings name them)} - a ticket's team is its application's."""
    applications = []
    with suppress(Exception):
        applications = (ctx.store.read_auto_assign() or {}).get("applications") or []
    resolver = teams.Resolver(ctx.hub or {}, None, applications)
    by_team = {}
    for r in requests:
        team = resolver.sla_team(teams.sla_system(r)) or teams.NO_TEAM
        by_team[team] = by_team.get(team, 0) + 1
    return {"open": len(requests), "by_team": by_team,
            "teams": [str(t).strip() for t in cfg.get("exclude_teams") or [] if str(t).strip()]}


def history_plan(ctx, cfg, now):
    """What the inquiry read covers this scan: {date_from, date_to (the recent window), slices, recent_day, older_day,
    index, cfg, ...}. Each scan searches TODAY, ONE other day of the last ``history_days`` days (``recent_day``, in
    turn) and ONE older day (at most ``HISTORY_REACH_DAYS`` back) that still has a request in progress - the oldest such
    day first, then the next one at the next scan, in turn - each its own search. A request is followed until it closes
    without the search ever covering many days at once: every result page costs the portal seconds (37 pages took the
    whole scan; 20 pages of a busy window took two minutes before a single request could be opened). None when the SLA
    history is switched off."""
    from .. import sla_history, sla_records
    if not cfg.get("history_enabled", True):
        return None
    state = ctx.store.read_sla_history_state() or {}
    today = now.date()
    days = max(1, int(cfg.get("history_days") or DEFAULTS["sla"]["history_days"]))
    date_from = today - timedelta(days=days - 1)
    floor = today - timedelta(days=HISTORY_REACH_DAYS - 1)
    rules = sla_records.clean_config((ctx.store.read_sla_records() or {}).get("config"))
    older, live_days, live_keys, live_spans = set(), set(), set(), {}
    still_open = [r.get("request_date") for r in (state.get("open") or {}).values()]
    still_open += [g.get("request_date") for g in (state.get("grid") or {}).values()   # a branch-only request has no
                   if sla_history.row_stage(g, rules) in ("open", "awaiting")]         # SLA yet
    for stamp in still_open:                              # follow every request that is not finished yet
        day = sla_history.parse_time(stamp)
        if day and floor <= day.date() < date_from:
            older.add(day.date())
    # and every ticket of the last Live scan (the Re-Assign page): its request date is at most a few days before the
    # day it was sent to the security team. A ticket whose last assignment is not read yet (the Live tab's clock) is
    # wanted first: its row is opened before any other (``live_keys``) and its days come every other scan
    with suppress(Exception):
        live = (((ctx.store.read_snapshot("sla") or {}).get("data") or {}).get("requests") or [])
        for r in live:
            if not isinstance(r, dict):
                continue
            sent = sla_history.parse_time(r.get("sent_at"))
            wanted = r.get("clock_from") != "assignment" and r.get("security_user") != UNASSIGNED
            key = sla_history.live_key(r)
            if wanted and key[0]:
                live_keys.add(key)
                if sent:                                  # the request days it can have («تاريخ الطلب»)
                    span = ((sent.date() - timedelta(days=LIVE_REACH_MARGIN_DAYS)).isoformat(), sent.date().isoformat())
                    old = live_spans.get(key)
                    live_spans[key] = (min(old[0], span[0]), max(old[1], span[1])) if old else span
            if sent:
                for back in range(LIVE_REACH_MARGIN_DAYS + 1):
                    day = sent.date() - timedelta(days=back)
                    if floor <= day < date_from:
                        older.add(day)
                        if wanted:
                            live_days.add(day)
    last = str(state.get("history_older_day") or "")
    pool = sorted(live_days) if live_days and last not in {d.isoformat() for d in live_days} else sorted(older)
    turn = [d for d in pool if d.isoformat() > last] or pool
    pick = turn[0] if turn else None
    # the other days of the recent window, one a scan in turn (oldest first); a day the last scan could not read keeps
    # its turn
    window = [date_from + timedelta(days=n) for n in range(days - 1)]
    last_recent = str(state.get("history_recent_day") or "")
    recent = next((d for d in window if d.isoformat() > last_recent), window[0] if window else None)
    slices = [(today, today)] + ([(recent, recent)] if recent else []) + ([(pick, pick)] if pick else [])
    return {"date_from": date_from, "date_to": today, "slices": slices,
            "recent_day": recent.isoformat() if recent else "", "older_day": pick.isoformat() if pick else "",
            "older_days": len(older), "live_days": len(live_days), "live_keys": live_keys, "live_spans": live_spans,
            "index": dict(state.get("index") or {}), "cfg": rules,
            "excluded": sla_excluder(ctx, cfg), "seconds": history_seconds(cfg, getattr(ctx, "general", None)),
            "day_start": day_start(cfg, getattr(ctx, "general", None)),
            "resume_page": int(((state.get("history_stats") or {}).get("resume_page")) or 0)}


def read_history_slices(browser, cfg, debug, plan, log=lambda m: None):
    """``read_history`` for each slice of the plan (today, a recent day, an older day), all within ONE time budget: the
    items of every slice, and ``covered`` = the request days whose rows were all read (a slice cut short keeps its
    other rows as the earlier scans read them). While one slice's details are opened, time is kept for the searches of
    the slices after it, so every day of the plan is counted."""
    started = time.monotonic()
    deadline = started + float(plan.get("seconds") or history_seconds(cfg))
    out = {"items": [], "rows": 0, "opened": 0, "left": 0, "failed": 0, "timing": {}, "portal_total": None,
           "resume_page": 0, "covered": [], "date_from": plan["date_from"].isoformat(),
           "recent_day": plan.get("recent_day", ""), "older_day": plan.get("older_day", "")}
    covered = set()
    slices = plan.get("slices") or [(plan["date_from"], plan["date_to"])]

    def not_read(dfrom):                                 # a day not read: its turn comes again
        for key in ("recent_day", "older_day"):
            if out[key] and out[key] == dfrom.isoformat():
                out[key] = ""
    for n, (dfrom, dto) in enumerate(slices):
        if time.monotonic() >= deadline - PAGE_RESERVE_S:
            for later, later_to in slices[n:]:
                log(f"SLA history: no time left for {later:%d/%m}"
                    + (f" – {later_to:%d/%m}" if later_to != later else "") + "; it is read at a next scan")
                not_read(later)
            break
        sub = {**plan, "date_from": dfrom, "date_to": dto, "deadline": deadline,
               "after_s": (len(slices) - n - 1) * SLICE_RESERVE_S,
               "resume_page": plan.get("resume_page", 0) if n == 0 else 0}
        got = read_history(browser, cfg, debug, sub, log)
        out["items"] += got["items"]
        for key in ("rows", "opened", "left", "failed"):
            out[key] += got[key]
        if n == 0:
            out.update(timing=got["timing"], portal_total=got["portal_total"], resume_page=got["resume_page"])
        if got.get("complete"):
            d = dfrom
            while d <= dto:
                covered.add(d.isoformat())
                d += timedelta(days=1)
        else:
            not_read(dfrom)
    out["covered"] = sorted(covered)
    return out


def day_start(cfg, general=None):
    """When the working day starts ("08:00"): the SLA's own active hours, else the general ones - where a period
    («الفترة من») given without a time starts."""
    hours = ((cfg or {}).get("schedule") or {}).get("active_hours") or (general or {}).get("active_hours") or ""
    first = str(hours).split("-")[0].strip()
    return first if re.fullmatch(r"\d{1,2}:\d{2}", first) else "08:00"


OPEN_STALE_DAYS = 45            # an open request not listed for this long (cancelled, deleted) leaves the open set


def fold_history_result(store, sink, now_iso, system_team=None):
    """Write one scan's inquiry read (``read_history``'s result) into the SLA history store (``OpsStore.sla_history``):
    a request the reviewer has confirmed goes to the file of its confirmation day, the others stay in the open set; the
    grid index and the latest grid (what is open / awaiting RIGHT NOW) are refreshed. The day files are written before
    the state, so a crash in between only means those rows are read again at the next scan. A failed read only records
    its error (the last good history stays). Closed days older than the recording's ``retention_days`` are dropped.
    Returns fold_history's counts, or None when nothing was folded."""
    from collections import defaultdict
    from .. import sla_history, sla_records
    cfg = sla_records.clean_config((store.read_sla_records() or {}).get("config"))
    items = sink.get("items") or []
    out = {}

    def change(d):
        for k in ("open", "index", "grid"):
            d.setdefault(k, {})
        if sink.get("error") and not items:
            d["history_error"] = {"at": now_iso, "message": sink["error"]}
            return
        # records read under older status rules are brought up to the current ones first (sla_history.rederive)
        index, records, loaded_from, days = d["index"], {k: sla_history.rederive(r, cfg) for k, r in d["open"].items()}, {}, {}
        for it in items:                                  # a closed request read again: load it from its day file
            rid, day = it.get("id"), ((index.get((it.get("grid") or {}).get("key")) or {}).get("day"))
            if it.get("history") is not None and rid and day and rid not in records:
                doc = days.setdefault(day, store.read_sla_closed_day(day))
                if rid in doc:
                    records[rid], loaded_from[rid] = dict(doc[rid]), day
        out.update(sla_history.fold_history(records, index, items, cfg, now_iso, system_team))
        new_open, per_day, removals = {}, defaultdict(dict), defaultdict(set)
        for rid, r in records.items():
            old = loaded_from.get(rid)
            if sla_history.is_finished(r):                # confirmed (or rejected): filed under the day it finished
                day = sla_history.closing_day(r)
                per_day[day][rid] = r
                if old and old != day:
                    removals[old].add(rid)
            else:
                new_open[rid] = r
                if old:
                    removals[old].add(rid)
        for day in sorted(set(per_day) | set(removals)):

            def put(doc, day=day):
                for rid in removals.get(day, ()):
                    doc.pop(rid, None)
                doc.update(per_day.get(day, {}))
            store.update_sla_closed_day(day, put)
        stale = (datetime.strptime(now_iso[:10], "%Y-%m-%d") - timedelta(days=OPEN_STALE_DAYS)).strftime("%Y-%m-%d")
        d["open"] = {rid: r for rid, r in new_open.items() if (r.get("last_seen") or now_iso)[:10] >= stale}
        oldest = (datetime.strptime(now_iso[:10], "%Y-%m-%d")
                  - timedelta(days=HISTORY_REACH_DAYS + 2)).strftime("%Y-%m-%d")
        d["index"] = {k: v for k, v in index.items()       # rows that can still be listed by the inquiry window
                      if (v.get("day") and v["day"] >= oldest) or (not v.get("day") and v.get("id") in d["open"])
                      or ((v.get("fail") or {}).get("at") or "") >= oldest}
        if items:
            fresh = sla_history.grid_state([it["grid"] for it in items], d.get("grid"), now_iso, cfg)
            keep_from = sink.get("date_from") or now_iso[:10]
            d["grid"] = sla_history.merge_grid(d.get("grid"), fresh, set(sink.get("covered") or ()), keep_from)
            d["grid_from"] = keep_from
            out["unknown_wordings"] = sla_history.unknown_wordings(fresh, cfg)
        if sink.get("older_day"):
            d["history_older_day"] = sink["older_day"]
        if sink.get("recent_day"):
            d["history_recent_day"] = sink["recent_day"]
        d["last_history"] = now_iso
        d["history_stats"] = {"at": now_iso, "rows": sink.get("rows", 0), "portal_total": sink.get("portal_total"),
                              "opened": sink.get("opened", 0), "resume_page": int(sink.get("resume_page") or 0),
                              "left": sink.get("left", 0), "failed": sink.get("failed", 0),
                              "timing": sink.get("timing") or {}, **out}
        d.pop("history_error", None)
    store.update_sla_history_state(change)
    keep = int(cfg.get("retention_days") or 0)
    if keep > 0:                                          # the recording's retention, by confirmation day
        cutoff = (datetime.strptime(now_iso[:10], "%Y-%m-%d") - timedelta(days=keep)).strftime("%Y-%m-%d")
        for day in store.sla_closed_days():
            if day < cutoff:
                store.drop_sla_closed_day(day)
    return out or None


def run(ctx, scraper=None):
    """One scan: read the portal, analyse, e-mail on escalation, return the snapshot data."""
    cfg = _config(ctx)
    username = str(cfg.get("username") or "").strip()
    if not username:
        return _fail(ctx, "No SLA account is set. Put the portal user in collector_config.json "
                          "(\"sla\": {\"username\": ...}) or in System > Settings.", "auth")
    try:
        password = ctx.password("sla", username)
        if not password:
            raise CredentialError("empty password")
    except CredentialError:
        return _fail(ctx, f"No password for the SLA portal account '{username}'. Fill sla.password in "
                          "collector_config.json (next to collector.py), or set the NBE_SSO_PASSWORD environment "
                          "variable, or run:  python collector.py --set-password sla", "auth")
    secrets = [password]

    sink = {}
    try:
        plan = None
        if scraper is None:
            with suppress(Exception):                     # a broken recording file must not stop the live scan
                plan = history_plan(ctx, cfg, ctx.now())
        headers, rows = scraper(ctx) if scraper else scrape(ctx, username, password, plan, sink if plan else None)
        now = ctx.now()
        columns, table = build_table(headers, rows)
        requests = analyze(columns, table, cfg, now)
        excluded = sla_excluder(ctx, cfg)
        left_out = None
        if excluded:                                      # Third Party (by default): not in the SLA, its alerts or totals
            kept = [r for r in requests if not excluded(teams.sla_system(r))]
            if len(kept) != len(requests):
                ctx.log(f"{len(requests) - len(kept)} request(s) of excluded applications left out of the SLA "
                        "(Settings › Ticketing › Left out of the SLA) - only counted")
            left_out = left_out_count(ctx, cfg, [r for r in requests if excluded(teams.sla_system(r))])
            requests = kept
    except Exception as exc:
        if ctx.debug:
            import traceback
            ctx.log(scrub(traceback.format_exc(), secrets))
        return _fail(ctx, clean_error(exc, secrets), "auth" if isinstance(exc, LoginError) else None)

    if sink and not ctx.dry_run:                          # the SLA Analysis history (exact, live): kept for good
        try:
            from .. import sla_records
            folded = fold_history_result(ctx.store, sink, now.strftime("%Y-%m-%dT%H:%M:%S"),
                                         sla_records.system_team_map(ctx.store.read_auto_assign())) or {}
            odd = folded.get("unknown_wordings") or []
            if odd:                                       # placed by «مستخدم السرية» until the settings name them
                ctx.log("SLA history: status wording(s) the settings do not name (System › Settings › SLA statuses): "
                        + ", ".join(f"«{w}» ×{n}" for w, n in odd)
                        + " - counted with the security team only when the inquiry names a security user")
        except Exception as exc:
            ctx.log(f"SLA history could not be recorded: {type(exc).__name__}: {scrub(str(exc), secrets)}")
    note_assignments(ctx, cfg, requests, secrets)
    previous, previous_at = [], None
    with suppress(Exception):                             # the last scan: who held each request then
        last = ctx.store.read_snapshot("sla") or {}
        previous, previous_at = (last.get("data") or {}).get("requests") or [], last.get("generated_at")
    clock = clock_requests(requests, previous, now, cfg, previous_at)
    if requests:
        ctx.log(f"{len(requests)} request(s) aged from their last «اسناد لمستخدم السرية»: {clock['assignment']} from "
                f"the status history, {clock['seen']} not read yet (since the scans first saw them with their user, at "
                f"least; read first at the next scans), {clock['sent']} not assigned yet (from «تاريخ الارسال للسرية») "
                f"- reminder from {fmt_hours(int(cfg['warning_minutes']))}, breached from "
                f"{fmt_hours(int(cfg['breach_minutes']))}")
    if sink and not ctx.dry_run:                          # today's numbers kept for the day-by-day view
        try:
            from .. import performance
            from ...storage import Store
            calls_store = Store(ctx.data_dir)
            performance.save_day_log(ctx.store, calls_store, ctx.hub or {}, calls_store.users() or {}, now)
        except Exception as exc:
            ctx.log(f"SLA day log could not be kept: {type(exc).__name__}: {scrub(str(exc), secrets)}")

    by_user = summarize(requests)
    data = snapshot_data(requests, by_user, cfg)
    if left_out is not None:                              # only how many: never an SLA status or an age
        data["left_out"] = left_out
    totals = data["totals"]
    ctx.log(f"{totals['open']} open requests - {totals['reminder']} reminder, {totals['breached']} over "
            f"{fmt_hours(int(cfg['breach_minutes']))}")

    try:
        sent, note = _notify(ctx, cfg, requests, by_user, columns, now)
    except Exception as exc:                  # the scan itself is good: keep its snapshot
        sent, note = False, f"report / e-mail step failed: {clean_error(exc, secrets)}"
        ctx.log(f"SLA {note}")
    ts = now.isoformat(timespec="seconds")
    summary = f"{totals['open']} open, {totals['reminder']} reminder, {totals['breached']} breached"
    return JobResult(
        ok=True, snapshot_data=data, summary=summary + (f" - {note}" if note else ""), emails_sent=int(sent),
        history_point={"ts": ts, "open": totals["open"], "within": totals["within"],
                       "reminder": totals["reminder"], "breached": totals["breached"],
                       "groups": teams.sla_groups(requests)},          # per application + user: a team's own history
        history_extra={"first_breach": {r["key"]: ts for r in requests if r["status"] == schemas.SLA_BREACHED}})
