"""The IP phone directory: which branch / department a phone number belongs to.

Source: an Active Directory export (Excel or CSV) with a phone number, a person's name and the branch / department name.
**Only the number and the unit (branch / department) are ever kept.** The person's name column is never read into anything
that is stored, logged or sent to a page; the uploaded file itself is not kept either.

Stored (``<data_dir>/ip_phones.json``, written by ``Store.replace_ip_phones``)::

    {"version": 1, "imported_at": ISO, "imported_by": uid, "source_name": "ip_phones.xlsx", "count": N,
     "units": ["...", "..."], "numbers": {"75942": 3, "73550": [4, 9]}, "report": {...}}

``numbers`` maps a phone number to an index into ``units``, or to a list of indexes when the number was re-used by
several units (most frequent first). Reading the file is the slow part, so it is done outside any lock: ``build`` turns
the rows into the dict above and only then does the store swap the file atomically.

Everything here is plain Python (no Flask), so ``manage.py`` and the tests use it directly.
"""
import csv
import io
import re
import unicodedata
import zipfile

from .ops.engines.common import norm_arabic

FILE_VERSION = 1
MIN_DIGITS, MAX_DIGITS = 3, 15
MAX_UNIT_CHARS = 120                  # also the length a Call Log record keeps in ``ip_unit``
MAX_ROWS = 400000                     # a real export has ~42,000
MAX_UNITS = 5000                      # a real export has ~700 - a lot more means the wrong column was picked
MAX_UNZIPPED_BYTES = 400 * 1024 * 1024
MAX_ALTERNATIVES = 5
SOURCE_NAME_CHARS = 120

_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")
_NO_NUMBER, _BAD_NUMBER = "no_number", "bad_number"

# Header titles (compared without spaces, '_' and '-', in lower case, Arabic spelling folded), best name first.
NUMBER_TITLES = ("telephonenumber", "phone", "ipphone", "extension", "ext", "phonenumber", "telephone", "ipphonenumber",
                 "number", "رقمالتليفون", "رقمالهاتف", "التليفون", "الهاتف")
UNIT_TITLES = ("msexchextensionattribute27", "branch", "department", "unit", "branchname", "departmentname", "dept",
               "الفرع", "القسم", "الإدارة", "جهةالعمل")


class IpPhoneError(ValueError):
    """The file cannot be used. The message is safe to show to the administrator (it never quotes file content)."""


# ============================================================================================ cleaning
def _title_key(value):
    return re.sub(r"[\s_\-]+", "", norm_arabic("" if value is None else str(value)))


def clean_number(value):
    """A phone number from a cell -> (digits, None) or (None, 'no_number' | 'bad_number').

    Arabic-Indic digits are converted, spaces / dashes / invisible marks dropped; what is left must be 3-15 digits."""
    if value is None:
        return None, _NO_NUMBER
    if isinstance(value, bool):
        return None, _BAD_NUMBER
    if isinstance(value, int):
        text = str(value)
    elif isinstance(value, float):
        text = str(int(value)) if value.is_integer() else str(value)
    else:
        text = str(value)
    if not text.isascii() or not text.isdigit():
        text = "".join(ch for ch in text.translate(_DIGITS) if not ch.isspace() and unicodedata.category(ch) != "Cf")
    text = re.sub(r"[-‐-―−]+", "", text)
    if re.fullmatch(r"\d+\.0+", text):                     # a number typed in Excel and saved as text: 75942.0
        text = text.split(".")[0]
    if not text:
        return None, _NO_NUMBER
    if not text.isascii() or not text.isdigit() or not MIN_DIGITS <= len(text) <= MAX_DIGITS:
        return None, _BAD_NUMBER
    return text, None


def clean_unit(value):
    """The unit name for display: trimmed, one space between words, a trailing '#' and digits glued to the last word
    ('...2000') removed - but only when a name is left; otherwise the text is kept as the administrator wrote it."""
    if value is None:
        return ""
    text = "".join(ch for ch in str(value) if unicodedata.category(ch) not in ("Cc", "Cf") or ch.isspace())
    text = re.sub(r"\s+", " ", text).strip()
    for _ in range(3):
        trimmed = re.sub(r"[#\s]+$", "", text)
        trimmed = re.sub(r"(?<=[^\W\d_])\d+$", "", trimmed)
        if trimmed == text or not re.search(r"[^\W\d_]", trimmed):
            break
        text = trimmed
    if not any(ch.isalnum() for ch in text):                # '#', '-' ... : nothing to show
        return ""
    return text[:MAX_UNIT_CHARS].rstrip()


# ============================================================================================ reading
def _csv_rows(raw):
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1256", errors="replace")       # "CSV" saved by Excel on an Arabic Windows
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return csv.reader(io.StringIO(text), dialect)


def _xlsx_rows(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED_BYTES:
                raise IpPhoneError("That Excel file is too large once opened.")
    except zipfile.BadZipFile:
        raise IpPhoneError("That file is not a valid .xlsx workbook.")
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:                                       # noqa: BLE001 - any damaged workbook is the same message
        raise IpPhoneError("That Excel file could not be read.")
    sheet = wb.worksheets[0]
    try:
        sheet.reset_dimensions()                             # some exporters write a wrong <dimension>; read what is there
    except AttributeError:
        pass
    try:
        yield from sheet.iter_rows(values_only=True)
    finally:
        wb.close()


def read_rows(raw, filename):
    """Rows (tuples / lists of cell values) of an uploaded .xlsx or .csv, read lazily."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        return _xlsx_rows(raw)
    if name.endswith(".csv"):
        return _csv_rows(raw)
    raise IpPhoneError("Upload a .csv or .xlsx file.")


def _find_columns(first_row):
    """(number_col, unit_col, header_used) from the first non-empty row; falls back to the first / last column."""
    titles = [_title_key(c) for c in first_row]
    number = next((i for t in NUMBER_TITLES for i, h in enumerate(titles) if h == _title_key(t)), None)
    unit = next((i for t in UNIT_TITLES for i, h in enumerate(titles) if h == _title_key(t)), None)
    header = number is not None or unit is not None
    width = len(first_row)
    if number is None:
        number = 0
    if unit is None:
        unit = width - 1
    if number == unit:
        raise IpPhoneError("Could not tell which columns hold the phone number and the branch / department. The first row "
                           "should have the titles telephoneNumber and msExchExtensionAttribute27 (or phone and branch / department).")
    return number, unit, header


# ============================================================================================ building
def build(rows, source_name="", imported_by="", now=""):
    """Rows -> the dict that is stored. Raises ``IpPhoneError`` when nothing usable is in the file.

    Counting: ``rows`` = data rows read (blank lines and the title row are not counted); a row is skipped when it has
    no number (``no_number`` - section headers, people without a phone), a number that is not 3-15 digits
    (``bad_number``) or no unit (``no_unit``)."""
    it = iter(rows)
    first = None
    for row in it:
        if row is not None and any(c not in (None, "") and str(c).strip() for c in row):
            first = list(row)
            break
    if first is None:
        raise IpPhoneError("The file is empty.")
    number_col, unit_col, header = _find_columns(first)

    units, unit_ids, folded = [], {}, {}           # unit text -> index (first seen first); folded spelling -> index
    seen = {}                                      # number -> {unit index: how many rows}
    counts = {"rows": 0, _NO_NUMBER: 0, _BAD_NUMBER: 0, "no_unit": 0, "kept": 0}
    need = max(number_col, unit_col) + 1

    def take(row):
        if not row or not any(c not in (None, "") and str(c).strip() for c in row):
            return
        counts["rows"] += 1
        if counts["rows"] > MAX_ROWS:
            raise IpPhoneError(f"The file has more than {MAX_ROWS:,} rows - is it the right file?")
        cells = row if len(row) >= need else list(row) + [None] * (need - len(row))
        number, problem = clean_number(cells[number_col])
        if problem:
            counts[problem] += 1
            return
        unit = clean_unit(cells[unit_col])
        if not unit:
            counts["no_unit"] += 1
            return
        idx = unit_ids.get(unit)
        if idx is None:                            # a new text: the same name spelled another way (taa marbuta ...) is one unit
            idx = folded.get(fold(unit))
            if idx is None:
                idx = folded[fold(unit)] = len(units)
                units.append(unit)
            unit_ids[unit] = idx
        per = seen.get(number)
        if per is None:
            seen[number] = {idx: 1}
        else:
            per[idx] = per.get(idx, 0) + 1
        counts["kept"] += 1

    if not header:
        take(first)
    for row in it:
        take(row)

    if not seen:
        raise IpPhoneError("No phone number with a branch / department was found. The first row should have the column "
                           "titles telephoneNumber and msExchExtensionAttribute27 (or phone and branch / department).")
    if len(units) > MAX_UNITS or (len(seen) >= 200 and len(units) * 2 > len(seen)):
        raise IpPhoneError("The branch / department column looks wrong (almost every phone has a different one - "
                           "are those names of people?). Nothing was imported.")

    numbers, ambiguous = {}, 0
    for number in sorted(seen):
        per = seen[number]
        if len(per) == 1:
            numbers[number] = next(iter(per))
        else:
            ambiguous += 1
            numbers[number] = [i for i, _ in sorted(per.items(), key=lambda kv: (-kv[1], kv[0]))]
    report = {"rows": counts["rows"], "kept": counts["kept"], "numbers": len(numbers), "units": len(units),
              "ambiguous": ambiguous, "no_number": counts[_NO_NUMBER], "bad_number": counts[_BAD_NUMBER],
              "no_unit": counts["no_unit"]}
    return {"version": FILE_VERSION, "imported_at": now, "imported_by": imported_by,
            "source_name": safe_source_name(source_name), "count": len(numbers), "units": units, "numbers": numbers,
            "report": report}


def safe_source_name(name):
    base = re.split(r"[\\/]", str(name or ""))[-1]
    base = "".join(ch for ch in base if unicodedata.category(ch) not in ("Cc", "Cf"))
    return re.sub(r"\s+", " ", base).strip()[:SOURCE_NAME_CHARS]


def build_from_bytes(raw, filename, imported_by="", now=""):
    return build(read_rows(raw, filename), source_name=filename, imported_by=imported_by, now=now)


# ============================================================================================ matching
_DASHES = re.compile(r"\s*[-‐-―−]\s*")
_PREFIX = re.compile(r"^(?:فرع|فروع|اداره|مكتب)\s+(?=\S)")      # already folded: إدارة -> اداره
_SUFFIX = re.compile(r"[\s#]*\d+[\s#]*$|[\s#]+$")


def fold(text):
    """One spelling for Arabic text: ``norm_arabic`` (tatweel, diacritics, alef / yaa / taa marbuta, digits, spaces) plus
    Persian digits, one kind of dash and no '#'. Used to compare names and for the search boxes."""
    s = norm_arabic(str(text or "").translate(_DIGITS)).replace("#", " ")
    return re.sub(r"\s+", " ", _DASHES.sub("-", s)).strip()


def match_key(text, loose=False):
    """The comparison form of a unit / branch name: folded, without a leading 'فرع' / 'فروع' / 'إدارة' / 'مكتب'. ``loose``
    also drops trailing digits ('... 2', '... 2000'), keeping the name when that would leave nothing."""
    s = fold(text)
    stripped = _PREFIX.sub("", s)
    s = stripped or s
    if loose:
        trimmed = _SUFFIX.sub("", s).strip()
        s = trimmed or s
    return s


class IpDirectory:
    """The parsed directory plus the Call Log branch list, ready for dictionary look-ups. Treat as read-only."""

    def __init__(self, data, branches):
        data = data if isinstance(data, dict) else {}
        units = data.get("units")
        self.units = [u if isinstance(u, str) else "" for u in units] if isinstance(units, list) else []
        numbers = data.get("numbers")
        self.numbers = numbers if isinstance(numbers, dict) else {}
        self.meta = {k: data.get(k) for k in ("version", "imported_at", "imported_by", "source_name", "count", "report")}
        self.branches = branches if isinstance(branches, list) else []
        self._idx = None                      # (exact, loose) name -> [branch]; built on first use
        self._branch_of = {}                  # unit index -> branch dict | None

    def __bool__(self):
        return bool(self.numbers)

    # ------------------------------------------------------------------ branches
    def _index(self):
        if self._idx is None:
            exact, loose = {}, {}
            for b in self.branches:
                if not isinstance(b, dict) or not b.get("name"):
                    continue
                exact.setdefault(match_key(b["name"]), []).append(b)
                loose.setdefault(match_key(b["name"], loose=True), []).append(b)
            self._idx = (exact, loose)
        return self._idx

    def branch_for(self, unit):
        """The one Call Log branch whose name equals ``unit`` (after normalisation), else None. Never guesses: no match or
        several matches -> None. An exact comparison is tried first; trailing digits are dropped only when it finds nothing."""
        exact, loose = self._index()
        hits = exact.get(match_key(unit))
        if hits is None:
            hits = loose.get(match_key(unit, loose=True))
        if hits and len(hits) == 1:
            b = hits[0]
            return {"id": b.get("id", ""), "code": b.get("code", ""), "name": b.get("name", ""), "region": b.get("region", "")}
        return None

    def _unit_branch(self, idx):
        if idx not in self._branch_of:
            self._branch_of[idx] = self.branch_for(self.units[idx])
        return self._branch_of[idx]

    # ------------------------------------------------------------------- look-up
    def _indexes(self, number):
        ref = self.numbers.get(number)
        if isinstance(ref, int) and not isinstance(ref, bool):
            ref = [ref]
        if not isinstance(ref, list):
            return []
        return [i for i in ref if isinstance(i, int) and not isinstance(i, bool) and 0 <= i < len(self.units) and self.units[i]]

    def lookup(self, number):
        """``{"number","unit","alternatives","branch"}`` or None when the directory does not know the number."""
        idx = self._indexes(number)
        if not idx:
            return None
        return {"number": number, "unit": self.units[idx[0]],
                "alternatives": [self.units[i] for i in idx[1:1 + MAX_ALTERNATIVES]], "branch": self._unit_branch(idx[0])}

    def status(self):
        """Numbers / units / ambiguity / how many units are Call Log branches - for the administrator's card."""
        if not self.numbers:
            return {"imported": False, "numbers": 0, "units": 0, "ambiguous": 0, "branch_units": 0}
        ambiguous = sum(1 for v in self.numbers.values() if isinstance(v, list) and len(v) > 1)
        branch_units = sum(1 for i in range(len(self.units)) if self._unit_branch(i))
        return {"imported": True, "numbers": len(self.numbers), "units": len(self.units), "ambiguous": ambiguous,
                "branch_units": branch_units}


def normalize_lookup_number(value):
    """The number a page asked about -> digits, or None when it cannot be a phone number at all."""
    number, problem = clean_number(value)
    return None if problem else number
