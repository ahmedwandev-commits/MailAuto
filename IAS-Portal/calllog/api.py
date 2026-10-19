"""JSON API used by the two dashboards."""
import csv
import io
import re
from collections import Counter

import hashlib
import json

from flask import Blueprint, current_app, g, jsonify, request

from . import auth, branches as branch_list, ip_phones, overview as overview_data, stats
from .ops import settings as hub_settings
from .ops.overlay import (UNCLASSIFIED, legacy_team, normalize_email, split_teams, user_teams,
                          valid_email)
from .ops.schemas import ALL_TEAMS_MARK
from .queries import (ApiError, clean_ip_phone, clean_note, clean_reason,
                      decorate, filtered_raw, json_body, paginate, parse_day, parse_range,
                      reasons_by_key)
from .security import generate_password, hash_password, password_problem
from .storage import (STATUSES, STATUS_LABELS, StorageError, iso_now, new_record_id, norm_uid,
                      parse_record_id, valid_uid)

bp = Blueprint("api", __name__)

ROLES = ("user", "manager", "admin")
ROLE_ORDER = {"admin": 0, "manager": 1, "user": 2}
MAX_SECURITY_NAME = 120
MAX_TEAMS_PER_ACCOUNT = 40


def public_user(u):
    """The profile as pages see it. ``teams`` is the list of teams (['__all__'] = every team); ``team`` is the older
    single value (first team, '__all__' or '') kept so existing pages and scripts keep working."""
    teams = user_teams(u)
    return {"user_id": u["user_id"], "name": u["name"], "department": u.get("department", ""),
            "role": u["role"], "email": u.get("email", ""), "team": legacy_team(teams), "teams": teams,
            "security_name": u.get("security_name", "")}


def admin_user_row(u):
    return {**public_user(u), "active": u.get("active", True),
            "must_change_password": bool(u.get("must_change_password")),
            "last_login": u.get("last_login"), "created_at": u.get("created_at")}


def edit_window_open(day):
    """May a non-admin still change a record from ``day``?"""
    limit = current_app.config["USER_EDIT_DAYS"]
    if limit is None:
        return True
    from datetime import date
    return (date.fromisoformat(auth.today()) - date.fromisoformat(day)).days <= int(limit)


# ====================================================================== common
@bp.get("/me")
@auth.login_required
def me():
    return jsonify(
        user=public_user(g.user),
        reasons=auth.store().reasons(),
        status_labels=STATUS_LABELS,
        csrf=auth.csrf_token(),
        today=auth.today(),
        server_time=iso_now(auth.now()),
        user_edit_days=current_app.config["USER_EDIT_DAYS"],
    )


# =================================================================== my records
def day_payload(day, uid):
    users = auth.store().users()
    reasons = reasons_by_key()
    recs = auth.store().day_records(day, uid)
    rows = [decorate(r, users, reasons) for r in reversed(recs)]    # reversed: among equal seconds the later call comes first
    rows.sort(key=lambda r: r["time"], reverse=True)
    editable = edit_window_open(day)
    for row in rows:
        row["editable"] = editable
    summary = stats.summarize(recs, reasons)
    return {
        "date": day, "today": auth.today(), "server_time": iso_now(auth.now()), "records": rows,
        "stats": {**{k: summary[k] for k in ("total", "first_call", "last_call", "by_reason", *STATUSES)},
                  "hourly": stats.buckets(recs, reasons, "hour")},
    }


# ------------------------------------------------------------------ branches
def branch_fields(branch_id):
    """The branch snapshot stored on a call. Empty = the caller is not a branch (sector / other)."""
    if not branch_id:
        return {"branch": "", "branch_name": "", "region": ""}
    entry = auth.store().branch_by_id(str(branch_id))
    if not entry:
        raise ApiError("That branch is not in the list. Reload the page and choose it again.")
    return {"branch": entry["code"], "branch_name": entry["name"], "region": entry["region"]}


def branch_label(rec):
    return f"{rec['branch']} · {rec['branch_name']}" if rec.get("branch") else ""


@bp.get("/branches")
@auth.login_required
def list_branches():
    items = auth.store().branches()
    body = json.dumps({"branches": items}, ensure_ascii=False, separators=(",", ":"))
    resp = current_app.response_class(body, mimetype="application/json")
    resp.set_etag(hashlib.sha1(body.encode("utf-8")).hexdigest())
    resp.headers["Cache-Control"] = "no-cache"      # the browser re-validates; unchanged list = tiny 304
    return resp.make_conditional(request)


# ------------------------------------------------------------- IP phone directory
def _same_branch(branch, entry):
    """Is the branch fields ``branch`` (code / name / region of a call) the directory's ``entry`` (a branch dict or None)?"""
    return bool(entry and branch.get("branch")) and (
        (entry["code"], entry["name"], entry["region"]) == (branch["branch"], branch["branch_name"], branch["region"]))


def ip_unit_of(ip_phone):
    """The branch / department the directory gives for ``ip_phone`` (looked up here - a page's own claim is never believed)
    and the branch it maps to. ('', None) when the number is not in the directory."""
    hit = auth.store().ip_directory().lookup(ip_phone)
    if not hit:
        return "", None
    return hit["unit"][:ip_phones.MAX_UNIT_CHARS], hit["branch"]


def ip_unit_auto(requested, branch, mapped):
    """May the call say its branch was filled in automatically? Only when the page said so *and* the saved branch really
    is the one the directory gives for the number."""
    return requested is True and _same_branch(branch, mapped)


@bp.get("/ip-phones/<number>")
@auth.login_required
def ip_phone_lookup(number):
    """Branch / department of an IP phone number (the call form asks while the employee types). Cheap: a dictionary read."""
    clean = ip_phones.normalize_lookup_number(number[:40])
    if not clean:
        raise ApiError("IP phone must be digits only, for example 75942.")
    hit = auth.store().ip_directory().lookup(clean)
    if not hit:
        return jsonify(found=False, number=clean)
    return jsonify(found=True, number=hit["number"], unit=hit["unit"], alternatives=hit["alternatives"], branch=hit["branch"])


@bp.get("/records")
@auth.login_required
def my_records():
    day = parse_day(request.args.get("date"), auth.today())
    if day > auth.today():
        raise ApiError("That date is in the future.")
    return jsonify(day_payload(day, g.user["user_id"]))


@bp.post("/records")
@auth.login_required
def create_record():
    if g.user["role"] != "user":
        raise ApiError("Administrator and manager accounts do not log calls. Use an employee account.", 403)
    body = json_body()
    ip_phone = clean_ip_phone(body.get("ip_phone"))
    reason = clean_reason(body.get("reason"))
    note = clean_note(body.get("note"))
    branch = branch_fields(body.get("branch"))
    unit, mapped = ip_unit_of(ip_phone)
    auto = ip_unit_auto(body.get("ip_unit_auto"), branch, mapped)
    stamp = auth.now().replace(microsecond=0)
    day = stamp.date().isoformat()
    rec = {
        "id": new_record_id(day, g.user["user_id"]), "user_id": g.user["user_id"],
        "date": day, "time": stamp.time().isoformat(), "ip_phone": ip_phone,
        "reason": reason, "note": note, **branch, "created_at": iso_now(stamp),
        "updated_at": None, "updated_by": None,
    }
    if unit:
        rec["ip_unit"] = unit
    if auto:
        rec["ip_unit_auto"] = True
    auth.store().add_record(rec)
    auth.audit("record_created", target=rec["id"],
               details={"ip_phone": ip_phone, "reason": reason, "note": note, "branch": branch_label(rec),
                        **({"ip_unit": unit} if unit else {}), **({"branch_auto": True} if auto else {})})
    return jsonify(record=decorate(rec, auth.store().users(), reasons_by_key())), 201


def _owned_record_ref(rid):
    """(day, owner) of ``rid`` if the signed-in user may touch it, else 404 / 403."""
    parsed = parse_record_id(rid)
    if not parsed:
        raise ApiError("Record not found.", 404)
    day, owner = parsed
    is_admin = g.user["role"] == "admin"
    if owner != g.user["user_id"] and not is_admin:
        raise ApiError("Record not found.", 404)
    if not is_admin and not edit_window_open(day):
        raise ApiError(
            f"Records older than {current_app.config['USER_EDIT_DAYS']} days are locked. "
            "Ask your administrator.", 403)
    return day, owner


@bp.put("/records/<rid>")
@auth.login_required
def update_record(rid):
    _, owner = _owned_record_ref(rid)
    body = json_body()
    changes = {}
    if "ip_phone" in body:
        changes["ip_phone"] = clean_ip_phone(body["ip_phone"])
    if "reason" in body:
        changes["reason"] = clean_reason(body["reason"])
    if "note" in body:
        changes["note"] = clean_note(body["note"])
    if "branch" in body:
        changes.update(branch_fields(body["branch"]))
    if not changes:
        raise ApiError("Nothing to change.")
    current = auth.store().get_record(rid)
    if not current:
        raise ApiError("Record not found.", 404)
    changes = {k: v for k, v in changes.items() if current.get(k, "") != v}
    if not changes:
        return jsonify(record=decorate(current, auth.store().users(), reasons_by_key()))
    changes.update(_ip_changes(current, changes, body))
    result = auth.store().update_record(rid, {
        **changes, "updated_at": iso_now(auth.now()), "updated_by": g.user["user_id"]})
    if not result:
        raise ApiError("Record not found.", 404)
    before, after = result
    audit_keys = {("branch" if k in ("branch", "branch_name", "region") else k) for k in changes} - {"ip_unit_auto"}
    view = lambda rec: {k: (branch_label(rec) if k == "branch" else rec.get(k)) for k in audit_keys}
    auth.audit("record_updated", target=rid, details={
        "owner": owner, "before": view(before), "after": view(after)})
    return jsonify(record=decorate(after, auth.store().users(), reasons_by_key()))


def _ip_changes(current, changes, body):
    """The IP-phone fields an edit changes. The unit is looked up again only when the IP phone itself changed (a call keeps
    the unit it had when it was made); 'filled automatically' is re-checked when the IP phone or the branch changed, against
    the branch of the unit the call carries."""
    out = {}
    if "ip_phone" not in changes and not any(k in changes for k in ("branch", "branch_name", "region")):
        return out
    if "ip_phone" in changes:
        unit, mapped = ip_unit_of(changes["ip_phone"])
        if unit != current.get("ip_unit", ""):
            out["ip_unit"] = unit
    else:
        unit = current.get("ip_unit", "")
        mapped = auth.store().ip_directory().branch_for(unit) if unit else None
    branch = {k: changes.get(k, current.get(k, "")) for k in ("branch", "branch_name", "region")}
    if ip_unit_auto(body.get("ip_unit_auto"), branch, mapped) != bool(current.get("ip_unit_auto")):
        out["ip_unit_auto"] = not bool(current.get("ip_unit_auto"))
    return out


@bp.delete("/records/<rid>")
@auth.login_required
def delete_record(rid):
    _, owner = _owned_record_ref(rid)
    removed = auth.store().delete_record(rid)
    if not removed:
        raise ApiError("Record not found.", 404)
    auth.audit("record_deleted", target=rid, details={
        "owner": owner, "branch": branch_label(removed), **{k: removed.get(k) for k in
                                                           ("date", "time", "ip_phone", "reason", "note")},
        **({"ip_unit": removed["ip_unit"]} if removed.get("ip_unit") else {})})
    return jsonify(ok=True)


# ===================================================================== admin
@bp.get("/admin/overview")
@auth.admin_required
def overview():
    d_from, d_to = parse_range()

    def compute():
        reasons = reasons_by_key()
        users = auth.store().users()
        data = overview_data.merged(d_from, d_to)
        totals = data["totals"]
        rows = stats.employee_rows(users, data["users"])
        kind = "hour" if d_from == d_to else "day"
        employees = [u for u in users.values() if u["role"] == "user" and u.get("active", True)]
        branch_rows = stats.branch_rows(data["branches"])
        with_branch = sum(r["total"] for r in branch_rows)
        return dict(
            range={"from": d_from, "to": d_to},
            totals={k: totals[k] for k in ("total", *STATUSES)} | {
                "active_employees": sum(1 for r in rows if r["total"] > 0), "employees": len(employees)},
            chart={"kind": kind, "buckets": stats.chart_buckets(data["rollups"], kind, d_from, d_to)},
            by_reason=data["by_reason"], employees=rows,
            recent=[decorate(r, users, reasons) for r in overview_data.recent_calls(d_from, d_to)],
            departments=sorted({u.get("department", "") for u in employees if u.get("department")}),
            branches={"by_region": stats.per_region(branch_rows), "top": branch_rows[:10],
                      "with_branch": with_branch, "without_branch": totals["total"] - with_branch,
                      "distinct": len(branch_rows)},
        )

    data = auth.store().memo(("overview", d_from, d_to), compute)
    return jsonify(**data, server_time=iso_now(auth.now()))


@bp.get("/admin/records")
@auth.admin_required
def admin_records():
    d_from, d_to = parse_range()
    rows = filtered_raw(d_from, d_to, request.args)
    page = paginate(rows)
    users, reasons = auth.store().users(), reasons_by_key()
    page["items"] = [{**decorate(r, users, reasons), "editable": True} for r in page["items"]]
    return jsonify(**page, range={"from": d_from, "to": d_to}, server_time=iso_now(auth.now()))


@bp.get("/admin/users")
@auth.admin_required
def list_users():
    users = sorted(auth.store().users().values(), key=lambda u: (ROLE_ORDER.get(u["role"], 9), u["name"]))
    return jsonify(users=[admin_user_row(u) for u in users])


def _hub_teams():
    return hub_settings.load(current_app.extensions["ops"])["email"]["teams"]


def _clean_teams(value, role, current_teams=()):
    """The teams of an account from a list or a text such as 'Flex Cube; Card Management'.

    Each name must be one of the teams in the Hub settings (any spelling of upper / lower case is accepted) - except a
    team the account already has, so a team renamed or removed in the settings never blocks editing the account.
    '__all__' (every team) is for managers only and stands alone. Employees cannot be put in "Unclassified": that
    bucket is for managers to sort out. Returns the list to store."""
    known = {t.casefold(): t for t in _hub_teams()}
    mine = {t.casefold(): t for t in current_teams}
    if isinstance(value, str) and value.strip().casefold() in {**known, **mine}:
        names = [value.strip()]                              # a team whose own name contains ',' or ';'
    else:
        names = split_teams(value)
    out = []
    for name in names:
        if len(name) > 80:
            raise ApiError("A team name is too long (max 80 characters).")
        if name.casefold() == ALL_TEAMS_MARK:
            if role == "user":
                raise ApiError("Only managers can follow every team. Tick the employee's own teams instead.")
            return [ALL_TEAMS_MARK]
        canonical = known.get(name.casefold()) or mine.get(name.casefold())
        if canonical is None:
            raise ApiError(f"'{name}' is not one of the teams. Choose teams from the list in the settings.")
        if role == "user" and canonical.casefold() == UNCLASSIFIED.casefold() and canonical.casefold() not in mine:
            raise ApiError(f"Employees cannot be put in '{canonical}': that bucket is for managers to sort out.")
        if canonical not in out:
            out.append(canonical)
    if len(out) > MAX_TEAMS_PER_ACCOUNT:
        raise ApiError(f"Too many teams (max {MAX_TEAMS_PER_ACCOUNT}).")
    return out


def _email_in_use(email, users, own_uid=None):
    return any(u.get("email") == email and uid != own_uid for uid, u in users.items())


def _clean_profile(body, partial=False, current=None):
    """Checked profile fields. ``current`` = the account being edited (its own e-mail may stay)."""
    current = current or {}
    out = {}
    if not partial or "name" in body:
        name = str(body.get("name") or "").strip()
        if not 1 <= len(name) <= 80:
            raise ApiError("Enter the employee's name (up to 80 characters).")
        out["name"] = name
    if not partial or "department" in body:
        dept = str(body.get("department") or "").strip()
        if len(dept) > 80:
            raise ApiError("Department is too long (max 80 characters).")
        out["department"] = dept
    if "role" in body:
        if body["role"] not in ROLES:
            raise ApiError("Role must be 'user', 'manager' or 'admin'.")
        out["role"] = body["role"]
    if "email" in body:
        email = normalize_email(body["email"])
        if email and (len(email) > 120 or not valid_email(email)):
            raise ApiError("That e-mail address is not valid.")
        out["email"] = email
    if "teams" in body or "team" in body:
        # `teams` (a list or 'A; B') wins when both are sent; the single `team` is the older way to say one team
        raw = body["teams"] if "teams" in body else body["team"]
        role = out.get("role") or current.get("role") or "user"
        teams = _clean_teams(raw, role, user_teams(current))
        out["teams"], out["team"] = teams, legacy_team(teams)
    if "security_name" in body:
        security_name = " ".join(str(body["security_name"] or "").split())
        if len(security_name) > MAX_SECURITY_NAME:
            raise ApiError(f"Security name is too long (max {MAX_SECURITY_NAME} characters).")
        out["security_name"] = security_name
    return out


def _new_user(uid, profile, password):
    return {"user_id": uid, "role": "user", "active": True, "must_change_password": True,
            "created_at": iso_now(auth.now()), "last_login": None,
            "password_hash": hash_password(password), **profile}


@bp.post("/admin/users")
@auth.admin_required
def create_user():
    body = json_body()
    uid = norm_uid(body.get("user_id"))
    if not valid_uid(uid):
        raise ApiError("User ID may use letters, digits, dot, dash and underscore (max 32).")
    profile = _clean_profile(body)
    if profile.get("email") and _email_in_use(profile["email"], auth.store().users()):
        raise ApiError("Another account already uses that e-mail address.", 409)
    password = str(body.get("password") or "")
    if password:
        problem = password_problem(password)
        if problem:
            raise ApiError(problem)
    else:
        password = generate_password()
    try:
        user = auth.store().create_user(_new_user(uid, {"role": "user", **profile}, password))
    except StorageError as exc:
        raise ApiError(str(exc), 409)
    teams = user_teams(user)
    auth.audit("user_created", target=uid, details={
        "name": user["name"], "role": user["role"], **({"teams": teams} if teams else {}),
        **{k: user[k] for k in ("email", "security_name") if user.get(k)}})
    return jsonify(user=admin_user_row(user), temp_password=password), 201


@bp.put("/admin/users/<uid>")
@auth.admin_required
def update_user(uid):
    uid = norm_uid(uid)
    body = json_body()
    changes = _clean_profile(body, partial=True, current=auth.store().get_user(uid))
    if "active" in body:
        changes["active"] = bool(body["active"])

    def guard(users, after):
        if after.get("email") and _email_in_use(after["email"], users, uid):
            raise ApiError("Another account already uses that e-mail address.", 409)
        if uid == g.user["user_id"] and (not after.get("active", True) or after["role"] != "admin"):
            raise ApiError("You cannot disable or demote your own account.", 409)
        if not any(u.get("active", True) and u["role"] == "admin"
                   for k, u in users.items() if k != uid) and not (
                after.get("active", True) and after["role"] == "admin"):
            raise ApiError("At least one active administrator is required.", 409)

    result = auth.store().update_user(uid, changes, guard)
    if not result:
        raise ApiError("User not found.", 404)
    before, after = result
    diff = {k: {"from": before.get(k), "to": after.get(k)} for k in changes
            if k not in ("team", "teams") and before.get(k) != after.get(k)}
    if "teams" in changes and user_teams(before) != user_teams(after):
        diff["teams"] = {"from": user_teams(before), "to": user_teams(after)}
    if diff:
        auth.audit("user_updated", target=uid, details=diff)
    return jsonify(user=admin_user_row(after))


@bp.post("/admin/users/<uid>/reset-password")
@auth.admin_required
def reset_password(uid):
    uid = norm_uid(uid)
    body = request.get_json(silent=True) or {}
    password = str(body.get("password") or "")
    if password:
        problem = password_problem(password)
        if problem:
            raise ApiError(problem)
    else:
        password = generate_password()
    result = auth.store().update_user(uid, {
        "password_hash": hash_password(password), "must_change_password": True})
    if not result:
        raise ApiError("User not found.", 404)
    auth.audit("password_reset", target=uid)
    return jsonify(temp_password=password)


# ------------------------------------------------------------- bulk import
_HEADER_ALIASES = {
    "user_id": {"user_id", "userid", "user id", "id", "employee id", "employee_id", "username"},
    "name": {"name", "full name", "employee name", "employee"},
    "department": {"department", "dept", "section"},
    "password": {"password", "pwd"},
    "role": {"role", "account type"},
    "email": {"email", "e-mail", "mail", "email address"},
    "team": {"team"},
    "teams": {"teams", "team list", "team names"},
    "security_name": {"security_name", "security name", "security user"},
}


def _read_rows(upload, prefer_sheets=()):
    """Rows of a .csv or .xlsx upload as lists of strings (for .xlsx: the preferred sheet, else the first)."""
    filename = (upload.filename or "").lower()
    raw = upload.read()
    if filename.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = next((ws for ws in wb.worksheets if ws.title.strip().lower() in prefer_sheets), wb.worksheets[0])
        return [["" if c is None else str(c).strip() for c in row] for row in sheet.iter_rows(values_only=True)]
    if filename.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    raise ApiError("Upload a .csv or .xlsx file.")


@bp.post("/admin/users/import")
@auth.admin_required
def import_users():
    upload = request.files.get("file")
    if not upload:
        raise ApiError("Choose a file to import.")
    rows = [r for r in _read_rows(upload) if any(r)]
    if not rows:
        raise ApiError("The file is empty.")
    header = [h.strip().lower() for h in rows[0]]
    col = {}
    for field, aliases in _HEADER_ALIASES.items():
        for i, h in enumerate(header):
            if h in aliases:
                col[field] = i
                break
    if "user_id" not in col or "name" not in col:
        raise ApiError("The first row must have column titles including 'user_id' and 'name' "
                       "(optional: 'department', 'password', 'role', 'email', 'teams', 'team', 'security_name').")

    def cell(row, field):
        i = col.get(field)
        return row[i].strip() if i is not None and i < len(row) else ""

    to_create, skipped, seen, credentials = [], [], set(), []
    taken = {u["email"] for u in auth.store().users().values() if u.get("email")}
    for line_no, row in enumerate(rows[1:], start=2):
        uid = norm_uid(cell(row, "user_id").removesuffix(".0"))  # Excel numbers arrive as 230966.0
        name, dept, password = cell(row, "name"), cell(row, "department"), cell(row, "password")
        problem, extra = None, {}
        try:
            fields = {k: cell(row, k) for k in ("role", "email", "team", "teams", "security_name") if k in col}
            if fields.get("role", "user") not in ("", "user", "manager"):
                raise ApiError("Role must be 'user' or 'manager'")
            if fields.get("teams"):                  # `teams` (A;B or A,B) wins over the older single `team` column
                fields.pop("team", None)
            extra = _clean_profile({k: v for k, v in fields.items() if k != "role" and v != ""}, partial=True,
                                   current={"role": fields.get("role") or "user"})
            if fields.get("role"):
                extra["role"] = fields["role"]
        except ApiError as exc:
            problem = exc.message.rstrip(".")
        if problem:
            pass
        elif not valid_uid(uid):
            problem = "Invalid User ID"
        elif not name or len(name) > 80 or len(dept) > 80:
            problem = "Missing or too-long name/department"
        elif uid in seen:
            problem = "Duplicate User ID in the file"
        elif extra.get("email") in taken:
            problem = "E-mail address already used by another account"
        elif password and password_problem(password):
            problem = password_problem(password)
        if problem:
            skipped.append({"row": line_no, "user_id": uid, "reason": problem})
            continue
        seen.add(uid)
        if extra.get("email"):
            taken.add(extra["email"])
        password = password or generate_password()
        to_create.append(_new_user(uid, {"name": name, "department": dept, **extra}, password))
        credentials.append({"user_id": uid, "name": name, "password": password})

    created, existing = auth.store().create_users(to_create)
    created_ids = {u["user_id"] for u in created}
    for uid in existing:
        skipped.append({"row": None, "user_id": uid, "reason": "User ID already exists"})
    credentials = [c for c in credentials if c["user_id"] in created_ids]
    auth.audit("users_imported", details={"created": len(created), "skipped": len(skipped)})
    return jsonify(created=len(created), skipped=skipped, credentials=credentials)


# -------------------------------------------------------------- branch list import
@bp.post("/admin/branches/import")
@auth.admin_required
def import_branches():
    upload = request.files.get("file")
    if not upload:
        raise ApiError("Choose a file to import.")
    rows = _read_rows(upload, prefer_sheets=("الفروع", "branches", "branch"))
    try:
        entries, report = branch_list.parse_rows(rows)
    except ValueError as exc:
        raise ApiError(str(exc))
    if not entries:
        raise ApiError("No branches were found in the file.")
    previous = len(auth.store().branches())
    auth.store().replace_branches(entries)
    auth.audit("branches_imported", details={"count": len(entries), "previous": previous, **report})
    return jsonify(count=len(entries), previous=previous, **report)


# ------------------------------------------------------- IP phone directory import
def ip_phone_status():
    """What the administrator's card shows: how many numbers / units, when and by whom, and the last import's counts."""
    directory = auth.store().ip_directory(strict=True)
    status = directory.status()
    meta = directory.meta
    return {**status, "imported_at": meta.get("imported_at") or "", "imported_by": meta.get("imported_by") or "",
            "source_name": meta.get("source_name") or "", "report": meta.get("report") or {},
            "has_previous": auth.store().ip_phones_previous_path.is_file()}


@bp.get("/admin/ip-phones")
@auth.admin_required
def ip_phones_status():
    return jsonify(**ip_phone_status())


@bp.post("/admin/ip-phones/import")
@auth.admin_required
def import_ip_phones():
    """Replace the IP phone directory with an uploaded Excel / CSV export (20 MB at most - see ``app.LARGE_UPLOADS``).
    Only the number and the branch / department are kept: the person's name column is never stored or logged."""
    upload = request.files.get("file")
    if not upload:
        raise ApiError("Choose a file to import.")
    try:
        data = ip_phones.build_from_bytes(upload.read(), upload.filename, imported_by=g.user["user_id"],
                                          now=iso_now(auth.now()))             # the slow part: no lock is held
    except ip_phones.IpPhoneError as exc:
        raise ApiError(str(exc))
    previous = auth.store().ip_directory().status()["numbers"]
    auth.store().replace_ip_phones(data)
    auth.audit("ip_phones_imported", target=data["source_name"], details={**data["report"], "previous": previous})
    return jsonify(**ip_phone_status(), previous=previous)


# ======================================================================== audit
@bp.get("/admin/audit")
@auth.admin_required
def audit_log():
    d_from, d_to = parse_range()
    actor = norm_uid(request.args.get("actor"))
    action = request.args.get("action") or ""
    q = (request.args.get("q") or "").strip().lower()
    entries = auth.store().read_audit(d_from, d_to)
    actions = sorted(Counter(e.get("action") for e in entries))
    out = []
    for e in entries:
        if actor and actor not in (e.get("actor") or "").lower() \
                and actor not in (e.get("actor_name") or "").lower():
            continue
        if action and e.get("action") != action:
            continue
        if q and q not in re.sub(r"\s+", " ", str(e)).lower():
            continue
        out.append(e)
    out.reverse()  # files are chronological; show newest first
    return jsonify(**paginate(out, default_size=100), actions=actions)
