"""
Oracle User Provisioning SQL Generator
=======================================
Reads a user list from an Excel file, pulls each user's full record from the
SOURCE Oracle DB (fcuatdb-scan/PRODPDB) and writes a ready-to-run SQL*Plus
script that re-provisions those users on the TARGET Oracle DB (fcshr-scan/PRODPDB).

HOW TO RUN
----------
  1. Edit the DB config below if needed.
  2. Place your user IDs in the Excel file (see EXCEL_FILE path below).
  3. Run the script:
         python provision_users.py
     or with optional overrides:
         python provision_users.py --date 2026-06-01
         python provision_users.py --excel "C:\\other\\path\\Users.xlsx"
         python provision_users.py --users ADMNMWUT,USER2 --date 2026-06-01
         python provision_users.py --users-file users.txt --date 2026-06-01
  4. Output appears next to this script:
       generated_sql/generated_users_part001.sql, part002.sql, ...
                                     <- run these on the TARGET DB, in order
       generated_users_report.csv   <- summary of what was found

TABLES CLONED
-------------
  Each user is cloned across the full set of user/PDATA tables listed in
  SINGLE_ROW_TABLES, MULTI_ROW_TABLES and ADDITIONAL_TABLES below. The
  "additional" tables are configured automatically at run time: the script
  reads each one's user-id column and primary key from the Oracle data
  dictionary (ALL_TAB_COLUMNS / ALL_CONSTRAINTS), so no column names have to
  be hard-coded and a table that does not exist in the source schema is simply
  skipped — same as the explicit tables.

PERFORMANCE NOTES
-----------------
  - Data is read from Oracle in BATCHES (one query per table per chunk of
    up to 900 users via a bound IN-list), not one query per table per user.
    Table existence is also checked once, not per user. This turns millions of
    per-user round-trips (tens of thousands of users x ~28 tables) into a few
    hundred.
  - Each chunk of users (default 900) is fetched, converted to SQL, and
    written to its own file immediately — only one chunk's data is ever
    held in memory, so this no longer runs out of RAM on large user lists.
  - Tune chunk size with --chunk-size (max 900, Oracle's IN-list limit).

DELETE-THEN-INSERT (every user)
-------------------------------
  - In the generated SQL, every user is ONE PL/SQL block (one transaction):
      0. check TARGET: DELETE the user's rows from all configured tables
         (children first, SSZB_USER last). If rows were found the user existed
         and is now removed; if 0 rows, the user was not in TARGET.
      1. INSERT the user fresh from SOURCE (same data / date rules as before).
      2. COMMIT.
    On ANY error the block ROLLS BACK, so that user stays in TARGET exactly
    as it was (never half-deleted), a "FAILED" line with the Oracle error is
    printed, and the script continues with the next user.
  - The spool log shows one line per user:  OK - existed in TARGET (N rows
    deleted) -> re-added   /   OK - not in TARGET -> added   /   FAILED ...
  - SMZB_MSGS_RIGHTS / SMZB_QUEUE_RIGHTS: only the user's own row
    (USER_ROLE_FLAG = 'U', see MSGS_RIGHTS_USER_FLAG) is deleted, never a
    role's row. For the additional role/queue-rights tables this filter is
    applied automatically whenever the table is keyed on USER_ROLE_ID and has
    a USER_ROLE_FLAG column.
  - The old "UPDATE if different" statements are gone: after the delete the
    INSERT always writes the full row, with every date column = RUN_DATE.

Key design rules
----------------
  - Never modifies Oracle data; only SELECT queries are executed against source.
    (The generated SQL DELETEs + INSERTs on TARGET — see above.)
  - The date is supplied as a command-line parameter (or prompted if omitted).
    ALL date columns in the generated SQL use Oracle TO_DATE with an explicit
    format mask (YYYY-MM-DD), so the output is NLS-safe.
  - Source schema (read from) and target schema (written into SQL) are
    completely independent and both configurable below.
  - Excel file is the PRIMARY source of user IDs; --users and --users-file
    are fallback/override options.

Requirements
------------
  pip install oracledb openpyxl
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import time
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, Iterator, List, Optional, Tuple

try:
    import oracledb
except ImportError as exc:
    raise SystemExit(
        "Missing dependency: oracledb\n"
        "Install with:  pip install oracledb"
    ) from exc

try:
    import openpyxl
except ImportError as exc:
    raise SystemExit(
        "Missing dependency: openpyxl\n"
        "Install with:  pip install openpyxl"
    ) from exc


# =============================================================================
# RUN DATE  — set interactively at startup (see main)
# Every date column in the generated SQL uses this value via TO_DATE().
# =============================================================================
RUN_DATE: date = date.today()   # overwritten in main() before any SQL is built


# =============================================================================
# ★  CONFIGURATION  — edit these values if anything changes  ★
# =============================================================================

# ── EXCEL INPUT ───────────────────────────────────────────────────────────────
EXCEL_FILE = Path(
    r"D:\Felex\access matrix  vs system pipeline\BS Scripts\Scripts\SA_USERS\Users.xlsx"
)
# Column name(s) to search for user IDs in the Excel file.
# The script tries each alias in order and uses the first match.
EXCEL_USER_ID_COLUMN_ALIASES = [
    "USER_ID", "USERID", "USER ID", "ID",
    "EMP_ID", "EMPID", "USERNAME", "USER NAME",
]
# Which sheet to read (0 = first sheet, or supply a name e.g. "Sheet1")
EXCEL_SHEET = 0

# ── SOURCE DB (read-only; UAT / PRODPDB on fcuatdb) ──────────────────────────
# Alternative environments (set SRC_DSN / SRC_USER / SRC_PWD accordingly):
#     PREPROD : fcpreprd-scan:1521/PRODPDB
#     BS/SHR  : fcshr-scan:1521/PRODPDB
#     UAT     : fcuatdb-scan:1521/PRODPDB   (default below)
#
# Credentials come from environment variables so NO password lives in this
# file (safe to commit / share). Set them before running, e.g.
#     set SRC_USER=your_db_user    (Windows cmd)
#     set SRC_PWD=your_password
#     set SRC_DSN=fcuatdb-scan:1521/PRODPDB
# or export them (Linux/macOS), or load a .env file. For a quick local-only run
# you MAY assign them literally here instead — but never commit a real password.
SRC_USER   = os.environ.get("SRC_USER", "")
SRC_PWD    = os.environ.get("SRC_PWD",  "")
SRC_DSN    = os.environ.get("SRC_DSN",  "fcuatdb-scan:1521/PRODPDB")  # UAT source
SRC_SCHEMA = os.environ.get("SRC_SCHEMA", "FCPROD")

# ── TARGET DB (written into generated SQL; BS / PRODPDB on fcshr) ────────────
#    The script only CONNECTS to SOURCE — target details are embedded in SQL.
TGT_DSN    = "fcshr-scan:1521/PRODPDB"  # informational (appears in header)
TGT_SCHEMA = "FCPROD"                     # schema name used in every INSERT/UPDATE

# ── Oracle Instant Client (for thick mode; leave as-is if not needed) ────────
ORA_CLIENT_DIR = r"C:\oracle\instantclient_23_9"

# ── Output file paths ─────────────────────────────────────────────────────────
HERE       = Path(__file__).parent
OUT_DIR    = HERE / "generated_sql"          # SQL is split into files in here
OUT_CSV    = HERE / "generated_users_report.csv"
SPOOL_FILE = r"D:\SA_USERS.TXT"

# ── Performance tuning ────────────────────────────────────────────────────────
# Users per output SQL file AND per DB round-trip batch. Oracle caps a bound
# IN-list at 1000 items (ORA-01795), so keep this at or below 900.
CHUNK_SIZE = 900


# =============================================================================
# Table catalogue
# =============================================================================

# (table_name, user_id_column, display_name_column, key_cols, update_cols)
SINGLE_ROW_TABLES: List[Tuple[str, str, str, List[str], List[str]]] = [
    (
        "SSZB_USER", "USER_ID", "USER_NAME",
        ["USER_ID"],
        ["USER_PASSWORD", "SALT", "PWD_CHANGED_ON", "USER_STATUS", "STATUS_CHANGED_ON"],
    ),
    (
        "SMZB_USER", "USER_ID", "USER_NAME",
        ["USER_ID"],
        ["TIME_LEVEL", "AUTO_AUTH", "MULTIBRANCH_ACCESS", "F10_REQD", "F11_REQD", "F12_REQD"],
    ),
    (
        "SMZB_USERLOG_DETAILS", "USER_ID", "USER_ID",
        ["USER_ID"],
        ["LAST_SIGNED_ON"],
    ),
    (
        "SMZB_USER_ALLOW_UDE_CUSTOM", "USER_ID", "USER_ID",
        ["USER_ID"],
        ["ALLOW_UDE"],
    ),
    (
        "SMZB_DASHBOARD_MASTER", "USER_ID", "USER_NAME",
        ["USER_ID"],
        [],
    ),
    (
        "COZM_USER_DETAILS", "USERID", "USER_NAME",
        ["USERID"],
        [],
    ),
    (
        "SMTB_USER_CUSTOM", "USER_ID", "USER_ID",
        ["USER_ID"],
        [],
    ),
]

# (table_name, where_column, key_cols, order_by)
MULTI_ROW_TABLES: List[Tuple[str, str, List[str], str]] = [
    ("SMZB_USER_ENTITY",        "USER_ID",      ["ENTITY_ID", "USER_ID"],              "ENTITY_ID, USER_ID"),
    ("SMZB_MSGS_RIGHTS",        "USER_ROLE_ID", ["USER_ROLE_ID", "USER_ROLE_FLAG"],    "USER_ROLE_ID, USER_ROLE_FLAG"),
    ("SMZB_USER_ROLE",          "USER_ID",      ["ROLE_ID", "USER_ID", "BRANCH_CODE"], "ROLE_ID, USER_ID, BRANCH_CODE"),
    ("SMZB_USER_CENTRAL_ROLE",  "USER_ID",      ["ROLE_ID", "USER_ID"],                "ROLE_ID, USER_ID"),
    ("SMZB_USER_CENTRAL_ROLES", "USER_ID",      ["ROLE_ID", "USER_ID"],                "ROLE_ID, USER_ID"),
    ("SMZB_USER_ACCCLASS",      "USER_ID",      ["USER_ID", "ACCOUNT_CLASS"],          "USER_ID, ACCOUNT_CLASS"),
    ("SMZB_USER_BRANCHES",      "USER_ID",      ["BRANCH", "USER_ID"],                 "BRANCH, USER_ID"),
]

# ── ADDITIONAL requested PDATA tables ────────────────────────────────────────
#    These are cloned with the SAME delete-then-insert rules as the explicit
#    tables above, but their user-id column and primary key are discovered
#    automatically from the Oracle data dictionary at run time (see
#    discover_additional_specs). Add/remove table names here only.
ADDITIONAL_TABLES: List[str] = [
    "SMZB_DASHBOARD_DETAILS",
    "SMZB_QUEUE_RIGHTS",
    "SMZB_USER_ACCESS_PRODUCTS",
    "SMZB_USER_FUNC_DISALLOW",
    "SMZB_USER_GLEXCEPT",
    "SMZB_USER_GLREST",
    "SMZB_USER_GROUP",
    "SMZB_USER_HOTKEY",
    "SMZB_USER_LIMITS_ROLE",
    "SMZB_USER_PRODUCTS",
    "SMZB_USERS_FUNCTIONS",
    "SMZB_USER_DETAIL",
    "SMZM_USER_NETWORK_DET_CUSTOM",
    "SMZM_USER_SOURCE_DET_CUSTOM",
]

# When auto-detecting the "which column holds the user id" column for an
# additional table, the first of these that the table actually has is used.
# USER_ROLE_ID is after the plain user columns so rights tables fall back to it
# (as SMZB_MSGS_RIGHTS does), only when no direct USER_ID/USERID exists.
USER_COL_PREFERENCE: List[str] = ["USER_ID", "USERID", "USER_ROLE_ID", "USER_NAME"]

# USER_ROLE_FLAG value that marks a USER row (not a ROLE row) in
# SMZB_MSGS_RIGHTS / SMZB_QUEUE_RIGHTS. Only rows with this flag are deleted.
MSGS_RIGHTS_USER_FLAG = "U"

# DELETE order used before re-inserting a user in TARGET:
# exact reverse of the INSERT order (children first, SSZB_USER last).
# (table, user-id column, extra WHERE condition)
DELETE_ORDER: List[Tuple[str, str, str]] = [
    ("SMZB_USER_BRANCHES",         "USER_ID",      ""),
    ("SMZB_USER_ACCCLASS",         "USER_ID",      ""),
    ("SMZB_USER_CENTRAL_ROLES",    "USER_ID",      ""),
    ("SMZB_USER_CENTRAL_ROLE",     "USER_ID",      ""),
    ("SMZB_USER_ROLE",             "USER_ID",      ""),
    ("SMZB_USER_ENTITY",           "USER_ID",      ""),
    ("SMZB_MSGS_RIGHTS",           "USER_ROLE_ID", f" AND USER_ROLE_FLAG = '{MSGS_RIGHTS_USER_FLAG}'"),
    ("SMTB_USER_CUSTOM",           "USER_ID",      ""),
    ("COZM_USER_DETAILS",          "USERID",       ""),
    ("SMZB_DASHBOARD_MASTER",      "USER_ID",      ""),
    ("SMZB_USER_ALLOW_UDE_CUSTOM", "USER_ID",      ""),
    ("SMZB_USERLOG_DETAILS",       "USER_ID",      ""),
    ("SMZB_USER",                  "USER_ID",      ""),
    ("SSZB_USER",                  "USER_ID",      ""),
]


# =============================================================================
# Data model
# =============================================================================

@dataclass
class UserReport:
    user_id:       str
    display_name:  str  = ""
    has_sszb_user: bool = False
    has_smzb_user: bool = False
    has_entity:    bool = False
    normal_roles:  int  = 0
    central_roles: int  = 0
    branches:      int  = 0
    accclasses:    int  = 0


@dataclass
class TableSpec:
    """Run-time-discovered copy rule for an ADDITIONAL_TABLES table."""
    table:       str
    user_col:    str                      # column that holds the user id
    key_cols:    List[str]                # primary key (for WHERE NOT EXISTS)
    order_by:    str                      # deterministic ORDER BY
    where_extra: str = ""                 # e.g. " AND USER_ROLE_FLAG = 'U'"


# Discovered at start-up in main(), read by the SQL builders. Kept as a module
# global (like RUN_DATE / OUT_DIR) to avoid threading it through every call.
ADDITIONAL_SPECS: List[TableSpec] = []


# =============================================================================
# Excel reader
# =============================================================================

def _normalize_cell(value: Any) -> Optional[str]:
    """Convert an Excel cell value to a clean string, or None if empty."""
    if value is None:
        return None
    if isinstance(value, str):
        return value.strip() or None
    if isinstance(value, bool):
        return None                          # Excel boolean — skip
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value).strip()
    if isinstance(value, (datetime, date)):
        return None                          # date cells are not user IDs
    return str(value).strip() or None


def _find_user_id_column(headers: List[Any]) -> Optional[int]:
    """
    Return the 0-based index of the column whose header matches one of the
    known user-ID aliases (case-insensitive).  Returns None if not found.
    """
    normalised = [
        str(h).strip().upper() if h is not None else ""
        for h in headers
    ]
    for alias in EXCEL_USER_ID_COLUMN_ALIASES:
        try:
            return normalised.index(alias.upper())
        except ValueError:
            continue
    return None


def read_user_ids_from_excel(path: Path, sheet=EXCEL_SHEET) -> List[str]:
    """
    Open the Excel workbook at *path*, locate the user-ID column and return
    a deduplicated list of non-empty user ID strings.

    Strategy
    --------
    1. Try to find a header row whose column matches EXCEL_USER_ID_COLUMN_ALIASES.
       The header row is assumed to be within the first 10 rows.
    2. If no recognised header is found, fall back to reading the FIRST column
       (column A) and skipping any cell that looks like a header word.
    """
    if not path.exists():
        raise SystemExit(
            f"❌  Excel file not found:\n   {path}\n"
            "   Update EXCEL_FILE in the script configuration section."
        )

    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise SystemExit(f"❌  Cannot open Excel file:\n   {exc}") from exc

    # Select sheet
    if isinstance(sheet, int):
        try:
            ws = wb.worksheets[sheet]
        except IndexError:
            raise SystemExit(
                f"❌  Excel file has no sheet at index {sheet}.\n"
                f"   Available sheets: {wb.sheetnames}"
            )
    else:
        if sheet not in wb.sheetnames:
            raise SystemExit(
                f"❌  Sheet '{sheet}' not found in workbook.\n"
                f"   Available sheets: {wb.sheetnames}"
            )
        ws = wb[sheet]

    # Read all rows into memory (read_only worksheet yields tuples of cells)
    all_rows: List[List[Any]] = []
    for row in ws.iter_rows(values_only=True):
        all_rows.append(list(row))

    wb.close()

    if not all_rows:
        raise SystemExit("❌  The Excel sheet appears to be empty.")

    # ── Step 1: locate header row ────────────────────────────────────────────
    header_row_idx: Optional[int] = None
    uid_col_idx:    Optional[int] = None

    for row_idx, row in enumerate(all_rows[:10]):          # scan first 10 rows
        col_idx = _find_user_id_column(row)
        if col_idx is not None:
            header_row_idx = row_idx
            uid_col_idx    = col_idx
            break

    # ── Step 2: extract user IDs ─────────────────────────────────────────────
    seen: set   = set()
    result: List[str] = []

    if header_row_idx is not None and uid_col_idx is not None:
        # Normal path: skip header row, read uid_col_idx column
        col_header = str(all_rows[header_row_idx][uid_col_idx]).strip()
        print(f"  ✔  Excel: found user-ID column '{col_header}' "
              f"at column {uid_col_idx + 1}, header row {header_row_idx + 1}")

        for row in all_rows[header_row_idx + 1:]:
            if uid_col_idx >= len(row):
                continue
            uid = _normalize_cell(row[uid_col_idx])
            if uid and uid not in seen:
                seen.add(uid)
                result.append(uid)
    else:
        # Fallback: read column A, skip header-like words
        print("  ⚠  No recognised user-ID header found in the first 10 rows.")
        print("     Falling back to reading Column A and skipping header words.")
        skip_words = {a.upper() for a in EXCEL_USER_ID_COLUMN_ALIASES}

        for row in all_rows:
            if not row:
                continue
            uid = _normalize_cell(row[0])
            if uid and uid.upper() not in skip_words and uid not in seen:
                seen.add(uid)
                result.append(uid)

    if not result:
        raise SystemExit(
            "❌  No user IDs could be extracted from the Excel file.\n"
            f"   File  : {path}\n"
            f"   Sheet : index {sheet}\n"
            "   Check that the column header matches one of:\n"
            + "\n".join(f"       {a}" for a in EXCEL_USER_ID_COLUMN_ALIASES)
        )

    return result


# =============================================================================
# User list helpers  (text / CLI fallback)
# =============================================================================

def parse_user_ids(raw: str) -> List[str]:
    """Parse user IDs from a comma/space/semicolon/newline separated string."""
    parts = re.split(r"[\s,;]+", raw.strip())
    seen: set = set()
    result: List[str] = []
    skip = {a.upper() for a in EXCEL_USER_ID_COLUMN_ALIASES}
    for item in parts:
        uid = _normalize_cell(item)
        if not uid:
            continue
        if uid.upper() in skip:
            continue
        if uid not in seen:
            seen.add(uid)
            result.append(uid)
    return result


def read_user_ids_from_text_file(path: Path) -> List[str]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return parse_user_ids(text)


def load_user_ids(
    excel_path: Optional[Path],
    users:       Optional[str],
    users_file:  Optional[str],
) -> List[str]:
    """
    Priority order
    --------------
    1. --users  (explicit CLI list)
    2. --users-file  (text file)
    3. --excel / EXCEL_FILE  (Excel workbook)  ← default source
    """
    if users:
        print("  ℹ  User source: --users flag (CLI list)")
        return parse_user_ids(users)

    if users_file:
        p = Path(users_file)
        if not p.exists():
            raise SystemExit(f"❌  Users file not found: {p}")
        print(f"  ℹ  User source: text file  →  {p}")
        return read_user_ids_from_text_file(p)

    # Default: Excel
    ep = excel_path or EXCEL_FILE
    print(f"  ℹ  User source: Excel file  →  {ep}")
    return read_user_ids_from_excel(ep)


# =============================================================================
# Oracle helpers  (all reads use SRC_SCHEMA)
# =============================================================================

def _init_thick_mode() -> None:
    try:
        oracledb.init_oracle_client(lib_dir=ORA_CLIENT_DIR)
    except Exception:
        pass   # Already initialised, or falling back to thin mode


def connect_oracle() -> "oracledb.Connection":
    if not SRC_USER or not SRC_PWD:
        raise SystemExit(
            "❌  Source DB credentials are not set.\n"
            "   Set the SRC_USER and SRC_PWD environment variables before "
            "running (SRC_DSN / SRC_SCHEMA optional),\n"
            "   e.g.  set SRC_USER=your_user & set SRC_PWD=your_password  (Windows)\n"
            "         export SRC_USER=your_user SRC_PWD=your_password      (Linux/macOS)"
        )
    _init_thick_mode()
    return oracledb.connect(user=SRC_USER, password=SRC_PWD, dsn=SRC_DSN)


def _fetch(
    conn:  "oracledb.Connection",
    sql:   str,
    binds: Optional[list] = None,
) -> List[Dict[str, Any]]:
    """Execute a SELECT and return rows as uppercase-keyed dicts."""
    with conn.cursor() as cur:
        cur.arraysize = 2000          # fewer network round-trips on big result sets
        cur.execute(sql, binds or [])
        if cur.description is None:
            return []
        cols = [d[0].upper() for d in cur.description]
        return [dict(zip(cols, row)) for row in cur.fetchall()]


def _chunked(seq: List[Any], size: int) -> Iterator[List[Any]]:
    for i in range(0, len(seq), size):
        yield seq[i:i + size]


def get_existing_tables(
    conn:   "oracledb.Connection",
    tables: List[str],
    schema: str,
) -> set:
    """
    Check which of *tables* exist in *schema* — ONE query for all of them,
    instead of one query per table per user.
    """
    if not tables:
        return set()
    names = sorted({t.upper() for t in tables})
    placeholders = ",".join(f":{i + 2}" for i in range(len(names)))
    rows = _fetch(
        conn,
        f"SELECT TABLE_NAME FROM ALL_TABLES "
        f"WHERE OWNER = :1 AND TABLE_NAME IN ({placeholders})",
        [schema.upper(), *names],
    )
    return {r["TABLE_NAME"] for r in rows}


def get_table_columns(
    conn:   "oracledb.Connection",
    table:  str,
    schema: str,
) -> set:
    """Return the set of (uppercased) column names for *table* in *schema*."""
    rows = _fetch(
        conn,
        "SELECT COLUMN_NAME FROM ALL_TAB_COLUMNS "
        "WHERE OWNER = :1 AND TABLE_NAME = :2",
        [schema.upper(), table.upper()],
    )
    return {r["COLUMN_NAME"] for r in rows}


def get_primary_key_cols(
    conn:   "oracledb.Connection",
    table:  str,
    schema: str,
) -> List[str]:
    """
    Return the primary-key column names of *table* (in key order), or an
    empty list if the table has no primary key.
    """
    rows = _fetch(
        conn,
        "SELECT cc.COLUMN_NAME "
        "FROM ALL_CONSTRAINTS c "
        "JOIN ALL_CONS_COLUMNS cc "
        "  ON c.OWNER = cc.OWNER AND c.CONSTRAINT_NAME = cc.CONSTRAINT_NAME "
        "WHERE c.OWNER = :1 AND c.TABLE_NAME = :2 "
        "  AND c.CONSTRAINT_TYPE = 'P' "
        "ORDER BY cc.POSITION",
        [schema.upper(), table.upper()],
    )
    return [r["COLUMN_NAME"] for r in rows]


def discover_additional_specs(
    conn:            "oracledb.Connection",
    tables:          List[str],
    existing_tables: set,
    schema:          str,
) -> List[TableSpec]:
    """
    Build a TableSpec for each ADDITIONAL_TABLES table that exists, by reading
    its columns and primary key from the Oracle data dictionary:

      - user_col    : first of USER_COL_PREFERENCE the table actually has
      - where_extra : " AND USER_ROLE_FLAG = 'U'" when the table is keyed on
                      USER_ROLE_ID and has a USER_ROLE_FLAG column (rights
                      tables — same rule as SMZB_MSGS_RIGHTS)
      - key_cols    : the table's primary key, else [user_col]
      - order_by    : the key columns (deterministic output)

    A table that exists but has none of the preferred user columns is skipped
    with a warning (so the run never breaks on it).
    """
    specs: List[TableSpec] = []
    for table in tables:
        if table not in existing_tables:
            continue   # already reported as missing by get_existing_tables

        cols = get_table_columns(conn, table, schema)
        user_col = next((c for c in USER_COL_PREFERENCE if c in cols), None)
        if not user_col:
            print(
                f"  ⚠  {table}: no user-id column found "
                f"(looked for {', '.join(USER_COL_PREFERENCE)}) — skipped."
            )
            continue

        where_extra = ""
        if user_col == "USER_ROLE_ID" and "USER_ROLE_FLAG" in cols:
            where_extra = f" AND USER_ROLE_FLAG = '{MSGS_RIGHTS_USER_FLAG}'"

        pk = get_primary_key_cols(conn, table, schema)
        key_cols = pk if pk else [user_col]
        order_by = ", ".join(key_cols)

        specs.append(
            TableSpec(
                table=table,
                user_col=user_col,
                key_cols=key_cols,
                order_by=order_by,
                where_extra=where_extra,
            )
        )
        flag_note = " (USER_ROLE_FLAG='U')" if where_extra else ""
        print(
            f"  ✔  {table}: key={user_col}{flag_note}, "
            f"pk=({', '.join(key_cols)})"
        )
    return specs


def fetch_table_batch(
    conn:        "oracledb.Connection",
    table:       str,
    uid_col:     str,
    user_ids:    List[str],
    order_by:    Optional[str] = None,
    where_extra: str = "",
) -> Dict[str, List[Dict[str, Any]]]:
    """
    Fetch rows from *table* for a whole batch of user_ids using a single
    'WHERE uid_col IN (:1, :2, ...)' query (chunked at 900 to respect
    Oracle's ORA-01795 bind-list limit), instead of one query per user.
    *where_extra* is appended to the WHERE clause verbatim (e.g. the
    USER_ROLE_FLAG filter). Returns {user_id: [rows...]}.
    """
    result: Dict[str, List[Dict[str, Any]]] = {uid: [] for uid in user_ids}
    order_clause = f" ORDER BY {order_by}" if order_by else ""
    uid_col_u = uid_col.upper()

    for chunk in _chunked(user_ids, 900):
        placeholders = ",".join(f":{i + 1}" for i in range(len(chunk)))
        sql = (
            f"SELECT * FROM {SRC_SCHEMA}.{table} "
            f"WHERE {uid_col} IN ({placeholders}){where_extra}{order_clause}"
        )
        for row in _fetch(conn, sql, list(chunk)):
            uid = row.get(uid_col_u)
            result.setdefault(uid, []).append(row)

    return result


def fetch_bundles_batch(
    conn:            "oracledb.Connection",
    user_ids:        List[str],
    existing_tables: set,
) -> List[Dict[str, Any]]:
    """
    Build the per-user bundles for a whole chunk of users at once, using
    batched (IN-list) queries — one or two queries per table for the whole
    chunk, instead of one query per table per user.
    """
    bundles: Dict[str, Dict[str, Any]] = {
        uid: {"user_id": uid, "display_name": ""} for uid in user_ids
    }

    for table, uid_col, _display_col, *_ in SINGLE_ROW_TABLES:
        if table not in existing_tables:
            for b in bundles.values():
                b[table] = []
            continue
        data = fetch_table_batch(conn, table, uid_col, user_ids)
        for uid, b in bundles.items():
            b[table] = data.get(uid, [])

    for table, where_col, _key_cols, order_by in MULTI_ROW_TABLES:
        if table not in existing_tables:
            for b in bundles.values():
                b[table] = []
            continue
        data = fetch_table_batch(conn, table, where_col, user_ids, order_by=order_by)
        for uid, b in bundles.items():
            b[table] = data.get(uid, [])

    # Additional (auto-discovered) tables
    for spec in ADDITIONAL_SPECS:
        if spec.table not in existing_tables:
            for b in bundles.values():
                b[spec.table] = []
            continue
        data = fetch_table_batch(
            conn, spec.table, spec.user_col, user_ids,
            order_by=spec.order_by, where_extra=spec.where_extra,
        )
        for uid, b in bundles.items():
            b[spec.table] = data.get(uid, [])

    # Resolve display name from the first single-row table that has one
    for uid, b in bundles.items():
        for table, _uid_col, display_col, *_ in SINGLE_ROW_TABLES:
            rows = b.get(table) or []
            if rows:
                row  = rows[0]
                name = str(
                    row.get(display_col)
                    or row.get("USER_ID")
                    or row.get("USERID")
                    or ""
                ).strip()
                if name:
                    b["display_name"] = name
                    break

    return list(bundles.values())


# =============================================================================
# SQL value serialisation
# =============================================================================

def _date_lit() -> str:
    """
    Return a TO_DATE literal for RUN_DATE — NLS_DATE_FORMAT independent.
    Format: TO_DATE('2026-06-10', 'YYYY-MM-DD')
    """
    return f"TO_DATE('{RUN_DATE.strftime('%Y-%m-%d')}', 'YYYY-MM-DD')"


def _lit(value: Any) -> str:
    """Serialise a Python value to an Oracle SQL literal."""
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "'Y'" if value else "'N'"
    if isinstance(value, (date, datetime)):
        # Always stamp with the user-supplied RUN_DATE, never SYSDATE
        return _date_lit()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    if isinstance(value, bytes):
        return f"HEXTORAW('{value.hex().upper()}')"
    return f"'{str(value).replace(chr(39), chr(39) * 2)}'"


# =============================================================================
# SQL statement builders  (all table refs use TGT_SCHEMA)
# =============================================================================

def _tgt(table: str) -> str:
    return f"{TGT_SCHEMA}.{table}"


def _insert_if_missing(
    table:    str,
    row:      Dict[str, Any],
    key_cols: List[str],
) -> str:
    cols     = list(row.keys())
    col_list = ",\n    ".join(cols)
    val_list = ",\n    ".join(_lit(row[c]) for c in cols)
    # Guard against a discovered key column that is not actually in the row
    # (shouldn't happen, but keeps the generated SQL valid if it does).
    key_cols = [k for k in key_cols if k in row] or cols
    where_pk = " AND ".join(f"{k} = {_lit(row[k])}" for k in key_cols)
    return (
        f"INSERT INTO {_tgt(table)} (\n"
        f"    {col_list}\n"
        f")\n"
        f"SELECT\n"
        f"    {val_list}\n"
        f"FROM DUAL\n"
        f"WHERE NOT EXISTS (\n"
        f"    SELECT 1 FROM {_tgt(table)}\n"
        f"    WHERE {where_pk}\n"
        f");"
    )


def _delete_user_lines(user_id: str, existing_tables: set) -> List[str]:
    """
    PL/SQL lines that remove every row the user has in TARGET (children
    first, SSZB_USER last), adding the deleted row count to v_deleted.
    Only tables that exist (same set used for the INSERTs) are touched.
    The auto-discovered ADDITIONAL tables are deleted first (they are all
    user-detail children; none must survive the core SMZB_USER/SSZB_USER).
    """
    uid = _lit(user_id)
    out: List[str] = []

    for spec in ADDITIONAL_SPECS:
        if spec.table not in existing_tables:
            continue
        out += [
            f"v_step := {_lit('delete ' + _tgt(spec.table))};",
            f"DELETE FROM {_tgt(spec.table)} "
            f"WHERE {spec.user_col} = {uid}{spec.where_extra};",
            "v_deleted := v_deleted + SQL%ROWCOUNT;",
        ]

    for table, col, extra in DELETE_ORDER:
        if table not in existing_tables:
            continue
        out += [
            f"v_step := {_lit('delete ' + _tgt(table))};",
            f"DELETE FROM {_tgt(table)} WHERE {col} = {uid}{extra};",
            "v_deleted := v_deleted + SQL%ROWCOUNT;",
        ]
    return out


def _found_in_source(bundle: Dict[str, Any], existing_tables: set) -> bool:
    """
    A user is only deleted + re-added when SOURCE really has it (its
    SSZB_USER row). Otherwise TARGET is left untouched — deleting a user we
    cannot re-create would just lose it.
    """
    if "SSZB_USER" in existing_tables:
        return bool(bundle.get("SSZB_USER"))
    return any(bundle.get(t) for t, *_ in DELETE_ORDER)


# =============================================================================
# Per-user SQL block builder
#   One PL/SQL block per user = one transaction:
#     0. check TARGET: delete the user everywhere if it exists
#     1. insert it fresh from SOURCE
#     COMMIT  — or ROLLBACK on any error (user left as it was in TARGET)
# =============================================================================

def build_user_sql(bundle: Dict[str, Any], existing_tables: set) -> Tuple[str, UserReport]:
    user_id      = bundle["user_id"]
    display_name = bundle.get("display_name") or ""
    report       = UserReport(user_id=user_id, display_name=display_name)
    lines: List[str] = []
    tag          = _lit(f"[{user_id}]")

    lines += [
        "-- " + "=" * 60,
        f"-- USER : {user_id}" + (f"  ({display_name})" if display_name else ""),
        "-- " + "=" * 60,
    ]

    # Advisory notes
    if not bundle.get("SMZB_USER_ROLE"):
        lines.append(f"-- NOTE: no NORMAL_ROLE rows found for {user_id}")
    if not (bundle.get("SMZB_USER_CENTRAL_ROLE") or bundle.get("SMZB_USER_CENTRAL_ROLES")):
        lines.append(f"-- NOTE: no CENTRAL_ROLE rows found for {user_id}")
    if not bundle.get("SMZB_USER_ENTITY"):
        lines.append(f"-- NOTE: no ENTITY rows found for {user_id}")

    if not _found_in_source(bundle, existing_tables):
        # Nothing to re-create from SOURCE -> do NOT delete the TARGET user.
        lines += [
            f"-- SKIPPED: {user_id} not found in SOURCE (no SSZB_USER row) — "
            f"TARGET left untouched.",
            f"PROMPT [{user_id}] SKIPPED - not found in SOURCE (TARGET left untouched)",
            "",
        ]
        return "\n".join(lines) + "\n", report

    lines += [
        "DECLARE",
        "  v_deleted NUMBER := 0;",
        "  v_step    VARCHAR2(200);",
        "BEGIN",
        f"-- 0. make sure {user_id} is NOT in TARGET: delete it first if it exists",
        *_delete_user_lines(user_id, existing_tables),
        "",
        f"-- 1. insert {user_id} fresh from SOURCE",
    ]

    def _ins(table: str, row: Dict[str, Any], key_cols: List[str], label: str = "") -> None:
        lines.append(f"v_step := {_lit('insert ' + _tgt(table) + label)};")
        lines.append(_insert_if_missing(table, row, key_cols))
        lines.append("")

    def _emit_single(table: str, key_cols: List[str]) -> None:
        rows = bundle.get(table) or []
        if rows:
            _ins(table, rows[0], key_cols)

    # 1. SSZB_USER
    _emit_single("SSZB_USER", ["USER_ID"])
    if bundle.get("SSZB_USER"):
        report.has_sszb_user = True

    # 2. SMZB_USER
    _emit_single("SMZB_USER", ["USER_ID"])
    if bundle.get("SMZB_USER"):
        report.has_smzb_user = True

    # 3. SMZB_USERLOG_DETAILS  (LAST_SIGNED_ON = RUN_DATE, like every date column)
    _emit_single("SMZB_USERLOG_DETAILS", ["USER_ID"])

    # 4. SMZB_USER_ALLOW_UDE_CUSTOM
    _emit_single("SMZB_USER_ALLOW_UDE_CUSTOM", ["USER_ID"])

    # 5. SMZB_DASHBOARD_MASTER
    _emit_single("SMZB_DASHBOARD_MASTER", ["USER_ID"])

    # 6. COZM_USER_DETAILS
    _emit_single("COZM_USER_DETAILS", ["USERID"])

    # 7. SMTB_USER_CUSTOM
    _emit_single("SMTB_USER_CUSTOM", ["USER_ID"])

    # 8. SMZB_MSGS_RIGHTS  (first row only — as before)
    if bundle.get("SMZB_MSGS_RIGHTS"):
        _ins("SMZB_MSGS_RIGHTS", bundle["SMZB_MSGS_RIGHTS"][0],
             ["USER_ROLE_ID", "USER_ROLE_FLAG"])

    # 9. Multi-row tables
    for table, _, key_cols, _ in MULTI_ROW_TABLES:
        if table == "SMZB_MSGS_RIGHTS":
            continue   # already handled above

        rows = bundle.get(table) or []
        if not rows:
            continue

        if   table == "SMZB_USER_ENTITY":
            report.has_entity    = True
        elif table == "SMZB_USER_ROLE":
            report.normal_roles  = len(rows)
        elif table in {"SMZB_USER_CENTRAL_ROLE", "SMZB_USER_CENTRAL_ROLES"}:
            report.central_roles += len(rows)
        elif table == "SMZB_USER_ACCCLASS":
            report.accclasses    = len(rows)
        elif table == "SMZB_USER_BRANCHES":
            report.branches      = len(rows)

        for i, row in enumerate(rows, 1):
            _ins(table, row, key_cols, f" row {i}/{len(rows)}")

    # 10. Additional (auto-discovered) PDATA tables — inserted after the core
    #     SMZB_USER / SSZB_USER rows exist, so any FK to them is satisfied.
    for spec in ADDITIONAL_SPECS:
        rows = bundle.get(spec.table) or []
        for i, row in enumerate(rows, 1):
            _ins(spec.table, row, spec.key_cols, f" row {i}/{len(rows)}")

    lines += [
        "COMMIT;",
        "IF v_deleted > 0 THEN",
        f"  DBMS_OUTPUT.PUT_LINE({tag} || ' OK - existed in TARGET (' || v_deleted"
        f" || ' rows deleted) -> re-added');",
        "ELSE",
        f"  DBMS_OUTPUT.PUT_LINE({tag} || ' OK - not in TARGET -> added');",
        "END IF;",
        "EXCEPTION",
        "  WHEN OTHERS THEN",
        "    ROLLBACK;   -- undo this user's delete + inserts: TARGET left as it was",
        f"    DBMS_OUTPUT.PUT_LINE({tag} || ' FAILED at ' || v_step"
        f" || ' -> ROLLED BACK (target unchanged): ' || SQLERRM);",
        "END;",
        "/",
    ]

    return "\n".join(lines).rstrip() + "\n", report


# =============================================================================
# Full script assembler
# =============================================================================

def build_script(
    bundles:         List[Dict[str, Any]],
    existing_tables: set,
    part_num:        int = 1,
    total_parts:     int = 1,
) -> Tuple[str, List[UserReport]]:
    """
    Build the SQL script for ONE chunk of users. Each chunk becomes its own
    file (see process_users_in_chunks), so this only ever holds a bounded
    number of user blocks (CHUNK_SIZE) in memory at a time.
    """
    reports:      List[UserReport] = []
    user_blocks:  List[str]        = []
    total_normal  = 0
    total_central = 0

    for bundle in bundles:
        block, report = build_user_sql(bundle, existing_tables)
        reports.append(report)
        total_normal  += report.normal_roles
        total_central += report.central_roles
        user_blocks.append(block)

    ts = datetime.now().strftime("%Y-%m-%d %H:%M")
    part_spool = f"{SPOOL_FILE.rsplit('.', 1)[0]}_part{part_num:03d}.TXT" \
        if "." in SPOOL_FILE else f"{SPOOL_FILE}_part{part_num:03d}"

    header = (
        "SET DEFINE OFF;\n"
        "SET SERVEROUTPUT ON SIZE UNLIMITED;\n"
        "WHENEVER SQLERROR EXIT SQL.SQLCODE ROLLBACK;\n"
        f"SPOOL {part_spool}\n"
        "\n"
        "-- ================================================================\n"
        "-- Oracle User Provisioning Script\n"
        f"-- Part            : {part_num} of {total_parts}\n"
        f"-- Generated       : {ts}\n"
        f"-- Source DB       : {SRC_DSN}   (data READ from here)\n"
        f"-- Source schema   : {SRC_SCHEMA}\n"
        f"-- Target DB       : {TGT_DSN}   (run this script here)\n"
        f"-- Target schema   : {TGT_SCHEMA}   <-- VERIFY before running!\n"
        "-- ----------------------------------------------------------------\n"
        f"-- Users in this part    : {len(bundles)}\n"
        f"-- NORMAL_ROLE rows      : {total_normal}\n"
        f"-- CENTRAL_ROLE rows     : {total_central}\n"
        "-- ----------------------------------------------------------------\n"
        f"-- All date columns use {_date_lit()} format — NLS-safe.\n"
        "-- ----------------------------------------------------------------\n"
        "-- EVERY USER: checked in TARGET first; if it exists it is DELETED\n"
        "--   from all tables and then INSERTED fresh from SOURCE.\n"
        "--   One PL/SQL block per user = one transaction. On any error that\n"
        "--   user is ROLLED BACK (left as it was) and the script continues.\n"
        "--   Search the spool for 'FAILED'.\n"
        "-- ================================================================\n"
        "\n"
    )

    footer = "\nCOMMIT;\nSPOOL OFF;\n"
    return header + "\n".join(user_blocks).rstrip() + footer, reports


# =============================================================================
# Chunked processing — bounded memory, batched DB reads, one file per chunk
# =============================================================================

def process_users_in_chunks(
    conn:            "oracledb.Connection",
    user_ids:        List[str],
    existing_tables: set,
    out_dir:         Path,
    chunk_size:      int = CHUNK_SIZE,
) -> List[UserReport]:
    out_dir.mkdir(parents=True, exist_ok=True)

    chunks      = list(_chunked(user_ids, chunk_size))
    total_parts = len(chunks)
    all_reports: List[UserReport] = []
    written_files: List[Path] = []

    t_start = time.time()
    for part_num, chunk in enumerate(chunks, 1):
        t0 = time.time()

        bundles = fetch_bundles_batch(conn, chunk, existing_tables)
        sql_text, reports = build_script(bundles, existing_tables, part_num, total_parts)

        out_path = out_dir / f"generated_users_part{part_num:03d}.sql"
        out_path.write_text(sql_text, encoding="utf-8")
        written_files.append(out_path)
        all_reports.extend(reports)

        elapsed   = time.time() - t0
        total_el  = time.time() - t_start
        avg       = total_el / part_num
        remaining = avg * (total_parts - part_num)
        print(
            f"  [Part {part_num:>3}/{total_parts}] {len(chunk):>4} users → "
            f"{out_path.name}  ({elapsed:.1f}s, ETA {remaining/60:.1f} min)"
        )

    total_el = time.time() - t_start
    print(
        f"\n  ✔  {len(user_ids)} users written to {total_parts} file(s) in "
        f"{out_dir.resolve()}  (total time: {total_el/60:.1f} min)"
    )
    return all_reports


# =============================================================================
# CSV report
# =============================================================================

_REPORT_FIELDS = [
    "user_id", "display_name",
    "has_sszb_user", "has_smzb_user", "has_entity",
    "normal_roles", "central_roles", "branches", "accclasses",
]


def write_report_csv(reports: List[UserReport], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=_REPORT_FIELDS)
        writer.writeheader()
        for r in reports:
            writer.writerow({f: getattr(r, f) for f in _REPORT_FIELDS})


# =============================================================================
# CLI helpers
# =============================================================================

def _parse_date_arg(raw: str) -> date:
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    raise ValueError("Use YYYY-MM-DD or DD/MM/YYYY")


def _prompt_run_date() -> date:
    today_str = date.today().strftime("%Y-%m-%d")
    print()
    print("─" * 60)
    print("  DATE STAMP FOR GENERATED SQL")
    print("  This date is written into EVERY date column in the output")
    print("  using TO_DATE with an explicit format mask.")
    print(f"  Press Enter to use today: {today_str}")
    print("─" * 60)

    while True:
        raw = input("  Enter date [YYYY-MM-DD or DD/MM/YYYY]: ").strip()
        if not raw:
            chosen = date.today()
            print(f"  ✔  Using today: {chosen.strftime('%Y-%m-%d')}\n")
            return chosen
        try:
            chosen = _parse_date_arg(raw)
            print(f"  ✔  Date set to: {chosen.strftime('%Y-%m-%d')}\n")
            return chosen
        except ValueError:
            print(f"  ❌  Could not parse '{raw}'. Use YYYY-MM-DD or DD/MM/YYYY.")


def _build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate Oracle provisioning SQL for a list of users.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "User-ID source priority:\n"
            "  1. --users  (explicit list, highest priority)\n"
            "  2. --users-file  (plain text file)\n"
            "  3. --excel / EXCEL_FILE constant  (default, lowest priority)\n"
        ),
    )
    parser.add_argument(
        "--excel",
        metavar="PATH",
        help=(
            f"Path to the Excel workbook containing user IDs "
            f"(default: {EXCEL_FILE})."
        ),
    )
    parser.add_argument(
        "--sheet",
        default=None,
        metavar="SHEET",
        help=(
            "Sheet name or 0-based index to read from the workbook "
            "(default: first sheet)."
        ),
    )
    parser.add_argument(
        "--users",
        metavar="U1,U2,...",
        help="Comma/space/semicolon separated user IDs — overrides Excel.",
    )
    parser.add_argument(
        "--users-file",
        metavar="FILE",
        help="Text file with user IDs (one per line or comma-separated).",
    )
    parser.add_argument(
        "--date",
        dest="run_date",
        metavar="DATE",
        help="Date stamp for all date columns (YYYY-MM-DD or DD/MM/YYYY).",
    )
    parser.add_argument(
        "--output-dir",
        default=str(OUT_DIR),
        metavar="DIR",
        help="Directory to write the split generated_users_partNNN.sql files into.",
    )
    parser.add_argument(
        "--output-csv",
        default=str(OUT_CSV),
        metavar="FILE",
        help="Path to write the summary CSV report.",
    )
    parser.add_argument(
        "--chunk-size",
        type=int,
        default=CHUNK_SIZE,
        metavar="N",
        help=(
            f"Users per output SQL file / per DB batch query (default: "
            f"{CHUNK_SIZE}; Oracle caps a bound IN-list at 1000)."
        ),
    )
    return parser


# =============================================================================
# Entry point
# =============================================================================

def main() -> int:
    global RUN_DATE, OUT_DIR, OUT_CSV, ADDITIONAL_SPECS

    parser = _build_arg_parser()
    args   = parser.parse_args()

    OUT_DIR    = Path(args.output_dir)
    OUT_CSV    = Path(args.output_csv)
    chunk_size = max(1, min(args.chunk_size, 900))   # Oracle IN-list cap

    # ── Banner ────────────────────────────────────────────────────────────────
    print("=" * 60)
    print("  Oracle User Provisioning SQL Generator")
    print("  Mode: user exists in TARGET -> DELETE first, then ADD")
    print("=" * 60)
    print(f"  Source DB    : {SRC_DSN}  (schema: {SRC_SCHEMA})")
    print(f"  Target DB    : {TGT_DSN}  (schema: {TGT_SCHEMA})")
    print(f"  Output dir   : {OUT_DIR}  (split SQL files)")
    print(f"  Output CSV   : {OUT_CSV}")
    print(f"  Chunk size   : {chunk_size} users/file")

    # ── Date ──────────────────────────────────────────────────────────────────
    if args.run_date:
        try:
            RUN_DATE = _parse_date_arg(args.run_date)
            print(f"  Date         : {RUN_DATE.strftime('%Y-%m-%d')}\n")
        except ValueError as exc:
            raise SystemExit(
                f"❌  Invalid --date value: {args.run_date}\n   {exc}"
            ) from exc
    else:
        RUN_DATE = _prompt_run_date()

    # ── Resolve Excel sheet override ──────────────────────────────────────────
    sheet = EXCEL_SHEET
    if args.sheet is not None:
        try:
            sheet = int(args.sheet)
        except ValueError:
            sheet = args.sheet   # treat as sheet name

    # ── Load user IDs ─────────────────────────────────────────────────────────
    excel_path = Path(args.excel) if args.excel else None
    user_ids   = load_user_ids(excel_path, args.users, args.users_file)

    if not user_ids:
        raise SystemExit(
            "❌  No user IDs were found.\n"
            "   Check the Excel file path and column headers, or use "
            "--users / --users-file."
        )
    print(f"  ✔  Loaded {len(user_ids)} unique user ID(s).\n")

    # Preview the first few IDs
    preview = user_ids[:8]
    print("  First user IDs:")
    for uid in preview:
        print(f"    • {uid}")
    if len(user_ids) > 8:
        print(f"    … and {len(user_ids) - 8} more")
    print()

    # ── Connect to Oracle ─────────────────────────────────────────────────────
    print(f"Connecting to Oracle source DB ({SRC_DSN}) …")
    try:
        conn = connect_oracle()
    except Exception as exc:
        raise SystemExit(
            f"❌  Could not connect to Oracle:\n   {exc}"
        ) from exc
    print("  ✔  Connected.\n")

    # ── Check which tables actually exist — ONCE, not per user ─────────────────
    all_tables = sorted(
        {t for t, *_ in SINGLE_ROW_TABLES}
        | {t for t, *_ in MULTI_ROW_TABLES}
        | set(ADDITIONAL_TABLES)
    )
    print("Checking table existence …")
    existing_tables = get_existing_tables(conn, all_tables, SRC_SCHEMA)
    missing = [t for t in all_tables if t not in existing_tables]
    if missing:
        print(f"  ⚠  Not found in {SRC_SCHEMA}, will be skipped: {', '.join(missing)}")
    print(f"  ✔  {len(existing_tables)}/{len(all_tables)} tables present.\n")

    # ── Auto-configure the ADDITIONAL requested tables ─────────────────────────
    print("Discovering key columns for additional PDATA tables …")
    ADDITIONAL_SPECS = discover_additional_specs(
        conn, ADDITIONAL_TABLES, existing_tables, SRC_SCHEMA
    )
    print(
        f"  ✔  {len(ADDITIONAL_SPECS)} additional table(s) configured.\n"
    )

    # ── Fetch + generate, in bounded-memory chunks ──────────────────────────────
    print(f"Fetching data & generating SQL in batches of {chunk_size} users …")
    try:
        reports = process_users_in_chunks(
            conn, user_ids, existing_tables, OUT_DIR, chunk_size
        )
    finally:
        conn.close()

    write_report_csv(reports, OUT_CSV)

    # ── Summary ───────────────────────────────────────────────────────────────
    print("\n" + "=" * 60)
    print("  ✔  Done!")
    print(f"  Date stamped : {RUN_DATE.strftime('%Y-%m-%d')}  (all date columns)")
    print(f"  SQL files    →  {OUT_DIR.resolve()}\\generated_users_part*.sql")
    print(f"  Report CSV   →  {OUT_CSV.resolve()}")
    print("=" * 60)

    col_w = (20, 25, 6, 8, 9)
    hdr   = (
        f"  {'USER_ID':<{col_w[0]}} {'NAME':<{col_w[1]}} "
        f"{'ROLES':>{col_w[2]}} {'CENTRAL':>{col_w[3]}} {'BRANCHES':>{col_w[4]}}"
    )
    sep   = (
        f"  {'-'*col_w[0]} {'-'*col_w[1]} "
        f"{'-'*col_w[2]} {'-'*col_w[3]} {'-'*col_w[4]}"
    )
    print(f"\n{hdr}\n{sep}")
    for r in reports:
        print(
            f"  {r.user_id:<{col_w[0]}} {r.display_name[:col_w[1]-1]:<{col_w[1]}} "
            f"{r.normal_roles:>{col_w[2]}} {r.central_roles:>{col_w[3]}} "
            f"{r.branches:>{col_w[4]}}"
        )

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
